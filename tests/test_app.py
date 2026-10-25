import hashlib
from flask.testing import FlaskClient
import io
import re
import zipfile
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import app as rotation_app


class CsrfAwareTestClient(FlaskClient):
    def post(self, *args, **kwargs):
        data = dict(kwargs.get("data") or {})
        if "_csrf_token" not in data:
            with self.session_transaction() as sess:
                token = sess.get("_csrf_token")
            if not token:
                self.get("/login")
                with self.session_transaction() as sess:
                    token = sess.get("_csrf_token")
            if token:
                data["_csrf_token"] = token
        kwargs["data"] = data
        return super().post(*args, **kwargs)


def setup_client(tmp_path, monkeypatch):
    rotation_app.app.test_client_class = CsrfAwareTestClient
    monkeypatch.setenv("ROTATION_INITIAL_PASSWORD", "test-strong-password-2026")
    monkeypatch.setattr(rotation_app, "DB", str(tmp_path / "test_rotation.db"))
    monkeypatch.setattr(rotation_app.app, "_ready", False, raising=False)
    rotation_app.app.config.update(TESTING=True, SECRET_KEY="test-only-secret")
    return rotation_app.app.test_client()


def login(client):
    return client.post("/login", data={"password": "test-strong-password-2026"}, follow_redirects=True)


def test_login_page_and_language_selector(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    response = client.get("/login")
    assert response.status_code == 200
    assert "Français".encode() in response.data
    assert "English".encode() in response.data
    assert b'name="_csrf_token"' in response.data



def test_login_branding_uses_resimanage_and_navy_cyan_identity(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    response = client.get("/login")
    assert response.status_code == 200
    assert b"ResiManage" in response.data
    assert b"ResiTrack" not in response.data
    assert b"#00dff755" in response.data
    assert b"#19dfbd" in response.data
    assert b"prefers-reduced-motion:reduce" in response.data
    assert b"Tous droits r\xc3\xa9serv\xc3\xa9s" in response.data
    assert b"All rights reserved" in response.data
    assert "جميع الحقوق محفوظة".encode() in response.data
    assert "✚".encode() not in response.data


def test_login_and_protected_dashboard(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    response = login(client)
    assert response.status_code == 200
    assert "Resident".encode() in response.data or "المقيمون".encode() in response.data
    client.get("/logout")
    protected = client.get("/residents")
    assert protected.status_code == 302
    assert "/login" in protected.headers["Location"]


def test_can_add_resident_without_touching_real_database(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    response = client.post(
        "/residents",
        data={"surname": "TestSurname", "firstname": "TestFirst", "year": "2", "phone": "000"},
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"TestSurname" in response.data
    with rotation_app.db() as conn:
        row = conn.execute("SELECT surname, firstname, year FROM residents WHERE surname=?", ("TestSurname",)).fetchone()
    assert tuple(row) == ("TestSurname", "TestFirst", 2)


def test_seed_has_nineteen_terrains_and_four_periods(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM terrains").fetchone()[0] == 19
        assert conn.execute("SELECT COUNT(*) FROM periods").fetchone()[0] == 4
        optional = {r[0] for r in conn.execute("SELECT name FROM terrains WHERE kind='optional'")}
    assert optional == {"Sédation", "Maxillo-faciale", "CAC DBK", "Brûlés", "Chirurgie cardiaque"}
    assert "Gynécologie" not in optional
    with rotation_app.db() as conn:
        assert conn.execute("SELECT kind FROM terrains WHERE name='Gynécologie'").fetchone()[0] != "optional"
        assert conn.execute("SELECT kind FROM terrains WHERE name='Chirurgie générale'").fetchone()[0] == "mandatory"


def test_assignment_requires_capacity_and_blocks_same_period_duplicate(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        resident_id = conn.execute("SELECT id FROM residents WHERE surname='TestSurname'").fetchone()
        if resident_id is None:
            conn.execute(
                "INSERT INTO residents(surname, firstname, year) VALUES(?,?,?)",
                ("AssignTest", "Resident", 2),
            )
            resident_id = conn.execute("SELECT id FROM residents WHERE surname='AssignTest'").fetchone()
        rid = resident_id[0]
        terrain_id = conn.execute("SELECT id FROM terrains WHERE name='Réanimation chirurgicale'").fetchone()[0]
        second_terrain_id = conn.execute("SELECT id FROM terrains WHERE name='Urologie'").fetchone()[0]
        period_id = conn.execute("SELECT id FROM periods ORDER BY id LIMIT 1").fetchone()[0]
        year_label = conn.execute("SELECT year_label FROM periods WHERE id=?", (period_id,)).fetchone()[0]

    response = client.post(
        "/assign",
        data={"resident": str(rid), "terrain": str(terrain_id), "period": str(period_id)},
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert "حدّد السعة أولاً".encode() in response.data

    with rotation_app.db() as conn:
        conn.execute(
            "INSERT INTO capacities(year_label, terrain_id, period_id, places) VALUES(?,?,?,1)",
            (year_label, terrain_id, period_id),
        )

    response = client.post(
        "/assign",
        data={"resident": str(rid), "terrain": str(terrain_id), "period": str(period_id)},
        follow_redirects=True,
    )
    assert response.status_code == 200
    with rotation_app.db() as conn:
        assigned = conn.execute(
            "SELECT status FROM assignments WHERE resident_id=? AND period_id=?",
            (rid, period_id),
        ).fetchone()
    assert assigned is not None
    assert assigned["status"] == "confirmed"

    response = client.post(
        "/assign",
        data={"resident": str(rid), "terrain": str(second_terrain_id), "period": str(period_id)},
        follow_redirects=True,
    )
    assert response.status_code == 200
    with rotation_app.db() as conn:
        count = conn.execute(
            "SELECT COUNT(*) FROM assignments WHERE resident_id=? AND period_id=?",
            (rid, period_id),
        ).fetchone()[0]
    assert count == 1


def test_capacity_limit_is_enforced_for_two_residents(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute(
            "INSERT INTO residents(surname, firstname, year) VALUES(?,?,?)",
            ("CapacityOne", "Resident", 3),
        )
        conn.execute(
            "INSERT INTO residents(surname, firstname, year) VALUES(?,?,?)",
            ("CapacityTwo", "Resident", 2),
        )
        ids = conn.execute(
            "SELECT id, surname FROM residents WHERE surname IN ('CapacityOne','CapacityTwo')"
        ).fetchall()
        resident_ids = {r["surname"]: r["id"] for r in ids}
        terrain_id = conn.execute("SELECT id FROM terrains WHERE name='Réanimation polyvalente'").fetchone()[0]
        period_id = conn.execute("SELECT id FROM periods ORDER BY id LIMIT 1").fetchone()[0]
        year_label = conn.execute("SELECT year_label FROM periods WHERE id=?", (period_id,)).fetchone()[0]
        conn.execute(
            "INSERT INTO capacities(year_label, terrain_id, period_id, places) VALUES(?,?,?,1)",
            (year_label, terrain_id, period_id),
        )

    first = client.post(
        "/assign",
        data={"resident": str(resident_ids["CapacityOne"]), "terrain": str(terrain_id), "period": str(period_id)},
        follow_redirects=True,
    )
    assert first.status_code == 200
    second = client.post(
        "/assign",
        data={"resident": str(resident_ids["CapacityTwo"]), "terrain": str(terrain_id), "period": str(period_id)},
        follow_redirects=True,
    )
    assert second.status_code == 200
    assert "السعة مكتملة".encode() in second.data
    with rotation_app.db() as conn:
        count = conn.execute(
            "SELECT COUNT(*) FROM assignments WHERE terrain_id=? AND period_id=? AND status='confirmed'",
            (terrain_id, period_id),
        ).fetchone()[0]
    assert count == 1


def test_repeat_is_blocked_until_all_core_terrains_are_completed(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute(
            "INSERT INTO residents(surname, firstname, year) VALUES(?,?,?)",
            ("RepeatRule", "Resident", 4),
        )
        rid = conn.execute("SELECT id FROM residents WHERE surname='RepeatRule'").fetchone()[0]
        target = conn.execute(
            "SELECT id FROM terrains WHERE name='Réanimation chirurgicale'"
        ).fetchone()[0]
        period_id = conn.execute("SELECT id FROM periods ORDER BY id LIMIT 1").fetchone()[0]
        year_label = conn.execute("SELECT year_label FROM periods WHERE id=?", (period_id,)).fetchone()[0]
        conn.execute(
            "INSERT INTO capacities(year_label, terrain_id, period_id, places) VALUES(?,?,?,2)",
            (year_label, target, period_id),
        )
        conn.execute(
            "INSERT INTO history(resident_id, terrain_id, year_label, period, date_done) VALUES(?,?,?,?,?)",
            (rid, target, "2025-2026", "01 Fév - 30 Avr", "2026-04-30"),
        )

    with rotation_app.app.test_request_context("/"):
        with rotation_app.db() as conn:
            allowed, message, status = rotation_app.check(conn, rid, target, period_id, year_label)
            assert allowed is False
            assert "التربص منجز سابقاً" in message

    with rotation_app.db() as conn:
        conn.execute(
            """
            INSERT INTO history(resident_id, terrain_id, year_label, period, date_done)
            SELECT ?, id, '2025-2026', 'historique', '2026-04-30'
            FROM terrains
            WHERE kind IN ('mandatory', 'basic')
              AND id NOT IN (SELECT terrain_id FROM history WHERE resident_id=?)
            """,
            (rid, rid),
        )

    with rotation_app.app.test_request_context("/"):
        with rotation_app.db() as conn:
            allowed, message, status = rotation_app.check(conn, rid, target, period_id, year_label)
    assert allowed is True
    assert status == "confirmed"
    assert "يمكن التكرار" in message

def test_resident_profile_and_objectives_are_saved_and_audited(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname, firstname, year) VALUES(?,?,?)", ("Profile", "Resident", 2))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Profile'").fetchone()[0]
        terrain_id = conn.execute("SELECT id FROM terrains WHERE name='Réanimation chirurgicale'").fetchone()[0]
    response = client.post("/profile/" + str(rid), data={
        "birth_place": "Tizi Ouzou", "birth_date": "1998-05-02",
        "phone": "0550000000", "photo_path": ""
    }, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        row = conn.execute("SELECT birth_place,birth_date,phone FROM residents WHERE id=?", (rid,)).fetchone()
        assert tuple(row) == ("Tizi Ouzou", "1998-05-02", "0550000000")
    response = client.post("/objectives", data={
        "terrain": str(terrain_id), "study_year": "2", "objective": "Manage airway safely"
    }, follow_redirects=True)
    assert response.status_code == 200
    assert b"Manage airway safely" in response.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='resident.profile_updated'").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM objectives WHERE objective='Manage airway safely'").fetchone()[0] == 1


def test_first_year_progression_requires_every_component_at_least_ten(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,1)", ("Pass", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Pass'").fetchone()[0]
    for module, theory, practical in [("Module A", 12, 11), ("Module B", 10, 10)]:
        client.post("/exams", data={"action":"record","resident":str(rid),"study_year":"1","module":module,"theory":str(theory),"practical":str(practical)})
    client.post("/exams", data={"action":"evaluate_first_year","resident":str(rid),"study_year":"1","all_modules_confirmed":"yes"})
    with rotation_app.db() as conn:
        assert conn.execute("SELECT year FROM residents WHERE id=?", (rid,)).fetchone()[0] == 2
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,1)", ("Fail", "Resident"))
        fail_id = conn.execute("SELECT id FROM residents WHERE surname='Fail'").fetchone()[0]
    for module, theory, practical in [("Module A", 15, 15), ("Module B", 9, 16)]:
        client.post("/exams", data={"action":"record","resident":str(fail_id),"study_year":"1","module":module,"theory":str(theory),"practical":str(practical)})
    client.post("/exams", data={"action":"evaluate_first_year","resident":str(fail_id),"study_year":"1","all_modules_confirmed":"yes"})
    with rotation_app.db() as conn:
        assert conn.execute("SELECT year FROM residents WHERE id=?", (fail_id,)).fetchone()[0] == 1


def test_dems_archive_keeps_resident_and_history(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    for username in ("benhocine", "drdris"):
        client.post("/users", data={"username": username, "password": "strong-test-password-2026", "role": "super_admin"}, follow_redirects=True)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,4)", ("Graduate", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Graduate'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Urologie'").fetchone()[0]
        conn.execute("INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",
                     (rid, tid, "2025-2026", "01 Mai - 30 Juillet", "2025-07-30"))
        conn.commit()

    with client.session_transaction() as sess:
        sess["user"] = "benhocine"
        sess["role"] = "super_admin"
    first_vote = client.post("/dems", data={"resident":str(rid),"success_date":"2026-10-01","outcome":"approve","reason":"DEMS passed"}, follow_redirects=True)
    assert first_vote.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT status FROM residents WHERE id=?", (rid,)).fetchone()[0] == "active"
        assert conn.execute("SELECT COUNT(*) FROM dems_archive WHERE resident_id=?", (rid,)).fetchone()[0] == 0

    with client.session_transaction() as sess:
        sess["user"] = "drdris"
        sess["role"] = "super_admin"
    second_vote = client.post("/dems", data={"resident":str(rid),"success_date":"2026-10-01","outcome":"approve","reason":"Reviewed and approved"}, follow_redirects=True)
    assert second_vote.status_code == 200
    with rotation_app.db() as conn:
        row = conn.execute("SELECT status,dems_status FROM residents WHERE id=?", (rid,)).fetchone()
        assert tuple(row) == ("graduated", "success")
        assert conn.execute("SELECT COUNT(*) FROM history WHERE resident_id=?", (rid,)).fetchone()[0] == 1
        archive = conn.execute("SELECT history_json FROM dems_archive WHERE resident_id=?", (rid,)).fetchone()
        assert "Urologie" in archive["history_json"]
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='dems.vote_recorded'").fetchone()[0] == 2


def test_standard_user_cannot_access_super_admin_pages(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    response = client.post("/users", data={
        "username":"ordinary-user","password":"ordinary-password-2026","role":"user"
    }, follow_redirects=True)
    assert response.status_code == 200
    client.get("/logout")
    response = client.post("/login", data={
        "username":"ordinary-user","password":"ordinary-password-2026"
    }, follow_redirects=True)
    assert response.status_code == 200
    assert client.get("/settings").status_code == 403
    assert client.get("/dems").status_code == 403
    assert client.get("/audit").status_code == 403

def test_backup_and_restore_preserve_database_and_resident_photos(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    photo_dir = tmp_path / "resident_photos"
    photo_dir.mkdir(exist_ok=True)
    photo_path = photo_dir / "test-resident.webp"
    photo_path.write_bytes(b"original-photo-bytes")
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,?)", ("BeforeBackup", "Resident", 2))
    backup_response = client.get("/backup")
    assert backup_response.status_code == 200
    assert backup_response.mimetype == "application/zip"
    backup_bytes = backup_response.get_data()
    with zipfile.ZipFile(io.BytesIO(backup_bytes)) as bundle:
        assert "rotation.db" in bundle.namelist()
        assert "resident_photos/test-resident.webp" in bundle.namelist()
        assert bundle.read("resident_photos/test-resident.webp") == b"original-photo-bytes"
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,?)", ("AddedAfterBackup", "Resident", 1))
    photo_path.write_bytes(b"changed-photo-bytes")
    restored = client.post("/restore", data={
        "backup_file": (io.BytesIO(backup_bytes), "resimanage_full_backup.zip")
    }, content_type="multipart/form-data", follow_redirects=True)
    assert restored.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM residents WHERE surname='BeforeBackup'").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM residents WHERE surname='AddedAfterBackup'").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='database.restored'").fetchone()[0] == 1
    assert (tmp_path / "resident_photos" / "test-resident.webp").read_bytes() == b"original-photo-bytes"
    assert list(tmp_path.glob("rotation_backup_before_restore_*.db"))
    assert list(tmp_path.glob("rotation_backup_before_restore_*.zip"))

def test_csv_import_requires_preview_and_does_not_invent_history(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    source = b"surname,firstname,year,phone\nImportSurname,ImportFirst,3,0551234567\n"
    preview = client.post("/import", data={
        "file": (io.BytesIO(source), "residents.csv")
    }, content_type="multipart/form-data", follow_redirects=True)
    assert preview.status_code == 200, preview.data.decode("utf-8", errors="replace")
    assert b"ImportSurname" in preview.data, preview.data.decode("utf-8", errors="replace")
    match = re.search(rb'name="preview_id" value="(\d+)"', preview.data)
    assert match
    confirmed = client.post("/import", data={
        "action":"confirm", "preview_id":match.group(1).decode()
    }, follow_redirects=True)
    assert confirmed.status_code == 200
    with rotation_app.db() as conn:
        rid = conn.execute("SELECT id FROM residents WHERE surname='ImportSurname'").fetchone()[0]
        assert conn.execute("SELECT COUNT(*) FROM history WHERE resident_id=?", (rid,)).fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='residents.imported'").fetchone()[0] == 1

def test_mandatory_year_override_requires_reason_and_records_audit(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,2)", ("Override", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Override'").fetchone()[0]
        conn.execute("INSERT INTO terrains(name,kind,mandatory_year) VALUES(?,?,3)", ("Year Three Only", "mandatory"))
        tid = conn.execute("SELECT id FROM terrains WHERE name='Year Three Only'").fetchone()[0]
        period_id = conn.execute("SELECT id FROM periods ORDER BY id LIMIT 1").fetchone()[0]
        year_label = conn.execute("SELECT year_label FROM periods WHERE id=?", (period_id,)).fetchone()[0]
        conn.execute("INSERT INTO capacities(year_label,terrain_id,period_id,places) VALUES(?,?,?,1)", (year_label,tid,period_id))
    response = client.post("/assign", data={"resident":str(rid),"terrain":str(tid),"period":str(period_id)}, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        aid = conn.execute("SELECT id FROM assignments WHERE resident_id=? AND status='pending'", (rid,)).fetchone()[0]
    client.post("/decisions", data={"aid":str(aid),"outcome":"approve","reason":"Authorized educational exception"}, follow_redirects=True)
    with rotation_app.db() as conn:
        assert conn.execute("SELECT status FROM assignments WHERE id=?", (aid,)).fetchone()[0] == "confirmed"
        assert conn.execute("SELECT COUNT(*) FROM decisions WHERE reason='Authorized educational exception'").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='assignment.override_approved'").fetchone()[0] == 1


def test_uploaded_resident_photo_is_stored_and_served_safely(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,1)", ("Photo", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Photo'").fetchone()[0]
    png = b"\x89PNG\r\n\x1a\n" + b"test"
    response = client.post("/profile/" + str(rid), data={
        "birth_place":"", "birth_date":"", "phone":"",
        "photo":(io.BytesIO(png),"portrait.png")
    }, content_type="multipart/form-data", follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        path = conn.execute("SELECT photo_path FROM residents WHERE id=?", (rid,)).fetchone()[0]
    assert path.startswith("/photos/resident_")
    photo_response = client.get(path)
    assert photo_response.status_code == 200
    assert photo_response.data.startswith(b"\x89PNG")

def test_assignment_resident_choices_are_sorted_by_year_priority(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,1)", ("FirstYear", "Resident"))
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,3)", ("ThirdYear", "Resident"))
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,4)", ("FourthYear", "Resident"))
    response = client.get("/assign")
    html = response.data.decode("utf-8")
    assert html.index("FourthYear") < html.index("ThirdYear") < html.index("FirstYear")


def test_academic_year_change_creates_four_new_periods_and_preserves_old_periods(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        old_year = conn.execute("SELECT year_label FROM periods ORDER BY id LIMIT 1").fetchone()[0]
    response = client.post("/settings", data={
        "lang":"fr","year":"2027-2028","director":"Dr Dris","new_password":""
    }, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM periods WHERE year_label='2027-2028'").fetchone()[0] == 4
        assert conn.execute("SELECT COUNT(*) FROM periods WHERE year_label=?", (old_year,)).fetchone()[0] == 4


def test_xlsx_import_preview(tmp_path, monkeypatch):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Nom","Prénom","Année","Téléphone"])
    sheet.append(["ExcelResident","ExcelFirst",2,"0551111111"])
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    preview = client.post("/import", data={"file":(stream,"residents.xlsx")},
                          content_type="multipart/form-data", follow_redirects=True)
    assert preview.status_code == 200
    assert b"ExcelResident" in preview.data, preview.data.decode("utf-8", errors="replace")

def test_docx_import_reads_table_headers_and_resident_rows(tmp_path, monkeypatch):
    from docx import Document
    document = Document()
    table = document.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Nom"
    table.rows[0].cells[1].text = "Prénom"
    table.rows[0].cells[2].text = "Année"
    row = table.add_row().cells
    row[0].text = "WordResident"
    row[1].text = "WordFirst"
    row[2].text = "4"
    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    preview = client.post("/import", data={"file":(stream,"residents.docx")},
                          content_type="multipart/form-data", follow_redirects=True)
    assert preview.status_code == 200
    assert b"WordResident" in preview.data


def test_csv_export_includes_exam_objective_archive_and_audit_tables_but_not_users(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    response = client.get("/export")
    assert response.status_code == 200
    content = response.get_data().decode("utf-8-sig")
    for table in ("exam_records", "objectives", "dems_archive", "leave_years", "leave_entries", "external_rotations", "promotion_votes", "dems_votes", "audit"):
        assert table in content
    assert "users" not in content
    assert "password_hash" not in content

def test_fresh_database_requires_admin_to_set_a_private_password(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    monkeypatch.delenv("ROTATION_INITIAL_PASSWORD", raising=False)
    page = client.get("/login")
    assert page.status_code == 200
    assert b"confirm_password" in page.data
    response = client.post("/login", data={
        "username":"admin",
        "password":"fresh-private-password-2026",
        "confirm_password":"fresh-private-password-2026"
    }, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        stored = conn.execute("SELECT value FROM settings WHERE key='password'").fetchone()[0]
        user = conn.execute("SELECT role,password_hash FROM users WHERE username='admin'").fetchone()
    assert stored.startswith("pbkdf2_sha256$")
    assert user["role"] == "super_admin"
    assert user["password_hash"].startswith("pbkdf2_sha256$")


def test_legacy_password_hash_is_upgraded_after_successful_login(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    client.get("/login")
    with rotation_app.app.test_request_context("/"):
        with rotation_app.db() as conn:
            conn.execute("UPDATE settings SET value=? WHERE key='password'",
                         (hashlib.sha256(b"legacy-test-password-2026").hexdigest(),))
    response = client.post("/login", data={
        "username":"admin","password":"legacy-test-password-2026"
    }, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        stored = conn.execute("SELECT value FROM settings WHERE key='password'").fetchone()[0]
        user = conn.execute("SELECT password_hash,role FROM users WHERE username='admin'").fetchone()
    assert stored.startswith("pbkdf2_sha256$")
    assert user["password_hash"].startswith("pbkdf2_sha256$")
    assert user["role"] == "super_admin"

def test_terrain_deactivation_preserves_history_and_hides_it_from_new_assignments(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,2)", ("TerrainHistory", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='TerrainHistory'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Urologie'").fetchone()[0]
        conn.execute("INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",
                     (rid,tid,"2025-2026","01 Mai - 30 Juillet","2025-07-30"))
    response = client.post("/terrains", data={"action":"deactivate","tid":str(tid)}, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT active FROM terrains WHERE id=?", (tid,)).fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM history WHERE resident_id=? AND terrain_id=?", (rid,tid)).fetchone()[0] == 1
    assert b"Urologie" not in client.get("/assign").data


def test_assignment_cancellation_is_audited_and_does_not_delete_record(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year) VALUES(?,?,3)", ("Cancel", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Cancel'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Urologie'").fetchone()[0]
        pid = conn.execute("SELECT id FROM periods ORDER BY id LIMIT 1").fetchone()[0]
        conn.execute("INSERT INTO assignments(resident_id,terrain_id,period_id,status) VALUES(?,?,?,'confirmed')",(rid,tid,pid))
        aid = conn.execute("SELECT id FROM assignments WHERE resident_id=?", (rid,)).fetchone()[0]
    response = client.post("/assignments/" + str(aid) + "/cancel", follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT status FROM assignments WHERE id=?", (aid,)).fetchone()[0] == "cancelled"
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='assignment.cancelled'").fetchone()[0] == 1


def test_user_deactivation_is_audited_and_blocks_login(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    client.post("/users", data={"username":"temporary-user","password":"temporary-user-password-2026","role":"user"})
    with rotation_app.db() as conn:
        uid = conn.execute("SELECT id FROM users WHERE username='temporary-user'").fetchone()[0]
    client.post("/users", data={"action":"deactivate","user_id":str(uid)})
    with rotation_app.db() as conn:
        assert conn.execute("SELECT active FROM users WHERE id=?", (uid,)).fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='user.deactivated'").fetchone()[0] == 1
    client.get("/logout")
    response = client.post("/login", data={"username":"temporary-user","password":"temporary-user-password-2026"}, follow_redirects=True)
    assert b"Temporary user" not in response.data
    assert client.get("/settings").status_code == 302


def test_email_registration_requires_verification_code(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    sent = {}
    monkeypatch.setattr(rotation_app, "send_verification_email", lambda recipient, code: sent.update(email=recipient, code=code))
    response = client.get("/register")
    assert response.status_code == 200
    assert b'type="email"' in response.data

    response = client.post("/register", data={
        "username": "emailuser",
        "email": "emailuser@example.com",
        "password": "strong-test-password-2026",
        "confirm_password": "strong-test-password-2026",
    })
    assert response.status_code == 302
    assert response.headers["Location"] == "/verify-email"
    assert sent["email"] == "emailuser@example.com"
    with rotation_app.db() as conn:
        assert conn.execute("SELECT 1 FROM users WHERE username='emailuser'").fetchone() is None
        pending = conn.execute("SELECT attempts FROM pending_registrations WHERE email=?", (sent["email"],)).fetchone()
        assert pending is not None
        assert pending["attempts"] == 0

    wrong = client.post("/verify-email", data={"code": "000000"})
    assert wrong.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT 1 FROM users WHERE username='emailuser'").fetchone() is None

    verified = client.post("/verify-email", data={"code": sent["code"]})
    assert verified.status_code == 302
    assert verified.headers["Location"] == "/login"
    with rotation_app.db() as conn:
        account = conn.execute("SELECT username,email,email_verified,role FROM users WHERE username='emailuser'").fetchone()
        assert tuple(account) == ("emailuser", "emailuser@example.com", 1, "user")
        assert conn.execute("SELECT 1 FROM pending_registrations WHERE email=?", (sent["email"],)).fetchone() is None


def test_annual_leave_records_days_and_carries_unused_balance(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Leave", "Test", 2, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Leave'").fetchone()[0]
        conn.commit()
    from datetime import date, timedelta
    start = date.today().replace(month=1, day=1)
    end = start + timedelta(days=4)
    response = client.post("/leave", data={
        "resident_id": str(rid), "start_date": start.isoformat(),
        "end_date": end.isoformat(), "notes": "test leave"
    }, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        entry = conn.execute("SELECT days FROM leave_entries WHERE resident_id=?", (rid,)).fetchone()
        assert entry[0] == 5
        earned = conn.execute("SELECT SUM(earned_days) FROM leave_years WHERE resident_id=?", (rid,)).fetchone()[0]
        used = conn.execute("SELECT SUM(days) FROM leave_entries WHERE resident_id=?", (rid,)).fetchone()[0]
        assert earned >= 30
        assert earned - used >= 25


def test_external_rotation_blocks_overlapping_local_period(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Away", "Resident", 2, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Away'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Réanimation chirurgicale'").fetchone()[0]
        period = conn.execute("SELECT id,year_label,name FROM periods WHERE name LIKE '%Août%' OR name LIKE '%Aout%'").fetchone()
        if period is None:
            period = conn.execute("SELECT id,year_label,name FROM periods ORDER BY id LIMIT 1").fetchone()
        conn.execute("INSERT INTO external_rotations(resident_id,supervisor,destination,hospital,country,rotation_type,start_date,end_date,status) VALUES(?,?,?,?,?,?,?,?,?)",
                     (rid, "Dr Test", "Paris", "CHU Test", "France", "Observation", "2026-08-01", "2026-10-30", "active"))
        conn.commit()
        allowed, message, status = rotation_app.check(conn, rid, tid, period["id"], period["year_label"])
        assert allowed is False
        assert "External rotation" in message


def test_external_rotation_form_requires_all_destination_details(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Trip", "Resident", 3, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Trip'").fetchone()[0]
        conn.commit()
    response = client.post("/external-rotations", data={
        "resident_id": str(rid), "supervisor": "Dr Test", "destination": "Paris",
        "hospital": "CHU Test", "country": "France", "rotation_type": "Observation",
        "start_date": "2026-11-01", "end_date": "2026-12-01", "notes": "test"
    }, follow_redirects=True)
    assert response.status_code == 200
    with rotation_app.db() as conn:
        row = conn.execute("SELECT supervisor,hospital,country,rotation_type FROM external_rotations WHERE resident_id=?", (rid,)).fetchone()
        assert tuple(row) == ("Dr Test", "CHU Test", "France", "Observation")

def test_leave_carries_unused_days_forward_and_blocks_overdraw(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Carry", "Leave", 2, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Carry'").fetchone()[0]
        conn.commit()

    first = client.post("/leave", data={"resident_id": rid, "start_date": "2024-01-01", "end_date": "2024-01-10"}, follow_redirects=True)
    assert first.status_code == 200
    second = client.post("/leave", data={"resident_id": rid, "start_date": "2025-01-01", "end_date": "2025-01-25"}, follow_redirects=True)
    assert second.status_code == 200
    with rotation_app.db() as conn:
        years = conn.execute("SELECT SUM(earned_days) FROM leave_years WHERE resident_id=? AND year IN (2024,2025)", (rid,)).fetchone()[0]
        used = conn.execute("SELECT SUM(days) FROM leave_entries WHERE resident_id=?", (rid,)).fetchone()[0]
    assert years == 60
    assert used == 35
    rejected = client.post("/leave", data={"resident_id": rid, "start_date": "2025-02-01", "end_date": "2025-02-26"}, follow_redirects=True)
    assert rejected.status_code == 200
    assert b"Insufficient accrued leave balance" in rejected.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT SUM(days) FROM leave_entries WHERE resident_id=?", (rid,)).fetchone()[0] == 35


def test_leave_rejects_overlapping_entries(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Overlap", "Leave", 1, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Overlap'").fetchone()[0]
        conn.commit()
    client.post("/leave", data={"resident_id": rid, "start_date": "2026-01-01", "end_date": "2026-01-10"}, follow_redirects=True)
    rejected = client.post("/leave", data={"resident_id": rid, "start_date": "2026-01-05", "end_date": "2026-01-12"}, follow_redirects=True)
    assert rejected.status_code == 200
    assert b"Leave dates overlap an existing entry" in rejected.data
    with rotation_app.db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM leave_entries WHERE resident_id=?", (rid,)).fetchone()[0] == 1


def test_external_rotation_returned_record_still_blocks_overlapping_period(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Returned", "Away", 2, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Returned'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Réanimation chirurgicale'").fetchone()[0]
        period = conn.execute("SELECT id,year_label,name FROM periods WHERE name LIKE '%Août%' OR name LIKE '%Aout%'").fetchone()
        assert period is not None
        conn.execute("INSERT INTO external_rotations(resident_id,supervisor,destination,hospital,country,rotation_type,start_date,end_date,status) VALUES(?,?,?,?,?,?,?,?,?)",
                     (rid, "Dr Test", "Paris", "CHU Test", "France", "Observation", "2026-08-01", "2026-10-30", "returned"))
        conn.commit()
        allowed, message, _ = rotation_app.check(conn, rid, tid, period["id"], period["year_label"])
    assert allowed is False
    assert "External rotation" in message


def test_assignment_resident_options_mark_external_period_conflicts(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Blocked", "Resident", 2, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Blocked'").fetchone()[0]
        period = conn.execute("SELECT id FROM periods WHERE name LIKE '%Août%' OR name LIKE '%Aout%'").fetchone()
        assert period is not None
        conn.execute("INSERT INTO external_rotations(resident_id,supervisor,destination,hospital,country,rotation_type,start_date,end_date,status) VALUES(?,?,?,?,?,?,?,?,?)",
                     (rid, "Dr Test", "Paris", "CHU Test", "France", "Observation", "2026-08-01", "2026-10-30", "returned"))
        conn.commit()
    response = client.get("/assign")
    assert response.status_code == 200
    assert f'data-blocked-periods="{period["id"]}"'.encode() in response.data
    assert b"o.disabled=blocked" in response.data


def test_external_rotation_can_be_edited_and_cancelled_with_audit(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)", ("Edit", "External", 3, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Edit'").fetchone()[0]
        conn.execute("INSERT INTO external_rotations(resident_id,supervisor,destination,hospital,country,rotation_type,start_date,end_date,status) VALUES(?,?,?,?,?,?,?,?,?)",
                     (rid, "Dr Old", "Paris", "CHU Old", "France", "Observation", "2026-11-01", "2026-12-01", "active"))
        xid = conn.execute("SELECT id FROM external_rotations WHERE resident_id=?", (rid,)).fetchone()[0]
        conn.commit()
    edit = client.post(f"/external-rotations/{xid}/edit", data={
        "resident_id": str(rid), "supervisor": "Dr New", "destination": "Lyon",
        "hospital": "CHU New", "country": "France", "rotation_type": "Clinical",
        "start_date": "2026-11-10", "end_date": "2026-12-10", "notes": "updated"
    }, follow_redirects=True)
    assert edit.status_code == 200
    with rotation_app.db() as conn:
        row = conn.execute("SELECT supervisor,destination,hospital,rotation_type,status FROM external_rotations WHERE id=?", (xid,)).fetchone()
        assert tuple(row) == ("Dr New", "Lyon", "CHU New", "Clinical", "active")
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='external_rotation.updated'").fetchone()[0] == 1
    cancelled = client.post(f"/external-rotations/{xid}/cancel", follow_redirects=True)
    assert cancelled.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT status FROM external_rotations WHERE id=?", (xid,)).fetchone()[0] == "cancelled"
        assert conn.execute("SELECT COUNT(*) FROM audit WHERE action='external_rotation.cancelled'").fetchone()[0] == 1

def test_post_without_csrf_token_is_rejected(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    client.get("/login")
    response = client.open("/login", method="POST", data={"username": "admin", "password": "not-a-real-password"})
    assert response.status_code == 400
    assert b"Invalid or missing CSRF token" in response.data


def test_official_admin_accounts_must_be_super_admin_and_cannot_be_disabled(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    for username in ("benhocine", "drdris"):
        invalid = client.post("/users", data={
            "username": username, "password": "strong-test-password-2026", "role": "user"
        }, follow_redirects=True)
        assert invalid.status_code == 200
        assert b"Official accounts must be Super Admins" in invalid.data
        with rotation_app.db() as conn:
            assert conn.execute("SELECT 1 FROM users WHERE lower(username)=?", (username,)).fetchone() is None

        created = client.post("/users", data={
            "username": username, "password": "strong-test-password-2026", "role": "super_admin"
        }, follow_redirects=True)
        assert created.status_code == 200
        with rotation_app.db() as conn:
            row = conn.execute("SELECT id,role,active FROM users WHERE lower(username)=?", (username,)).fetchone()
            assert tuple(row[1:]) == ("super_admin", 1)
            uid = row["id"]
        deactivated = client.post("/users", data={"action": "deactivate", "user_id": str(uid)}, follow_redirects=True)
        assert deactivated.status_code == 200
        assert b"Official accounts cannot be deactivated" in deactivated.data
        with rotation_app.db() as conn:
            assert conn.execute("SELECT active FROM users WHERE id=?", (uid,)).fetchone()[0] == 1

def test_year_progression_requires_approval_from_both_official_admins(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    for username in ("benhocine", "drdris"):
        client.post("/users", data={"username": username, "password": "strong-test-password-2026", "role": "super_admin"}, follow_redirects=True)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,2,'active')", ("Promotion", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Promotion'").fetchone()[0]
        conn.commit()

    unauthorized = client.post("/exams", data={"action":"approve_advance","resident":str(rid),"study_year":"2","outcome":"approve","reason":"ok"}, follow_redirects=False)
    assert unauthorized.status_code == 403

    with client.session_transaction() as sess:
        sess["user"] = "benhocine"
        sess["role"] = "super_admin"
    first = client.post("/exams", data={"action":"approve_advance","resident":str(rid),"study_year":"2","outcome":"approve","reason":"First official approves"}, follow_redirects=True)
    assert first.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT year FROM residents WHERE id=?", (rid,)).fetchone()[0] == 2
        assert conn.execute("SELECT COUNT(*) FROM promotion_votes WHERE resident_id=?", (rid,)).fetchone()[0] == 1

    with client.session_transaction() as sess:
        sess["user"] = "drdris"
        sess["role"] = "super_admin"
    second = client.post("/exams", data={"action":"approve_advance","resident":str(rid),"study_year":"2","outcome":"approve","reason":"Second official approves"}, follow_redirects=True)
    assert second.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT year FROM residents WHERE id=?", (rid,)).fetchone()[0] == 3
        assert conn.execute("SELECT COUNT(*) FROM promotion_votes WHERE resident_id=?", (rid,)).fetchone()[0] == 0
        decision = conn.execute("SELECT decision,actor FROM decisions WHERE details LIKE '%Promotion Resident — year 2%'").fetchone()
        assert decision["decision"] == "Approved by both official administrators"
        assert decision["actor"] == "benhocine+drdris"


def test_year_progression_is_not_approved_if_either_official_rejects(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    for username in ("benhocine", "drdris"):
        client.post("/users", data={"username": username, "password": "strong-test-password-2026", "role": "super_admin"}, follow_redirects=True)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,3,'active')", ("Reject", "Resident"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='Reject'").fetchone()[0]
        conn.commit()
    for username,outcome in (("benhocine","approve"),("drdris","reject")):
        with client.session_transaction() as sess:
            sess["user"] = username
            sess["role"] = "super_admin"
        response = client.post("/exams", data={"action":"approve_advance","resident":str(rid),"study_year":"3","outcome":outcome,"reason":username+" reviewed"}, follow_redirects=True)
        assert response.status_code == 200
    with rotation_app.db() as conn:
        assert conn.execute("SELECT year FROM residents WHERE id=?", (rid,)).fetchone()[0] == 3
        decision = conn.execute("SELECT decision FROM decisions WHERE details LIKE '%Reject Resident — year 3%'").fetchone()
        assert decision[0] == "Rejected by official vote"



def test_selected_postgraduate_year_filters_history_summary_and_pending_decisions(tmp_path, monkeypatch):
    client = setup_client(tmp_path, monkeypatch)
    login(client)
    with rotation_app.db() as conn:
        conn.execute("INSERT INTO residents(surname,firstname,year,status) VALUES(?,?,?,?)",
                     ("YearFilter", "Resident", 2, "active"))
        rid = conn.execute("SELECT id FROM residents WHERE surname='YearFilter'").fetchone()[0]
        tid = conn.execute("SELECT id FROM terrains WHERE name='Urologie'").fetchone()[0]
        conn.execute("INSERT INTO periods(year_label,name) VALUES(?,?)", ("2025-2026", "01 Fév - 30 Avr"))
        current_pid = conn.execute("SELECT id FROM periods WHERE year_label=? AND name=?",
                                   ("2025-2026", "01 Fév - 30 Avr")).fetchone()[0]
        conn.execute("INSERT INTO periods(year_label,name) VALUES(?,?)", ("2024-2025", "01 Fév - 30 Avr"))
        previous_pid = conn.execute("SELECT id FROM periods WHERE year_label=? AND name=?",
                                   ("2024-2025", "01 Fév - 30 Avr")).fetchone()[0]
        conn.execute("INSERT INTO settings(key,value) VALUES('year',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                     ("2025-2026",))
        conn.execute("INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",
                     (rid, tid, "2025-2026", "01 Fév - 30 Avr", "2025-02-10"))
        conn.execute("INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",
                     (rid, tid, "2024-2025", "01 Fév - 30 Avr", "2024-02-10"))
        conn.execute("INSERT INTO assignments(resident_id,terrain_id,period_id,status) VALUES(?,?,?,?)",
                     (rid, tid, current_pid, "pending"))
        conn.execute("INSERT INTO assignments(resident_id,terrain_id,period_id,status) VALUES(?,?,?,?)",
                     (rid, tid, previous_pid, "pending"))
        conn.commit()

    history_page = client.get("/history")
    assert history_page.status_code == 200
    assert b"2025-02-10" in history_page.data
    assert b"2024-02-10" not in history_page.data

    summary_page = client.get("/summary")
    assert summary_page.status_code == 200
    assert b"2025-2026" in summary_page.data
    assert b"2024-2025" not in summary_page.data

    decisions_page = client.get("/decisions")
    assert decisions_page.status_code == 200
    assert "01 Fév - 30 Avr".encode() in decisions_page.data
    assert decisions_page.data.count(b"YearFilter Resident") == 1
