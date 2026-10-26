# -*- coding: utf-8 -*-
from flask import Flask, request, redirect, session, render_template_string, flash, send_file, send_from_directory, jsonify
import sqlite3, os, sys, hashlib, hmac, io, csv, shutil, json, tempfile, re, subprocess, secrets, smtplib, time, zipfile
import warnings
from PIL import Image, ImageOps
from datetime import datetime, date
from werkzeug.utils import secure_filename

BASE=os.path.dirname(os.path.abspath(sys.executable)) if getattr(sys,"frozen",False) else os.path.dirname(os.path.abspath(__file__))
DB=os.path.join(BASE,"rotation.db")
def local_secret_key():
 path=os.path.join(BASE,".rotation_secret")
 try:
  with open(path,"r",encoding="ascii") as handle: return handle.read().strip()
 except FileNotFoundError:
  value=secrets.token_hex(32)
  try:
   fd=os.open(path,os.O_WRONLY|os.O_CREAT|os.O_EXCL,0o600)
   with os.fdopen(fd,"w",encoding="ascii") as handle: handle.write(value)
   return value
  except FileExistsError:
   with open(path,"r",encoding="ascii") as handle: return handle.read().strip()
app=Flask(__name__)
app.config["MAX_CONTENT_LENGTH"]=25*1024*1024
app.secret_key=os.environ.get("ROTATION_SECRET") or local_secret_key()
LANGS=("ar","fr","en")
D={
"ar":{"app":"ResiManage","home":"الرئيسية","residents":"المقيمون","manage_residents":"إضافة وإدارة المقيمين","import_residents":"استيراد المقيمين","promotion":"قرارات الترقية والامتحانات","dems_success":"قائمة الناجحين في DEMS","settings_page":"إعدادات التطبيق","terrains":"المصالح","capacities":"المقاعد","assign":"التوزيع","history":"التاريخ","summary":"ملخص المقيمين","decisions":"القرارات","settings":"الإعدادات","logout":"خروج","login":"دخول المسؤول","password":"كلمة المرور","signin":"دخول","badpass":"كلمة المرور غير صحيحة","year":"سنة ما بعد التدرج","resident":"المقيم","surname":"اللقب","firstname":"الاسم","study_year":"السنة","phone":"الهاتف","repeat":"يعيد السنة","status":"الحالة","active":"نشط","inactive":"غير نشط","add":"إضافة","save":"حفظ","terrain":"المصلحة","kind":"النوع","mandatory":"إجباري","basic":"أساسي","optional":"اختياري","mandatory_year":"السنة الإلزامية","any":"حسب البرنامج","period":"الفترة","places":"عدد المقاعد","used":"المستعمل","remaining":"المتبقي","date":"تاريخ الإنجاز","academic_year":"سنة ما بعد التدرج","choose":"اختيار","assign_now":"تعيين مقيم","check_save":"التحقق والحفظ","confirmed":"مؤكد","pending":"بانتظار القرار","rejected":"مرفوض","history_add":"تسجيل تربص سابق","reason":"سبب القرار","approve":"موافقة","reject":"رفض","language":"لغة الواجهة","director":"المسؤول","new_password":"كلمة مرور جديدة (اختياري)","change":"حفظ الإعدادات","note":"تغيير اللغة لا يغيّر البيانات المحفوظة.","priority":"الأولوية للسنة الرابعة ثم الثالثة والثانية والأولى.","none":"لا توجد بيانات","capacity_first":"حدّد السعة أولاً لهذه المصلحة والفترة","full":"السعة مكتملة","duplicate":"المقيم مسجّل في نفس الفترة بمصلحة: ","same_terrain_year":"سبق للمقيم إنجاز هذه المصلحة خلال سنة ما بعد التدرج في الفترة: ","history_block":"التربص منجز سابقاً؛ يمنع التكرار حتى إكمال جميع المصالح.","history_repeat":"يمكن التكرار لأن المقيم أكمل جميع المصالح.","mismatch":"السنة لا تطابق السنة الإلزامية؛ التعيين بانتظار قرار المسؤول","done":"تم الحفظ","overview":"ملخص التوزيع السنوي","count":"عدد التعيينات","type":"التصنيف" ,"export":"تصدير CSV","backup":"تنزيل نسخة احتياطية","all":"كل المصالح","four":"التعيينات المؤكدة / 4 فترات","waiting":"معلّق","users":"المستخدمون","import":"استيراد","objectives":"الأهداف","exams":"الامتحانات","dems":"أرشيف DEMS","audit":"سجل التدقيق","restore":"استرجاع","username":"اسم المستخدم","role":"الصلاحية","standard_user":"مستخدم عادي","super_admin":"مسؤول عام","created":"تاريخ الإنشاء","active_col":"نشط","actor":"الفاعل","action":"الإجراء","details":"التفاصيل","module":"المادة","theory":"نظري","practical":"تطبيقي","success_date":"تاريخ النجاح","file":"الملف","objective":"الهدف","archive":"أرشفة","users_help":"أنشئ حسابًا مستقلًا لكل مسؤول، وامنح الأستاذ Benhocine والدكتور Dris صلاحية المسؤول العام.","import_help":"ارفع ملفًا يحتوي على اللقب والاسم والسنة الدراسية. راجع المعاينة قبل التأكيد؛ لن تُختلق أي تربصات سابقة.","restore_help":"سيتم التحقق من سلامة قاعدة البيانات وإنشاء نسخة احتياطية من القاعدة الحالية قبل الاسترجاع.","confirm_password":"تأكيد كلمة المرور","confirm_import":"تأكيد الاستيراد","cancel":"إلغاء","upload_preview":"رفع ومعاينة","archive_note":"تُحفظ بيانات المقيم وتاريخه ولا تُحذف."},
"fr":{"app":"ResiManage","home":"Accueil","residents":"Résidents","manage_residents":"Gérer les résidents","import_residents":"Importer les résidents","promotion":"Décisions de promotion et examens","dems_success":"Lauréats du DEMS","settings_page":"Paramètres de l’application","terrains":"Services","capacities":"Capacités","assign":"Affectations","history":"Historique","summary":"Bilan résidents","decisions":"Décisions","settings":"Paramètres","logout":"Déconnexion","login":"Connexion administrateur","password":"Mot de passe","signin":"Connexion","badpass":"Mot de passe incorrect","year":"Année de post-graduation","resident":"Résident","surname":"Nom","firstname":"Prénom","study_year":"Année","phone":"Téléphone","repeat":"Redouble","status":"Statut","active":"Actif","inactive":"Inactif","add":"Ajouter","save":"Enregistrer","terrain":"Service","kind":"Type","mandatory":"Obligatoire","basic":"De base","optional":"Optionnel","mandatory_year":"Année requise","any":"Selon le programme","period":"Période","places":"Nombre de places","used":"Utilisées","remaining":"Restantes","date":"Date de réalisation","academic_year":"Année de post-graduation","choose":"Choisir","assign_now":"Affecter un résident","check_save":"Vérifier et enregistrer","confirmed":"Confirmé","pending":"En attente","rejected":"Refusé","history_add":"Enregistrer un stage antérieur","reason":"Motif de décision","approve":"Approuver","reject":"Refuser","language":"Langue de l’interface","director":"Responsable","new_password":"Nouveau mot de passe (facultatif)","change":"Enregistrer les paramètres","note":"Changer de langue ne modifie pas les données.","priority":"Priorité aux 4e, puis 3e, 2e et 1re années.","none":"Aucune donnée","capacity_first":"Définissez d’abord la capacité du service et de la période","full":"Capacité complète","duplicate":"Résident déjà affecté pendant cette période au service : ","same_terrain_year":"Ce résident a déjà effectué ce service cette année de post-graduation, pendant la période : ","history_block":"Stage déjà effectué ; répétition bloquée avant d’avoir terminé tous les services.","history_repeat":"Répétition autorisée : tous les services ont été effectués.","mismatch":"Année différente de l’année requise ; décision du responsable nécessaire","done":"Enregistré","overview":"Vue annuelle des affectations","count":"Nombre d’affectations","type":"Catégorie" ,"export":"Exporter CSV","backup":"Télécharger la sauvegarde","all":"Tous les services","four":"Affectations confirmées / 4 périodes","waiting":"En attente","users":"Utilisateurs","import":"Importer","objectives":"Objectifs","exams":"Examens","dems":"Archive DEMS","audit":"Journal d’audit","restore":"Restaurer","username":"Identifiant","role":"Rôle","standard_user":"Utilisateur standard","super_admin":"Super Admin","created":"Créé le","active_col":"Actif","actor":"Acteur","action":"Action","details":"Détails","module":"Matière","theory":"Théorie","practical":"Pratique","success_date":"Date de réussite","file":"Fichier","objective":"Objectif","archive":"Archiver","users_help":"Créez un compte distinct pour chaque responsable et attribuez le rôle Super Admin à Pr Benhocine et Dr Dris.","import_help":"Importez une liste contenant le nom, le prénom et l’année d’étude. Vérifiez l’aperçu avant confirmation ; aucun stage passé ne sera inventé.","restore_help":"L’intégrité de la base sera vérifiée et une copie de la base actuelle sera créée avant la restauration.","confirm_password":"Confirmer le mot de passe","confirm_import":"Confirmer l’importation","cancel":"Annuler","upload_preview":"Charger et prévisualiser","archive_note":"Les données et l’historique du résident sont conservés, pas supprimés."},
"en":{"app":"ResiManage","home":"Home","residents":"Residents","manage_residents":"Manage residents","import_residents":"Import residents","promotion":"Promotion decisions & exams","dems_success":"DEMS successful graduates","settings_page":"App settings","terrains":"Departments","capacities":"Capacity","assign":"Assignments","history":"History","summary":"Resident summary","decisions":"Decisions","settings":"Settings","logout":"Log out","login":"Administrator login","password":"Password","signin":"Sign in","badpass":"Incorrect password","year":"Postgraduate year","resident":"Resident","surname":"Surname","firstname":"First name","study_year":"Year","phone":"Phone","repeat":"Repeating year","status":"Status","active":"Active","inactive":"Inactive","add":"Add","save":"Save","terrain":"Department","kind":"Type","mandatory":"Mandatory","basic":"Core","optional":"Optional","mandatory_year":"Required year","any":"As per curriculum","period":"Period","places":"Places","used":"Used","remaining":"Remaining","date":"Completion date","academic_year":"Postgraduate year","choose":"Choose","assign_now":"Assign resident","check_save":"Validate and save","confirmed":"Confirmed","pending":"Pending decision","rejected":"Rejected","history_add":"Record previous rotation","reason":"Decision reason","approve":"Approve","reject":"Reject","language":"Interface language","director":"Administrator","new_password":"New password (optional)","change":"Save settings","note":"Changing language does not change saved data.","priority":"Priority: year 4, then 3, 2 and 1.","none":"No data","capacity_first":"Set capacity for this department and period first","full":"Capacity is full","duplicate":"Resident already assigned during this period to: ","same_terrain_year":"This resident already completed this department this postgraduate year, during: ","history_block":"Rotation completed previously; repeats blocked until all departments are completed.","history_repeat":"Repeat allowed because all departments have been completed.","mismatch":"Year does not match required year; administrator decision needed","done":"Saved","overview":"Annual assignment overview","count":"Assignment count","type":"Category" ,"export":"Export CSV","backup":"Download backup","all":"All departments","four":"Confirmed assignments / 4 periods","waiting":"Pending","users":"Users","import":"Import","objectives":"Objectives","exams":"Exams","dems":"DEMS archive","audit":"Audit log","restore":"Restore","username":"Username","role":"Role","standard_user":"Standard user","super_admin":"Super Admin","created":"Created","active_col":"Active","actor":"Actor","action":"Action","details":"Details","module":"Module","theory":"Theory","practical":"Practical","success_date":"Success date","file":"File","objective":"Objective","archive":"Archive","users_help":"Create a separate account for each administrator and assign Super Admin to Pr Benhocine and Dr Dris.","import_help":"Upload a list containing surname, first name and study year. Review the preview before confirming; no past rotations are invented.","restore_help":"The database integrity will be checked and a copy of the current database will be created before restoring.","confirm_password":"Confirm password","confirm_import":"Confirm import","cancel":"Cancel","upload_preview":"Upload and preview","archive_note":"Resident data and history are retained, not deleted."}
}
def current_language():
 try: lang=session.get("lang","fr")
 except RuntimeError: lang="fr"
 return lang if lang in LANGS else "fr"
def L(fr,en,ar):
 return {"fr":fr,"en":en,"ar":ar}[current_language()]
def period_label(name):
 value=str(name or "").strip()
 aliases={"01 Fév - 30 Avr":1,"01 Fév–30 Avr":1,"01 Mai - 30 Juillet":2,"01 Août - 30 Oct":3,"01 Aout - 30 Oct":3,"01 Nov - 30 Jan":4}
 number=aliases.get(value)
 if number is None:
  import re as _re
  match=_re.fullmatch(r"(?:Période|Period|الفترة)\\s*(\\d+)",value,flags=_re.IGNORECASE)
  if match:
   number=int(match.group(1))
 if number in (1,2,3,4):
  return L(f"Période {number}",f"Period {number}",f"الفترة {number}")
 return value
def tr(k):
 lang=current_language()
 return D.get(lang,D["fr"]).get(k,D["fr"].get(k,k))
def db():
 c=sqlite3.connect(DB,timeout=20); c.row_factory=sqlite3.Row; c.execute("PRAGMA foreign_keys=ON"); return c
def setting(c,k,default=""):
 r=c.execute("SELECT value FROM settings WHERE key=?",(k,)).fetchone(); return r["value"] if r else default
def available_postgrad_years(c):
 # Offer every postgraduate year from 2022-2023 through the configured current
 # year, while also retaining years found in historical records.
 current=setting(c,"year")
 try:
  current_start=int(current.split("-")[0]) if re.fullmatch(r"\\d{4}-\\d{4}",current) else date.today().year
 except (ValueError,IndexError):
  current_start=date.today().year
 years={f"{start}-{start+1}" for start in range(2022,max(2022,current_start)+1)}
 if current: years.add(current)
 # Recover years from every feature that stores its own postgraduate-year label.
 # Older databases may not have every newer table/column yet, so each source is
 # deliberately queried independently and missing schema is tolerated.
 for table,column in (("periods","year_label"),("history","year_label"),
                      ("exam_records","postgrad_year"),("decisions","postgrad_year")):
  try:
   years.update(str(row[0]) for row in c.execute(
    "SELECT DISTINCT "+column+" FROM "+table+" WHERE "+column+" IS NOT NULL AND "+column+"<>''"
   ).fetchall())
  except sqlite3.Error:
   pass
 return sorted((y for y in years if re.fullmatch(r"\d{4}-\d{4}",y) and int(y.split("-")[1])==int(y.split("-")[0])+1),reverse=True)
def active_postgrad_year(c):
 candidate=str(session.get("selected_postgrad_year","")).strip()
 if re.fullmatch(r"\d{4}-\d{4}",candidate) and int(candidate.split("-")[1])==int(candidate.split("-")[0])+1:
  return candidate
 return setting(c,"year")
def hash_password(password):
 salt=os.urandom(16).hex()
 digest=hashlib.pbkdf2_hmac("sha256",password.encode("utf-8"),bytes.fromhex(salt),310000).hex()
 return "pbkdf2_sha256$"+salt+"$"+digest
def verify_password(password, stored):
 if stored.startswith("pbkdf2_sha256$"):
  try:
   _,salt,digest=stored.split("$",2)
   candidate=hashlib.pbkdf2_hmac("sha256",password.encode("utf-8"),bytes.fromhex(salt),310000).hex()
   return hmac.compare_digest(candidate,digest)
  except (ValueError,TypeError):
   return False
 return hmac.compare_digest(hashlib.sha256(password.encode("utf-8")).hexdigest(),stored)
def ensure_column(c, table, column, spec):
 existing={row[1] for row in c.execute("PRAGMA table_info("+table+")")}
 if column not in existing:
  c.execute("ALTER TABLE "+table+" ADD COLUMN "+column+" "+spec)
def audit_event(c, action, details):
 c.execute("INSERT INTO audit(action,details,actor) VALUES(?,?,?)",
           (action,str(details)[:1000],str(session.get("user","system"))))


# Transactional change journal groundwork for future authenticated device sync.
# Authentication/transport endpoints and offline client queues are intentionally separate phases.
_SYNC_ENTITY_FIELDS={
 "residents":("id","surname","firstname","year","phone","phone2","repeat_year","status","birth_place","birth_date","address","email","resident_card_number","national_id_number","medical_order_registered","medical_order_number","social_security_number","registration_date","additional_info","resident_type","origin_hospital","external_start_date","external_end_date","archive_type","archive_reason","archive_return_date","archived_at","photo_path","dems_status"),
 "terrains":("id","name","kind","mandatory_year","active"),
 "periods":("id","year_label","name"),
 "capacities":("id","year_label","terrain_id","period_id","places"),
 "assignments":("id","resident_id","terrain_id","period_id","status"),
 "history":("id","resident_id","terrain_id","year_label","period","date_done"),
 "decisions":("id","details","decision","reason","actor","created_at"),
 "exam_records":("id","resident_id","study_year","module","theory","practical","exam_type","attempt","created_at"),
 "objectives":("id","terrain_id","study_year","objective","completed"),
 "dems_archive":("id","resident_id","success_date","profile_json","history_json","archived_at"),
 "promotion_votes":("id","resident_id","from_year","decision","reason","actor","created_at"),
 "dems_votes":("id","resident_id","success_date","decision","reason","actor","created_at"),
 "leave_years":("id","resident_id","year","earned_days","created_by"),
 "leave_entries":("id","resident_id","start_date","end_date","days","notes","created_by","created_at"),
 "external_rotations":("id","resident_id","supervisor","destination","hospital","country","rotation_type","start_date","end_date","notes","status","request_status","created_by","created_at","updated_at"),
 "audit":("id","action","details","actor","created_at"),
}
def install_sync_change_triggers(c):
 c.execute("""CREATE TABLE IF NOT EXISTS sync_change_log(
  seq INTEGER PRIMARY KEY AUTOINCREMENT,
  entity TEXT NOT NULL,
  entity_id TEXT NOT NULL,
  operation TEXT NOT NULL CHECK(operation IN ('upsert','delete')),
  payload_json TEXT,
  base_seq INTEGER NOT NULL DEFAULT 0,
  changed_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
 )""")
 ensure_column(c,"sync_change_log","base_seq","INTEGER NOT NULL DEFAULT 0")
 c.execute("CREATE INDEX IF NOT EXISTS idx_sync_change_log_entity_seq ON sync_change_log(entity,seq)")
 c.execute("""CREATE TABLE IF NOT EXISTS sync_row_versions(
  entity TEXT NOT NULL, entity_id TEXT NOT NULL, server_seq INTEGER NOT NULL DEFAULT 0,
  PRIMARY KEY(entity,entity_id)
 )""")
 c.execute("""CREATE TABLE IF NOT EXISTS sync_control(
  id INTEGER PRIMARY KEY CHECK(id=1), suppress_log INTEGER NOT NULL DEFAULT 0
 )""")
 c.execute("INSERT OR IGNORE INTO sync_control(id,suppress_log) VALUES(1,0)")
 c.execute("""CREATE TABLE IF NOT EXISTS sync_devices(
  device_id TEXT PRIMARY KEY,
  label TEXT NOT NULL,
  token_hash TEXT NOT NULL UNIQUE,
  active INTEGER NOT NULL DEFAULT 1,
  mode TEXT NOT NULL DEFAULT 'read_only',
  created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
  last_seen TEXT DEFAULT ''
 )""")
 ensure_column(c,"sync_devices","mode","TEXT NOT NULL DEFAULT 'read_only'")
 c.execute("""CREATE TABLE IF NOT EXISTS sync_applied_ops(
  op_id TEXT PRIMARY KEY,
  device_id TEXT NOT NULL,
  entity TEXT NOT NULL,
  entity_id TEXT NOT NULL,
  result_json TEXT NOT NULL,
  applied_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
 )""")
 c.execute("""CREATE TABLE IF NOT EXISTS sync_photo_uploads(
  device_id TEXT NOT NULL, resident_id INTEGER NOT NULL, sha256 TEXT NOT NULL,
  photo_path TEXT NOT NULL, server_seq INTEGER NOT NULL, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
  PRIMARY KEY(device_id,resident_id,sha256)
 )""")
 for table,fields in _SYNC_ENTITY_FIELDS.items():
  pairs=", ".join("'"+field+"', NEW."+field for field in fields)
  for suffix,operation in (("insert","upsert"),("update","upsert")):
   name="sync_"+table+"_"+suffix
   c.execute("DROP TRIGGER IF EXISTS "+name)
   c.execute(f"""CREATE TRIGGER {name} AFTER {suffix.upper()} ON {table}
    WHEN COALESCE((SELECT suppress_log FROM sync_control WHERE id=1),0)=0
    BEGIN INSERT INTO sync_change_log(entity,entity_id,operation,payload_json,base_seq)
    VALUES('{table}',CAST(NEW.id AS TEXT),'{operation}',json_object({pairs}),
     COALESCE((SELECT server_seq FROM sync_row_versions WHERE entity='{table}' AND entity_id=CAST(NEW.id AS TEXT)),0)); END""")
  name="sync_"+table+"_delete"
  c.execute("DROP TRIGGER IF EXISTS "+name)
  c.execute(f"""CREATE TRIGGER {name} AFTER DELETE ON {table}
   WHEN COALESCE((SELECT suppress_log FROM sync_control WHERE id=1),0)=0
   BEGIN INSERT INTO sync_change_log(entity,entity_id,operation,payload_json,base_seq)
   VALUES('{table}',CAST(OLD.id AS TEXT),'delete',json_object('id',OLD.id),
    COALESCE((SELECT server_seq FROM sync_row_versions WHERE entity='{table}' AND entity_id=CAST(OLD.id AS TEXT)),0)); END""")

def init():
 c=db()
 c.executescript("""
 CREATE TABLE IF NOT EXISTS residents(id INTEGER PRIMARY KEY,surname TEXT NOT NULL,firstname TEXT NOT NULL,year INTEGER NOT NULL CHECK(year BETWEEN 1 AND 4),phone TEXT DEFAULT '',repeat_year INTEGER DEFAULT 0,status TEXT DEFAULT 'active');
 CREATE TABLE IF NOT EXISTS terrains(id INTEGER PRIMARY KEY,name TEXT UNIQUE NOT NULL,kind TEXT DEFAULT 'basic',mandatory_year INTEGER DEFAULT 0,active INTEGER NOT NULL DEFAULT 1);
 CREATE TABLE IF NOT EXISTS periods(id INTEGER PRIMARY KEY,year_label TEXT NOT NULL,name TEXT NOT NULL,start_date TEXT NOT NULL DEFAULT '',end_date TEXT NOT NULL DEFAULT '',UNIQUE(year_label,name));
 CREATE TABLE IF NOT EXISTS capacities(id INTEGER PRIMARY KEY,year_label TEXT NOT NULL,terrain_id INTEGER,period_id INTEGER,places INTEGER NOT NULL,UNIQUE(year_label,terrain_id,period_id));
 CREATE TABLE IF NOT EXISTS assignments(id INTEGER PRIMARY KEY,resident_id INTEGER,terrain_id INTEGER,period_id INTEGER,status TEXT DEFAULT 'confirmed',UNIQUE(resident_id,period_id));
 CREATE TABLE IF NOT EXISTS history(id INTEGER PRIMARY KEY,resident_id INTEGER,terrain_id INTEGER,year_label TEXT,period TEXT,date_done TEXT);
 CREATE TABLE IF NOT EXISTS decisions(id INTEGER PRIMARY KEY,details TEXT,decision TEXT,reason TEXT,actor TEXT,postgrad_year TEXT NOT NULL DEFAULT '',created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY,action TEXT,details TEXT,actor TEXT DEFAULT '',created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT NOT NULL);
 """)
 # Additive migrations preserve existing resident and rotation-history records.
 for col, spec in {"birth_place":"TEXT DEFAULT ''","birth_date":"TEXT DEFAULT ''","registration_date":"TEXT DEFAULT ''","photo_path":"TEXT DEFAULT ''","dems_status":"TEXT DEFAULT ''","resident_type":"TEXT NOT NULL DEFAULT 'local'","origin_hospital":"TEXT DEFAULT ''","external_start_date":"TEXT DEFAULT ''","external_end_date":"TEXT DEFAULT ''","address":"TEXT DEFAULT ''","email":"TEXT DEFAULT ''","phone2":"TEXT DEFAULT ''","resident_card_number":"TEXT DEFAULT ''","national_id_number":"TEXT DEFAULT ''","medical_order_number":"TEXT DEFAULT ''","medical_order_registered":"TEXT DEFAULT ''","social_security_number":"TEXT DEFAULT ''","additional_info":"TEXT DEFAULT ''"}.items():
  ensure_column(c,"residents",col,spec)
 ensure_column(c,"terrains","active","INTEGER NOT NULL DEFAULT 1")
 ensure_column(c,"periods","start_date","TEXT NOT NULL DEFAULT ''")
 ensure_column(c,"periods","end_date","TEXT NOT NULL DEFAULT ''")
 # Replace old hard-coded date labels with neutral period names; dates are now configurable per postgraduate year.
 for old_name,new_name in (("01 Fév - 30 Avr","Période 1"),("01 Mai - 30 Juillet","Période 2"),("01 Août - 30 Oct","Période 3"),("01 Nov - 30 Jan","Période 4")):
  c.execute("UPDATE periods SET name=? WHERE name=?",(new_name,old_name))
 ensure_column(c,"audit","actor","TEXT DEFAULT ''")
 ensure_column(c,"decisions","postgrad_year","TEXT NOT NULL DEFAULT ''")
 c.executescript("""
 CREATE TABLE IF NOT EXISTS exam_records(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL,study_year INTEGER NOT NULL,module TEXT NOT NULL,theory REAL NOT NULL,practical REAL NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS objectives(id INTEGER PRIMARY KEY,terrain_id INTEGER NOT NULL,study_year INTEGER NOT NULL,objective TEXT NOT NULL,completed INTEGER NOT NULL DEFAULT 0);
 CREATE TABLE IF NOT EXISTS dems_archive(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL UNIQUE,success_date TEXT NOT NULL,profile_json TEXT NOT NULL,history_json TEXT NOT NULL,archived_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS promotion_votes(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL,from_year INTEGER NOT NULL,decision TEXT NOT NULL CHECK(decision IN ('approve','reject')),reason TEXT NOT NULL,actor TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP,UNIQUE(resident_id,from_year,actor));
 CREATE TABLE IF NOT EXISTS dems_votes(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL,success_date TEXT NOT NULL,decision TEXT NOT NULL CHECK(decision IN ('approve','reject')),reason TEXT NOT NULL,actor TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP,UNIQUE(resident_id,success_date,actor));
 CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY,username TEXT UNIQUE NOT NULL,password_hash TEXT NOT NULL,role TEXT NOT NULL CHECK(role IN ('super_admin','user')),active INTEGER NOT NULL DEFAULT 1,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS pending_registrations(id INTEGER PRIMARY KEY,username TEXT NOT NULL,email TEXT NOT NULL UNIQUE,password_hash TEXT NOT NULL,code_hash TEXT NOT NULL,expires_at INTEGER NOT NULL,attempts INTEGER NOT NULL DEFAULT 0,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS import_previews(id INTEGER PRIMARY KEY,owner TEXT NOT NULL,rows_json TEXT NOT NULL,errors_json TEXT NOT NULL,filename TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS leave_years(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL,year INTEGER NOT NULL,earned_days INTEGER NOT NULL DEFAULT 30,created_by TEXT DEFAULT '',UNIQUE(resident_id,year));
 CREATE TABLE IF NOT EXISTS leave_entries(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL,start_date TEXT NOT NULL,end_date TEXT NOT NULL,days INTEGER NOT NULL,notes TEXT DEFAULT '',created_by TEXT DEFAULT '',created_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS external_rotations(id INTEGER PRIMARY KEY,resident_id INTEGER NOT NULL,supervisor TEXT NOT NULL,destination TEXT NOT NULL,hospital TEXT NOT NULL,country TEXT NOT NULL,rotation_type TEXT NOT NULL,start_date TEXT NOT NULL,end_date TEXT NOT NULL,notes TEXT DEFAULT '',status TEXT NOT NULL DEFAULT 'active',request_status TEXT NOT NULL DEFAULT 'approved',created_by TEXT DEFAULT '',created_at TEXT DEFAULT CURRENT_TIMESTAMP,updated_at TEXT DEFAULT CURRENT_TIMESTAMP);
 CREATE TABLE IF NOT EXISTS trash_items(id INTEGER PRIMARY KEY,entity_type TEXT NOT NULL,entity_id TEXT NOT NULL,label TEXT NOT NULL,reason TEXT NOT NULL,snapshot_json TEXT NOT NULL,deleted_at TEXT NOT NULL,deleted_by TEXT NOT NULL,UNIQUE(entity_type,entity_id));
 """)
 ensure_column(c,"exam_records","exam_type","TEXT NOT NULL DEFAULT 'actual'")
 ensure_column(c,"exam_records","attempt","INTEGER NOT NULL DEFAULT 1")
 ensure_column(c,"exam_records","postgrad_year","TEXT NOT NULL DEFAULT ''")
 # Migrate user email fields only after the users table exists (fresh DB and legacy DB).
 ensure_column(c,"users","email","TEXT DEFAULT ''")
 ensure_column(c,"users","email_verified","INTEGER NOT NULL DEFAULT 0")
 ensure_column(c,"import_previews","import_type","TEXT NOT NULL DEFAULT 'residents'")
 ensure_column(c,"objectives","completed","INTEGER NOT NULL DEFAULT 0")
 ensure_column(c,"external_rotations","request_status","TEXT NOT NULL DEFAULT 'approved'")
 for col,spec in {"archive_type":"TEXT NOT NULL DEFAULT ''","archive_reason":"TEXT NOT NULL DEFAULT ''","archive_return_date":"TEXT NOT NULL DEFAULT ''","archived_at":"TEXT NOT NULL DEFAULT ''"}.items(): ensure_column(c,"residents",col,spec)
 y=str(datetime.now().year)+"-"+str(datetime.now().year+1)
 c.execute("INSERT OR IGNORE INTO settings VALUES('year',?)",(y,))
 initial_password=os.environ.get("ROTATION_INITIAL_PASSWORD")
 if initial_password and not setting(c,"password"):
  c.execute("INSERT OR IGNORE INTO settings VALUES('password',?)",(hash_password(initial_password),))
 c.execute("INSERT OR IGNORE INTO settings VALUES('director',?)",("Professor Benhocine",))
 for n,k in [("Réanimation chirurgicale","basic"),("Réanimation polyvalente","basic"),("Chirurgie générale","mandatory"),("Urologie","basic"),("Orthopédie","basic"),("Neurochirurgie","basic"),("Sédation","optional"),("Maxillo-faciale","optional"),("CCI","basic"),("Obstétrique","basic"),("Chirurgie cardiaque","optional"),("PU Médecine","basic"),("PU Chirurgie","basic"),("ORL","basic"),("Ophtalmologie","basic"),("Gynécologie","basic"),("Chirurgie thoracique","basic"),("Brûlés","optional"),("CAC DBK","optional")]:
  c.execute("INSERT OR IGNORE INTO terrains(name,kind) VALUES(?,?)",(n,k))
 # One-time correction for databases created before the agreed terrain classification.
 # Do not reapply on every startup, so later authorized admin edits remain intact.
 if not setting(c,"terrain_classification_migration_20261017"):
  for name in ("Sédation","Maxillo-faciale","Chirurgie cardiaque","Brûlés","CAC DBK"):
   c.execute("UPDATE terrains SET kind='optional' WHERE name=?",(name,))
  c.execute("UPDATE terrains SET kind='basic' WHERE name='Gynécologie'")
  c.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",("terrain_classification_migration_20261017","1"))
 y=setting(c,"year",y)
 # Fill any missing canonical periods in legacy/partially initialized databases.
 for n in ["Période 1","Période 2","Période 3","Période 4"]:
  c.execute("INSERT OR IGNORE INTO periods(year_label,name) VALUES(?,?)",(y,n))
 install_sync_change_triggers(c)
 c.commit(); c.close()
@app.before_request
def boot():
 if not getattr(app,"_ready",False): init(); app._ready=True
 # Enforce six-calendar-month trash retention automatically, not only when a
 # Super Admin opens the trash page. Run once per process-local calendar day.
 today_key=date.today().isoformat()
 if getattr(app,"_trash_cleanup_date","") != today_key:
  c=db()
  try:
   now=datetime.now()
   for item in c.execute("SELECT id,deleted_at FROM trash_items").fetchall():
    try:
     deleted=datetime.fromisoformat(item["deleted_at"])
     month=deleted.month+6
     year=deleted.year+(month-1)//12
     month=(month-1)%12+1
     day=min(deleted.day,__import__("calendar").monthrange(year,month)[1])
     if now>=deleted.replace(year=year,month=month,day=day):
      c.execute("DELETE FROM trash_items WHERE id=?",(item["id"],))
      audit_event(c,"trash.expired",f"trash_id={item['id']}; purged_after_six_calendar_months")
    except (ValueError,TypeError):
     continue
   c.commit()
   app._trash_cleanup_date=today_key
  finally:
   c.close()

@app.before_request
def verify_csrf():
 if request.method=="POST" and request.path.startswith("/api/sync/") and request.headers.get("Authorization","").startswith("Bearer "):
  # API writes authenticate with a per-device bearer token; browser session CSRF is not used.
  return None
 if request.method=="POST":
  expected=session.get("_csrf_token","")
  supplied=request.form.get("_csrf_token","")
  if not expected or not supplied or not hmac.compare_digest(expected,supplied):
   return "Invalid or missing CSRF token",400

def page(body,**kw):
 lang=session.get("lang","fr")
 csrf_token=session.setdefault("_csrf_token",secrets.token_urlsafe(32))
 body=re.sub(r'(<form\b(?=[^>]*\bmethod\s*=\s*["\']post["\'])[^>]*>)',lambda match:match.group(1)+'<input type="hidden" name="_csrf_token" value="{{ csrf_token }}">',body,flags=re.IGNORECASE)
 shell="""<!doctype html><html lang="{{lang}}" dir="{{'rtl' if lang=='ar' else 'ltr'}}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="theme-color" content="#103c66"><meta name="application-name" content="ResiManage"><link rel="icon" type="image/jpeg" href="/static/resimanage-3d-icon.jpg"><title>{{t('app')}} — {{t('year')}}</title></head><style>*{box-sizing:border-box}html{width:100%;max-width:100%;overflow-x:hidden}body{font-family:Arial,sans-serif;background:#f1f5f9;margin:0;color:#182b3d;width:100%;max-width:100%;overflow-x:hidden;line-height:1.5}header{background:linear-gradient(135deg,#103c66 0%,#17689a 58%,#23a6a5 100%);color:white;padding:clamp(10px,3vw,18px);overflow-wrap:anywhere}.brand-lockup{display:flex;align-items:center;gap:14px;max-width:1050px;margin:0 auto}.brand-icon{width:clamp(48px,12vw,70px);height:auto;flex:0 0 auto;filter:drop-shadow(0 5px 5px #001b3380)}.brand-lockup h2{margin:0;font-size:clamp(21px,4vw,30px);letter-spacing:.2px}.brand-tagline{margin:2px 0 5px;opacity:.93;font-size:clamp(12px,2.5vw,15px)}.brand-year{font-size:12px;opacity:.9}main{width:100%;max-width:1050px;margin:0 auto;padding:clamp(7px,2.5vw,16px);min-width:0}nav{display:flex;align-items:flex-start;flex-wrap:wrap;gap:6px;overflow:visible;max-width:100%;background:white;padding:10px}.nav-group{position:relative;min-width:0}.nav-group summary{list-style:none;background:#164e83;color:white;padding:9px 12px;border-radius:7px;cursor:pointer;white-space:normal;overflow-wrap:anywhere;max-width:100%;line-height:1.3;text-align:center}.nav-group summary::-webkit-details-marker{display:none}.nav-group summary:after{content:" ▾";font-size:.85em}.nav-menu{display:flex;flex-direction:column;gap:7px;min-width:0;width:max-content;max-width:min(92vw,340px);padding:9px;margin-top:6px;background:#fff;border:1px solid #d7e1eb;border-radius:10px;box-shadow:0 6px 18px #182b3d22}.nav-menu a{display:block;width:100%;min-width:0;white-space:normal;overflow-wrap:anywhere;word-break:normal;line-height:1.45;padding:10px 12px;text-align:start}.main-nav>a{flex:0 0 auto}.nav-group[open]{z-index:20}.nav-group[open] .nav-menu{position:absolute;top:100%;inset-inline-start:0;max-width:min(92vw,340px);max-height:70vh;overflow-y:auto;overflow-x:hidden}.nav-group[open] .nav-menu a{white-space:normal;min-height:40px}a,button{background:#164e83;color:white;padding:9px 12px;border:0;border-radius:7px;text-decoration:none;display:inline-block;white-space:normal;overflow-wrap:anywhere;max-width:100%;cursor:pointer}section{background:white;padding:clamp(9px,2.5vw,15px);border:1px solid #d7e1eb;border-radius:10px;margin:10px 0;min-width:0;max-width:100%;overflow-wrap:anywhere}input,select,textarea{font:inherit;padding:9px;border:1px solid #bccada;border-radius:7px;width:100%;max-width:100%;margin:4px 0 9px}label{display:block;min-width:0}table{border-collapse:collapse;width:100%;min-width:540px}td,th{text-align:start;border-bottom:1px solid #ddd;padding:8px;overflow-wrap:anywhere}.wrap{max-width:100%;overflow-x:auto;-webkit-overflow-scrolling:touch}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(170px,100%),1fr));gap:8px;min-width:0}.msg{background:#fff0c2;padding:10px;border-radius:7px;overflow-wrap:anywhere}.good{color:green}.bad{color:#b21d1d}.muted{color:#526274}img,video,iframe{max-width:100%;height:auto}@media(max-width:600px){main{padding:7px}header{font-size:clamp(13px,3.5vw,16px)}nav{padding:7px;gap:5px}.nav-group summary,nav>a{font-size:13px;padding:8px 9px}.nav-menu{min-width:min(220px,82vw)}section{margin:7px 0;padding:10px}.grid{grid-template-columns:minmax(0,1fr)}table{min-width:540px}} [dir=rtl]{text-align:right}[dir=rtl] nav{direction:rtl}
/* ResiManage unified interior visual identity */
:root{--rt-navy:#142b49;--rt-blue:#2768a8;--rt-mint:#36cfae;--rt-cyan:#45c7e8;--rt-coral:#ff8b72;--rt-amber:#ffc36c;--rt-violet:#9b8cf5;--rt-ink:#21334b;--rt-line:#dce6f0;--rt-glass:#ffffffd9}
body{background:radial-gradient(ellipse at 8% 0%,#ffd2ad55,transparent 28%),radial-gradient(ellipse at 95% 22%,#7ce8d533,transparent 26%),linear-gradient(145deg,#f3f7fc,#edf3fa 58%,#f8f1f8);color:var(--rt-ink)}
header.brand-header{background:radial-gradient(ellipse at 85% 0%,#ffb66c45,transparent 34%),linear-gradient(115deg,#101d36,#1c4265 58%,#28566d);box-shadow:0 12px 32px #172d4825;border-bottom:3px solid transparent;border-image:linear-gradient(90deg,var(--rt-mint),var(--rt-cyan),var(--rt-coral),var(--rt-violet)) 1}
.brand-icon{filter:drop-shadow(0 7px 12px #0005)}main{padding-top:18px;padding-bottom:26px}
.main-nav{margin:0 auto;position:sticky;top:0;z-index:20;justify-content:center;align-items:center;gap:8px;background:#ffffffdf;backdrop-filter:blur(16px);border:1px solid #ffffff;border-bottom-color:#dce6f0;border-radius:0 0 18px 18px;box-shadow:0 9px 25px #17365b12}
.main-nav>a,.nav-group summary{border:1px solid #ffffff25;background:linear-gradient(135deg,#214e79,#286da2);box-shadow:0 4px 10px #1c42651a;transition:transform .18s ease,box-shadow .18s ease,filter .18s ease}
.main-nav>a:hover,.nav-group summary:hover{transform:translateY(-2px);box-shadow:0 7px 15px #1c42652a;filter:saturate(1.12)}
.main-nav>a:nth-of-type(1){background:linear-gradient(135deg,#2768a8,#3986c9)}
.main-nav>a:nth-of-type(2){background:linear-gradient(135deg,#168f88,#35c7a9)}
.main-nav>a:nth-of-type(3){background:linear-gradient(135deg,#cf7547,#ffab68)}
.main-nav>a:nth-of-type(4){background:linear-gradient(135deg,#6d62bf,#9c8df5)}
.nav-group summary{background:linear-gradient(135deg,#254a70,#477e9e)}
.nav-menu{background:#fffffff2;backdrop-filter:blur(15px);border:1px solid #d9e7f0;border-radius:14px;box-shadow:0 18px 45px #17365b25}
.nav-menu a{background:#f1f7fb;color:#244461;border:1px solid transparent;border-inline-start:3px solid var(--rt-cyan);border-radius:9px;transition:background .18s ease,transform .18s ease}
.nav-menu a:hover{background:#e3f7f1;transform:translateX(-2px)}
section{background:var(--rt-glass);backdrop-filter:blur(10px);border:1px solid #ffffff;box-shadow:0 9px 28px #2444610d;border-radius:17px;transition:box-shadow .2s ease}
section h2,section h3{color:#204c70}.nav-icon{display:inline-flex;align-items:center;justify-content:center;min-width:20px;margin-inline-end:5px;font-size:16px;line-height:1;color:#b9ffec}.nav-group summary .nav-icon{color:#ffcf91}.rt-quick-actions{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(210px,100%),1fr));gap:12px;margin:14px 0}.rt-action-card{display:flex;align-items:center;gap:13px;min-width:0;padding:15px;border-radius:16px;background:linear-gradient(135deg,#ffffffef,#f1f8fb);border:1px solid #fff;box-shadow:0 8px 22px #24446112;color:#244461;text-decoration:none}.rt-action-card:hover{transform:translateY(-3px)}.rt-action-icon{display:flex;align-items:center;justify-content:center;flex:0 0 46px;width:46px;height:46px;border-radius:14px;font-size:24px;color:#15364d;background:linear-gradient(135deg,#b9f6e8,#a8e8f7)}.rt-action-card:nth-child(2) .rt-action-icon{background:linear-gradient(135deg,#ffd6b2,#ffad8e)}.rt-action-card:nth-child(3) .rt-action-icon{background:linear-gradient(135deg,#d9d1ff,#b7b0fa)}.rt-action-card:nth-child(4) .rt-action-icon{background:linear-gradient(135deg,#fff0b9,#ffd18a)}.rt-action-card strong,.rt-action-card small{display:block}.rt-action-card strong{color:#214c6b}.rt-action-card small{font-size:12px;color:#657b8e}.rt-stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(145px,100%),1fr));gap:10px;margin:14px 0}.rt-stat{padding:15px;border-radius:15px;background:linear-gradient(135deg,#fff,#f0f7fc);border:1px solid #fff;box-shadow:0 6px 18px #2444610d}.rt-stat span{display:block;font-size:12px;color:#617b8f}.rt-stat strong{display:block;font-size:26px;color:#214f73;line-height:1.3}.rt-stat:nth-child(2) strong{color:#15977f}.rt-stat:nth-child(3) strong{color:#c8754f}.rt-stat:nth-child(4) strong{color:#7668c5}.rt-dashboard-title{display:flex;align-items:center;gap:10px}.rt-dashboard-title .rt-dot{width:11px;height:11px;border-radius:50%;background:linear-gradient(135deg,var(--rt-mint),var(--rt-cyan));box-shadow:0 0 0 5px #42d9bd1b}.rt-year-chip{display:inline-flex;min-width:28px;justify-content:center;padding:3px 8px;border-radius:9px;background:#e4f0ff;color:#285a92;font-weight:700}.rt-progress{display:inline-block;vertical-align:middle;width:66px;height:7px;overflow:hidden;border-radius:10px;background:#e5edf4}.rt-progress span{display:block;height:100%;border-radius:10px;background:linear-gradient(90deg,#36cfae,#45c7e8)}.rt-resident-link{background:transparent;color:#275b7f;padding:2px 0;box-shadow:none}
section h2:after,section h3:after{content:"";display:block;width:54px;height:3px;margin-top:7px;border-radius:5px;background:linear-gradient(90deg,var(--rt-mint),var(--rt-cyan),var(--rt-coral))}
a,button{background:linear-gradient(135deg,#24659c,#2b80b1);box-shadow:0 3px 9px #214e7918;transition:filter .18s ease,transform .18s ease,box-shadow .18s ease}
a:hover,button:hover{filter:brightness(1.07) saturate(1.1);box-shadow:0 6px 14px #214e7925}
button:active,a:active{transform:translateY(1px)}
input,select,textarea{background:#ffffffdc;border-color:#c8d8e7;border-radius:11px;transition:border-color .15s ease,box-shadow .15s ease}
input:focus,select:focus,textarea:focus{outline:none;border-color:#37bfae;box-shadow:0 0 0 3px #37cfae26}
table{border-collapse:separate;border-spacing:0;border:1px solid var(--rt-line);border-radius:12px;overflow:hidden;background:#fff}
th{background:linear-gradient(100deg,#e7f4f7,#eef0fc);color:#234b68;font-weight:700}
td,th{border-bottom:1px solid #e7edf4;padding:10px}
tr:last-child td{border-bottom:0}
tbody tr:nth-child(even),table tr:nth-child(even){background:#f8fbfe}
.msg{background:linear-gradient(100deg,#fff2d8,#fff8eb);border:1px solid #f4d49a;color:#72501c;border-radius:12px}
.good{color:#13856d;font-weight:700}.bad{color:#c34e50;font-weight:700}.muted{color:#60758a}
.rt-language{display:flex;align-items:center;gap:5px;margin-inline-start:auto}
.rt-language a{padding:7px 9px;border-radius:20px;background:#e9f5f7;color:#244c66;border:1px solid #d5e8ef;box-shadow:none;font-size:12px}
.rt-language a[aria-current=page]{background:linear-gradient(100deg,#2ac7a7,#55c9e8);color:#102f47;font-weight:700}
.rt-footer{text-align:center;padding:17px 10px 10px;color:#637b91;font-size:12px}
.rt-footer strong{color:#2d5877}.rt-footer .rt-sep{color:#ff986c;padding:0 4px}
@media(max-width:600px){.main-nav{position:relative;border-radius:0 0 14px 14px;justify-content:flex-start}.rt-language{width:100%;justify-content:center;margin:3px 0 0}.main-nav>a,.nav-group summary{font-size:12px}.rt-footer{line-height:1.9}}
@media(prefers-reduced-motion:reduce){*,*:before,*:after{scroll-behavior:auto!important;transition:none!important;animation:none!important}}
/* ResiManage dark navy mockup theme */
:root{color-scheme:dark;--rt-cyan:#19d9e8;--rt-mint:#19dfbd;--rt-text:#e8f7ff}
html,body{background:#041426!important;color:#e8f7ff!important}
body{background:radial-gradient(ellipse at 10% 0%,#0a467055,transparent 34%),linear-gradient(150deg,#020d1c,#061a30 58%,#041426)!important;color:#e8f7ff}
header.brand-header{background:linear-gradient(115deg,#020d1c,#07375b 58%,#075a70)!important;border-bottom:2px solid var(--rt-cyan)!important}
main{max-width:1120px}
.main-nav{background:#061a2be8!important;border:1px solid #0c4564!important;border-radius:0 0 18px 18px!important}
.main-nav>a,.nav-group summary{background:linear-gradient(135deg,#0b3557,#0b5572)!important;border:1px solid #17668a!important;color:#e9fbff!important;border-radius:11px!important}
.nav-menu{background:#071d32f5!important;border-color:#155a80!important}
.nav-menu a{background:#0a2b47!important;color:#dff8ff!important;border-color:#104765!important}
section,.rt-stat,.rt-action-card,.login-brand,.login-panel{background:linear-gradient(145deg,#08223aee,#061a30f5)!important;border:1px solid #125173!important;box-shadow:0 10px 28px #0004!important;color:#e8f7ff!important;border-radius:17px!important}
h1,h2,h3,h4,strong,th{color:#eafaff}
p,label,td,small,.muted{color:#a9ccdf}
a,button{background:linear-gradient(110deg,#087f9f,#09b9bd)!important;color:#f3feff!important;border:1px solid #24cce0!important;border-radius:11px!important}
input,select,textarea{background:#04182b!important;color:#e8f7ff!important;border:1px solid #176083!important;border-radius:10px!important}
th{background:#0b3552!important;color:#70eaf2!important}td{background:#061d32;border-bottom:1px solid #12415e!important}
.rt-dot{background:var(--rt-mint)!important}.rt-year-chip{background:#0b4c69!important;color:#8ff5f4!important}
.rt-progress{background:#123750!important}.rt-progress span{background:linear-gradient(90deg,#16d8bc,#16c8ed)!important}
.login-emblem{color:#20e2e7!important;text-shadow:0 0 25px #16d8ed!important}
@media(max-width:600px){main{padding:8px}.main-nav{position:relative;top:auto;justify-content:flex-start}.rt-stats,.rt-quick-actions{grid-template-columns:repeat(2,minmax(0,1fr))}}

/* Functional navigation redesign: task-based sidebar, clear active states, responsive mobile groups */
.rt-app-shell{display:grid;grid-template-columns:260px minmax(0,1fr);align-items:start;gap:18px;width:100%;max-width:1500px;margin:0 auto;padding:16px;direction:inherit}
.rt-app-shell>main{width:100%;max-width:none;min-width:0;margin:0;padding:0}
.main-nav{position:sticky;top:12px;z-index:30;display:flex;flex-direction:column;align-items:stretch;justify-content:flex-start;gap:7px;width:100%;min-width:0;padding:12px;background:#071a2beF!important;border:1px solid #15506b!important;border-radius:18px!important;box-shadow:0 12px 30px #00101f30!important;backdrop-filter:blur(18px)}
.main-nav>a,.nav-group summary{display:flex;align-items:center;gap:9px;width:100%;min-height:44px;padding:10px 12px!important;border:1px solid #174c68!important;border-radius:11px!important;background:#0a2b47!important;color:#dff8ff!important;box-shadow:none!important;text-align:start;font-size:14px;font-weight:600;transition:background .16s ease,border-color .16s ease,transform .16s ease}
.nav-group{width:100%;min-width:0}
.nav-group summary{list-style:none;cursor:pointer;justify-content:space-between}
.nav-group summary::-webkit-details-marker{display:none}
.nav-group summary:after{content:'⌄';font-size:14px;color:#74e7e8;transition:transform .16s ease}
.nav-group[open] summary:after{transform:rotate(180deg)}
.nav-group summary .nav-icon,.main-nav .nav-icon{display:inline-flex;align-items:center;justify-content:center;flex:0 0 24px;min-width:24px;height:24px;margin:0;color:#71e5e4;font-size:17px}
.nav-menu{position:static!important;display:flex;flex-direction:column;gap:4px;min-width:0;width:100%;max-width:none!important;max-height:none!important;overflow:visible!important;margin:5px 0 3px;padding:5px 7px;background:#041728!important;border:1px solid #123e57!important;border-radius:12px!important;box-shadow:none!important}
.nav-menu a{display:flex;align-items:center;width:100%;min-height:38px;padding:8px 10px!important;background:transparent!important;color:#b8d5e4!important;border:1px solid transparent!important;border-inline-start:3px solid transparent!important;border-radius:8px!important;box-shadow:none!important;font-size:13px;font-weight:500}
.main-nav a:hover,.nav-group summary:hover,.nav-menu a:hover{background:#0e3a56!important;color:#fff!important;border-color:#1a6a83!important;transform:none!important;box-shadow:none!important}
.main-nav a.is-active,.nav-menu a.is-active{background:linear-gradient(100deg,#0a6b7a,#0d536b)!important;color:#fff!important;border-color:#22c9c7!important;border-inline-start-color:#39e0c1!important;font-weight:700}
.rt-language{display:flex;align-items:center;justify-content:center;gap:6px;margin:8px 0 0;padding-top:12px;border-top:1px solid #17435b}
.rt-language a{min-width:40px;text-align:center}
.rt-language a[aria-current=page]{background:#19b9b5!important;color:#031b2c!important;border-color:#51e4d9!important}
.rt-app-shell section{margin:0 0 14px}
.rt-app-shell .rt-quick-actions{margin-top:16px}
:focus-visible{outline:3px solid #39dfcf!important;outline-offset:2px}
 .rt-app-shell:not(:has(.main-nav)){display:block;max-width:1120px} @media(max-width:900px){.rt-app-shell{grid-template-columns:220px minmax(0,1fr);gap:12px;padding:10px}.main-nav{padding:9px}}
@media(max-width:680px){.rt-app-shell{display:flex;flex-direction:column;gap:12px;padding:8px}.main-nav{position:relative;top:auto;display:grid;grid-template-columns:repeat(2,minmax(0,1fr));align-items:start;gap:7px;padding:9px;border-radius:14px!important}.main-nav>.nav-home{grid-column:1/-1}.main-nav>.nav-group{align-self:start}.nav-group summary{min-height:44px;font-size:12px!important;padding:8px!important}.nav-group[open] .nav-menu{position:static!important;grid-column:1/-1;max-height:42vh!important;overflow:auto!important}.nav-menu a{font-size:12px}.rt-language{grid-column:1/-1;margin-top:2px;padding-top:8px}.rt-app-shell>main{padding:0}.rt-app-shell section{margin-bottom:10px}}
@media(prefers-reduced-motion:reduce){.main-nav *{transition:none!important}}
.rt-profile-panel{padding:clamp(16px,3vw,28px);background:linear-gradient(145deg,#ffffffef,#f1f8fcf2);border:1px solid #fff;box-shadow:0 16px 38px #17365b14}.rt-profile-head{display:flex;align-items:center;gap:20px;padding:4px 0 22px;border-bottom:1px solid #dce7f0;margin-bottom:20px}.rt-profile-avatar[hidden]{display:none!important}.rt-avatar-wrap{flex:0 0 auto;width:clamp(220px,30vw,280px)}.rt-profile-avatar{display:block;width:100%;height:clamp(220px,30vw,280px);object-fit:cover;border-radius:24px;border:4px solid white;box-shadow:0 10px 24px #17365b24;background:#e8f2f7}.rt-avatar-placeholder{display:grid;place-items:center;font-size:48px}.rt-profile-heading{min-width:0}.rt-profile-heading h2{margin:4px 0 8px;font-size:clamp(22px,4vw,32px);overflow-wrap:anywhere}.rt-eyebrow{font-size:11px;letter-spacing:1.3px;font-weight:800;color:#168c91}.rt-chip{display:inline-flex;align-items:center;padding:4px 10px;margin-inline-end:6px;border-radius:999px;background:#e5efff;color:#285b8c;font-size:12px;font-weight:700}.rt-chip-active{background:#dcf8eb;color:#167b58}.rt-chip-muted{background:#eef1f5;color:#596a7c}.rt-profile-form{display:grid;gap:16px}.rt-field-grid{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:6px 16px}.rt-field-grid label{font-weight:700;color:#34536d}.rt-field-grid input{margin-top:7px;background:#fff}.rt-photo-picker{padding:15px;border:1px dashed #8cb7ca;border-radius:15px;background:linear-gradient(135deg,#f0fbfc,#f7f6ff)}.rt-photo-picker label{font-weight:800;color:#285c79}.rt-photo-picker input[type=file]{background:#fff;border:1px solid #d6e4ee;padding:10px}.rt-photo-picker small{display:block;color:#61778a}.rt-savebar{display:flex;flex-wrap:wrap;align-items:center;gap:10px}.rt-savebar button{min-width:190px;font-weight:800;padding:12px 18px}.rt-secondary-action{background:#edf4f9;color:#31536f;box-shadow:none;border:1px solid #dce7f0}.rt-profile-links{display:flex;flex-wrap:wrap;gap:10px;margin-top:18px}.rt-profile-links a{background:transparent;color:#28668e;box-shadow:none;padding:4px 0}.rt-profile-form input:focus-visible,button:focus-visible,a:focus-visible{outline:3px solid #39bda9;outline-offset:3px}@media(max-width:620px){.rt-profile-head{align-items:stretch;flex-direction:column;gap:14px}.rt-avatar-wrap{width:min(260px,100%);flex-basis:auto;align-self:center}.rt-profile-avatar{height:260px;border-radius:20px}.rt-field-grid{grid-template-columns:minmax(0,1fr)}.rt-savebar{align-items:stretch;flex-direction:column}.rt-savebar button,.rt-savebar a{width:100%;text-align:center}.rt-profile-panel{padding:14px}}@media(prefers-reduced-motion:reduce){*,*:before,*:after{scroll-behavior:auto!important;transition:none!important;animation:none!important}}
/* Selectable 3D interior themes */
html[data-interior-theme="emerald"]{color-scheme:dark}
html[data-interior-theme="emerald"],html[data-interior-theme="emerald"] body{background:#061c1b!important;color:#e3fff6!important}
html[data-interior-theme="emerald"] body{background:radial-gradient(ellipse at 8% 0%,#0b625044,transparent 34%),radial-gradient(ellipse at 92% 20%,#1b7d7040,transparent 28%),linear-gradient(145deg,#041716,#082a29 58%,#071e24)!important}
html[data-interior-theme="emerald"] header.brand-header{background:linear-gradient(125deg,#031b1b,#07534b 58%,#14746a)!important;border-bottom:2px solid #48e0bd!important}
html[data-interior-theme="emerald"] .main-nav{background:#052421ed!important;border-color:#0f6558!important}
html[data-interior-theme="emerald"] .main-nav>a,html[data-interior-theme="emerald"] .nav-group summary{background:linear-gradient(135deg,#0b453f,#116b60)!important;color:#e3fff6!important;border-color:#168775!important}
html[data-interior-theme="emerald"] .nav-icon,html[data-interior-theme="emerald"] .nav-group summary:after{color:#7cf3cf!important}
html[data-interior-theme="emerald"] .nav-menu{background:#052421f5!important;border-color:#12665a!important}
html[data-interior-theme="emerald"] .nav-menu a{background:#08332f!important;color:#d5fff1!important;border-color:#10554c!important}
html[data-interior-theme="emerald"] section,html[data-interior-theme="emerald"] .rt-stat,html[data-interior-theme="emerald"] .rt-action-card,html[data-interior-theme="emerald"] .login-brand,html[data-interior-theme="emerald"] .login-panel{background:linear-gradient(145deg,#0b302cf2,#082321f5)!important;border-color:#14594f!important;box-shadow:0 12px 32px #0004!important;color:#e3fff6!important}
html[data-interior-theme="emerald"] h1,html[data-interior-theme="emerald"] h2,html[data-interior-theme="emerald"] h3,html[data-interior-theme="emerald"] h4,html[data-interior-theme="emerald"] strong,html[data-interior-theme="emerald"] th{color:#e6fff8!important}
html[data-interior-theme="emerald"] p,html[data-interior-theme="emerald"] label,html[data-interior-theme="emerald"] td,html[data-interior-theme="emerald"] small,html[data-interior-theme="emerald"] .muted{color:#b0d8cc!important}
html[data-interior-theme="emerald"] a,html[data-interior-theme="emerald"] button{background:linear-gradient(120deg,#087c6b,#168a86)!important;color:#fff!important;border-color:#2ac5a6!important}
html[data-interior-theme="emerald"] input,html[data-interior-theme="emerald"] select,html[data-interior-theme="emerald"] textarea{background:#05211f!important;color:#e3fff6!important;border-color:#17685c!important}
html[data-interior-theme="emerald"] th{background:#10483f!important;color:#9ef5d7!important}
html[data-interior-theme="emerald"] td{background:#082a27!important;border-bottom-color:#14594f!important}
html[data-interior-theme="emerald"] table tr:nth-child(even){background:#0b332f!important}
html[data-interior-theme="emerald"] .msg{background:#46381d!important;color:#ffe8a8!important;border-color:#806a31!important}
html[data-interior-theme="emerald"] .rt-footer,html[data-interior-theme="emerald"] .rt-footer strong{color:#a6d3c5!important}
.theme-chooser{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(190px,100%),1fr));gap:12px;margin:12px 0}
.theme-choice{display:flex;flex-direction:column;align-items:stretch;gap:8px;text-align:start;min-height:112px;padding:12px!important;border:2px solid transparent!important;border-radius:15px!important;white-space:normal!important}
.theme-choice[aria-pressed="true"]{outline:3px solid #50d8c0;outline-offset:2px}
.theme-swatch{display:flex;gap:5px;height:27px;border-radius:9px;padding:4px;width:100%}
.theme-swatch i{display:block;flex:1;border-radius:5px;box-shadow:inset 0 1px 2px #ffffff55,0 2px 4px #0002}
html[data-interior-theme="holographic"] body{background:radial-gradient(ellipse at 10% 5%,#5524a966,transparent 38%),linear-gradient(145deg,#09051c,#160d35 55%,#07172d)!important}html[data-interior-theme="holographic"] section,html[data-interior-theme="holographic"] .rt-stat,html[data-interior-theme="holographic"] .rt-action-card,html[data-interior-theme="holographic"] .login-brand,html[data-interior-theme="holographic"] .login-panel{background:linear-gradient(145deg,#211044ee,#100a29f5)!important;border-color:#7748b9!important;box-shadow:0 12px 32px #0006,0 0 18px #8b5cf633!important}html[data-interior-theme="holographic"] a,html[data-interior-theme="holographic"] button{background:linear-gradient(120deg,#6d28d9,#c026d3)!important;border-color:#d879f9!important}html[data-interior-theme="titanium"] body{background:radial-gradient(ellipse at 10% 0%,#64748b44,transparent 38%),linear-gradient(145deg,#080b11,#1b2432 55%,#080b11)!important}html[data-interior-theme="titanium"] section,html[data-interior-theme="titanium"] .rt-stat,html[data-interior-theme="titanium"] .rt-action-card,html[data-interior-theme="titanium"] .login-brand,html[data-interior-theme="titanium"] .login-panel{background:linear-gradient(145deg,#252d3bee,#101722f5)!important;border-color:#64748b!important}html[data-interior-theme="titanium"] a,html[data-interior-theme="titanium"] button{background:linear-gradient(120deg,#475569,#334155)!important;border-color:#94a3b8!important}html[data-interior-theme="aurora"] body{background:radial-gradient(ellipse at 15% 0%,#0e749055,transparent 38%),linear-gradient(145deg,#021a25,#063542 55%,#041426)!important}html[data-interior-theme="aurora"] section,html[data-interior-theme="aurora"] .rt-stat,html[data-interior-theme="aurora"] .rt-action-card,html[data-interior-theme="aurora"] .login-brand,html[data-interior-theme="aurora"] .login-panel{background:linear-gradient(145deg,#073744ee,#05252ff5)!important;border-color:#0e7490!important}html[data-interior-theme="aurora"] a,html[data-interior-theme="aurora"] button{background:linear-gradient(120deg,#0891b2,#0d9488)!important;border-color:#2dd4bf!important}html[data-interior-theme="ruby"] body{background:radial-gradient(ellipse at 15% 0%,#9f123955,transparent 38%),linear-gradient(145deg,#160713,#300d24 55%,#100817)!important}html[data-interior-theme="ruby"] section,html[data-interior-theme="ruby"] .rt-stat,html[data-interior-theme="ruby"] .rt-action-card,html[data-interior-theme="ruby"] .login-brand,html[data-interior-theme="ruby"] .login-panel{background:linear-gradient(145deg,#40152fee,#210d20f5)!important;border-color:#9f456e!important}html[data-interior-theme="ruby"] a,html[data-interior-theme="ruby"] button{background:linear-gradient(120deg,#9f1239,#be557c)!important;border-color:#fb7185!important}.rt-year-select{display:flex;flex-direction:column;text-decoration:none;cursor:pointer}.rt-year-select small{font-size:11px;opacity:.8}.rt-stat{min-width:0}.rt-quick-actions .rt-action-card{min-width:0;overflow-wrap:anywhere}.rt-action-card strong{overflow-wrap:anywhere;word-break:normal}.login-emblem{animation:heartbeat 1.8s ease-in-out infinite!important;filter:drop-shadow(0 0 10px #20e2e788)}

/* ResiManage cinematic medical UI: deep navy, cyan pulse, glass/3D surfaces */
:root{--rm-bg:#03101f;--rm-panel:#071d35;--rm-panel2:#0b2946;--rm-cyan:#18e6e0;--rm-blue:#08a9ff;--rm-line:#1477a6;--rm-text:#eafaff;--rm-muted:#9cc7dc}
html{background:#020b18!important;color-scheme:dark}
body{background:radial-gradient(ellipse at 8% 8%,#087cba35 0,transparent 34%),radial-gradient(ellipse at 90% 0%,#00e5d425 0,transparent 30%),linear-gradient(145deg,#020815 0%,#06182c 48%,#020b18 100%)!important;color:var(--rm-text)!important;background-attachment:fixed!important}
body:before{content:"";position:fixed;inset:0;pointer-events:none;z-index:-1;opacity:.17;background:repeating-linear-gradient(0deg,transparent 0 44px,#16c8e51a 45px 46px),repeating-linear-gradient(90deg,transparent 0 44px,#16c8e51a 45px 46px);mask-image:linear-gradient(to bottom,black,transparent 78%)}
.brand-header{background:linear-gradient(115deg,#03162bea,#06406bea 58%,#007f91e8)!important;border-bottom:1px solid #20dff080!important;box-shadow:0 12px 36px #00bde52b,inset 0 -1px #4dfaff35!important;backdrop-filter:blur(16px)}
.brand-lockup{position:relative}.brand-icon{border-radius:17px;box-shadow:0 0 0 1px #30e9ff8c,0 0 24px #00cfff72,0 12px 25px #0009;transform:perspective(500px) rotateY(-7deg);animation:rmLogoFloat 5s ease-in-out infinite}
.brand-header h2{color:#f1feff!important;text-shadow:0 0 15px #00d9ff80;letter-spacing:.02em}.brand-tagline,.brand-year{color:#a8e8f5!important}
.rt-app-shell{max-width:1280px;margin:auto}
.main-nav{background:linear-gradient(180deg,#061b31f2,#031324f5)!important;border:1px solid #0f729b!important;border-radius:18px!important;box-shadow:0 14px 34px #0006,inset 0 1px #ffffff12!important;backdrop-filter:blur(16px)}
main, .rt-content{color:var(--rm-text)!important}
section,.rt-stat,.rt-action-card,.login-brand,.login-panel,.rt-card,.panel,.form-card,.resident-card,.dashboard-card{background:linear-gradient(145deg,#0b2b49ed 0%,#06182cf5 65%,#071d34f5)!important;color:var(--rm-text)!important;border:1px solid #137eaa!important;border-radius:18px!important;box-shadow:0 12px 28px #0005,0 0 0 1px #16d9ee12,inset 0 1px #c2f9ff18!important;backdrop-filter:blur(12px);transition:transform .22s ease,box-shadow .22s ease,border-color .22s ease}
.rt-action-card:hover,.rt-stat:hover,section:hover{transform:translateY(-3px);border-color:#26e7ee!important;box-shadow:0 18px 38px #0008,0 0 24px #00c9e52b,inset 0 1px #d2ffff2b!important}
h1,h2,h3,h4,label,strong{color:#eafaff!important}p,small,.muted{color:#a6cde0}
button,input,select,textarea,.btn,a.button{border-radius:12px!important}
button,.btn,a.button,input[type=submit]{background:linear-gradient(120deg,#00b9d9,#087fe0 58%,#05d8c8)!important;color:#001827!important;border:1px solid #56f8f0!important;box-shadow:0 5px 16px #00b9d936,inset 0 1px #ffffff6b!important;font-weight:700!important;transition:transform .18s ease,filter .18s ease,box-shadow .18s ease}
button:hover,.btn:hover,a.button:hover,input[type=submit]:hover{transform:translateY(-2px);filter:brightness(1.13);box-shadow:0 9px 24px #00d8ef50,inset 0 1px #fff9!important}
a{color:#54e7f3}input,select,textarea{background:#031427!important;color:#eaffff!important;border:1px solid #17658b!important;box-shadow:inset 0 3px 8px #0005!important}input:focus,select:focus,textarea:focus{outline:2px solid #14d9e8!important;outline-offset:1px;box-shadow:0 0 18px #00d5e94a!important}
table{color:#eafaff!important;border-color:#155b7b!important}th{background:#0b3554!important;color:#9ffaff!important}td{border-color:#144361!important}tr:nth-child(even){background:#082039a8!important}
.badge,.chip,.status-pill{box-shadow:0 0 12px #00c9d52b}
.login-emblem{filter:drop-shadow(0 0 18px #1ce8ffb5)!important;animation:rmPulse 2.4s ease-in-out infinite!important}
@keyframes rmLogoFloat{0%,100%{transform:perspective(500px) rotateY(-7deg) translateY(0)}50%{transform:perspective(500px) rotateY(7deg) translateY(-4px)}}
@keyframes rmPulse{0%,100%{filter:drop-shadow(0 0 10px #1ce8ff7a)}50%{filter:drop-shadow(0 0 25px #1ce8ffe0)}}
@media(prefers-reduced-motion:reduce){*,*:before,*:after{animation:none!important;transition:none!important;scroll-behavior:auto!important}}

/* ResiManage 3D visual pass inspired by the user's interior reference board.
   Keep the existing medical imagery, navigation, labels and controls intact. */
:root{--rm-depth:0 24px 52px #0008;--rm-glass:linear-gradient(145deg,#0e3557f5,#06162cf7 62%,#09213af2)}
body{perspective:1200px}
body:after{content:"";position:fixed;inset:-12%;z-index:-1;pointer-events:none;opacity:.32;background:radial-gradient(ellipse at 20% 25%,#0ae5ff18,transparent 25%),radial-gradient(ellipse at 78% 32%,#20f4d018,transparent 22%),radial-gradient(ellipse at 52% 90%,#347dff18,transparent 30%);animation:rmAmbientDrift 18s ease-in-out infinite alternate}
.brand-header{position:relative;overflow:hidden;isolation:isolate}
.brand-header:after{content:"";position:absolute;inset:0;z-index:-1;pointer-events:none;background:linear-gradient(115deg,transparent 15%,#67f9ff13 43%,transparent 68%);transform:translateX(-100%);animation:rmLightSweep 9s ease-in-out infinite}
.brand-lockup{transform-style:preserve-3d;gap:16px}
.brand-icon{object-fit:cover;border:1px solid #77fbffad;box-shadow:0 0 0 2px #0b6c8b70,0 0 28px #00dfff85,0 16px 30px #000b,inset 0 1px #fff9}
.brand-header h2{font-weight:850;text-shadow:0 2px 0 #00415c,0 0 22px #13dfff9c;letter-spacing:.015em}
.rt-app-shell{transform-style:preserve-3d}
.main-nav{box-shadow:0 24px 42px #0008,inset 0 1px #ffffff20,inset 0 -1px #00eaff20!important}
.main-nav>a,.nav-group summary{position:relative;overflow:hidden;box-shadow:inset 0 1px #ffffff22,0 5px 0 #03101f,0 9px 14px #0005!important;transform:translateY(0);transition:transform .22s ease,box-shadow .22s ease,filter .22s ease!important}
.main-nav>a:before,.nav-group summary:before{content:"";position:absolute;inset:0;pointer-events:none;background:linear-gradient(120deg,#ffffff13,transparent 38%,#00e5ff08 70%,#ffffff04)}
.main-nav>a:hover,.nav-group summary:hover{transform:translateY(-2px)!important;box-shadow:inset 0 1px #ffffff35,0 7px 0 #03101f,0 14px 22px #0008,0 0 18px #00dfff22!important}
.main-nav a.is-active,.nav-menu a.is-active{box-shadow:inset 0 1px #ffffff50,0 5px 0 #034257,0 10px 22px #00dfff24!important}
section,.rt-stat,.rt-action-card,.rt-card,.panel,.form-card,.resident-card,.dashboard-card{position:relative;isolation:isolate;overflow:hidden;transform-style:preserve-3d;box-shadow:0 18px 34px #0007,0 0 0 1px #26eaff10,inset 0 1px #e7ffff22,inset 0 -1px #00bde51b!important}
section:before,.rt-stat:before,.rt-action-card:before,.rt-card:before,.panel:before,.form-card:before,.resident-card:before,.dashboard-card:before{content:"";position:absolute;inset:0;z-index:-1;pointer-events:none;border-radius:inherit;background:linear-gradient(130deg,#ffffff0b,transparent 28%,transparent 70%,#11dff10a)}
.rt-stat,.rt-action-card{transition:transform .25s cubic-bezier(.2,.8,.2,1),box-shadow .25s ease,border-color .25s ease!important}
.rt-action-card:hover,.rt-stat:hover{transform:translateY(-5px) rotateX(1deg)!important;box-shadow:0 25px 45px #0009,0 0 25px #00dfff30,inset 0 1px #ffffff35!important}
button,.btn,a.button,input[type=submit]{position:relative;overflow:hidden;box-shadow:0 7px 0 #03476a,0 12px 24px #00101f80,inset 0 1px #ffffff8c,inset 0 -2px #002f4e55!important;transform:translateY(0);transition:transform .16s ease,filter .16s ease,box-shadow .16s ease!important}
button:active,.btn:active,a.button:active,input[type=submit]:active{transform:translateY(4px)!important;box-shadow:0 3px 0 #03476a,0 5px 12px #00101f80,inset 0 1px #ffffff55!important}
button:after,.btn:after,a.button:after,input[type=submit]:after{content:"";position:absolute;inset:-80% auto -80% -45%;width:28%;pointer-events:none;transform:rotate(22deg);background:linear-gradient(90deg,transparent,#ffffff55,transparent);transition:left .55s ease}
button:hover:after,.btn:hover:after,a.button:hover:after,input[type=submit]:hover:after{left:125%}
input,select,textarea{box-shadow:inset 0 4px 12px #0008,0 1px #ffffff0d!important}
.login-hero{border:1px solid #35e9ff80;box-shadow:0 34px 90px #000a,0 0 0 1px #18c8ff18,inset 0 1px #ffffff40, inset 0 -2px #00dfff36!important;transform-style:preserve-3d}
.login-hero:before{box-shadow:inset 0 0 40px #d5ffff18,0 18px 50px #00dfff25}
.login-emblem{border:1px solid #8effffb0!important;box-shadow:0 0 0 3px #0c72945c,0 0 34px #17e8ffb0,0 18px 28px #000b!important;transform-style:preserve-3d;animation:rmEmblem3D 5s ease-in-out infinite!important}
.login-panel{border-color:#5beeff91!important;box-shadow:0 25px 55px #0008,inset 0 1px #ffffff2b,inset 0 -1px #00dfff30!important}
.login-submit{box-shadow:0 7px 0 #00607a,0 14px 26px #00cde944,inset 0 1px #ffffffbd!important}
.login-feature-row span{border-color:#28dff177!important;box-shadow:0 8px 18px #0005,inset 0 1px #ffffff24;background:linear-gradient(145deg,#0b3654cc,#061a2bcc)!important}
.rt-profile-avatar{box-shadow:0 0 0 3px #0b5575,0 16px 32px #0008,0 0 26px #00cfff35!important;transform:perspective(800px) rotateY(-2deg)}
.rt-photo-picker{box-shadow:inset 0 1px #ffffff24,0 12px 22px #0004!important}
table{border-radius:12px;overflow:hidden;box-shadow:0 14px 28px #0005}
th{box-shadow:inset 0 1px #ffffff1e}
.msg{box-shadow:0 8px 20px #0004,inset 0 1px #ffffff25}
@keyframes rmAmbientDrift{0%{transform:translate3d(-1%,0,0) scale(1)}100%{transform:translate3d(1.5%,1%,0) scale(1.04)}}
@keyframes rmLightSweep{0%,55%{transform:translateX(-100%)}80%,100%{transform:translateX(100%)}}
@keyframes rmEmblem3D{0%,100%{transform:perspective(700px) rotateY(-4deg) rotateX(1deg) translateY(0)}50%{transform:perspective(700px) rotateY(4deg) rotateX(-1deg) translateY(-5px)}}
@media(max-width:680px){.brand-icon{width:52px!important;height:52px!important;border-radius:14px!important}.main-nav>a,.nav-group summary{box-shadow:inset 0 1px #ffffff1f,0 4px 0 #03101f,0 7px 12px #0005!important}.rt-action-card:hover,.rt-stat:hover{transform:translateY(-2px)!important}}
@media(prefers-reduced-motion:reduce){body:after,.brand-header:after,.login-emblem{animation:none!important}button:after,.btn:after,a.button:after,input[type=submit]:after{display:none!important}}

.rt-history-yearbar{display:flex;align-items:end;justify-content:space-between;gap:14px;flex-wrap:wrap;margin-bottom:12px}.rt-year-inline{display:flex;align-items:center;gap:8px;flex-wrap:wrap}.rt-year-inline label{font-weight:800}.rt-year-inline select{min-width:150px}@media(max-width:600px){.rt-history-yearbar{align-items:stretch;flex-direction:column}.rt-year-inline{align-items:stretch;flex-direction:column}.rt-year-inline select{width:100%}}\n@media(max-width:680px){html{overflow-x:clip;overflow-y:auto;scroll-behavior:auto!important;touch-action:pan-y}body{overflow-x:clip;min-height:100%;touch-action:pan-y}.main-nav{backdrop-filter:none!important;-webkit-backdrop-filter:none!important;transform:none!important;align-items:stretch}.main-nav>a:hover,.nav-group summary:hover{transform:none!important}.wrap{-webkit-overflow-scrolling:touch;overscroll-behavior-x:contain}.nav-group[open]{flex:1 0 100%;min-width:0}.nav-group[open] .nav-menu{position:static;inset:auto;width:100%;max-width:100%;max-height:none;overflow:visible;overscroll-behavior:auto;-webkit-overflow-scrolling:auto;box-shadow:none;margin:5px 0 0}.nav-group[open] .nav-menu a{min-height:44px}.nav-group summary{touch-action:manipulation}}\n</style><header class="brand-header"><div class="brand-lockup"><img class="brand-icon" src="/static/resimanage-3d-icon.jpg" width="64" height="64" alt="ResiManage"><div><h2>{{t('app')}}</h2><p class="brand-tagline">{% if lang=='ar' %}تسيير تربصات الأطباء المقيمين{% elif lang=='fr' %}Gestion des stages des médecins résidents{% else %}Resident rotation management{% endif %}</p><span class="brand-year">{{t('year')}}: {{year}}</span></div></div></header><div class="rt-app-shell">{% if user %}<nav class="main-nav" aria-label="{{L('Navigation principale','Main navigation','التنقل الرئيسي')}}">
<a class="nav-home {% if request.path=='/' %}is-active{% endif %}" href="/"><span class="nav-icon">⌂</span><span>{{t('home')}}</span></a>
<details class="nav-group" {% if request.path in ['/residents','/objectives','/import','/exams','/dems'] %}open{% endif %}><summary><span class="nav-icon">♙</span><span>{{t('residents')}}</span></summary><div class="nav-menu"><a href="/residents" class="{% if request.path=='/residents' %}is-active{% endif %}">{{t('manage_residents')}}</a><a href="/objectives" class="{% if request.path=='/objectives' %}is-active{% endif %}">{{t('objectives')}}</a><a href="/import" class="{% if request.path=='/import' %}is-active{% endif %}">{{t('import_residents')}}</a><a href="/templates" class="{% if request.path=='/templates' %}is-active{% endif %}">{{L('Formulaires vierges','Blank forms','استمارات فارغة')}}</a><a href="/exams" class="{% if request.path=='/exams' %}is-active{% endif %}">{{t('promotion')}}</a><a href="/dems" class="{% if request.path=='/dems' %}is-active{% endif %}">{{t('dems_success')}}</a></div></details>
<details class="nav-group" {% if request.path in ['/terrains','/assign','/history','/leave','/external-rotations','/rotation-board'] %}open{% endif %}><summary><span class="nav-icon">⇄</span><span>{% if lang=='ar' %}التربصات والتخطيط{% elif lang=='fr' %}Stages et planification{% else %}Rotations & planning{% endif %}</span></summary><div class="nav-menu"><a href="/terrains" class="{% if request.path=='/terrains' %}is-active{% endif %}">{{t('terrains')}}</a><a href="/rotation-board" class="{% if request.path=='/rotation-board' %}is-active{% endif %}">{% if lang=='ar' %}جدول التوزيع{% elif lang=='fr' %}Tableau de répartition{% else %}Rotation board{% endif %}</a><a href="/assign" class="{% if request.path=='/assign' %}is-active{% endif %}">{{t('assign')}}</a><a href="/history" class="{% if request.path=='/history' %}is-active{% endif %}">{{t('history')}}</a><a href="/templates" class="{% if request.path=='/templates' %}is-active{% endif %}">{{L('Formulaires vierges','Blank forms','استمارات فارغة')}}</a><a href="/leave" class="{% if request.path=='/leave' %}is-active{% endif %}">{% if lang=='ar' %}الإجازات{% elif lang=='fr' %}Congés{% else %}Annual leave{% endif %}</a><a href="/external-rotations" class="{% if request.path.startswith('/external-rotations') %}is-active{% endif %}">{% if lang=='ar' %}التربصات الخارجية{% elif lang=='fr' %}Stages externes{% else %}External rotations{% endif %}</a></div></details>
<details class="nav-group" {% if request.path in ['/decisions','/summary'] %}open{% endif %}><summary><span class="nav-icon">▤</span><span>{% if lang=='ar' %}المتابعة والتقارير{% elif lang=='fr' %}Suivi et rapports{% else %}Monitoring & reports{% endif %}</span></summary><div class="nav-menu"><a href="/decisions" class="{% if request.path=='/decisions' %}is-active{% endif %}">{{t('decisions')}}</a><a href="/summary" class="{% if request.path=='/summary' %}is-active{% endif %}">{{t('summary')}}</a></div></details>
<details class="nav-group" {% if request.path in ['/settings','/users','/sync-devices','/sync','/audit','/backup','/export','/restore'] %}open{% endif %}><summary><span class="nav-icon">⚙</span><span>{{t('settings')}}</span></summary><div class="nav-menu"><a href="/settings" class="{% if request.path=='/settings' %}is-active{% endif %}">{{t('settings_page')}}</a><a href="/users" class="{% if request.path=='/users' %}is-active{% endif %}">{{t('users')}}</a><a href="/sync-devices" class="{% if request.path=='/sync-devices' %}is-active{% endif %}">{{L('Appareils synchronisés','Sync devices','أجهزة المزامنة')}}</a><a href="/sync" class="{% if request.path=='/sync' %}is-active{% endif %}">{{L('Synchroniser cet appareil','Sync this device','مزامنة هذا الجهاز')}}</a><a href="/audit" class="{% if request.path=='/audit' %}is-active{% endif %}">{{t('audit')}}</a><a href="/backup" class="{% if request.path=='/backup' %}is-active{% endif %}">{{t('backup')}}</a><a href="/export" class="{% if request.path=='/export' %}is-active{% endif %}">{{t('export')}}</a><a href="/restore" class="{% if request.path=='/restore' %}is-active{% endif %}">{{t('restore')}}</a><a href="/trash" class="{% if request.path=='/trash' %}is-active{% endif %}">{{L('Corbeille (6 mois)','Trash (6 months)','سلة المهملات (6 أشهر)')}}</a><a href="/logout">{{t('logout')}}</a></div></details>
<div class="rt-language" aria-label="{{L('Langue','Language','اللغة')}}"><a href="/lang/ar" {% if lang=='ar' %}aria-current="page"{% endif %} title="العربية">🇩🇿</a><a href="/lang/fr" {% if lang=='fr' %}aria-current="page"{% endif %} title="Français">🇫🇷</a><a href="/lang/en" {% if lang=='en' %}aria-current="page"{% endif %} title="English">🇬🇧</a></div></nav>{% endif %}<main>{% for m in get_flashed_messages() %}<p class="msg">{{m}}</p>{% endfor %}{{body|safe}}</main></div><footer class="rt-footer">© <strong>{{L('Tous droits réservés · ResiManage · Dr Dris','All rights reserved · ResiManage · Dr Dris','جميع الحقوق محفوظة · ResiManage · Dr Dris')}}</strong></footer><script>
(function(){var allowed=["futuristic","emerald","holographic","titanium","aurora","ruby"];function apply(theme){if(!allowed.includes(theme))theme="futuristic";document.documentElement.setAttribute("data-interior-theme",theme);document.querySelectorAll("[data-theme-choice]").forEach(function(b){b.setAttribute("aria-pressed",String(b.getAttribute("data-theme-choice")===theme));});}var saved="futuristic";try{saved=localStorage.getItem("resimanage-interior-theme")||saved;}catch(e){}apply(saved);window.setInteriorTheme=function(theme){if(!allowed.includes(theme))return;apply(theme);try{localStorage.setItem("resimanage-interior-theme",theme);}catch(e){}};document.addEventListener("DOMContentLoaded",function(){apply(saved);});})();
</script></html>"""
 c=db()
 try:
  selected_year=active_postgrad_year(c)
  current_year=setting(c,"year")
  year_options=available_postgrad_years(c)
 finally:
  c.close()
 context={"t":tr,"L":L,"lang":lang,"period_label":period_label,"year":selected_year,"current_year":current_year,"year_options":year_options}
 context.update(kw)
 context.setdefault("role",session.get("role",""))
 context.setdefault("csrf_token",csrf_token)
 inner=render_template_string(body,**context)
 return render_template_string(shell,body=inner,lang=lang,t=tr,L=L,year=selected_year,current_year=current_year,year_options=year_options,user=bool(session.get("user")),role=session.get("role",""))
def guard():
 username=session.get("user")
 if not username: return redirect("/login")
 c=db()
 account=c.execute("SELECT role,active FROM users WHERE username=?",(username,)).fetchone()
 c.close()
 if not account or not account["active"]:
  session.clear()
  return redirect("/login")
 # Reflect the current role from the database instead of trusting stale session data.
 session["role"]=account["role"]
def super_guard():
 denied=guard()
 if denied: return denied
 if session.get("role")!="super_admin": return ("Forbidden — Super Admin only",403)

@app.before_request
def purge_expired_trash_automatically():
 # Expired trash is purged on normal app traffic, not only when an admin opens /trash.
 try:
  c=db()
  if not c.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='trash_items'").fetchone():
   c.close(); return
  for item in c.execute("SELECT id,deleted_at,snapshot_json FROM trash_items").fetchall():
   try:
    deleted=datetime.fromisoformat(item["deleted_at"])
    month=deleted.month+6; year=deleted.year+(month-1)//12; month=(month-1)%12+1
    day=min(deleted.day,__import__("calendar").monthrange(year,month)[1])
    if datetime.now()<deleted.replace(year=year,month=month,day=day): continue
    snapshot=json.loads(item["snapshot_json"] or "{}")
    photo_path=(snapshot.get("resident") or {}).get("photo_path","")
    if photo_path:
     photo_name=os.path.basename(photo_path)
     photo_file=os.path.join(os.path.dirname(os.path.abspath(DB)),"resident_photos",photo_name)
     if photo_name.startswith("resident_") and os.path.isfile(photo_file):
      try: os.remove(photo_file)
      except OSError: app.logger.warning("Could not purge expired resident photo: %s",photo_file)
    c.execute("DELETE FROM trash_items WHERE id=?",(item["id"],))
    audit_event(c,"trash.expired",f"trash_id={item['id']}; purged_after_six_calendar_months")
   except (ValueError,TypeError,KeyError): continue
  c.commit(); c.close()
 except sqlite3.Error:
  try: c.close()
  except Exception: pass
@app.get("/lang/<lang>")
def lang(lang):
 if lang in LANGS: session["lang"]=lang
 return redirect(request.referrer or "/login")
@app.route("/login",methods=["GET","POST"])
def login():
 c=db(); p=setting(c,"password")
 if request.method=="POST":
  username=request.form.get("username","admin").strip()
  password=request.form.get("password","")
  account=c.execute("SELECT * FROM users WHERE (username=? OR (lower(email)=? AND email_verified=1)) AND active=1",(username,username.lower())).fetchone()
  if not p:
   confirm=request.form.get("confirm_password","")
   if username!="admin" or len(password)<12 or password!=confirm:
    c.close(); flash(L("La première configuration exige l’identifiant admin et un mot de passe confirmé d’au moins 12 caractères.","First setup requires username admin and a matching password of at least 12 characters.","يتطلب الإعداد الأول اسم المستخدم admin وكلمة مرور مؤكدة من 12 حرفًا على الأقل."))
    return redirect("/login")
   p=hash_password(password)
   c.execute("INSERT INTO settings(key,value) VALUES('password',?)",(p,))
   c.execute("INSERT OR IGNORE INTO users(username,password_hash,role) VALUES('admin',?,'super_admin')",(p,))
   c.commit(); c.close(); lang=session.get("lang","fr"); session.clear(); session["lang"]=lang; session["user"]="admin"; session["role"]="super_admin"; return redirect("/")
  if account and verify_password(password,account["password_hash"]):
   c.close(); lang=session.get("lang","fr"); session.clear(); session["lang"]=lang; session["user"]=account["username"]; session["role"]=account["role"]; return redirect("/")
  if username=="admin" and not account and verify_password(password,p):
   secure_hash=hash_password(password)
   c.execute("UPDATE settings SET value=? WHERE key='password'",(secure_hash,))
   c.execute("INSERT OR IGNORE INTO users(username,password_hash,role) VALUES('admin',?,'super_admin')",(secure_hash,))
   c.commit(); c.close(); lang=session.get("lang","fr"); session.clear(); session["lang"]=lang; session["user"]="admin"; session["role"]="super_admin"; return redirect("/")
  c.close(); flash(tr("badpass"))
 else:
  c.close()
 setup_required=not p
 return page("""<section class="login-hero"><div class="login-language"><span>🌐</span> <a href="/lang/ar">🇩🇿 العربية</a><a href="/lang/fr">🇫🇷 Français</a><a href="/lang/en">🇬🇧 English</a></div><div class="login-heartbeat" role="img" aria-label="{{L('Cœur battant avec tracé de pouls','Beating heart with pulse trace','قلب نابض مع تخطيط النبض')}}"><svg viewBox="0 0 180 86" xmlns="http://www.w3.org/2000/svg"><path class="login-heart-shape" d="M90 73S18 35 32 17C43 2 66 8 90 32C114 8 137 2 148 17C162 35 90 73 90 73Z"/><path class="login-ecg-glow" d="M8 47H48L61 35L74 59L91 18L108 55L120 43H172"/><path class="login-ecg-line" d="M8 47H48L61 35L74 59L91 18L108 55L120 43H172"/></svg></div><div class="login-brand"><img class="login-emblem" src="/static/resimanage-3d-icon.jpg" alt="{{L('ResiManage — médecin et stéthoscope','ResiManage — doctor and stethoscope','ResiManage — طبيب وسماعة طبية')}}"><p class="login-kicker">CHU Nedir Mohamed · Tizi Ouzou</p><h1>Resi<span>Manage</span></h1><p class="login-subtitle">{{L('Gestion des rotations des résidents','Resident rotation management','تسيير تربصات الأطباء المقيمين')}}</p><p class="login-motto">{{L('Aujourd’hui, un pas de plus vers votre avenir','One step closer to your future, every day','اليوم خطوة إضافية نحو مستقبلك')}}</p></div><div class="login-panel"><h2>{{t('login')}}</h2><p class="login-hint">{{L('Accédez à votre espace sécurisé','Access your secure workspace','ادخل إلى مساحتك الآمنة')}}</p><form method="post"><label>{{L('Identifiant ou e-mail','Username or email','اسم المستخدم أو البريد الإلكتروني')}}</label><input name="username" value="admin" autocomplete="username" required><label>{{t('password')}}</label><div class="password-wrap"><input id="login-password" name="password" type="password" autocomplete="current-password" minlength="{{12 if setup_required else 1}}" required><button class="password-toggle" type="button" onclick="const p=document.getElementById('login-password');p.type=p.type==='password'?'text':'password';this.setAttribute('aria-label',p.type==='password'?'{{L('Afficher le mot de passe','Show password','إظهار كلمة المرور')}}':'{{L('Masquer le mot de passe','Hide password','إخفاء كلمة المرور')}}')">◉</button></div>{% if setup_required %}<label>{{L('Confirmer le mot de passe','Confirm password','تأكيد كلمة المرور')}}</label><div class="password-wrap"><input id="first-confirm-password" name="confirm_password" type="password" autocomplete="new-password" minlength="12" required><button class="password-toggle" type="button" aria-label="{{L('Afficher le mot de passe','Show password','إظهار كلمة المرور')}}" onclick="const p=document.getElementById('first-confirm-password');p.type=p.type==='password'?'text':'password';this.setAttribute('aria-label',p.type==='password'?'{{L('Afficher le mot de passe','Show password','إظهار كلمة المرور')}}':'{{L('Masquer le mot de passe','Hide password','إخفاء كلمة المرور')}}')">◉</button></div><p>{{L('Créez le mot de passe du premier administrateur (12 caractères minimum).','Create the first administrator password (12+ characters).','أنشئ كلمة مرور المسؤول الأول (12 حرفًا على الأقل).')}}</p>{% endif %}<button class="login-submit">{{L('Créer l’administrateur','Create administrator','إنشاء المسؤول') if setup_required else t('signin')}} <span>➜</span></button></form><p class="login-security">♧ {{L('Connexion sécurisée','Secure sign-in','دخول آمن')}}</p><p class="login-security"><a href="/register" style="color:#8df2df;text-decoration:underline">{{L('Créer un compte','Create account','إنشاء حساب')}}</a></p></div><div class="login-feature-row"><span>🩺 {{L('Médecine','Medicine','الطب')}}</span><span>🩺 {{L('Réanimation','Critical care','الإنعاش')}}</span><span>✨ {{L('Anesthésie','Anesthesia','التخدير')}}</span><span>📋 {{L('Formation','Training','التكوين')}}</span></div><footer class="login-rights">{{L('Tous droits réservés · ResiManage · Dr Dris','All rights reserved · ResiManage · Dr Dris','جميع الحقوق محفوظة · ResiManage · Dr Dris')}}</footer></section><style>.login-hero{position:relative;isolation:isolate;overflow:hidden;max-width:980px;margin:12px auto;padding:clamp(16px,4vw,38px);border-radius:28px;color:#f8fbff;background:radial-gradient(ellipse at 12% 15%,#00dff755,transparent 32%),radial-gradient(ellipse at 88% 24%,#3278ff55,transparent 30%),radial-gradient(ellipse at 80% 85%,#19dfbd55,transparent 30%),linear-gradient(135deg,#020b22,#073d85 48%,#041426);box-shadow:0 24px 60px #152d4c35}.login-hero:before,.login-hero:after{content:'';position:absolute;z-index:-1;border-radius:50%;filter:blur(2px);animation:orbit 12s ease-in-out infinite alternate}.login-hero:before{width:230px;height:230px;right:-55px;top:-75px;background:linear-gradient(135deg,#47e4d8aa,#4c8bff44)}.login-hero:after{width:190px;height:190px;left:-65px;bottom:90px;background:linear-gradient(135deg,#00eaffaa,#087dff55);animation-delay:-4s}.login-language{display:flex;justify-content:center;align-items:center;gap:clamp(8px,2vw,20px);flex-wrap:wrap}.login-language a{color:#fff;text-decoration:none;border:1px solid #ffffff45;background:#ffffff12;border-radius:22px;padding:7px 12px}.login-heartbeat{position:relative;display:flex;justify-content:center;align-items:center;width:min(180px,55vw);height:78px;margin:0 auto -4px;transform:translateY(-8px);filter:drop-shadow(0 0 12px #ff436b66);animation:heartbeat 1.65s ease-in-out infinite;pointer-events:none}.login-heartbeat svg{display:block;width:100%;height:100%;overflow:visible;background:transparent;border:0;box-shadow:none}.login-heart-shape{fill:#ff416b;fill-opacity:.92;stroke:#ff9caf;stroke-width:1.5}.login-ecg-glow{fill:none;stroke:#37f4ed;stroke-width:7;stroke-linejoin:round;stroke-linecap:round;opacity:.22}.login-ecg-line{fill:none;stroke:#9ffff7;stroke-width:2.8;stroke-linejoin:round;stroke-linecap:round;stroke-dasharray:260;stroke-dashoffset:260;animation:ecg-trace 1.65s linear infinite;filter:drop-shadow(0 0 4px #38fff0)}.login-brand{text-align:center;padding:12px 0 14px}.login-emblem{display:block;width:clamp(76px,16vw,112px);height:clamp(76px,16vw,112px);object-fit:cover;border-radius:24px;margin:0 auto 12px;border:1px solid #ffffff55;box-shadow:0 0 28px #53e2c477;animation:none}.login-brand h1{font-size:clamp(38px,8vw,68px);line-height:1;margin:4px 0;font-weight:800;letter-spacing:-2px}.login-brand h1 span{background:linear-gradient(90deg,#19d9e8,#58a7ff,#19dfbd);-webkit-background-clip:text;background-clip:text;color:transparent}.login-kicker,.login-subtitle,.login-motto{margin:7px 0}.login-kicker{color:#b8d8f0}.login-subtitle{font-size:clamp(15px,3vw,20px)}.login-motto{color:#70e8f2;font-style:italic}.login-panel{max-width:480px;margin:16px auto 10px;padding:clamp(18px,4vw,30px);border:1px solid #ffffff55;border-radius:22px;background:#ffffff13;backdrop-filter:blur(14px);box-shadow:0 15px 35px #050d2040}.login-panel h2{text-align:center;margin:0;font-size:26px}.login-hint,.login-security{text-align:center;color:#d0e3f4}.login-panel label{display:block;margin-top:12px}.login-panel input{width:100%;min-width:0;background:#ffffff17;color:white;border:1px solid #ffffff48;border-radius:14px;padding:13px}.login-panel input::placeholder{color:#d4e2f1}.password-wrap{display:flex;align-items:center;gap:8px}.password-wrap input{flex:1}.password-toggle{border:1px solid #ffffff55;background:#ffffff18;color:white;padding:10px 14px;border-radius:12px;cursor:pointer}.login-submit{width:100%;margin-top:18px;border:0;border-radius:15px;padding:14px;background:linear-gradient(100deg,#5ce7b0,#22c9e5 55%,#6697ff);color:#08223d;font-weight:800;font-size:17px;box-shadow:0 7px 22px #31d5cf44;cursor:pointer}.login-submit span{margin-inline-start:10px}.login-feature-row{display:flex;justify-content:center;gap:10px;flex-wrap:wrap;margin:22px 0 12px}.login-feature-row span{border:1px solid #ffffff35;background:#ffffff10;border-radius:12px;padding:9px 13px}.login-rights{text-align:center;color:#d6e5f3;border-top:1px solid #ffffff24;padding-top:13px;font-size:12px}.login-rights span{color:#45c7e8;padding:0 5px}@keyframes orbit{to{transform:translate3d(-18px,22px,0) scale(1.08)}}@keyframes pulse{50%{transform:scale(1.08);filter:drop-shadow(0 0 12px #5ce7b0)}}@keyframes heartbeat{0%,100%{transform:scale(1)}14%{transform:scale(1.22)}28%{transform:scale(1)}42%{transform:scale(1.14)}56%{transform:scale(1)}}@keyframes ecg-trace{0%{stroke-dashoffset:260}85%,100%{stroke-dashoffset:0}}@media(prefers-reduced-motion:reduce){.login-hero:before,.login-hero:after,.login-emblem,.login-heartbeat,.login-ecg-line{animation:none}.login-ecg-line{stroke-dashoffset:0}}@media(max-width:520px){.login-hero{margin:4px;padding:14px;border-radius:18px}.login-feature-row span{font-size:12px;padding:7px}.login-rights span{padding:0 2px}}</style>""",setup_required=setup_required)

def _smtp_setting(name, default=""):
 # Prefer current product-specific names, but accept older deployments during migration.
 return os.environ.get("RESIMANAGE_SMTP_"+name, os.environ.get("RESITRACK_SMTP_"+name, default)).strip()

def send_verification_email(recipient, code, lang="fr"):
 host=_smtp_setting("HOST")
 port_text=_smtp_setting("PORT","587")
 sender=_smtp_setting("FROM")
 username=_smtp_setting("USER")
 password=os.environ.get("RESIMANAGE_SMTP_PASSWORD",os.environ.get("RESITRACK_SMTP_PASSWORD",""))
 try:
  port=int(port_text)
 except (TypeError,ValueError):
  raise RuntimeError("SMTP port must be a valid integer.") from None
 if not host or not sender or not username or not password:
  raise RuntimeError("Email verification is not configured. Set RESIMANAGE_SMTP_HOST, RESIMANAGE_SMTP_PORT, RESIMANAGE_SMTP_FROM, RESIMANAGE_SMTP_USER and RESIMANAGE_SMTP_PASSWORD in the environment; do not put credentials in the app or APK.")
 from email.message import EmailMessage
 message=EmailMessage()
 language=lang if lang in LANGS else "fr"
 subjects={"fr":"ResiManage — code de vérification de l’e-mail","en":"ResiManage — email verification code","ar":"ResiManage — رمز التحقق من البريد الإلكتروني"}
 bodies={
  "fr":"Votre code de vérification ResiManage est : "+code+"\n\nIl expire dans 10 minutes. Si vous n’avez pas demandé ce code, ignorez ce message.",
  "en":"Your ResiManage verification code is: "+code+"\n\nIt expires in 10 minutes. If you did not request this, ignore this message.",
  "ar":"رمز التحقق الخاص بك في ResiManage هو: "+code+"\n\nتنتهي صلاحيته خلال 10 دقائق. إذا لم تطلب هذا الرمز، فتجاهل هذه الرسالة."
 }
 message["Subject"]=subjects[language]
 message["From"]=sender
 message["To"]=recipient
 message.set_content(bodies[language])
 with smtplib.SMTP(host,port,timeout=15) as server:
  server.starttls()
  server.login(username,password)
  server.send_message(message)

@app.route("/register",methods=["GET","POST"])
def register():
 if session.get("user"): return redirect("/")
 c=db()
 if request.method=="POST":
  username=request.form.get("username","").strip()
  email=request.form.get("email","").strip().lower()
  password=request.form.get("password","")
  confirm=request.form.get("confirm_password","")
  if not re.fullmatch(r"[A-Za-z0-9_.-]{3,40}",username):
   flash(L("Le nom d’utilisateur doit contenir de 3 à 40 caractères : lettres, chiffres, points, tirets bas ou tirets.","Username must be 3–40 characters: letters, numbers, dots, underscores or hyphens.","يجب أن يتكون اسم المستخدم من 3 إلى 40 حرفًا: أحرف أو أرقام أو نقاط أو شرطات سفلية أو شرطات."))
  elif not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+",email) or len(email)>254:
   flash(L("Saisissez une adresse e-mail valide.","Enter a valid email address.","أدخل عنوان بريد إلكتروني صالحًا."))
  elif len(password)<12 or password!=confirm:
   flash(L("Le mot de passe doit contenir au moins 12 caractères et correspondre à la confirmation.","Password must be at least 12 characters and match confirmation.","يجب أن تتكون كلمة المرور من 12 حرفًا على الأقل وأن تطابق التأكيد."))
  elif c.execute("SELECT 1 FROM users WHERE username=?",(username,)).fetchone() or c.execute("SELECT 1 FROM users WHERE lower(email)=?",(email,)).fetchone() or c.execute("SELECT 1 FROM pending_registrations WHERE username=? OR email=?",(username,email)).fetchone():
   flash(L("Ce nom d’utilisateur ou cette adresse e-mail est déjà enregistré.","Username or email is already registered.","اسم المستخدم أو البريد الإلكتروني مسجل مسبقًا."))
  else:
   code=("%06d" % secrets.randbelow(1000000))
   c.execute("DELETE FROM pending_registrations WHERE expires_at<?",(int(time.time()),))
   try:
    c.execute("INSERT INTO pending_registrations(username,email,password_hash,code_hash,expires_at) VALUES(?,?,?,?,?)",(username,email,hash_password(password),hashlib.sha256(code.encode()).hexdigest(),int(time.time())+600))
    c.commit()
    try:
     send_verification_email(email,code,session.get("lang","fr"))
    except Exception as exc:
     c.execute("DELETE FROM pending_registrations WHERE email=?",(email,)); c.commit()
     app.logger.error("ResiManage verification email failed: %s",exc)
     flash(L("L’e-mail n’a pas pu être envoyé ; la configuration SMTP est nécessaire.","Email could not be sent; SMTP setup is required.","تعذر إرسال البريد الإلكتروني؛ يلزم إعداد SMTP."))
    else:
     session["pending_email"]=email
     flash(L("Code de vérification envoyé ; validité de 10 minutes.","Verification code sent; valid for 10 minutes.","تم إرسال رمز التحقق؛ صلاحيته 10 دقائق."))
     c.close()
     return redirect("/verify-email")
   except sqlite3.IntegrityError:
    flash(L("Ce nom d’utilisateur ou cette adresse e-mail est déjà enregistré.","Username or email is already registered.","اسم المستخدم أو البريد الإلكتروني مسجل مسبقًا."))
 c.close()
 return page("""<section class="login-hero"><div class="login-brand"><img class="login-emblem" src="/static/resimanage-3d-icon.jpg" alt="ResiManage — médecin et stéthoscope"><h1>Resi<span>Manage</span></h1><h2>{{L('Créer un compte','Create account','إنشاء حساب')}}</h2><p class="login-subtitle">{{L('Une adresse e-mail valide est nécessaire pour confirmer le compte.','A valid email is required to confirm the account.','يلزم بريد إلكتروني صالح لتأكيد الحساب.')}}</p></div><div class="login-panel"><form method="post"><label>{{L('Nom d’utilisateur','Username','اسم المستخدم')}}</label><input name="username" minlength="3" maxlength="40" pattern="[A-Za-z0-9_.-]{3,40}" autocomplete="username" required><label>{{L('E-mail','Email','البريد الإلكتروني')}}</label><input name="email" type="email" maxlength="254" autocomplete="email" required><label>{{L('Mot de passe (12 caractères minimum)','Password (12+ characters)','كلمة المرور (12 حرفًا على الأقل)')}}</label><div class="password-wrap"><input id="register-password" name="password" type="password" minlength="12" autocomplete="new-password" required><button class="password-toggle" type="button" aria-label="Show password" onclick="const p=document.getElementById('register-password');p.type=p.type==='password'?'text':'password';this.setAttribute('aria-label',p.type==='password'?'Show password':'Hide password')">◉</button></div><label>{{L('Confirmer le mot de passe','Confirm password','تأكيد كلمة المرور')}}</label><div class="password-wrap"><input id="register-confirm-password" name="confirm_password" type="password" minlength="12" autocomplete="new-password" required><button class="password-toggle" type="button" aria-label="Show password" onclick="const p=document.getElementById('register-confirm-password');p.type=p.type==='password'?'text':'password';this.setAttribute('aria-label',p.type==='password'?'Show password':'Hide password')">◉</button></div><button class="login-submit">{{L('Créer le compte','Create account','إنشاء الحساب')}}</button></form><p class="login-security"><a href="/login">{{L('Retour à la connexion','Back to sign in','العودة إلى تسجيل الدخول')}}</a></p></div></section>""")

@app.route("/verify-email",methods=["GET","POST"])
def verify_email():
 email=session.get("pending_email","").strip().lower()
 if not email: return redirect("/register")
 c=db()
 if request.method=="POST":
  code=request.form.get("code","").strip()
  pending=c.execute("SELECT * FROM pending_registrations WHERE email=?",(email,)).fetchone()
  if not pending:
   session.pop("pending_email",None); flash(L("L’inscription a expiré ; veuillez recommencer.","Registration expired; please register again.","انتهت صلاحية التسجيل؛ يرجى إعادة التسجيل.")); c.close(); return redirect("/register")
  if pending["expires_at"]<int(time.time()):
   c.execute("DELETE FROM pending_registrations WHERE email=?",(email,)); c.commit()
   session.pop("pending_email",None); flash(L("Le code a expiré ; veuillez recommencer l’inscription.","Code expired; please register again.","انتهت صلاحية الرمز؛ يرجى إعادة التسجيل.")); c.close(); return redirect("/register")
  if pending["attempts"]>=5:
   c.execute("DELETE FROM pending_registrations WHERE email=?",(email,)); c.commit()
   session.pop("pending_email",None); flash(L("Trop de tentatives ; veuillez recommencer l’inscription.","Too many attempts; please register again.","محاولات كثيرة؛ يرجى إعادة التسجيل.")); c.close(); return redirect("/register")
  if not re.fullmatch(r"\d{6}",code) or not hmac.compare_digest(hashlib.sha256(code.encode()).hexdigest(),pending["code_hash"]):
   c.execute("UPDATE pending_registrations SET attempts=attempts+1 WHERE email=?",(email,)); c.commit()
   flash(L("Code de vérification incorrect.","Incorrect verification code.","رمز التحقق غير صحيح."))
  else:
   try:
    c.execute("INSERT INTO users(username,password_hash,role,email,email_verified) VALUES(?,?,'user',?,1)",(pending["username"],pending["password_hash"],pending["email"]))
    c.execute("DELETE FROM pending_registrations WHERE email=?",(email,))
    audit_event(c,"user.self_registered",f"username={pending['username']}; email_verified=1")
    c.commit()
    session.pop("pending_email",None)
    flash(L("E-mail vérifié et compte créé. Vous pouvez vous connecter.","Email verified and account created. Please sign in.","تم تأكيد البريد وإنشاء الحساب. يرجى تسجيل الدخول."))
    c.close()
    return redirect("/login")
   except sqlite3.IntegrityError:
    c.execute("DELETE FROM pending_registrations WHERE email=?",(email,)); c.commit()
    flash(L("Ce compte existe déjà.","Account already exists.","الحساب موجود مسبقًا."))
 c.close()
 return page("""<section class="login-hero"><div class="login-brand"><img class="login-emblem" src="/static/resimanage-3d-icon.jpg" alt="ResiManage — médecin et stéthoscope"><h1>Resi<span>Manage</span></h1><h2>{{L('Vérification de l’e-mail','Email verification','تأكيد البريد الإلكتروني')}}</h2><p class="login-subtitle">{{L('Code envoyé à','Code sent to','أُرسل الرمز إلى')}} {{email}}</p></div><div class="login-panel"><form method="post"><label>{{L('Code à 6 chiffres','6-digit code','رمز التأكيد المكوّن من 6 أرقام')}}</label><input name="code" inputmode="numeric" pattern="[0-9]{6}" minlength="6" maxlength="6" autocomplete="one-time-code" required><button class="login-submit">{{L('Vérifier l’e-mail','Verify email','تأكيد البريد الإلكتروني')}}</button></form><p class="login-security">{{L('Le code expire après 10 minutes.','Code expires after 10 minutes.','تنتهي صلاحية الرمز بعد 10 دقائق.')}}</p></div></section>""",email=email)

@app.get("/logout")
def logout(): session.clear(); return redirect("/login")
@app.post("/select-year")
def select_postgrad_year():
 if guard(): return redirect("/login")
 year=request.form.get("year","").strip()
 if not re.fullmatch(r"\d{4}-\d{4}",year) or int(year.split("-")[1])!=int(year.split("-")[0])+1:
  flash(L("Année de post-graduation invalide.","Invalid postgraduate year.","سنة ما بعد التدرج غير صالحة."))
  return redirect("/")
 session["selected_postgrad_year"]=year
 target=request.form.get("next","").strip()
 if target.startswith("/") and not target.startswith("//") and not target.startswith("/\\"):
  return redirect(target)
 referrer=request.referrer or "/"
 if referrer.startswith("/") and not referrer.startswith("//"):
  return redirect(referrer)
 return redirect("/")

@app.get("/")
def home():
 if guard(): return redirect("/login")
 c=db(); y=active_postgrad_year(c); rs=c.execute("SELECT * FROM residents WHERE status='active' ORDER BY year DESC,surname").fetchall()
 count=c.execute("SELECT COUNT(*) FROM terrains").fetchone()[0]; confirmed=c.execute("SELECT COUNT(*) FROM assignments a JOIN periods p ON p.id=a.period_id WHERE p.year_label=? AND a.status='confirmed'",(y,)).fetchone()[0]
 rows=[]
 for r in rs:
  n=c.execute("SELECT COUNT(*) FROM assignments a JOIN periods p ON p.id=a.period_id WHERE a.resident_id=? AND p.year_label=? AND a.status='confirmed'",(r["id"],y)).fetchone()[0]; rows.append((r,n))
 c.close()
 return page("""<section><h3 class="rt-dashboard-title"><span class="rt-dot"></span>{{t('overview')}}</h3><p>{{t('priority')}}</p><div class="rt-stats"><div class="rt-stat"><span>{{t('residents')}}</span><strong>{{rows|length}}</strong></div><div class="rt-stat"><span>{{t('terrains')}}</span><strong>{{count}}</strong></div><div class="rt-stat"><span>{{t('confirmed')}}</span><strong>{{confirmed}}</strong></div><form class="rt-stat rt-year-select" action="/select-year" method="post" aria-label="{{L('Choisir l’année de post-graduation','Choose the postgraduate year','اختر سنة ما بعد التدرج')}}"><span>{{t('year')}}</span><select name="year" aria-label="{{L('Année de post-graduation','Postgraduate year','سنة ما بعد التدرج')}}" onchange="this.form.submit()">{% set past_years = year_options|reject('equalto', current_year)|list %}<optgroup label="{{L('Année en cours','Current year','السنة الحالية')}}"><option value="{{current_year}}" {% if year==current_year %}selected{% endif %}>{{current_year}} — {{L('en cours','current','الحالية')}}</option></optgroup>{% if past_years %}<optgroup label="{{L('Historique des années précédentes','Previous years history','سجل السنوات السابقة')}}">{% for yy in past_years %}<option value="{{yy}}" {% if yy==year %}selected{% endif %}>{{yy}}</option>{% endfor %}</optgroup>{% endif %}</select><small>{{L('Choisir une année pour consulter ses données','Choose a year to view its data','اختر سنة لعرض بياناتها')}}</small></form></div></section><div class="rt-quick-actions"><a class="rt-action-card" href="/residents"><span class="rt-action-icon">♙</span><span><strong>{{t('manage_residents')}}</strong><small>{{L('Profils · années · coordonnées','Profiles · years · contact details','الملفات · السنوات · معلومات الاتصال')}}</small></span></a><a class="rt-action-card" href="/assign"><span class="rt-action-icon">⇄</span><span><strong>{{t('assign')}}</strong><small>{{L('Répartition des stages — affecter un résident','Rotation allocation — assign a resident','توزيع التربصات — تعيين مقيم')}}</small></span></a><a class="rt-action-card" href="/rotation-board"><span class="rt-action-icon">▦</span><span><strong>{% if lang=='ar' %}جدول التوزيع{% elif lang=='fr' %}Tableau général de répartition{% else %}Rotation board{% endif %}</strong><small>{{L('Services × périodes · capacités · résidents · PDF / impression','Departments × periods · capacity · residents · PDF / print','المصالح × الفترات · السعة · المقيمون · PDF / طباعة')}}</small></span></a><a class="rt-action-card" href="/terrains"><span class="rt-action-icon">▥</span><span><strong>{{t('terrains')}}</strong><small>{{L('Services obligatoires et optionnels','Mandatory and optional departments','المصالح الإلزامية والاختيارية')}}</small></span></a><a class="rt-action-card" href="/import"><span class="rt-action-icon">⇧</span><span><strong>{{t('import_residents')}}</strong><small>{{L('Importation avec aperçu et validation','Import with preview and validation','استيراد مع معاينة وتحقق')}}</small></span></a><a class="rt-action-card" href="/leave"><span class="rt-action-icon">◷</span><span><strong>{{L('Congés annuels','Annual leave','الإجازات السنوية')}}</strong><small>{{L('30 jours par an · report des jours non utilisés','30 days yearly · unused days carry forward','30 يومًا سنويًا · ترحيل الأيام غير المستعملة')}}</small></span></a><a class="rt-action-card" href="/external-rotations"><span class="rt-action-icon">✈</span><span><strong>{{L('Stages à l’étranger','External rotations','التربصات بالخارج')}}</strong><small>{{L('Encadrant · hôpital · pays · dates','Supervisor · hospital · country · dates','المشرف · المستشفى · البلد · التواريخ')}}</small></span></a></div><section><h3 class="rt-dashboard-title"><span class="rt-dot"></span>{{t('four')}}</h3><div class="wrap"><table><tr><th>{{t('resident')}}</th><th>{{t('study_year')}}</th><th>{{t('count')}}</th></tr>{% for r,n in rows %}<tr><td><a class="rt-resident-link" href="/profile/{{r.id}}">{{r.surname}} {{r.firstname}}</a></td><td><span class="rt-year-chip">{{r.year}}</span></td><td><span class="rt-progress"><span style="width:{{[n*25,100]|min}}%"></span></span> {{n}} / 4</td></tr>{% else %}<tr><td colspan="3">{{t('none')}}</td></tr>{% endfor %}</table></div></section>""",rows=rows,count=count,confirmed=confirmed,year=y)
@app.route("/residents",methods=["GET","POST"])
def residents():
 if guard(): return redirect("/login")
 c=db()
 if request.method=="POST":
  action=request.form.get("action","create")
  if action in ("edit","archive","restore_archive","delete") and session.get("role")!="super_admin": c.close(); return ("Forbidden — Super Admin only",403)
  if action=="archive":
   try:
    rid=int(request.form["rid"]); archive_type=request.form.get("archive_type","temporary"); reason=request.form.get("archive_reason","").strip(); return_date=request.form.get("archive_return_date","").strip()
    if archive_type not in ("temporary","permanent") or not reason: raise ValueError(L("Choisissez un type d’archivage et indiquez le motif.","Choose an archive type and provide a reason.","اختر نوع الأرشفة وأدخل السبب."))
    if archive_type=="temporary":
     if not return_date: raise ValueError(L("La date de retour est obligatoire pour un archivage temporaire.","A return date is required for temporary archiving.","تاريخ العودة إجباري للأرشفة المؤقتة."))
     date.fromisoformat(return_date)
    else: return_date=""
    target=c.execute("SELECT surname,firstname,status FROM residents WHERE id=?",(rid,)).fetchone()
    if not target or target["status"] in ("graduated","archived"): raise ValueError(L("Ce dossier ne peut pas être archivé depuis cet écran.","This record cannot be archived from this screen.","لا يمكن أرشفة هذا الملف من هذه الشاشة."))
    c.execute("UPDATE residents SET status='archived',archive_type=?,archive_reason=?,archive_return_date=?,archived_at=? WHERE id=?",(archive_type,reason,return_date,datetime.now().isoformat(timespec="seconds"),rid))
    audit_event(c,"resident.archived",f"resident_id={rid}; type={archive_type}; reason={reason}; return_date={return_date}; previous_status={target['status']}"); c.commit()
    flash(L("Dossier archivé. Le résident ne participe plus aux affectations.","Record archived. The resident is excluded from assignments.","تمت أرشفة الملف. لن يشارك المقيم في التعيينات."))
   except (ValueError,TypeError,KeyError) as exc: c.rollback(); flash(str(exc) or L("Données d’archivage invalides.","Invalid archive data.","بيانات الأرشفة غير صالحة."))
  elif action=="restore_archive":
   rid=int(request.form.get("rid","0")); target=c.execute("SELECT archive_type,archive_reason,status FROM residents WHERE id=?",(rid,)).fetchone()
   if not target or target["status"]!="archived" or target["archive_type"]!="temporary": flash(L("Seul un archivage temporaire peut être réactivé depuis cet écran.","Only a temporary archive can be restored from this screen.","يمكن إلغاء الأرشفة المؤقتة فقط من هذه الشاشة."))
   else:
    c.execute("UPDATE residents SET status='active',archive_type='',archive_reason='',archive_return_date='',archived_at='' WHERE id=?",(rid,)); audit_event(c,"resident.archive_restored",f"resident_id={rid}; reason={target['archive_reason']}"); c.commit(); flash(L("Dossier réactivé ; le résident revient dans les listes d’affectation.","Record restored; the resident is back in assignment lists.","تمت إعادة تفعيل الملف؛ عاد المقيم إلى قوائم التعيين."))
  elif action=="delete":
   rid=int(request.form.get("rid","0")); reason=request.form.get("delete_reason","").strip(); target=c.execute("SELECT surname,firstname,status FROM residents WHERE id=?",(rid,)).fetchone()
   if not target or not reason: flash(L("Dossier introuvable ou motif de suppression manquant.","Record not found or deletion reason missing.","الملف غير موجود أو سبب الحذف مفقود."))
   elif target["status"]=="graduated": flash(L("Un diplômé archivé par le processus DEMS ne peut pas être supprimé depuis cet écran.","A DEMS graduate cannot be deleted from this screen.","لا يمكن حذف خريج مؤرشف عبر مسار DEMS من هذه الشاشة."))
   else:
    tables=("assignments","history","exam_records","leave_years","leave_entries","external_rotations","promotion_votes","dems_votes","dems_archive")
    snapshot={"resident":dict(c.execute("SELECT * FROM residents WHERE id=?",(rid,)).fetchone()),"related":{}}
    for table in tables:
     columns={row[1] for row in c.execute("PRAGMA table_info("+table+")")}
     if "resident_id" in columns:
      snapshot["related"][table]=[dict(x) for x in c.execute("SELECT * FROM "+table+" WHERE resident_id=?",(rid,)).fetchall()]
    deleted_at=datetime.now().replace(microsecond=0)
    c.execute("INSERT INTO trash_items(entity_type,entity_id,label,reason,snapshot_json,deleted_at,deleted_by) VALUES(?,?,?,?,?,?,?)",("resident",str(rid),target["surname"]+" "+target["firstname"],reason,json.dumps(snapshot,ensure_ascii=False),deleted_at.isoformat(),session.get("username","")))
    for table in tables:
     columns={row[1] for row in c.execute("PRAGMA table_info("+table+")")}
     if "resident_id" in columns: c.execute("DELETE FROM "+table+" WHERE resident_id=?",(rid,))
    c.execute("DELETE FROM residents WHERE id=?",(rid,))
    audit_event(c,"resident.trashed",f"resident_id={rid}; name={target['surname']} {target['firstname']}; reason={reason}; retained_for_six_months")
    c.commit()
    flash(L("Dossier déplacé dans la corbeille pour six mois. Les données associées sont conservées.","Record moved to the trash for six months. Associated data is retained.","تم نقل الملف إلى سلة المهملات لمدة ستة أشهر مع الاحتفاظ بالبيانات المرتبطة."))
  elif action=="edit":
   try:
    rid=int(request.form["rid"])
    year=int(request.form.get("year","0"))
    repeat_year=request.form.get("repeat_year","0")
    status=request.form.get("status","")
    change_reason=request.form.get("change_reason","").strip()
   except (TypeError,ValueError):
    flash(L("Modification du résident invalide.","Invalid resident update.","تعديل بيانات المقيم غير صالح."))
   else:
    target=c.execute("SELECT status,year FROM residents WHERE id=?",(rid,)).fetchone()
    if not target:
     flash(L("Résident introuvable.","Resident not found.","لم يتم العثور على المقيم."))
    elif target["status"]=="graduated":
     flash(L("Les résidents diplômés sont archivés ; corrigez les dossiers DEMS depuis la gestion des archives.","Graduated residents are archived; correct DEMS records from archive management.","يُؤرشف المقيمون المتخرجون؛ صحّح ملفات DEMS من إدارة الأرشيف."))
    elif year not in (1,2,3,4) or repeat_year not in ("0","1") or status not in ("active","inactive"):
     flash(L("Année, redoublement ou statut invalide.","Invalid year, repeat status or account status.","السنة أو حالة الإعادة أو الحالة غير صالحة."))
    elif year!=target["year"] and not change_reason:
     flash(L("Tout changement d’année doit passer par le processus officiel d’examens et de promotion ; indiquez un motif pour toute correction administrative autorisée.","Study-year changes must use the official exam/promotion workflow; provide a reason for an authorized administrative correction.","يجب أن تمر تغييرات السنة عبر مسار الامتحانات والترقية الرسمي؛ قدّم سببًا لأي تصحيح إداري معتمد."))
    else:
     c.execute("UPDATE residents SET year=?,repeat_year=?,status=? WHERE id=?",(year,int(repeat_year),status,rid))
     audit_event(c,"resident.updated",f"resident_id={rid}; year_before={target['year']}; year_after={year}; repeat_year={repeat_year}; status_before={target['status']}; status_after={status}; reason={change_reason}; source=super_admin_correction")
     c.commit(); flash(L("Dossier du résident corrigé et opération consignée dans le journal.","Resident record corrected and action logged.","تم تصحيح ملف المقيم وتسجيل العملية في سجل التدقيق."))
  elif action=="create":
   if session.get("role")!="super_admin":
    c.close(); return ("Forbidden — Super Admin only",403)
   surname=request.form.get("surname","").strip()
   firstname=request.form.get("firstname","").strip()
   phone=request.form.get("phone","").strip()
   resident_type=request.form.get("resident_type","local")
   origin_choice=request.form.get("origin_hospital","").strip()
   origin_hospital=origin_choice
   if origin_choice.casefold() in ("autre","other"):
    origin_hospital=request.form.get("origin_hospital_other","").strip()
   external_start=request.form.get("external_start_date","").strip()
   external_end=request.form.get("external_end_date","").strip()
   repeat_year=request.form.get("repeat_year","0")
   try: year=int(request.form.get("year","0"))
   except (TypeError,ValueError): year=0
   external_dates_valid=True
   if resident_type=="external":
    try: external_dates_valid=bool(external_start and external_end and date.fromisoformat(external_end)>=date.fromisoformat(external_start))
    except ValueError: external_dates_valid=False
   if resident_type not in ("local","external"):
    flash(L("Type de résident invalide.","Invalid resident type.","نوع المقيم غير صالح."))
   elif not surname or not firstname or len(surname)>100 or len(firstname)>100 or year not in (1,2,3,4):
    flash(L("Le nom et le prénom sont obligatoires (100 caractères maximum) et l’année d’étude doit être valide.","Surname, first name and a valid study year are required (maximum 100 characters for each name).","اللقب والاسم إجباريان (100 حرف كحد أقصى) ويجب تحديد سنة دراسية صالحة."))
   elif len(phone)>40 or repeat_year not in ("0","1"):
    flash(L("Le téléphone doit contenir au maximum 40 caractères et le champ redoublement doit valoir 0 ou 1.","Phone must be at most 40 characters and repeat flag must be 0 or 1.","يجب ألا يتجاوز الهاتف 40 حرفًا، ويجب أن تكون قيمة الإعادة 0 أو 1."))
   elif resident_type=="external" and (not origin_hospital or not external_start or not external_end):
    flash(L("Pour un résident externe, l’établissement d’origine et les dates de début et de fin sont obligatoires.","For an external resident, the home institution and start/end dates are required.","بالنسبة للمقيم الخارجي، المؤسسة الأصلية وتاريخا البداية والنهاية إجبارية."))
   elif resident_type=="external" and origin_choice.casefold() not in ("ehs cardio dbk","cac dbk","autre","other"):
    flash(L("Choisissez EHS Cardio DBK, CAC DBK ou « Autre » et précisez l’établissement d’origine.","Choose EHS Cardio DBK, CAC DBK, or “Other” and specify the home institution.","اختر EHS Cardio DBK أو CAC DBK أو «أخرى» وحدد المؤسسة الأصلية."))
   elif resident_type=="external" and not external_dates_valid:
    flash(L("La date de fin du stage doit être égale ou postérieure à la date de début.","Rotation end date must be on or after the start date.","يجب أن يكون تاريخ نهاية التربص بعد تاريخ بدايته أو مساويًا له."))
   elif c.execute("SELECT 1 FROM residents WHERE lower(surname)=lower(?) AND lower(firstname)=lower(?) AND year=? AND COALESCE(resident_type,'local')=? AND COALESCE(origin_hospital,'')=?",(surname,firstname,year,resident_type,origin_hospital)).fetchone():
    flash(L("Un résident portant le même nom, la même année et le même établissement existe déjà.","A resident with the same name, year and institution already exists.","يوجد مقيم بالاسم والسنة والمؤسسة نفسها."))
   else:
    c.execute("INSERT INTO residents(surname,firstname,year,phone,repeat_year,status,resident_type,origin_hospital,external_start_date,external_end_date) VALUES(?,?,?,?,?,?,?,?,?,?)",(surname,firstname,year,phone,int(repeat_year),"active",resident_type,origin_hospital if resident_type=="external" else "",external_start if resident_type=="external" else "",external_end if resident_type=="external" else ""))
    audit_event(c,"resident.created",surname+" "+firstname+"; type="+resident_type+"; origin="+origin_hospital)
    c.commit(); flash(tr("done"))
 rows=c.execute("SELECT * FROM residents ORDER BY CASE status WHEN 'active' THEN 0 WHEN 'inactive' THEN 1 WHEN 'archived' THEN 2 ELSE 3 END,year DESC,surname").fetchall(); due_archives=[r for r in rows if r['status']=='archived' and r['archive_type']=='temporary' and r['archive_return_date'] and r['archive_return_date']<=date.today().isoformat()]; c.close()
 return page("""<section><h3>{{t('add')}} — {{t('resident')}}</h3><details class="rt-add-resident"><summary>{{L('Ajouter un résident','Add a resident','إضافة مقيم')}}</summary><form method="post"><input type="hidden" name="action" value="create"><div class="grid"><div><label>{{t('surname')}}</label><input name="surname" required></div><div><label>{{t('firstname')}}</label><input name="firstname" required></div><div><label>{{t('study_year')}}</label><select name="year">{% for y in range(1,5) %}<option>{{y}}</option>{% endfor %}</select></div><div><label>{{t('phone')}}</label><input name="phone"></div><div><label>{{L('Catégorie','Resident category','فئة المقيم')}}</label><select name="resident_type" id="resident-type"><option value="local">{{L('Résident du CHU Nedir Mohamed','Nedir Mohamed hospital resident','مقيم بمستشفى نذير محمد')}}</option><option value="external">{{L('Résident externe d’un autre hôpital','External resident from another hospital','مقيم خارجي من مستشفى آخر')}}</option></select></div></div><div id="external-resident-fields" class="grid"><div><label>{{L('Établissement d’origine','Home institution','المؤسسة الأصلية')}}</label><select name="origin_hospital" id="origin-hospital" required><option value="">{{L('Choisir un établissement…','Choose an institution…','اختر مؤسسة…')}}</option><option value="EHS Cardio DBK">{{L('EHS Cardio — établissement spécialisé de chirurgie cardiaque, DBK','EHS Cardio — specialized cardiac surgery hospital, DBK','EHS Cardio — المؤسسة المتخصصة في جراحة القلب، DBK')}}</option><option value="CAC DBK">{{L('CAC — Centre AntiCancer, DBK','CAC — Cancer Center, DBK','CAC — مركز مكافحة السرطان، DBK')}}</option><option value="Autre">{{L('Autre établissement','Other institution','مؤسسة أخرى')}}</option></select><div id="other-hospital-fields" hidden><label>{{L('Préciser l’hôpital d’origine','Specify home hospital','حدد المستشفى الأصلي')}}<input name="origin_hospital_other" id="origin-hospital-other" maxlength="160"></label></div><small>{{L('Sélectionnez EHS Cardio DBK, CAC DBK ou « Autre » puis précisez l’établissement.','Select EHS Cardio DBK, CAC DBK or “Other” and specify the institution.','اختر EHS Cardio DBK أو CAC DBK أو «أخرى» ثم حدد المؤسسة.')}}</small></div><div><label>{{L('Début du stage au CHU','Start date at our hospital','بداية التربص بمستشفانا')}}</label><input type="date" name="external_start_date" id="external-start"></div><div><label>{{L('Fin du stage au CHU','End date at our hospital','نهاية التربص بمستشفانا')}}</label><input type="date" name="external_end_date" id="external-end"></div></div><p class="muted">{{L('Les mêmes règles de rotation, de priorité, de capacité et de conflit s’appliquent aux résidents externes et aux résidents du CHU.','External residents follow the same rotation, priority, capacity and conflict rules as local residents.','تُطبق على المقيمين الخارجيين نفس قواعد التربص والأولوية والسعة والتعارض المطبقة على مقيمي المستشفى.')}}</p><label><input type="checkbox" name="repeat_year" value="1"> {{t('repeat')}}</label><button>{{t('add')}}</button></form><script>(function(){var type=document.getElementById('resident-type'),box=document.getElementById('external-resident-fields'),origin=document.getElementById('origin-hospital'),otherBox=document.getElementById('other-hospital-fields'),other=document.getElementById('origin-hospital-other');function update(){var external=type.value==='external',isOther=['autre','other'].includes((origin.value||'').trim().toLocaleLowerCase());box.hidden=!external;otherBox.hidden=!(external&&isOther);other.required=external&&isOther;['origin-hospital','external-start','external-end'].forEach(function(id){document.getElementById(id).required=external;});}type.addEventListener('change',update);origin.addEventListener('input',update);origin.addEventListener('change',update);update();})();</script></details></section>{% if due_archives %}<section><strong>{{L('Retours d’archivage à traiter','Archive returns due','حالات عودة من الأرشيف تستوجب المعالجة')}}</strong><ul>{% for r in due_archives %}<li>{{r.surname}} {{r.firstname}} — {{r.archive_return_date}} — {{r.archive_reason}}</li>{% endfor %}</ul><p>{{L('Un Super Admin peut réactiver les dossiers temporaires ci-dessous.','A Super Admin can restore temporary records below.','يمكن للمسؤول العام إعادة تفعيل الملفات المؤقتة أدناه.')}}</p></section>{% endif %}<section><div class="wrap"><table><tr><th>{{t('resident')}}</th><th>{{t('study_year')}}</th><th>{{t('save')}}</th></tr>{% for r in rows %}<tr><td><a href="/profile/{{r.id}}">{{r.surname}} {{r.firstname}}</a><br><small>{% if r.resident_type=='external' %}{{L('Résident externe —','External resident —','مقيم خارجي —')}} {{r.origin_hospital}} · {% endif %}{{L('Statut :','Status:','الحالة:')}} {{r.status}}{% if r.archive_type %} — {{L('Archivage','Archive','الأرشفة')}} {{L('temporaire' if r.archive_type=='temporary' else 'définitif','temporary' if r.archive_type=='temporary' else 'permanent','مؤقت' if r.archive_type=='temporary' else 'نهائي')}}{% endif %}</small>{% if r.archive_reason %}<p>{{L('Motif :','Reason:','السبب:')}} {{r.archive_reason}}</p>{% endif %}{% if r.archive_return_date %}<p>{{L('Retour prévu :','Expected return:','العودة المتوقعة:')}} {{r.archive_return_date}}{% if r.archive_return_date<=today %} ⚠ {{L('Date atteinte','Due','حان الموعد')}}{% endif %}</p>{% endif %}</td><td>{{r.year}}</td><td>{% if r.status=='graduated' %}<span>{{L('Diplômé — archive DEMS','Graduated — DEMS archive','متخرج — أرشيف DEMS')}}</span>{% elif r.status=='archived' %}{% if role=='super_admin' and r.archive_type=='temporary' %}<form method="post"><input type="hidden" name="action" value="restore_archive"><input type="hidden" name="rid" value="{{r.id}}"><button>{{L('Désarchiver / réactiver','Unarchive / reactivate','إلغاء الأرشفة / إعادة التفعيل')}}</button></form>{% endif %}{% else %}<form method="post"><input type="hidden" name="action" value="edit"><input type="hidden" name="rid" value="{{r.id}}"><label>{{L('Année de post-graduation','Postgraduate year','سنة ما بعد التدرج')}}<select name="year" required>{% for yy in range(1,5) %}<option value="{{yy}}" {% if r.year==yy %}selected{% endif %}>{{yy}}</option>{% endfor %}</select></label><label><input type="checkbox" name="repeat_year" value="1" {% if r.repeat_year %}checked{% endif %}> {{t('repeat')}}</label><select name="status"><option value="active" {% if r.status=='active' %}selected{% endif %}>{{t('active')}}</option><option value="inactive" {% if r.status=='inactive' %}selected{% endif %}>{{t('inactive')}}</option></select><label>{{L('Motif du changement d’année','Reason for year change','سبب تغيير السنة')}}<input name="change_reason" maxlength="500" placeholder="{{L('Passage, redoublement ou correction','Promotion, repeat or correction','ترقية أو إعادة أو تصحيح')}}"></label><button>{{t('save')}}</button></form><details><summary>{{L('Archiver ce dossier','Archive this record','أرشفة هذا الملف')}}</summary><form method="post"><input type="hidden" name="action" value="archive"><input type="hidden" name="rid" value="{{r.id}}"><label>{{L('Type d’archivage','Archive type','نوع الأرشفة')}}<select name="archive_type" required><option value="temporary">{{L('Temporaire et réversible','Temporary and reversible','مؤقت وقابل للاسترجاع')}}</option><option value="permanent">{{L('Définitif','Permanent','نهائي')}}</option></select></label><label>{{L('Motif obligatoire','Required reason','السبب إجباري')}}<textarea name="archive_reason" maxlength="500" required></textarea></label><label>{{L('Date de retour (obligatoire si temporaire)','Return date (required if temporary)','تاريخ العودة (إجباري إذا كانت الأرشفة مؤقتة)')}}<input type="date" name="archive_return_date"></label><button type="submit">{{L('Confirmer l’archivage','Confirm archive','تأكيد الأرشفة')}}</button></form></details><details><summary>{{L('Supprimer le dossier (corbeille pendant 6 mois)','Delete record (kept in trash for 6 months)','حذف الملف (الاحتفاظ في سلة المهملات 6 أشهر)')}}</summary><form method="post" onsubmit="return confirm('{{L('Confirmer le déplacement vers la corbeille ? Le dossier pourra être restauré pendant six mois.','Move this record to the trash? It can be restored for six months.','تأكيد نقل الملف إلى سلة المهملات؟ يمكن استعادته خلال ستة أشهر.')}}')"><input type="hidden" name="action" value="delete"><input type="hidden" name="rid" value="{{r.id}}"><label>{{L('Motif obligatoire','Required reason','السبب إجباري')}}<textarea name="delete_reason" maxlength="500" required></textarea></label><button type="submit">{{L('Déplacer vers la corbeille','Move to trash','نقل إلى سلة المهملات')}}</button></form>{% endif %}</td></tr>{% endfor %}</table></div></section>""",rows=rows,due_archives=due_archives,today=date.today().isoformat())
@app.route("/terrains",methods=["GET","POST"])
def terrains():
 if guard(): return redirect("/login")
 c=db()
 if request.method=="POST":
  if session.get("role")!="super_admin": c.close(); return ("Forbidden — Super Admin only",403)
  action=request.form.get("action","create")
  try:
   if action=="edit":
    tid=int(request.form["tid"]); name=request.form["name"].strip(); kind=request.form["kind"]; mandatory_year=int(request.form["mandatory_year"])
    if not name or kind not in ("mandatory","basic","optional") or mandatory_year not in (0,1,2,3,4): raise ValueError("Invalid terrain fields")
    current=c.execute("SELECT * FROM terrains WHERE id=?",(tid,)).fetchone()
    if not current: raise ValueError("Terrain not found.")
    has_records=c.execute("SELECT (SELECT COUNT(*) FROM history WHERE terrain_id=?)+(SELECT COUNT(*) FROM assignments WHERE terrain_id=?)",(tid,tid)).fetchone()[0]>0
    if has_records and (kind!=current["kind"] or mandatory_year!=current["mandatory_year"]):
     raise ValueError("Terrain classification/year cannot be changed after rotation history or assignments exist. Preserve the audit trail and correct it before use or via a reviewed migration.")
    c.execute("UPDATE terrains SET name=?,kind=?,mandatory_year=? WHERE id=?",(name,kind,mandatory_year,tid))
    audit_event(c,"terrain.updated",f"terrain_id={tid}; name={name}; kind={kind}; mandatory_year={mandatory_year}")
   elif action in ("deactivate","activate"):
    tid=int(request.form["tid"]); active=0 if action=="deactivate" else 1
    current=c.execute("SELECT * FROM terrains WHERE id=?",(tid,)).fetchone()
    if not current: raise ValueError("Terrain not found.")
    c.execute("UPDATE terrains SET active=? WHERE id=?",(active,tid))
    audit_event(c,"terrain.deactivated" if not active else "terrain.activated",f"terrain_id={tid}")
   else:
    name=request.form["name"].strip(); kind=request.form["kind"]; mandatory_year=int(request.form["mandatory_year"])
    if not name or kind not in ("mandatory","basic","optional") or mandatory_year not in (0,1,2,3,4): raise ValueError("Invalid terrain fields")
    c.execute("INSERT INTO terrains(name,kind,mandatory_year) VALUES(?,?,?)",(name,kind,mandatory_year))
    audit_event(c,"terrain.created",name+" kind="+kind)
   c.commit(); flash(tr("done"))
  except sqlite3.IntegrityError: c.rollback(); flash(L("Cet élément existe déjà.","Already exists.","هذا العنصر موجود مسبقًا."))
  except ValueError as exc: c.rollback(); flash(str(exc))
 rows=c.execute("SELECT * FROM terrains ORDER BY CASE WHEN kind='optional' THEN 1 ELSE 0 END, name").fetchall(); c.close()
 return page("""<section><h3>{{t('terrains')}}</h3><details class="rt-add-terrain"><summary>{{L('Ajouter un service','Add a department','إضافة مصلحة')}}</summary><form method="post"><input type="hidden" name="action" value="create"><label>{{t('terrain')}}</label><input name="name" required><div class="grid"><div><label>{{t('kind')}}</label><select name="kind"><option value="mandatory">{{t('mandatory')}}</option><option value="basic">{{t('basic')}}</option><option value="optional">{{t('optional')}}</option></select></div><div><label>{{t('mandatory_year')}}</label><select name="mandatory_year"><option value="0">{{t('any')}}</option>{% for y in range(1,5) %}<option value="{{y}}">{{y}}</option>{% endfor %}</select></div></div><button>{{L('Enregistrer le service','Save department','حفظ المصلحة')}}</button></form></details></section><section><div class="wrap"><table><tr><th>{{t('terrain')}}</th><th>{{t('kind')}}</th><th>{{t('mandatory_year')}}</th><th>{{t('status')}}</th><th>{{L('Actions administratives','Admin actions','الإجراءات الإدارية')}}</th></tr>{% for r in rows %}<tr><td colspan="5"><form method="post" class="grid"><input type="hidden" name="action" value="edit"><input type="hidden" name="tid" value="{{r.id}}"><input name="name" value="{{r.name}}" required><select name="kind"><option value="mandatory" {% if r.kind=='mandatory' %}selected{% endif %}>{{t('mandatory')}}</option><option value="basic" {% if r.kind=='basic' %}selected{% endif %}>{{t('basic')}}</option><option value="optional" {% if r.kind=='optional' %}selected{% endif %}>{{t('optional')}}</option></select><select name="mandatory_year"><option value="0" {% if not r.mandatory_year %}selected{% endif %}>{{t('any')}}</option>{% for y in range(1,5) %}<option value="{{y}}" {% if r.mandatory_year==y %}selected{% endif %}>{{y}}</option>{% endfor %}</select><span>{{L('Actif' if r.active else 'Inactif','Active' if r.active else 'Inactive','نشط' if r.active else 'غير نشط')}}</span><button>{{t('save')}}</button></form>{% if role=='super_admin' %}<form method="post"><input type="hidden" name="tid" value="{{r.id}}"><button name="action" value="{{'deactivate' if r.active else 'activate'}}">{{L('Désactiver' if r.active else 'Réactiver','Deactivate' if r.active else 'Reactivate','تعطيل' if r.active else 'إعادة التفعيل')}}</button></form>{% endif %}</td></tr>{% endfor %}</table></div><p>{{L('La désactivation retire un service des nouvelles affectations sans supprimer son historique.','Deactivation hides a department from new assignments but preserves its history.','يؤدي تعطيل المصلحة إلى منع التعيينات الجديدة دون حذف سجلها السابق.')}}</p></section>""",rows=rows)
@app.route("/capacities",methods=["GET","POST"])
def capacities():
 if guard(): return redirect("/login")
 c=db(); y=active_postgrad_year(c)
 if request.method=="POST":
  if session.get("role")!="super_admin": c.close(); return ("Forbidden — Super Admin only",403)
  try:
   tid=int(request.form["terrain"]); pid=int(request.form["period"]); places=int(request.form["places"])
   if places<0: raise ValueError(L("La capacité ne peut pas être négative.","Capacity cannot be negative.","لا يمكن أن تكون السعة سالبة."))
   terrain=c.execute("SELECT id FROM terrains WHERE id=? AND active=1",(tid,)).fetchone()
   period=c.execute("SELECT id FROM periods WHERE id=? AND year_label=?",(pid,y)).fetchone()
   if not terrain: raise ValueError(L("Choisissez un service actif.","Choose an active department.","اختر مصلحة نشطة."))
   if not period: raise ValueError(L("Choisissez une période de l’année de post-graduation sélectionnée.","Choose a period from the selected postgraduate year.","اختر فترة من سنة ما بعد التدرج المحددة."))
   used=c.execute("SELECT COUNT(*) FROM assignments WHERE terrain_id=? AND period_id=? AND status IN ('confirmed','pending')",(tid,pid)).fetchone()[0]
   if places<used: flash(tr("full"))
   else:
    c.execute("INSERT INTO capacities(year_label,terrain_id,period_id,places) VALUES(?,?,?,?) ON CONFLICT(year_label,terrain_id,period_id) DO UPDATE SET places=excluded.places",(y,tid,pid,places))
    audit_event(c,"capacity.changed",f"year={y}; terrain_id={tid}; period_id={pid}; places={places}")
    c.commit(); flash(tr("done"))
  except (ValueError,TypeError,KeyError):
   c.rollback(); flash(L("Saisie de capacité invalide. Choisissez un service actif, une période de l’année sélectionnée et un nombre entier positif ou nul.","Invalid capacity input. Choose an active department, a period in the selected postgraduate year, and a non-negative whole number of places.","إدخال السعة غير صالح. اختر مصلحة نشطة وفترة من سنة ما بعد التدرج المحددة وعددًا صحيحًا غير سالب للمقاعد."))
 ts=c.execute("SELECT * FROM terrains WHERE active=1 ORDER BY CASE WHEN kind='optional' THEN 1 ELSE 0 END, name").fetchall(); ps=c.execute("SELECT * FROM periods WHERE year_label=? ORDER BY id",(y,)).fetchall()
 rows=c.execute("SELECT t.name terrain,p.name period,cap.places,(SELECT COUNT(*) FROM assignments a WHERE a.terrain_id=t.id AND a.period_id=p.id AND a.status IN ('confirmed','pending')) used FROM capacities cap JOIN terrains t ON t.id=cap.terrain_id JOIN periods p ON p.id=cap.period_id WHERE cap.year_label=?",(y,)).fetchall(); c.close()
 return page("""<section><h3>{{t('capacities')}}</h3><form method="post"><div class="grid"><div><label>{{t('terrain')}}</label><select name="terrain">{% for x in ts %}<option value="{{x.id}}">{{x.name}}</option>{% endfor %}</select></div><div><label>{{t('period')}}</label><select name="period">{% for x in ps %}<option value="{{x.id}}">{{period_label(x.name)}}</option>{% endfor %}</select></div><div><label>{{t('places')}}</label><input type="number" min="0" name="places" required></div></div><button>{{t('save')}}</button></form></section><section><div class="wrap"><table><tr><th>{{t('period')}}</th><th>{{t('terrain')}}</th><th>{{t('places')}}</th><th>{{t('used')}}</th><th>{{t('remaining')}}</th></tr>{% for r in rows %}<tr><td>{{r.period}}</td><td>{{r.terrain}}</td><td>{{r.places}}</td><td>{{r.used}}</td><td>{{r.places-r.used}}</td></tr>{% endfor %}</table></div></section>""",ts=ts,ps=ps,rows=rows)
@app.route("/rotation-board", methods=["GET","POST"])
def rotation_board():
 if guard(): return redirect("/login")
 c=db(); y=request.form.get("year",request.args.get("year",active_postgrad_year(c)))
 if request.method=="POST":
  if session.get("role")!="super_admin":
   c.close(); return ("Forbidden — Super Admin only",403)
  action=request.form.get("action","")
  try:
   if action=="period_dates":
    pid=int(request.form.get("period_id","0"))
    period=c.execute("SELECT id,name FROM periods WHERE id=? AND year_label=?",(pid,y)).fetchone()
    if not period: raise ValueError("Période introuvable pour cette année de post-graduation")
    start_date=request.form.get("start_date","").strip(); end_date=request.form.get("end_date","").strip()
    if start_date: start_date=date.fromisoformat(start_date).isoformat()
    if end_date: end_date=date.fromisoformat(end_date).isoformat()
    if start_date and end_date and end_date<start_date: raise ValueError("La date de fin doit être postérieure ou égale à la date de début")
    if period["name"] in ("Période 1","Period 1","الفترة 1") and start_date:
     start_year=int(y.split("-")[0]); previous_year=f"{start_year-1}-{start_year}"
     previous_end=c.execute("SELECT end_date FROM periods WHERE year_label=? AND name IN ('Période 4','Period 4','الفترة 4') ORDER BY id DESC LIMIT 1",(previous_year,)).fetchone()
     if previous_end and previous_end["end_date"] and start_date<=previous_end["end_date"]: raise ValueError(L("Le début du nouveau roulement doit être postérieur à la fin du dernier stage de l’année précédente.","The new rotation must start after the final rotation of the previous postgraduate year.","يجب أن يبدأ التناوب الجديد بعد نهاية آخر تربص في سنة ما بعد التدرج السابقة."))
    c.execute("UPDATE periods SET start_date=?,end_date=? WHERE id=? AND year_label=?",(start_date,end_date,pid,y))
    audit_event(c,"period.dates_updated","year="+y+"; period_id="+str(pid)+"; start="+start_date+"; end="+end_date)
    c.commit(); flash(L("Dates de la période enregistrées.","Period dates saved.","تم حفظ تواريخ الفترة."))
    c.close(); return redirect("/rotation-board?year="+y.replace("&","%26").replace(" ","%20"))
   tid=int(request.form.get("terrain_id","0")); pid=int(request.form.get("period_id","0"))
   period=c.execute("SELECT id,name FROM periods WHERE id=? AND year_label=?",(pid,y)).fetchone()
   terrain=c.execute("SELECT id,name FROM terrains WHERE id=? AND active=1",(tid,)).fetchone()
   if not period or not terrain: raise ValueError("Terrain or period not found")
   if action=="unassign":
    rid=int(request.form.get("resident_id","0"))
    assignment=c.execute("SELECT a.*,r.surname,r.firstname,t.name terrain,p.name period FROM assignments a JOIN residents r ON r.id=a.resident_id JOIN terrains t ON t.id=a.terrain_id JOIN periods p ON p.id=a.period_id WHERE a.resident_id=? AND a.terrain_id=? AND a.period_id=? AND a.status IN ('confirmed','pending')",(rid,tid,pid)).fetchone()
    if not assignment: raise ValueError("Aucune affectation active à retirer pour ce résident")
    reason=L("Affectation retirée du tableau par le Super Admin","Assignment removed from the board by Super Admin","أزال المسؤول العام التعيين من الجدول")
    c.execute("INSERT INTO trash_items(entity_type,entity_id,label,reason,snapshot_json,deleted_at,deleted_by) VALUES(?,?,?,?,?,?,?)",("assignment",str(assignment["id"]),assignment["surname"]+" "+assignment["firstname"]+" — "+assignment["terrain"]+" / "+assignment["period"],reason,json.dumps({"record":{k:assignment[k] for k in ("id","resident_id","terrain_id","period_id","status")}},ensure_ascii=False),datetime.now().replace(microsecond=0).isoformat(),session.get("username","")))
    c.execute("UPDATE assignments SET status='cancelled' WHERE id=?",(assignment["id"],))
    audit_event(c,"assignment.cancelled","assignment_id="+str(assignment["id"])+"; resident_id="+str(rid)+"; terrain_id="+str(tid)+"; period_id="+str(pid)+"; year="+y+"; reason=user_removed_from_board; retained_in_trash_for_six_months=true")
    c.commit(); flash(L("Affectation retirée du tableau et conservée dans la corbeille pendant six mois. Vous pouvez la restaurer avant l’échéance.","Assignment removed from the board and kept in the trash for six months. You can restore it before expiry.","تمت إزالة التعيين من الجدول وحفظه في سلة المهملات لمدة ستة أشهر. يمكنك استعادته قبل انتهاء المدة."))
   elif action=="capacity":
    places=int(request.form.get("places",""))
    if places<0: raise ValueError(L("La capacité ne peut pas être négative.","Capacity cannot be negative.","لا يمكن أن تكون السعة سالبة."))
    used=c.execute("SELECT COUNT(*) FROM assignments WHERE terrain_id=? AND period_id=? AND status IN ('confirmed','pending')",(tid,pid)).fetchone()[0]
    if places<used:
     raise ValueError(L("La capacité ne peut pas être inférieure aux affectations actives déjà enregistrées ("+str(used)+").","Capacity cannot be lower than the active assignments already recorded ("+str(used)+").","لا يمكن أن تكون السعة أقل من التعيينات النشطة المسجلة بالفعل ("+str(used)+")."))
    c.execute("INSERT INTO capacities(year_label,terrain_id,period_id,places) VALUES(?,?,?,?) ON CONFLICT(year_label,terrain_id,period_id) DO UPDATE SET places=excluded.places",(y,tid,pid,places))
    audit_event(c,"capacity.updated","year="+y+"; terrain_id="+str(tid)+"; period_id="+str(pid)+"; places="+str(places))
    c.commit(); flash(L("Capacité enregistrée.","Capacity saved.","تم حفظ السعة."))
   elif action=="move":
    if session.get("role")!="super_admin": raise ValueError("Only a Super Admin may move an existing assignment")
    rid=int(request.form.get("resident_id","0"))
    old=c.execute("SELECT id,terrain_id,status FROM assignments WHERE resident_id=? AND period_id=? AND status IN ('confirmed','pending')",(rid,pid)).fetchone()
    if not old: raise ValueError("Resident has no active assignment in this period")
    c.execute("UPDATE assignments SET terrain_id=?,status='confirmed' WHERE id=?",(tid,old["id"]))
    audit_event(c,"assignment.moved","assignment_id="+str(old["id"])+"; resident_id="+str(rid)+"; from_terrain_id="+str(old["terrain_id"])+"; to_terrain_id="+str(tid)+"; period_id="+str(pid)+"; super_admin_override=1")
    c.commit(); flash(L("Affectation déplacée par le Super Admin.","Assignment moved by Super Admin.","تم نقل التعيين بواسطة المسؤول العام."))
   elif action=="assign":
    rid=int(request.form.get("resident_id","0"))
    force=request.form.get("force_rules")=="1"
    if force and session.get("role")!="super_admin":
     raise ValueError("Only a Super Admin may bypass assignment rules")
    if force:
     resident=c.execute("SELECT id FROM residents WHERE id=? AND status='active'",(rid,)).fetchone()
     if not resident: raise ValueError("Resident is not active")
     ok,msg,status=True,"Affectation dérogatoire Super Admin","confirmed"
    else:
     ok,msg,status=check(c,rid,tid,pid,y)
    if not ok:
     flash(msg)
    else:
     old=c.execute("SELECT id,status,terrain_id FROM assignments WHERE resident_id=? AND period_id=?",(rid,pid)).fetchone()
     if old and old["status"] not in ("rejected","cancelled"):
      flash(tr("duplicate"))
     elif old:
      c.execute("UPDATE assignments SET terrain_id=?,status=? WHERE id=?",(tid,status,old["id"]))
      audit_event(c,"assignment.replaced","assignment_id="+str(old["id"])+"; resident_id="+str(rid)+"; terrain_id="+str(tid)+"; period_id="+str(pid)+"; status="+str(status)+( "; super_admin_override=1" if force else ""))
      c.commit(); flash(msg)
     else:
      c.execute("INSERT INTO assignments(resident_id,terrain_id,period_id,status) VALUES(?,?,?,?)",(rid,tid,pid,status))
      audit_event(c,"assignment.created","resident_id="+str(rid)+"; terrain_id="+str(tid)+"; period_id="+str(pid)+"; status="+str(status)+( "; super_admin_override=1" if force else ""))
      c.commit(); flash(msg)
   else:
    raise ValueError("Unknown board action")
  except Exception as exc:
   c.rollback(); flash("Modification impossible : "+str(exc)[:180])
  c.close()
  return redirect("/rotation-board?year="+y.replace("&","%26").replace(" ","%20"))
 years=[x["year_label"] for x in c.execute("SELECT DISTINCT year_label FROM periods ORDER BY year_label DESC").fetchall()]
 if y not in years and years: y=years[0]
 periods=c.execute("SELECT * FROM periods WHERE year_label=? ORDER BY id",(y,)).fetchall()
 terrains=c.execute("SELECT * FROM terrains WHERE active=1 ORDER BY CASE WHEN kind='optional' THEN 1 ELSE 0 END,name").fetchall()
 assignments=c.execute("""SELECT a.resident_id,a.terrain_id,a.period_id,a.status,r.surname,r.firstname,r.year,r.resident_type,r.origin_hospital,t.name AS assigned_terrain
 FROM assignments a JOIN residents r ON r.id=a.resident_id JOIN terrains t ON t.id=a.terrain_id
 JOIN periods p ON p.id=a.period_id
 WHERE p.year_label=? AND a.status IN ('confirmed','pending')
 ORDER BY a.period_id,r.year DESC,r.surname,r.firstname""",(y,)).fetchall()
 capacities=c.execute("SELECT terrain_id,period_id,places FROM capacities WHERE year_label=?",(y,)).fetchall()
 capmap={(row["terrain_id"],row["period_id"]):row["places"] for row in capacities}
 residents=c.execute("SELECT id,surname,firstname,year,resident_type,origin_hospital FROM residents WHERE status='active' ORDER BY year DESC,surname,firstname").fetchall()
 assigned={}
 for a in assignments: assigned.setdefault((a["terrain_id"],a["period_id"]),[]).append(a)
 board=[]
 for terrain in terrains:
  cells=[]
  for period in periods:
   names=assigned.get((terrain["id"],period["id"]),[])
   capacity=capmap.get((terrain["id"],period["id"]))
   already_assigned={row["resident_id"] for row in c.execute("SELECT resident_id FROM assignments WHERE period_id=? AND status IN ('confirmed','pending')",(period["id"],)).fetchall()}
   pickable=[r for r in residents if r["id"] not in already_assigned]
   move_options=[a for a in assignments if a["period_id"]==period["id"] and a["terrain_id"]!=terrain["id"]]
   cells.append({"period_id":period["id"],"names":names,"used":len(names),"capacity":capacity,"remaining":max(0,capacity-len(names)) if capacity is not None else None,"pickable":pickable,"move_options":move_options})
  board.append({"id":terrain["id"],"name":terrain["name"],"cells":cells})
 if request.args.get("format")=="pdf":
  try:
   from reportlab.lib import colors
   from reportlab.lib.pagesizes import A4, landscape
   from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
   from reportlab.lib.enums import TA_CENTER, TA_LEFT
   from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
   from reportlab.lib.units import mm
   output=io.BytesIO()
   doc=SimpleDocTemplate(output,pagesize=landscape(A4),rightMargin=9*mm,leftMargin=9*mm,topMargin=10*mm,bottomMargin=10*mm,title="ResiManage - Répartition des stages")
   styles=getSampleStyleSheet()
   title_style=ParagraphStyle("BoardTitle",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=15,leading=18,alignment=TA_CENTER,textColor=colors.HexColor("#183c5a"),spaceAfter=5)
   cell_style=ParagraphStyle("BoardCell",parent=styles["BodyText"],fontName="Helvetica",fontSize=7,leading=9,alignment=TA_LEFT,wordWrap="CJK")
   head_style=ParagraphStyle("BoardHead",parent=cell_style,fontName="Helvetica-Bold",alignment=TA_CENTER,textColor=colors.white)
   story=[Paragraph("Centre Hospitalo-Universitaire Nedir Mohamed — Tizi Ouzou",title_style),Paragraph("Département d’Anesthésie-Réanimation",title_style),Paragraph("Chef de service : Professeur Benhocine",styles["Normal"]),Paragraph("Tableau général de répartition des stages — année de post-graduation : "+str(y),styles["Normal"]),Spacer(1,5*mm)]
   data=[[Paragraph("Terrain de stage",head_style)]+[Paragraph(str(p["name"]),head_style) for p in periods]]
   for item in board:
    row=[Paragraph(item["name"],cell_style)]
    for cell in item["cells"]:
     roster="<br/>".join(("• "+(str(a["surname"] or "")+" "+str(a["firstname"] or "")).strip()+" (A"+str(a["year"])+(" — en attente" if a["status"]=="pending" else "")+")") for a in cell["names"])
     cap="Capacité : "+str(cell["capacity"]) if cell["capacity"] is not None else "Capacité : non définie"
     remaining=("Places restantes : "+str(cell["remaining"])) if cell["remaining"] is not None else "Places restantes : —"
     content=(roster+"<br/>" if roster else "")+"Occupées : "+str(cell["used"])+" / "+(str(cell["capacity"]) if cell["capacity"] is not None else "—")+"<br/>"+remaining
     row.append(Paragraph(content,cell_style))
    data.append(row)
   widths=[42*mm]+[(landscape(A4)[0]-18*mm-42*mm)/max(1,len(periods))]*len(periods)
   table=Table(data,colWidths=widths,repeatRows=1,hAlign="LEFT")
   table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#214e79")),("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#cbd8e4")),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f0f6fb")])]))
   story.append(table); doc.build(story); output.seek(0)
   return send_file(output,mimetype="application/pdf",as_attachment=True,download_name="ResiManage_repartition_"+str(y).replace("/","-")+".pdf")
  except ImportError:
   return ("PDF export requires the reportlab package. Install requirements.txt and retry.",500)
 current_year_label=setting(c,"year")
 c.close()
 return page("""<section class="rt-board-head"><div class="rt-board-print-heading"><h2>{{L('Centre Hospitalio-universitaire Nedir Mohamed, Tizi Ouzou','Nedir Mohamed University Hospital, Tizi Ouzou','المركز الاستشفائي الجامعي نذير محمد، تيزي وزو')}}</h2><h3>{{L('Département d’Anesthésie-Réanimation','Anesthesia and Critical Care Department','قسم التخدير والإنعاش')}}</h3><p><strong>{{L('Chef de service : Professeur Benhocine','Department head: Professor Benhocine','رئيس المصلحة: الأستاذ بن حسين')}}</strong></p></div><h2>{{L('Tableau général de répartition des stages','General rotation allocation board','الجدول العام لتوزيع التربصات')}}</h2><p>{{L('Année de post-graduation :','Postgraduate year:','سنة ما بعد التدرج:')}} <strong>{{year_label}}</strong>. {{L('Les affectations confirmées et en attente sont affichées. Une capacité non définie est indiquée par « — ».','Confirmed and pending assignments are shown. Undefined capacity is indicated by “—”.','تظهر التعيينات المؤكدة والمعلقة. تُشار السعة غير المحددة بالرمز «—».')}}</p><form method="get" class="rt-board-controls"><label>{{L('Année de post-graduation','Postgraduate year','سنة ما بعد التدرج')}}<select name="year">{% set past_years = years|reject('equalto', current_year_label)|list %}<optgroup label="{{L('Année en cours','Current year','السنة الحالية')}}"><option value="{{current_year_label}}" {% if year_label==current_year_label %}selected{% endif %}>{{current_year_label}} — {{L('en cours','current','الحالية')}}</option></optgroup>{% if past_years %}<optgroup label="{{L('Historique des années précédentes','Previous years history','سجل السنوات السابقة')}}">{% for yy in past_years %}<option value="{{yy}}" {% if yy==year_label %}selected{% endif %}>{{yy}}</option>{% endfor %}</optgroup>{% endif %}</select></label><button type="submit">{{L('Afficher','Display','عرض')}}</button></form>{% if session.get("role")=="super_admin" %}<section class="rt-period-date-editor"><h3>{{L('Dates des périodes','Period dates','تواريخ الفترات')}} — {{year_label}}</h3><p>{{L('Les dates ne sont plus prédéfinies. Saisissez les dates de début et de fin de chaque période. Le début de la période 1 doit être postérieur à la fin du dernier stage de l’année de post-graduation précédente.','Dates are not predefined. Enter start and end dates for each period. Period 1 must start after the final rotation of the previous postgraduate year.','التواريخ غير محددة مسبقًا. أدخل تاريخ بداية ونهاية كل فترة. يجب أن تبدأ الفترة الأولى بعد نهاية آخر تربص في سنة ما بعد التدرج السابقة.')}}</p><div class="rt-period-date-grid">{% for p in periods %}<form method="post" class="rt-period-date-form"><input type="hidden" name="action" value="period_dates"><input type="hidden" name="year" value="{{year_label}}"><input type="hidden" name="period_id" value="{{p.id}}"><strong>{{period_label(p.name)}}</strong><label>{{L('Début','Start','البداية')}} <input type="date" name="start_date" value="{{p.start_date or ''}}"></label><label>{{L('Fin','End','النهاية')}} <input type="date" name="end_date" value="{{p.end_date or ''}}"></label><button type="submit">{{L('Enregistrer les dates','Save dates','حفظ التواريخ')}}</button></form>{% endfor %}</div></section>{% endif %}<div class="rt-board-actions"><a href="/rotation-board?year={{year_label|urlencode}}&format=pdf">{{L('Exporter en PDF','Export PDF','تصدير PDF')}}</a><button type="button" onclick="window.print()">{{L('Imprimer le tableau','Print board','طباعة الجدول')}}</button><a href="/assign">{{L('Retour aux affectations','Back to assignments','العودة إلى التعيينات')}}</a></div></section><section class="rt-board-table"><div class="wrap"><table><thead><tr><th>{{L('Service de stage','Rotation department','مصلحة التربص')}}</th>{% for p in periods %}<th>{{period_label(p.name)}}{% if p.start_date or p.end_date %}<br><small>{{p.start_date or L('Début à définir','Start to be set','تحديد البداية')}} – {{p.end_date or L('Fin à définir','End to be set','تحديد النهاية')}}</small>{% endif %}</th>{% endfor %}</tr></thead><tbody>{% for item in board %}<tr><th>{{item.name}}</th>{% for cell in item.cells %}<td class="{% if cell.capacity is not none and cell.remaining==0 %}rt-board-full{% elif cell.used>0 %}rt-board-partial{% else %}rt-board-free{% endif %}"><div class="rt-board-roster">{% if cell.names %}{% for n in cell.names %}<div class="rt-board-name">{{n.surname}} {{n.firstname}} <small>({{n.year}}e année{% if n.resident_type=='external' %} — {{L('externe','external','خارجي')}}{% if n.origin_hospital %}, {{n.origin_hospital}}{% endif %}{% endif %}{% if n.status=='pending' %}, {{L('en attente','pending','قيد الانتظار')}}{% endif %})</small>{% if session.get('role')=='super_admin' %}<form method="post" class="rt-board-remove-form" onsubmit="return confirm('Supprimer ce choix du tableau ? Cette action peut être corrigée en réaffectant le résident.');"><input type="hidden" name="action" value="unassign"><input type="hidden" name="year" value="{{year_label}}"><input type="hidden" name="terrain_id" value="{{item.id}}"><input type="hidden" name="period_id" value="{{cell.period_id}}"><input type="hidden" name="resident_id" value="{{n.resident_id}}"><button type="submit" aria-label="Supprimer le choix de {{n.surname}} {{n.firstname}}">Retirer</button></form>{% endif %}</div>{% endfor %}{% else %}<span class="muted">{{L('Aucun résident affecté','No resident assigned','لا يوجد مقيم معيّن')}}</span>{% endif %}</div><div class="rt-board-count">{{L('Places occupées','Occupied places','المقاعد المشغولة')}} : <strong>{{cell.used}}</strong> / {{cell.capacity if cell.capacity is not none else "—"}}</div><div class="rt-board-remaining">{{L('Places restantes','Remaining places','المقاعد المتبقية')}} : <strong>{{cell.remaining if cell.remaining is not none else L('capacité non définie','capacity undefined','السعة غير محددة')}}</strong></div>{% if session.get('role')=='super_admin' %}<form method="post" class="rt-board-cell-form"><input type="hidden" name="action" value="capacity"><input type="hidden" name="year" value="{{year_label}}"><input type="hidden" name="terrain_id" value="{{item.id}}"><input type="hidden" name="period_id" value="{{cell.period_id}}"><label>{{L('Nombre de places','Number of places','عدد المقاعد')}}<input type="number" min="0" step="1" name="places" value="{{cell.capacity if cell.capacity is not none else ''}}" placeholder="{{L('À définir','To be set','يُحدد لاحقًا')}}" required></label><button type="submit">{{L('Enregistrer','Save','حفظ')}}</button></form><form method="post" class="rt-board-cell-form"><input type="hidden" name="action" value="assign"><input type="hidden" name="year" value="{{year_label}}"><input type="hidden" name="terrain_id" value="{{item.id}}"><input type="hidden" name="period_id" value="{{cell.period_id}}"><label>{{L('Ajouter un résident','Add resident','إضافة مقيم')}}<select name="resident_id" required><option value="">{{L('Choisir un résident…','Choose a resident…','اختر مقيمًا…')}}</option>{% for year_group, residents_in_year in cell.pickable|groupby('year')|reverse %}<optgroup label="{{year_group}} {{L('e année','year','سنة')}}">{% for resident in residents_in_year %}<option value="{{resident.id}}">{{resident.surname}} {{resident.firstname}}{% if resident.resident_type=='external' %} — {{L('externe','external','خارجي')}}{% if resident.origin_hospital %} ({{resident.origin_hospital}}){% endif %}{% endif %}</option>{% endfor %}</optgroup>{% endfor %}</select></label><label class="rt-board-override"><input type="checkbox" name="force_rules" value="1"> {{L('Dérogation aux règles (Super Admin)','Override rules (Super Admin)','تجاوز القواعد (المسؤول العام)')}}</label><button type="submit">{{L('Affecter','Assign','تعيين')}}</button></form>{% if cell.move_options %}<form method="post" class="rt-board-cell-form"><input type="hidden" name="action" value="move"><input type="hidden" name="year" value="{{year_label}}"><input type="hidden" name="terrain_id" value="{{item.id}}"><input type="hidden" name="period_id" value="{{cell.period_id}}"><label>Déplacer ici un résident déjà affecté<select name="resident_id" required><option value="">{{L('Choisir…','Choose…','اختر…')}}</option>{% for assigned_resident in cell.move_options %}<option value="{{assigned_resident.resident_id}}">{{assigned_resident.surname}} {{assigned_resident.firstname}}{% if assigned_resident.resident_type=='external' %} — {{L('externe','external','خارجي')}}{% if assigned_resident.origin_hospital %} ({{assigned_resident.origin_hospital}}){% endif %}{% endif %} — {{assigned_resident.assigned_terrain}}</option>{% endfor %}</select></label><button type="submit">{{L('Déplacer (dérogation)','Move (override)','نقل (استثناء)')}}</button></form>{% endif %}{% endif %}</td>{% endfor %}</tr>{% endfor %}</tbody></table></div><p class="muted">{{L('Les capacités se modifient directement dans chaque cellule du tableau. Les résidents externes sont inclus dans les effectifs et signalés avec leur établissement d’origine.','Edit capacity directly in each table cell. External residents are included in counts and labeled with their home institution.','يمكن تعديل السعة مباشرة في كل خلية من الجدول. يُحتسب المقيمون الخارجيون ضمن العدد ويُشار إلى مؤسساتهم الأصلية.')}}</p></section><style>.rt-board-print-heading{display:none}.rt-board-head h2{margin-top:0}.rt-board-controls{display:flex;align-items:end;gap:12px;flex-wrap:wrap}.rt-board-controls label{min-width:210px;flex:1}.rt-board-actions{display:flex;gap:8px;flex-wrap:wrap}.rt-board-table table{min-width:950px;table-layout:fixed}.rt-board-table th:first-child{width:175px}.rt-board-table th{background:#214e79;color:white}.rt-board-table td{vertical-align:top;min-width:160px}.rt-board-name{padding:4px 0;border-bottom:1px solid #dce6f0}.rt-board-name small{color:#526274}.rt-board-count{margin-top:8px;padding-top:5px;border-top:1px solid #dce6f0}.rt-board-remaining{font-weight:600;margin-top:4px}.rt-board-cell-form{display:grid;gap:4px;margin-top:8px;padding-top:7px;border-top:1px solid #dce6f0}.rt-board-cell-form label{font-size:12px;font-weight:600}.rt-board-cell-form input,.rt-board-cell-form select{width:100%;min-width:0;padding:5px;border:1px solid #b7c7d7;border-radius:5px;background:white}.rt-board-cell-form button{padding:5px 7px;font-size:12px}.rt-board-full{background:#ffe8e6}.rt-board-partial{background:#fff5db}.rt-board-free{background:#e7f7ee}@page{size:A4 landscape;margin:8mm}@media print{html,body{width:100%!important;max-width:none!important;margin:0!important;padding:0!important;background:white!important;color:black!important}body{zoom:1!important}main,.rt-app-shell,.rt-app-shell main,.rt-board-head,.rt-board-table,.rt-board-table .wrap{width:100%!important;max-width:none!important;min-width:0!important;margin:0!important;padding:0!important;box-sizing:border-box!important}body{background:white!important;color:black}header,nav,.main-nav,.rt-board-controls,.rt-board-actions,.rt-board-cell-form,footer{display:none!important}.rt-board-print-heading{display:block;text-align:center;margin-bottom:12px}.rt-board-print-heading h2{margin:0;font-size:17pt}.rt-board-print-heading h3{margin:3px 0;font-size:12pt}.rt-board-print-heading p{margin:3px 0}main{max-width:none;padding:0!important}.rt-board-head,.rt-board-table{border:0;box-shadow:none;break-inside:avoid}.rt-board-table .wrap{overflow:visible}.rt-board-table table{min-width:0;width:100%;font-size:8pt;table-layout:fixed}.rt-board-table th:first-child{width:18%}.rt-board-table th,.rt-board-table td{padding:4px;border:1px solid #777;overflow-wrap:anywhere;break-inside:avoid}.rt-board-table th{background:#e8eef4!important;color:#111!important;-webkit-print-color-adjust:exact;print-color-adjust:exact}.rt-board-table tr{break-inside:avoid}}@media(max-width:600px){.rt-board-actions>*{flex:1;text-align:center}.rt-board-remove-form{display:inline-block;margin-left:6px}.rt-board-remove-form button{padding:2px 6px;border:1px solid #b42318;border-radius:5px;background:#fff1f0;color:#9f1c13;font-size:11px;cursor:pointer}.rt-period-date-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(210px,1fr));gap:10px}.rt-period-date-form{display:grid;gap:7px;padding:12px;border:1px solid #cbd8e4;border-radius:10px;background:#f8fbff}.rt-period-date-form label{font-size:13px;font-weight:600}.rt-period-date-form input{width:100%;padding:7px;border:1px solid #8aa3ba;border-radius:6px;background:#fff;color:#10263a}.rt-board-table td{color:#10263a!important;background-color:#f8fbff}.rt-board-table .rt-board-roster,.rt-board-table .rt-board-name,.rt-board-table .rt-board-count,.rt-board-table .rt-board-remaining{color:#10263a!important}.rt-board-table .rt-board-name small{color:#334e68!important}.rt-board-table .muted{color:#40566d!important}.rt-board-table .rt-board-full{background:#ffe3e3!important;color:#641b1b!important}.rt-board-table .rt-board-partial{background:#fff0c2!important;color:#5b3a00!important}.rt-board-table .rt-board-free{background:#dcf8e7!important;color:#14532d!important}.rt-board-cell-form{color:#10263a!important}.rt-board-cell-form label{color:#17324d!important}.rt-board-cell-form input,.rt-board-cell-form select{color:#10263a!important;background:#fff!important;border-color:#8aa3ba!important}.rt-board-cell-form button{color:#fff!important}@media print{.rt-board-table td,.rt-board-table .rt-board-roster,.rt-board-table .rt-board-name,.rt-board-table .rt-board-count,.rt-board-table .rt-board-remaining{color:#111!important}}</style>""",year_label=y,current_year_label=current_year_label,years=years,periods=periods,board=board)

def check(c,rid,tid,pid,y):
 r=c.execute("SELECT * FROM residents WHERE id=? AND status='active'",(rid,)).fetchone(); t=c.execute("SELECT * FROM terrains WHERE id=?",(tid,)).fetchone(); p=c.execute("SELECT * FROM periods WHERE id=? AND year_label=?",(pid,y)).fetchone()
 if not r or not t or not t["active"] or not p: return False,tr("none"),None
 start_day, end_day = rotation_period_bounds(p["year_label"], p["name"], p["start_date"], p["end_date"])
 if r["resident_type"]=="external":
  try:
   ext_start=date.fromisoformat(r["external_start_date"] or "")
   ext_end=date.fromisoformat(r["external_end_date"] or "")
  except ValueError:
   return False,L("Les dates du résident externe doivent être renseignées avant l’affectation.","External resident dates must be completed before assignment.","يجب إدخال تواريخ المقيم الخارجي قبل التعيين."),None
  if ext_start>end_day or ext_end<start_day:
   return False,L("Cette période est en dehors des dates du stage du résident externe.","This period falls outside the external resident's rotation dates.","هذه الفترة خارج تواريخ تربص المقيم الخارجي."),None
 external = c.execute("SELECT hospital,country,start_date,end_date FROM external_rotations WHERE resident_id=? AND request_status='approved' AND status<>'cancelled' AND start_date<=? AND end_date>=? LIMIT 1",(rid,end_day.isoformat(),start_day.isoformat())).fetchone()
 if external: return False,(L("Un stage à l’étranger approuvé bloque l’affectation locale : ","An approved external rotation blocks the local assignment: ","يمنع تربص خارجي معتمد التعيين المحلي: ")+external["hospital"]+", "+external["country"]+" ("+external["start_date"]+" "+L("au","to","إلى")+" "+external["end_date"]+")"),None
 dup=c.execute("SELECT t.name FROM assignments a JOIN terrains t ON t.id=a.terrain_id WHERE a.resident_id=? AND a.period_id=? AND a.status IN ('confirmed','pending')",(rid,pid)).fetchone()
 if dup:return False,tr("duplicate")+dup["name"],None
 repeat = c.execute(
  "SELECT t.name,p.name AS period,p.year_label FROM assignments a JOIN terrains t ON t.id=a.terrain_id JOIN periods p ON p.id=a.period_id WHERE a.resident_id=? AND a.terrain_id=? AND a.period_id<>? AND a.status IN ('confirmed','pending') LIMIT 1",
  (rid,tid,pid)
 ).fetchone()
 total=c.execute(
  "SELECT COUNT(*) FROM terrains WHERE kind IN ('mandatory','basic') AND active=1"
 ).fetchone()[0]
 # Historical rows without an explicit completion date still need a completed period.
 # A future academic year with a blank date is not proof that the rotation happened.
 today=date.today()
 completed_terrains=set()
 history_rows=c.execute("""
  SELECT h.terrain_id,h.year_label,h.period,h.date_done,
         p.start_date AS period_start_date,p.end_date AS period_end_date
  FROM history h JOIN terrains t ON t.id=h.terrain_id
  LEFT JOIN periods p ON p.year_label=h.year_label AND p.name=h.period
  WHERE h.resident_id=? AND t.kind IN ('mandatory','basic') AND t.active=1
 """,(rid,)).fetchall()
 for history_row in history_rows:
  done_date=str(history_row["date_done"] or "").strip()
  if done_date:
   try:
    if date.fromisoformat(done_date)<=today:
     completed_terrains.add(history_row["terrain_id"])
   except ValueError:
    continue
  else:
   try:
    _,history_period_end=rotation_period_bounds(history_row["year_label"],history_row["period"],history_row["period_start_date"],history_row["period_end_date"])
    if history_period_end<today:
     completed_terrains.add(history_row["terrain_id"])
   except (ValueError,TypeError):
    continue
 # A confirmed assignment is completion evidence only after its rotation period has ended.
 # Future/current assignments must not unlock repetition of mandatory terrains.
 completed_assignments=c.execute("""
  SELECT a.terrain_id,p.year_label,p.name,p.start_date,p.end_date FROM assignments a
  JOIN terrains t ON t.id=a.terrain_id JOIN periods p ON p.id=a.period_id
  WHERE a.resident_id=? AND a.status='confirmed'
    AND t.kind IN ('mandatory','basic') AND t.active=1
 """,(rid,)).fetchall()
 for completed in completed_assignments:
  _,period_end=rotation_period_bounds(completed["year_label"],completed["name"],completed["start_date"],completed["end_date"])
  if period_end<today:
   completed_terrains.add(completed["terrain_id"])
 done=len(completed_terrains)
 if repeat and done<total:
  return False,tr("same_terrain_year")+repeat["name"]+" — "+repeat["period"],None
 cap=c.execute("SELECT places FROM capacities WHERE year_label=? AND terrain_id=? AND period_id=?",(y,tid,pid)).fetchone()
 if not cap:return False,tr("capacity_first"),None
 used=c.execute("SELECT COUNT(*) FROM assignments WHERE terrain_id=? AND period_id=? AND status IN ('confirmed','pending')",(tid,pid)).fetchone()[0]
 if used>=cap["places"]:return False,tr("full"),None
 h=c.execute("SELECT year_label,period,date_done FROM history WHERE resident_id=? AND terrain_id=? ORDER BY id DESC LIMIT 1",(rid,tid)).fetchone()
 if h and done<total:return False,tr("history_block")+f" ({h['year_label']} / {h['period']} / {h['date_done']})",None
 if t["mandatory_year"] and t["mandatory_year"]!=r["year"]:return True,tr("mismatch"),"pending"
 return True,(tr("history_repeat") if h else tr("done")),"confirmed"
@app.route("/assign",methods=["GET","POST"])
def assign():
 if guard():return redirect("/login")
 c=db(); y=active_postgrad_year(c)
 if request.method=="POST":
  if session.get("role")!="super_admin":
   c.close(); return ("Forbidden — Super Admin only",403)
  try:
   rid,tid,pid=map(int,(request.form["resident"],request.form["terrain"],request.form["period"]))
   override=request.form.get("override")=="1"
   override_reason=request.form.get("override_reason","").strip()
   resident_check=c.execute("SELECT id,status FROM residents WHERE id=?",(rid,)).fetchone()
   terrain_check=c.execute("SELECT id,name FROM terrains WHERE id=? AND active=1",(tid,)).fetchone()
   period_check=c.execute("SELECT id,name,year_label FROM periods WHERE id=? AND year_label=?",(pid,y)).fetchone()
   if not resident_check or not terrain_check or not period_check:
    raise ValueError("Resident, terrain or period does not exist in the selected postgraduate year.")
   if override and (not override_reason or session.get("role")!="super_admin"):
    flash(L("Une dérogation exige un motif et le rôle Super Admin.","An override requires a reason and the Super Admin role.","يتطلب الاستثناء سببًا وصلاحية المسؤول العام."))
   else:
    ok,msg,status=check(c,rid,tid,pid,y)
    if not ok and not override:
     flash(msg)
    else:
     if override:
      status="confirmed"
      msg="Dérogation Super Admin enregistrée."
     old=c.execute("SELECT id,status,terrain_id FROM assignments WHERE resident_id=? AND period_id=?",(rid,pid)).fetchone()
     if old and old["status"] not in ("rejected","cancelled"):
      if not override:
       flash(tr("duplicate"))
      else:
       c.execute("UPDATE assignments SET terrain_id=?,status=? WHERE id=?",(tid,status,old["id"]))
       audit_event(c,"assignment.super_admin_override",f"assignment_id={old['id']}; resident_id={rid}; period_id={pid}; year={y}; terrain_before={old['terrain_id']}; terrain_after={tid}; status_before={old['status']}; status_after={status}; reason={override_reason}")
       c.commit(); flash(msg)
     elif old:
      c.execute("UPDATE assignments SET terrain_id=?,status=? WHERE id=?",(tid,status,old["id"]))
      audit_event(c,"assignment.replaced",f"assignment_id={old['id']}; resident_id={rid}; period_id={pid}; year={y}; terrain_id={tid}; status={status}; override={override}; reason={override_reason}")
      c.commit(); flash(msg)
     else:
      c.execute("INSERT INTO assignments(resident_id,terrain_id,period_id,status) VALUES(?,?,?,?)",(rid,tid,pid,status))
      audit_event(c,"assignment.created",f"resident_id={rid}; terrain_id={tid}; period_id={pid}; year={y}; status={status}; override={override}; reason={override_reason}")
      c.commit(); flash(msg)
  except Exception: c.rollback(); flash(L("Affectation invalide.","Invalid assignment.","تعيين غير صالح."))
 rs=c.execute("SELECT * FROM residents WHERE status='active' ORDER BY year DESC,surname").fetchall(); ts=c.execute("SELECT * FROM terrains WHERE active=1 ORDER BY CASE WHEN kind='optional' THEN 1 ELSE 0 END, name").fetchall(); ps=c.execute("SELECT * FROM periods WHERE year_label=? ORDER BY id",(y,)).fetchall()
 external_records=c.execute("SELECT resident_id,start_date,end_date FROM external_rotations WHERE request_status='approved' AND status<>'cancelled'").fetchall()
 blocked_periods_by_resident={}
 for r in rs:
  blocked=[]
  for p in ps:
   pstart,pend=rotation_period_bounds(p["year_label"],p["name"],p["start_date"],p["end_date"])
   if any(e["resident_id"]==r["id"] and e["start_date"]<=pend.isoformat() and e["end_date"]>=pstart.isoformat() for e in external_records):
    blocked.append(str(p["id"]))
  blocked_periods_by_resident[r["id"]]=",".join(blocked)
 rows=c.execute("SELECT a.*,r.surname,r.firstname,r.year,t.name terrain,p.name period FROM assignments a JOIN residents r ON r.id=a.resident_id JOIN terrains t ON t.id=a.terrain_id JOIN periods p ON p.id=a.period_id WHERE p.year_label=? ORDER BY p.id,r.year DESC",(y,)).fetchall(); c.close()
 return page("""<section class="rt-history-yearbar"><h3>{{L('Année de post-graduation sélectionnée','Selected postgraduate year','سنة ما بعد التدرج المختارة')}} : {{year}}</h3><form method="post" action="/select-year" class="rt-year-inline"><input type="hidden" name="next" value="/assign"><label for="assignment-postgrad-year">{{t('year')}}</label><select id="assignment-postgrad-year" name="year" onchange="this.form.submit()">{% for yy in year_options %}<option value="{{yy}}" {% if yy==year %}selected{% endif %}>{{yy}}{% if yy==current_year %} — {{L('en cours','current','الحالية')}}{% endif %}</option>{% endfor %}</select><noscript><button type="submit">{{t('choose')}}</button></noscript></form></section><section class="rt-assignment-hero"><h2>{% if lang=='ar' %}تخطيط التربصات وتعيين المقيمين{% elif lang=='fr' %}Planification et affectation des résidents{% else %}Rotation planning & resident assignment{% endif %}</h2><p>{{t('priority')}}</p><div class="rt-assignment-stats"><div><span>{% if lang=='ar' %}المقيمون النشطون{% elif lang=='fr' %}Résidents actifs{% else %}Active residents{% endif %}</span><strong>{{rs|length}}</strong></div><div><span>{% if lang=='ar' %}التخصصات / الأراضي{% elif lang=='fr' %}Terrains actifs{% else %}Active rotations{% endif %}</span><strong>{{ts|length}}</strong></div><div><span>{% if lang=='ar' %}الفترات المتاحة{% elif lang=='fr' %}Périodes{% else %}Periods{% endif %}</span><strong>{{ps|length}}</strong></div><div><span>{% if lang=='ar' %}التعيينات في السنة{% elif lang=='fr' %}Affectations cette année{% else %}Assignments this year{% endif %}</span><strong>{{rows|length}}</strong></div></div></section><section class="rt-board-entry"><h3>{% if lang=='ar' %}جدول توزيع التربصات{% elif lang=='fr' %}Tableau général de répartition{% else %}Rotation distribution board{% endif %}</h3><p>{% if lang=='ar' %}عرض جميع الأراضي والفترات والمقيمين والمقاعد المتبقية.{% elif lang=='fr' %}Visualisez tous les terrains, les quatre périodes, les résidents affectés et les places restantes. Export PDF ou impression disponibles.{% else %}View all terrains, four periods, assigned residents and remaining places. PDF export and printing available.{% endif %}</p><a href="/rotation-board">{% if lang=='ar' %}فتح الجدول{% elif lang=='fr' %}Ouvrir le tableau{% else %}Open distribution board{% endif %}</a></section>{% if role=='super_admin' %}<section class="rt-assignment-form"><h3>{% if lang=='ar' %}تعيين جديد{% elif lang=='fr' %}Nouvelle affectation{% else %}Create an assignment{% endif %}</h3><p class="muted">{% if lang=='ar' %}الأولوية حسب السنة: الرابعة، ثم الثالثة، الثانية، فالأولى. يتم إظهار تعارضات التربص الخارجي باللون الأحمر.{% elif lang=='fr' %}Priorité : 4e année, puis 3e, 2e et 1re année. Les conflits de stages externes sont signalés en rouge.{% else %}Priority order: Year 4, then Year 3, Year 2, and Year 1. External rotation conflicts are marked in red.{% endif %}</p><label for="assignment-search">{% if lang=='ar' %}ابحث عن مقيم{% elif lang=='fr' %}Rechercher un résident{% else %}Find a resident{% endif %}</label><input id="assignment-search" type="search" placeholder="{% if lang=='ar' %}الاسم أو اللقب...{% elif lang=='fr' %}Nom ou prénom...{% else %}Name or surname...{% endif %}"><label for="assignment-year-filter">{% if lang=='ar' %}تصفية حسب السنة{% elif lang=='fr' %}Filtrer par année{% else %}Filter by year{% endif %}</label><select id="assignment-year-filter"><option value="all">{% if lang=='ar' %}كل السنوات{% elif lang=='fr' %}Toutes les années{% else %}All years{% endif %}</option>{% for y in [4,3,2,1] %}<option value="{{y}}">{% if lang=='ar' %}السنة{% elif lang=='fr' %}Année{% else %}Year{% endif %} {{y}}</option>{% endfor %}</select><form method="post" id="assignment-form"><label for="assignment-resident">{{t('resident')}}</label><select name="resident" id="assignment-resident" required>{% for r in rs %}<option value="{{r.id}}" data-year="{{r.year}}" data-blocked-periods="{{blocked_periods_by_resident[r.id]}}" data-base-label="{{r.year}} — {{r.surname}} {{r.firstname}}{% if r.resident_type=='external' %} — {{L('Externe','External','خارجي')}} · {{r.origin_hospital}}{% endif %}">{{r.year}} — {{r.surname}} {{r.firstname}}{% if r.resident_type=='external' %} — {{L('Externe','External','خارجي')}} · {{r.origin_hospital}}{% endif %}</option>{% endfor %}</select><p id="assignment-conflict" class="rt-assignment-notice" role="status" aria-live="polite"></p><label for="assignment-terrain">{{t('terrain')}}</label><select name="terrain" id="assignment-terrain" required>{% for x in ts %}<option value="{{x.id}}" data-kind="{{x.kind}}">{{x.name}}{% if x.kind=='optional' %} — {% if lang=='ar' %}اختياري{% elif lang=='fr' %}Optionnel{% else %}Optional{% endif %}{% endif %}</option>{% endfor %}</select><label for="assignment-period">{{t('period')}}</label><select name="period" id="assignment-period" required>{% for x in ps %}<option value="{{x.id}}">{{period_label(x.name)}}</option>{% endfor %}</select><fieldset class="rt-admin-override" style="border:1px solid #c38a28;border-radius:10px;padding:10px;margin:12px 0"><legend>{{L('Dérogation Super Admin','Super Admin override','استثناء المسؤول العام')}}</legend><label><input type="checkbox" name="override" value="1"> {{L('Forcer ou remplacer une affectation malgré un conflit','Force or replace an assignment despite a conflict','فرض تعيين أو استبداله رغم وجود تعارض')}}</label><label for="override-reason">{{L('Motif obligatoire','Required reason','السبب إلزامي')}}</label><textarea id="override-reason" name="override_reason" rows="2" placeholder="{{L('Expliquez le motif de la dérogation','Explain the reason for the override','اشرح سبب الاستثناء')}}"></textarea></fieldset><button type="submit">{{t('check_save')}}</button></form><script>(function(){const p=document.getElementById("assignment-period"),r=document.getElementById("assignment-resident"),search=document.getElementById("assignment-search"),year=document.getElementById("assignment-year-filter"),notice=document.getElementById("assignment-conflict");if(!p||!r)return;const options=Array.from(r.options);function update(){const pid=p.value,term=(search.value||"").toLocaleLowerCase(),yr=year.value;let visible=0;options.forEach(o=>{const blocked=(o.dataset.blockedPeriods||"").split(",").includes(pid),matchesName=(o.dataset.baseLabel||"").toLocaleLowerCase().includes(term),matchesYear=yr==="all"||o.dataset.year===yr;o.disabled=blocked||!matchesName||!matchesYear;o.hidden=!matchesName||!matchesYear;o.textContent=o.dataset.baseLabel+(blocked?" — ⚠ {% if lang=='ar' %}تعارض تربص خارجي{% elif lang=='fr' %}Conflit de stage externe{% else %}External rotation conflict{% endif %}":"");if(!o.disabled)visible++;});if(r.selectedOptions.length&&r.selectedOptions[0].disabled){const next=options.find(o=>!o.disabled);if(next)r.value=next.value;}const selected=r.selectedOptions[0],blocked=selected&&(selected.dataset.blockedPeriods||"").split(",").includes(pid);if(blocked){notice.className="rt-assignment-notice is-conflict";notice.textContent="{% if lang=='ar' %}تعارض: هذا المقيم لديه تربص خارجي في هذه الفترة.{% elif lang=='fr' %}Conflit : ce résident a un stage externe pendant cette période.{% else %}Conflict: this resident has an external rotation during this period.{% endif %}";}else{notice.className="rt-assignment-notice";notice.textContent=visible===0?"{% if lang=='ar' %}لا يوجد مقيم مطابق للفلاتر في هذه الفترة.{% elif lang=='fr' %}Aucun résident ne correspond aux filtres pour cette période.{% else %}No residents match these filters for this period.{% endif %}":"";}}p.addEventListener("change",update);search.addEventListener("input",update);year.addEventListener("change",update);update();})();</script></section>{% else %}<section><p>{% if lang=='ar' %}إنشاء التعيينات متاح للمسؤول العام فقط.{% elif lang=='fr' %}La création des affectations est réservée aux Super Admins.{% else %}Only Super Admins can create assignments.{% endif %}</p></section>{% endif %}<section><h3>{% if lang=='ar' %}التعيينات الحالية{% elif lang=='fr' %}Affectations actuelles{% else %}Current assignments{% endif %}</h3><div class="wrap"><table class="rt-assignment-table"><thead><tr><th>{{t('period')}}</th><th>{{t('resident')}}</th><th>{{t('study_year')}}</th><th>{{t('terrain')}}</th><th>{{t('status')}}</th><th>{% if lang=='ar' %}إجراء إداري{% elif lang=='fr' %}Action admin{% else %}Admin action{% endif %}</th></tr></thead><tbody>{% for a in rows %}<tr class="{% if a.status=='pending' %}is-pending{% elif a.status=='rejected' or a.status=='cancelled' %}is-inactive{% else %}is-confirmed{% endif %}"><td>{{period_label(a.period)}}</td><td>{{a.surname}} {{a.firstname}}</td><td><span class="rt-year-badge">{{a.year}}</span></td><td>{{a.terrain}}</td><td><span class="rt-status-pill status-{{a.status}}">{{t(a.status) if a.status in ['confirmed','pending','rejected'] else a.status}}</span></td><td>{% if role=='super_admin' and a.status not in ['cancelled','rejected'] %}<form method="post" action="/assignments/{{a.id}}/cancel"><button class="rt-danger-button" type="submit">{% if lang=='ar' %}إلغاء{% elif lang=='fr' %}Annuler{% else %}Cancel{% endif %}</button></form>{% else %}—{% endif %}</td></tr>{% else %}<tr><td colspan="6">{% if lang=='ar' %}لا توجد تعيينات بعد.{% elif lang=='fr' %}Aucune affectation pour le moment.{% else %}No assignments yet.{% endif %}</td></tr>{% endfor %}</tbody></table></div></section><style>.rt-assignment-hero h2{margin-top:0}.rt-assignment-stats{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:10px;margin-top:16px}.rt-assignment-stats>div{padding:12px;border:1px solid #155173;border-radius:13px;background:#071e33}.rt-assignment-stats span{display:block;font-size:12px;color:#a9ccdf}.rt-assignment-stats strong{display:block;font-size:25px;color:#27d8ce}.rt-assignment-notice{min-height:0;margin:0;padding:0;color:#a9ccdf}.rt-assignment-notice.is-conflict{margin:6px 0;padding:10px;border:1px solid #ff6d77;border-radius:10px;background:#531c2a;color:#ffe5e7;font-weight:700}.rt-assignment-table .is-pending td{border-top:2px solid #ffbd61}.rt-assignment-table .is-inactive td{opacity:.7}.rt-year-badge{display:inline-block;min-width:30px;padding:4px 8px;border-radius:8px;background:#0b4c69;color:#8ff5f4;text-align:center;font-weight:700}.rt-status-pill{display:inline-block;padding:4px 8px;border-radius:20px;font-size:12px;font-weight:700}.status-confirmed{background:#064c3e;color:#a7f3d0}.status-pending{background:#624013;color:#ffe0a3}.status-rejected,.status-cancelled{background:#5b1f2b;color:#ffc1c8}.rt-danger-button{background:#762536!important;border-color:#d64f67!important}@media(max-width:680px){.rt-assignment-stats{grid-template-columns:repeat(2,minmax(0,1fr))}.rt-assignment-stats strong{font-size:21px}} </style>""",rs=rs,ts=ts,ps=ps,rows=rows,blocked_periods_by_resident=blocked_periods_by_resident)
@app.post("/assignments/<int:aid>/cancel")
def cancel_assignment(aid):
 if super_guard(): return super_guard()
 c=db()
 row=c.execute("SELECT status FROM assignments WHERE id=?",(aid,)).fetchone()
 if not row:
  c.close(); return ("Assignment not found",404)
 if row["status"]=="cancelled":
  c.close(); flash("Assignment already cancelled."); return redirect("/assign")
 c.execute("UPDATE assignments SET status='cancelled' WHERE id=?",(aid,))
 audit_event(c,"assignment.cancelled",f"assignment_id={aid}; previous_status={row['status']}")
 c.commit(); c.close(); flash("Assignment cancelled; historical record retained.")
 return redirect("/assign")
@app.route("/history",methods=["GET","POST"])
def history():
 if guard():return redirect("/login")
 c=db()
 if request.method=="POST":
  if session.get("role")!="super_admin": c.close(); return ("Forbidden — Super Admin only",403)
  if request.form.get("action")=="period_dates":
   try:
    pid=int(request.form.get("period_id","0"))
    period_row=c.execute("SELECT id,name FROM periods WHERE id=? AND year_label=?",(pid,active_postgrad_year(c))).fetchone()
    if not period_row: raise ValueError("Période introuvable pour l’année sélectionnée.")
    start_value=request.form.get("start_date","").strip(); end_value=request.form.get("end_date","").strip()
    if start_value: start_value=date.fromisoformat(start_value).isoformat()
    if end_value: end_value=date.fromisoformat(end_value).isoformat()
    if start_value and end_value and end_value<start_value: raise ValueError("La date de fin doit être postérieure ou égale à la date de début.")
    c.execute("UPDATE periods SET start_date=?,end_date=? WHERE id=? AND year_label=?",(start_value,end_value,pid,active_postgrad_year(c)))
    audit_event(c,"period.dates_updated",f"year={active_postgrad_year(c)}; period_id={pid}; start={start_value}; end={end_value}; source=history")
    c.commit(); c.close(); flash(L("Dates de période mises à jour.","Period dates updated.","تم تحديث تواريخ الفترة.")); return redirect("/history")
   except (ValueError,sqlite3.Error) as exc:
    c.rollback(); flash(str(exc)); c.close(); return redirect("/history")
  try:
   rid=int(request.form["resident"]); tid=int(request.form["terrain"])
   year_label=request.form.get("year_label","").strip()
   period=request.form.get("period","").strip()
   date_done=request.form.get("date_done","").strip()
   if not re.fullmatch(r"\d{4}-\d{4}",year_label) or int(year_label.split("-")[1])!=int(year_label.split("-")[0])+1:
    raise ValueError(L("L’année de post-graduation doit suivre le format AAAA-AAAA avec deux années consécutives.","Postgraduate year must be YYYY-YYYY with consecutive years.","يجب أن تكون سنة ما بعد التدرج بالتنسيق YYYY-YYYY مع سنتين متتاليتين."))
   valid_periods={"01 Fév - 30 Avr","01 Mai - 30 Juillet","01 Août - 30 Oct","01 Nov - 30 Jan"}
   if period and period not in valid_periods:
    raise ValueError(L("Choisissez l’une des quatre périodes officielles.","Choose one of the four official rotation periods.","اختر إحدى الفترات الرسمية الأربع."))
   if date_done:
    parsed_date=date.fromisoformat(date_done)
    if parsed_date>date.today():
     raise ValueError(L("Une date de stage historique ne peut pas être dans le futur.","A historical rotation date cannot be in the future.","لا يمكن أن يكون تاريخ التربص السابق في المستقبل."))
   if not c.execute("SELECT 1 FROM residents WHERE id=?",(rid,)).fetchone():
    raise ValueError(L("Ce résident n’existe pas.","Resident does not exist.","هذا المقيم غير موجود."))
   if not c.execute("SELECT 1 FROM terrains WHERE id=?",(tid,)).fetchone():
    raise ValueError(L("Ce service de stage n’existe pas.","Rotation department does not exist.","مصلحة التربص هذه غير موجودة."))
   duplicate=c.execute("SELECT 1 FROM history WHERE resident_id=? AND terrain_id=? AND year_label=? AND period=? AND COALESCE(date_done,'')=?",(rid,tid,year_label,period,date_done)).fetchone()
   if duplicate: raise ValueError(L("Ce stage historique existe déjà.","This historical rotation already exists.","هذا التربص السابق موجود مسبقًا."))
   c.execute("INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",(rid,tid,year_label,period,date_done))
   audit_event(c,"history.created",f"resident_id={rid}; terrain_id={tid}; year={year_label}; period={period}")
   c.commit(); flash(tr("done"))
  except (TypeError,ValueError,sqlite3.IntegrityError) as exc:
   c.rollback(); flash(str(exc) or L("Stage historique invalide.","Invalid historical rotation.","التربص السابق غير صالح."))
 selected_year=active_postgrad_year(c); year_options=available_postgrad_years(c); current_year=setting(c,"year")
 # Historical resident choices: include anyone who had joined by the end of the
 # selected postgraduate year and had not yet been archived before it began.
 year_start=f"{selected_year.split('-')[0]}-02-01"; year_end=f"{selected_year.split('-')[1]}-01-31"
 rs=c.execute("SELECT * FROM residents WHERE (registration_date IS NULL OR registration_date='' OR registration_date<=?) AND (archived_at IS NULL OR archived_at='' OR substr(archived_at,1,10)>=?) ORDER BY year DESC,surname",(year_end,year_start)).fetchall()
 ts=c.execute("SELECT * FROM terrains WHERE active=1 ORDER BY CASE WHEN kind='optional' THEN 1 ELSE 0 END, name").fetchall()
 period_rows=c.execute("SELECT * FROM periods WHERE year_label=? ORDER BY id",(selected_year,)).fetchall()
 rows=c.execute("SELECT h.*,r.surname,r.firstname,t.name terrain FROM history h JOIN residents r ON r.id=h.resident_id JOIN terrains t ON t.id=h.terrain_id WHERE h.year_label=? ORDER BY h.id DESC",(selected_year,)).fetchall(); c.close()
 return page("""<section><div class="rt-history-yearbar"><h3>{{L('Historique de l’année de post-graduation','History for postgraduate year','سجل سنة ما بعد التدرج')}}</h3><form method="post" action="/select-year" class="rt-year-inline"><label for="history-postgrad-year">{{t('year')}}</label><select id="history-postgrad-year" name="year" onchange="this.form.submit()">{% for yy in year_options %}<option value="{{yy}}" {% if yy==year %}selected{% endif %}>{{yy}}{% if yy==current_year %} — {{L('en cours','current','الحالية')}}{% endif %}</option>{% endfor %}</select><noscript><button type="submit">{{t('choose')}}</button></noscript></form></div><p>{{L('Les lignes affichées ci-dessous correspondent uniquement à l’année sélectionnée. Choisissez une année précédente pour saisir son historique.','The rows below show only the selected year. Choose a previous year to enter its history.','السطور أدناه تخص السنة المختارة فقط. اختر سنة سابقة لإدخال سجلها.')}}</p><h3>{{L('Dates des périodes de cette année','Period dates for this year','تواريخ فترات هذه السنة')}}</h3><p>{{L('Ces dates sont partagées avec le tableau des affectations pour la même année.','These dates are shared with the assignment board for the same year.','هذه التواريخ مشتركة مع جدول التعيينات للسنة نفسها.')}}</p>{% for p in period_rows %}<form method="post" class="grid" style="grid-template-columns:repeat(auto-fit,minmax(150px,1fr));align-items:end"><input type="hidden" name="action" value="period_dates"><input type="hidden" name="period_id" value="{{p.id}}"><strong>{{period_label(p.name)}}</strong><label>{{L('Début','Start','البداية')}}<input type="date" name="start_date" value="{{p.start_date}}" required></label><label>{{L('Fin','End','النهاية')}}<input type="date" name="end_date" value="{{p.end_date}}" required></label><button>{{L('Enregistrer les dates','Save dates','حفظ التواريخ')}}</button></form>{% endfor %}<h3>{{t('history_add')}}</h3><form method="post"><label>{{t('resident')}}</label><select name="resident">{% for r in rs %}<option value="{{r.id}}">{{r.surname}} {{r.firstname}} — {{r.year}}</option>{% endfor %}</select><label>{{t('terrain')}}</label><select name="terrain">{% for x in ts %}<option value="{{x.id}}">{{x.name}}</option>{% endfor %}</select><label>{{t('academic_year')}}</label><input name="year_label" required value="{{year}}" placeholder="2025-2026"><label>{{t('period')}}</label><select name="period"><option value="">—</option>{% for p in ["01 Fév - 30 Avr","01 Mai - 30 Juillet","01 Août - 30 Oct","01 Nov - 30 Jan"] %}<option value="{{p}}">{{period_label(p)}}</option>{% endfor %}</select><label>{{t('date')}}</label><input type="date" name="date_done"><button>{{t('save')}}</button></form></section><section><div class="wrap"><table><tr><th>{{t('resident')}}</th><th>{{t('terrain')}}</th><th>{{t('academic_year')}}</th><th>{{t('period')}}</th><th>{{t('date')}}</th></tr>{% for h in rows %}<tr><td>{{h.surname}} {{h.firstname}}</td><td>{{h.terrain}}</td><td>{{h.year_label}}</td><td>{{period_label(h.period)}}</td><td>{{h.date_done}}</td></tr>{% endfor %}</table></div></section>""",rs=rs,ts=ts,rows=rows,period_rows=period_rows)
@app.get("/summary")
def summary():
 if guard():return redirect("/login")
 c=db(); y=active_postgrad_year(c); year_start=date(int(y.split("-")[0]),2,1); year_end=date(int(y.split("-")[1]),1,31); rs=c.execute("SELECT * FROM residents ORDER BY year DESC,surname").fetchall(); rows=[]
 for r in rs:
  a=c.execute("SELECT p.name period,t.name terrain,a.status FROM assignments a JOIN periods p ON p.id=a.period_id JOIN terrains t ON t.id=a.terrain_id WHERE a.resident_id=? AND p.year_label=? ORDER BY p.id",(r["id"],y)).fetchall()
  h=c.execute("SELECT t.name terrain,h.year_label,h.period,h.date_done FROM history h JOIN terrains t ON t.id=h.terrain_id WHERE h.resident_id=? AND h.year_label=?",(r["id"],y)).fetchall()
  x=c.execute("SELECT hospital,country,rotation_type,start_date,end_date,status FROM external_rotations WHERE resident_id=? AND status<>'cancelled' AND start_date<=? AND end_date>=? ORDER BY start_date DESC",(r["id"],year_end.isoformat(),year_start.isoformat())).fetchall()
  rows.append((r,a,h,x))
 c.close()
 return page("""<section><h3>{{t('summary')}} {{'—'}} {{year}}</h3><div class="wrap"><table><tr><th>{{t('resident')}}</th><th>{{t('study_year')}}</th><th>{{t('assign')}}</th><th>{{t('history')}}</th><th>{{L('Stage à l’étranger','External rotation','تربص خارجي')}}</th></tr>{% for r,a,h,x in rows %}<tr><td>{{r.surname}} {{r.firstname}}</td><td>{{r.year}}</td><td>{% for item in a %}<div>{{period_label(item.period)}} — {{item.terrain}} ({{item.status}})</div>{% else %}—{% endfor %}</td><td>{% for item in h %}<div>{{item.terrain}} — {{item.year_label}} / {{period_label(item.period)}} / {{item.date_done}}</div>{% else %}—{% endfor %}</td><td>{% for item in x %}<div><strong>{{item.hospital}}, {{item.country}}</strong><br>{{item.rotation_type}} · {{item.start_date}} {{'→'}} {{item.end_date}} · {{item.status}}</div>{% else %}—{% endfor %}</td></tr>{% endfor %}</table></div></section>""",rows=rows,year=y)
@app.route("/decisions",methods=["GET","POST"])
def decisions():
 if super_guard():return super_guard()
 c=db(); selected_year=active_postgrad_year(c)
 if request.method=="POST":
  aid=request.form.get("aid",type=int); outcome=request.form.get("outcome",""); reason=request.form.get("reason","").strip()
  if outcome not in ("approve","reject"): c.close(); return ("Invalid decision",400)
  a=c.execute("SELECT a.*,r.surname,r.firstname,r.status resident_status,t.name terrain,t.active terrain_active,p.name period,p.year_label,p.start_date,p.end_date FROM assignments a JOIN residents r ON r.id=a.resident_id JOIN terrains t ON t.id=a.terrain_id JOIN periods p ON p.id=a.period_id WHERE a.id=? AND a.status='pending' AND p.year_label=?",(aid,selected_year)).fetchone()
  if a and reason:
   if outcome=="approve":
    if a["resident_status"]!="active":
     flash("Cannot approve: the resident is no longer active.")
    elif not a["terrain_active"]:
     flash("Cannot approve: the terrain has been deactivated.")
    else:
     start_day,end_day=rotation_period_bounds(a["year_label"],a["period"],a["start_date"],a["end_date"])
     external=c.execute("SELECT hospital,country,start_date,end_date FROM external_rotations WHERE resident_id=? AND request_status='approved' AND status<>'cancelled' AND start_date<=? AND end_date>=? LIMIT 1",(a["resident_id"],end_day.isoformat(),start_day.isoformat())).fetchone()
     if external:
      flash("Cannot approve: an external rotation overlaps this period.")
     else:
      cap=c.execute("SELECT places FROM capacities WHERE year_label=? AND terrain_id=? AND period_id=?",(a["year_label"],a["terrain_id"],a["period_id"])).fetchone()
      used=c.execute("SELECT COUNT(*) FROM assignments WHERE terrain_id=? AND period_id=? AND status='confirmed'",(a["terrain_id"],a["period_id"])).fetchone()[0]
      dup=c.execute("SELECT id FROM assignments WHERE resident_id=? AND period_id=? AND id<>? AND status IN ('confirmed','pending')",(a["resident_id"],a["period_id"],aid)).fetchone()
      if not cap or used>=cap["places"] or dup: flash(tr("full") if not dup else tr("duplicate"))
      else:
       c.execute("UPDATE assignments SET status='confirmed' WHERE id=?",(aid,)); c.execute("INSERT INTO decisions(details,decision,reason,actor,postgrad_year) VALUES(?,?,?,?,?)",(a["surname"]+" "+a["firstname"]+" — "+a["terrain"],tr("approve"),reason,session["user"],a["year_label"]))
       audit_event(c,"assignment.override_approved",f"assignment_id={aid}; reason={reason}")
       c.commit(); flash(tr("done"))
   else:
    c.execute("UPDATE assignments SET status='rejected' WHERE id=?",(aid,)); c.execute("INSERT INTO decisions(details,decision,reason,actor,postgrad_year) VALUES(?,?,?,?,?)",(a["surname"]+" "+a["firstname"]+" — "+a["terrain"],tr("reject"),reason,session["user"],a["year_label"]))
    audit_event(c,"assignment.override_rejected",f"assignment_id={aid}; reason={reason}")
    c.commit(); flash(tr("done"))
 pending=c.execute("SELECT a.id,r.surname,r.firstname,r.year,t.name terrain,p.name period FROM assignments a JOIN residents r ON r.id=a.resident_id JOIN terrains t ON t.id=a.terrain_id JOIN periods p ON p.id=a.period_id WHERE a.status='pending' AND p.year_label=?",(selected_year,)).fetchall(); ds=c.execute("SELECT * FROM decisions WHERE postgrad_year=? ORDER BY id DESC",(selected_year,)).fetchall(); legacy_year_start=f"{selected_year.split('-')[0]}-02-01"; legacy_year_end=f"{selected_year.split('-')[1]}-01-31"; legacy_ds=c.execute("SELECT * FROM decisions WHERE COALESCE(postgrad_year,'')='' AND date(created_at)>=? AND date(created_at)<=? ORDER BY id DESC",(legacy_year_start,legacy_year_end)).fetchall(); c.close()
 return page("""<section><h3>{{t('decisions')}} — {{year}}</h3>{% for a in pending %}<p><b>{{a.surname}} {{a.firstname}}</b> — {{a.terrain}} / {{period_label(a.period)}}</p><form method="post"><input type="hidden" name="aid" value="{{a.id}}"><input name="reason" placeholder="{{t('reason')}}" required><button name="outcome" value="approve">{{t('approve')}}</button> <button name="outcome" value="reject">{{t('reject')}}</button></form>{% else %}<p>{{t('none')}}</p>{% endfor %}</section><section><div class="wrap"><table><tr><th>Date</th><th>{{t('assign')}}</th><th>{{t('status')}}</th><th>{{t('reason')}}</th></tr>{% for d in ds %}<tr><td>{{d.created_at}}</td><td>{{d.details}}</td><td>{{d.decision}}</td><td>{{d.reason}}</td></tr>{% endfor %}</table></div></section>{% if legacy_ds %}<section><h4>Décisions historiques sans année de post-graduation associée</h4><div class="wrap"><table><tr><th>Date</th><th>Décision</th><th>Résultat</th><th>Motif</th></tr>{% for d in legacy_ds %}<tr><td>{{d.created_at}}</td><td>{{d.details}}</td><td>{{d.decision}}</td><td>{{d.reason}}</td></tr>{% endfor %}</table></div></section>{% endif %}""",pending=pending,ds=ds,legacy_ds=legacy_ds,year=selected_year)

@app.route("/settings",methods=["GET","POST"])
def settings():
 if super_guard():return super_guard()
 c=db()
 if request.method=="POST":
  lang=request.form.get("lang","ar")
  if lang in LANGS:session["lang"]=lang
  y=request.form.get("year","").strip()
  if y and (not re.fullmatch(r"\d{4}-\d{4}",y) or int(y.split("-")[1])!=int(y.split("-")[0])+1):
   flash("Academic year must be formatted as YYYY-YYYY with consecutive years.")
   y=setting(c,"year")
  if y:
   c.execute("INSERT INTO settings(key,value) VALUES('year',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(y,))
   for n in ["01 Fév - 30 Avr","01 Mai - 30 Juillet","01 Août - 30 Oct","01 Nov - 30 Jan"]:
    c.execute("INSERT OR IGNORE INTO periods(year_label,name) VALUES(?,?)",(y,n))
  d=request.form.get("director","").strip()
  if d:c.execute("INSERT INTO settings(key,value) VALUES('director',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(d,))
  pw=request.form.get("new_password","")
  if pw and len(pw)<12:
   flash("New password must be at least 12 characters.")
   pw=""
  if pw:
   secure_hash=hash_password(pw)
   c.execute("INSERT INTO settings(key,value) VALUES('password',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(secure_hash,))
   c.execute("UPDATE users SET password_hash=? WHERE username=?",(secure_hash,session.get("user","admin")))
   audit_event(c,"security.password_changed","Administrator password changed")
  audit_event(c,"settings.updated",f"year={y}; director_updated={bool(d)}")
  c.commit(); flash(tr("done"))
 vals={r["key"]:r["value"] for r in c.execute("SELECT key,value FROM settings")}; c.close()
 return page("""<section><h3>{{t('settings')}}</h3><form method="post"><label>{{t('language')}}</label><select name="lang"><option value="ar" {% if lang=='ar' %}selected{% endif %}>العربية</option><option value="fr" {% if lang=='fr' %}selected{% endif %}>Français</option><option value="en" {% if lang=='en' %}selected{% endif %}>English</option></select><label>{{t('year')}}</label><input name="year" value="{{vals.get('year','')}}" required><label>{{t('director')}}</label><input name="director" value="{{vals.get('director','')}}"><label>{{t('new_password')}}</label><div class="password-wrap"><input id="settings-new-password" type="password" name="new_password" minlength="12" autocomplete="new-password"><button class="password-toggle" type="button" aria-label="Show password" onclick="const p=document.getElementById('settings-new-password');p.type=p.type==='password'?'text':'password';this.setAttribute('aria-label',p.type==='password'?'Show password':'Hide password')">◉</button></div><button>{{t('change')}}</button></form><section aria-labelledby="interior-theme-title"><h3 id="interior-theme-title">{{L('Thèmes de l’interface','Interface themes','سمات الواجهة')}}</h3><p class="muted">{{L('Choisissez un thème 3D élégant. Il est enregistré sur cet appareil et appliqué immédiatement.','Choose an elegant 3D theme. Your choice is saved on this device and applied immediately.','اختر مظهرًا ثلاثي الأبعاد أنيقًا. سيتم حفظ اختيارك على هذا الجهاز وتطبيقه فورًا.')}}</p><div class="theme-chooser"><button type="button" class="theme-choice" data-theme-choice="futuristic" aria-pressed="false" onclick="setInteriorTheme('futuristic')"><span class="theme-swatch" style="background:#041426"><i style="background:#087f9f"></i><i style="background:#19dfbd"></i><i style="background:#19d9e8"></i></span><strong>{{L('Futuriste médical','Medical futuristic','طبي مستقبلي')}}</strong><small>{{L('Bleu nuit · cyan néon · verre 3D','Deep navy · neon cyan · 3D glass','أزرق داكن · سماوي نيون · زجاج ثلاثي الأبعاد')}}</small></button><button type="button" class="theme-choice" data-theme-choice="emerald" aria-pressed="false" onclick="setInteriorTheme('emerald')"><span class="theme-swatch" style="background:#061c1b"><i style="background:#087c6b"></i><i style="background:#48e0bd"></i><i style="background:#14746a"></i></span><strong>{{L('Émeraude futuriste','Futuristic emerald','زمردي مستقبلي')}}</strong><small>{{L('Vert profond · reflets lumineux · contraste','Deep green · luminous reflections · contrast','أخضر داكن · انعكاسات مضيئة · تباين')}}</small></button><button type="button" class="theme-choice" data-theme-choice="holographic" aria-pressed="false" onclick="setInteriorTheme('holographic')"><span class="theme-swatch" style="background:#160d35"><i style="background:#8a5cf6"></i><i style="background:#e879f9"></i><i style="background:#38bdf8"></i></span><strong>{{L('Holographique violet','Violet holographic','هولوغرافي بنفسجي')}}</strong><small>{{L('Violet 3D · halos néon','3D violet · neon halos','بنفسجي ثلاثي الأبعاد · هالات نيون')}}</small></button><button type="button" class="theme-choice" data-theme-choice="titanium" aria-pressed="false" onclick="setInteriorTheme('titanium')"><span class="theme-swatch" style="background:#171d29"><i style="background:#64748b"></i><i style="background:#cbd5e1"></i><i style="background:#38bdf8"></i></span><strong>{{L('Titane premium','Premium titanium','تيتانيوم فاخر')}}</strong><small>{{L('Noir graphite · argent métallique','Graphite black · metallic silver','أسود غرافيتي · فضي معدني')}}</small></button><button type="button" class="theme-choice" data-theme-choice="aurora" aria-pressed="false" onclick="setInteriorTheme('aurora')"><span class="theme-swatch" style="background:#042b36"><i style="background:#06b6d4"></i><i style="background:#2dd4bf"></i><i style="background:#818cf8"></i></span><strong>{{L('Aurore turquoise','Turquoise aurora','شفق فيروزي')}}</strong><small>{{L('Reflets lumineux · profondeur','Luminous reflections · depth','انعكاسات مضيئة · عمق')}}</small></button><button type="button" class="theme-choice" data-theme-choice="ruby" aria-pressed="false" onclick="setInteriorTheme('ruby')"><span class="theme-swatch" style="background:#300d24"><i style="background:#9f1239"></i><i style="background:#fb7185"></i><i style="background:#f9a8d4"></i></span><strong>{{L('Rubis et or rose','Ruby and rose gold','ياقوتي وذهبي وردي')}}</strong><small>{{L('Bordeaux · rose doré','Burgundy · rose gold','عنابي · ذهبي وردي')}}</small></button></div></section><p class="muted">{{t('note')}}</p><p><a href="/sync-devices">{{L('Gérer les appareils de synchronisation','Manage trusted sync devices','إدارة أجهزة المزامنة الموثوقة')}}</a></p><p><a href="/sync">{{L('Configurer cet appareil Android','Configure this Android device','إعداد هذا الجهاز أندرويد')}}</a></p></section>""",vals=vals,lang=session.get("lang","fr"))
@app.route("/users",methods=["GET","POST"])
def manage_users():
 if super_guard(): return super_guard()
 c=db()
 if request.method=="POST":
  action=request.form.get("action","create")
  if action in ("deactivate","activate","change_role","reset_password"):
   uid=request.form.get("user_id",type=int)
   target=c.execute("SELECT * FROM users WHERE id=?",(uid,)).fetchone()
   if not target:
    flash("User not found.")
   elif action=="deactivate" and target["username"].lower() in ("benhocine","drdris"):
    flash("Official accounts cannot be deactivated.")
   elif action=="deactivate" and target["username"]==session.get("user"):
    flash("You cannot deactivate your own account while signed in.")
   elif action=="deactivate" and target["role"]=="super_admin" and target["active"] and c.execute("SELECT COUNT(*) FROM users WHERE role='super_admin' AND active=1").fetchone()[0]<=1:
    flash("At least one active Super Admin must remain.")
   elif action=="change_role":
    new_role=request.form.get("new_role","")
    if new_role not in ("super_admin","user"):
     flash("Invalid role.")
    elif target["username"].lower() in ("benhocine","drdris") and new_role!="super_admin":
     flash("Official accounts must be Super Admins.")
    elif target["role"]=="super_admin" and target["active"] and new_role!="super_admin" and c.execute("SELECT COUNT(*) FROM users WHERE role='super_admin' AND active=1").fetchone()[0]<=1:
     flash("At least one active Super Admin must remain.")
    else:
     c.execute("UPDATE users SET role=? WHERE id=?",(new_role,uid))
     audit_event(c,"user.role_changed",f"username={target['username']}; from={target['role']}; to={new_role}")
     c.commit(); flash(tr("done"))
   elif action=="reset_password":
    new_password=request.form.get("new_password","")
    if len(new_password)<12:
     flash("New password must be at least 12 characters.")
    else:
     c.execute("UPDATE users SET password_hash=? WHERE id=?",(hash_password(new_password),uid))
     audit_event(c,"user.password_reset",f"username={target['username']}; reset_by={session.get('user','')}")
     c.commit(); flash("Password reset. Share the new password securely with the account owner.")
   else:
    active=0 if action=="deactivate" else 1
    c.execute("UPDATE users SET active=? WHERE id=?",(active,uid))
    audit_event(c,"user.deactivated" if not active else "user.activated",f"username={target['username']}")
    c.commit(); flash(tr("done"))
  else:
   username=request.form.get("username","").strip().lower()
   password=request.form.get("password","")
   role=request.form.get("role","user")
   if username in ("benhocine","drdris") and role!="super_admin":
    flash("Official accounts must be Super Admins.")
   elif len(username)<3 or len(password)<12 or role not in ("super_admin","user"):
    flash("Username must have 3+ characters and password 12+ characters.")
   else:
    try:
     c.execute("INSERT INTO users(username,password_hash,role) VALUES(?,?,?)",(username,hash_password(password),role))
     audit_event(c,"user.created",f"username={username}; role={role}")
     c.commit(); flash(tr("done"))
    except sqlite3.IntegrityError:
     flash(L("Ce nom d’utilisateur existe déjà.","Username already exists.","اسم المستخدم موجود مسبقًا."))
 rows=c.execute("SELECT id,username,role,active,created_at FROM users ORDER BY role DESC,username").fetchall(); c.close()
 return page("""<section><h3>{{t('users')}}</h3><p>{{t('users_help')}}</p><form method="post"><input type="hidden" name="action" value="create"><label>{{t('username')}}</label><input name="username" minlength="3" required><label>{{t('password')}} (12+)</label><div class="password-wrap"><input id="new-user-password" name="password" type="password" minlength="12" required><button class="password-toggle" type="button" aria-label="Show password" onclick="const p=document.getElementById('new-user-password');p.type=p.type==='password'?'text':'password';this.setAttribute('aria-label',p.type==='password'?'Show password':'Hide password')">◉</button></div><label>{{t('role')}}</label><select name="role"><option value="user">{{t('standard_user')}}</option><option value="super_admin">{{t('super_admin')}}</option></select><button>{{t('add')}}</button></form></section><section><div class="wrap"><table><tr><th>{{t('username')}}</th><th>{{t('role')}}</th><th>{{t('active_col')}}</th><th>{{t('created')}}</th><th>{{L('Gestion du compte','Account controls','إدارة الحساب')}}</th></tr>{% for u in rows %}<tr><td>{{u.username}}</td><td>{{u.role}}</td><td>{{u.active}}</td><td>{{u.created_at}}</td><td><form method="post"><input type="hidden" name="user_id" value="{{u.id}}"><label>{{L('Rôle','Role','الصلاحية')}}</label><select name="new_role"><option value="user" {% if u.role=='user' %}selected{% endif %}>{{L('Utilisateur standard','Standard user','مستخدم عادي')}}</option><option value="super_admin" {% if u.role=='super_admin' %}selected{% endif %}>{{L('Super administrateur','Super Admin','مسؤول عام')}}</option></select><button name="action" value="change_role">{{L('Enregistrer le rôle','Save role','حفظ الصلاحية')}}</button></form><form method="post"><input type="hidden" name="user_id" value="{{u.id}}"><label>{{L('Nouveau mot de passe (12 caractères minimum)','New password (12+ characters)','كلمة مرور جديدة (12 حرفًا على الأقل)')}}</label><input name="new_password" type="password" minlength="12" required><button name="action" value="reset_password">{{L('Réinitialiser le mot de passe','Reset password','إعادة تعيين كلمة المرور')}}</button></form>{% if u.username!=session.get('user') or u.active==0 %}<form method="post"><input type="hidden" name="user_id" value="{{u.id}}"><button name="action" value="{{'deactivate' if u.active else 'activate'}}">{{L('Désactiver','Deactivate','تعطيل') if u.active else L('Activer','Activate','تفعيل')}}</button></form>{% else %}{{L('Compte actuel','Current account','الحساب الحالي')}}{% endif %}</td></tr>{% endfor %}</table></div></section>""",rows=rows)

@app.get("/sync")
def sync_page():
 if super_guard(): return super_guard()
 return page("""<section><h3>{{L('Synchronisation des appareils','Device sync','مزامنة الأجهزة')}}</h3><p><strong>{{L('Synchronisation entre plusieurs smartphones','Multi-smartphone sync','المزامنة بين عدة هواتف ذكية')}}</strong></p><p>{{L('Utilisez le même serveur central ResiManage sous Windows pour les deux téléphones Android. Un Super Admin enregistre chaque téléphone séparément et délivre un jeton différent à chacun. Saisissez la même adresse HTTPS du serveur sur les deux téléphones, avec leur jeton respectif. Gardez le serveur accessible uniquement via le VPN privé ; ne l’exposez jamais directement à Internet.','Use the same central Windows ResiManage server on both Android phones. A Super Admin enrolls each phone separately and issues a different token for each device. Enter the same HTTPS server address on both phones, using each device’s own token. Keep the server reachable only through the private VPN; never expose it directly to the public Internet.','استعمل خادم ResiManage المركزي نفسه على هاتفي أندرويد. يسجل المسؤول العام كل هاتف على حدة ويمنحه رمزًا مختلفًا. أدخل عنوان HTTPS نفسه للخادم على الهاتفين مع استعمال رمز خاص بكل هاتف. أبقِ الخادم متاحًا عبر شبكة VPN خاصة فقط ولا تعرضه مباشرة على الإنترنت.') }}</p><p>{{L('Procédure : 1) sauvegardez les données ; 2) enregistrez chaque téléphone dans « Appareils » ; 3) configurez la même adresse HTTPS avec un jeton différent sur chaque téléphone ; 4) lancez « Synchroniser maintenant » sur chacun. Les modifications du téléphone A sont récupérées par B lors de sa synchronisation avec le serveur central.','Steps: 1) back up data; 2) enroll each phone in “Devices”; 3) configure the same HTTPS server address with a different token on each phone; 4) tap “Synchronize now” on each phone. Changes from phone A are retrieved by B when it syncs with the central server.','الخطوات: 1) أنشئ نسخة احتياطية؛ 2) سجّل كل هاتف في «الأجهزة»؛ 3) اضبط عنوان HTTPS نفسه مع رمز مختلف لكل هاتف؛ 4) اضغط «مزامنة الآن» على كل هاتف. يحصل الهاتف B على تغييرات الهاتف A عند المزامنة مع الخادم المركزي.')}}</p><p class="muted">{{L('Important : le premier appairage exige une base locale vide pour les résidents et les stages afin d’éviter les conflits d’identifiants ; faites une sauvegarde avant. La synchronisation est manuelle, pas continue en arrière-plan. Les votes de promotion/DEMS, l’archive des diplômés, la résolution interactive des conflits et la file hors ligne complète ne sont pas encore terminés. N’utilisez pas cette synchronisation partielle avec des données réelles avant la fin des tests fonctionnels et de sécurité.','Important: first pairing requires an empty local resident/stage database to avoid ID collisions; back up first. Sync is manual, not continuous/background. Promotion/DEMS votes, graduate archive, interactive conflict resolution and full offline queueing are not complete. Do not use this partial sync with real resident data until end-to-end and security tests pass.','مهم: يتطلب الربط الأول قاعدة محلية فارغة لبيانات المقيمين والتربصات لتجنب تعارض المعرّفات؛ أنشئ نسخة احتياطية أولًا. المزامنة يدوية حاليًا وليست مستمرة في الخلفية. لم تكتمل بعد أصوات الترقية وDEMS وأرشيف المتخرجين وحل التعارضات التفاعلي وقائمة العمل دون اتصال. لا تستخدم هذه المزامنة الجزئية مع بيانات حقيقية قبل اجتياز اختبارات الوظائف والأمان.') }}</p><label>{{L('Adresse HTTPS du serveur central','Central server HTTPS address','عنوان HTTPS للخادم المركزي')}}</label><input id="sync-server-url" type="url" placeholder="https://private-vpn-ip:5000" autocomplete="url"><label>{{L('Jeton d’accès de l’appareil enregistré','Enrolled device bearer token','رمز وصول الجهاز المسجل')}}</label><input id="sync-device-token" type="password" autocomplete="off" placeholder="{{L('Collez le jeton fourni par le Super Admin central','Paste the one-time token from the central Super Admin','ألصق الرمز الذي منحه المسؤول العام المركزي')}}"><label>{{L('Certificat CA privé (PEM ; facultatif pour un serveur reconnu publiquement)','Private CA certificate (PEM; optional for a publicly trusted server)','شهادة CA خاصة (PEM؛ اختيارية لخادم موثوق علنًا)')}}</label><textarea id="sync-ca-certificate" rows="5" placeholder="-----BEGIN CERTIFICATE-----&#10;{{L('Collez le certificat CA ayant signé le certificat du serveur','Paste the CA certificate that signed the server certificate','ألصق شهادة CA التي وقعت شهادة الخادم')}}&#10;-----END CERTIFICATE-----"></textarea><small>{{L('Pour un certificat de serveur privé/VPN, collez le certificat de l’autorité qui l’a émis. Il est stocké chiffré avec le jeton de synchronisation.','For a private/VPN server certificate, paste its issuing CA certificate. The certificate is stored encrypted with the sync token.','لشهادة خادم خاص أو VPN، ألصق شهادة سلطة التصديق التي أصدرتها. تُخزّن الشهادة مشفرة مع رمز المزامنة.')}}</small><div style="display:flex;flex-wrap:wrap;gap:8px"><button type="button" onclick="saveSyncConfig()">{{L('Enregistrer en sécurité','Save securely','حفظ بأمان')}}</button><button type="button" onclick="runSyncNow()">{{L('Synchroniser maintenant','Synchronize now','مزامنة الآن')}}</button><button type="button" onclick="clearSyncConfig()">{{L('Oublier le jeton','Forget token','نسيان الرمز')}}</button></div><p id="sync-status" class="msg" aria-live="polite">{{L('Le module Android natif affichera ici son état.','Android native sync bridge will report its status here.','ستعرض وحدة المزامنة الأصلية في أندرويد حالتها هنا.')}}</p><pre id="sync-result" style="white-space:pre-wrap;overflow-wrap:anywhere"></pre></section><script>
(function(){const status=document.getElementById('sync-status');const result=document.getElementById('sync-result');const lg="{{lang}}";const msg=(fr,en,ar)=>lg==="fr"?fr:lg==="ar"?ar:en;if(!window.ResiManageSync){status.textContent=msg('Cette configuration est disponible uniquement dans l’application Android. L’enrôlement serveur est géré par un Super Admin.','This sync setup is available inside the Android app only. Windows/server enrollment is managed by a Super Admin.','إعداد المزامنة متاح داخل تطبيق أندرويد فقط. يتولى المسؤول العام تسجيل الأجهزة على الخادم.');return;}status.textContent=window.ResiManageSync.isConfigured()?msg('Une configuration sécurisée du serveur est enregistrée sur cet appareil.','A secure server configuration is saved on this device.','تم حفظ إعداد آمن للخادم على هذا الجهاز.'):msg('Aucun serveur de synchronisation n’est configuré.','No sync server configured yet.','لم يتم إعداد خادم مزامنة بعد.');})();
function saveSyncConfig(){const status=document.getElementById('sync-status');const token=document.getElementById('sync-device-token');try{const answer=JSON.parse(window.ResiManageSync.saveConfig(document.getElementById('sync-server-url').value,token.value,document.getElementById('sync-ca-certificate').value));status.textContent=answer.message||({fr:'Configuration enregistrée.',en:'Configuration saved.',ar:'تم حفظ الإعدادات.'})["{{lang}}"];if(answer.ok)token.value='';}catch(e){status.textContent=({fr:'Impossible d’enregistrer : ',en:'Could not save: ',ar:'تعذر الحفظ: '})["{{lang}}"]+e;}}
function runSyncNow(){const status=document.getElementById('sync-status');status.textContent=({fr:'Synchronisation en cours… gardez l’application ouverte et connectée au VPN.',en:'Sync in progress… keep the app open and connected to the VPN.',ar:'المزامنة جارية… أبقِ التطبيق مفتوحًا ومتصلًا بشبكة VPN.'})["{{lang}}"];const started=window.ResiManageSync.runSync();if(!started&& !window.ResiManageSync.isConfigured())status.textContent=({fr:'Enregistrez d’abord l’adresse du serveur et le jeton.',en:'Save the server URL and token first.',ar:'احفظ عنوان الخادم والرمز أولًا.'})["{{lang}}"];}
function clearSyncConfig(){window.ResiManageSync.clearConfig();document.getElementById('sync-device-token').value='';document.getElementById('sync-ca-certificate').value='';document.getElementById('sync-status').textContent=({fr:'Le jeton et la configuration CA ont été supprimés de cet appareil.',en:'Saved sync token and CA configuration removed from this device.',ar:'تم حذف رمز المزامنة وإعداد شهادة CA من هذا الجهاز.'})["{{lang}}"];document.getElementById('sync-result').textContent='';}
window.resimanageSyncFinished=function(raw){try{const r=JSON.parse(raw);document.getElementById('sync-status').textContent=r.message||(r.ok?({fr:'Synchronisation terminée.',en:'Sync completed.',ar:'اكتملت المزامنة.'})["{{lang}}"]:({fr:'Échec de la synchronisation.',en:'Sync failed.',ar:'فشلت المزامنة.'})["{{lang}}"]);document.getElementById('sync-result').textContent=JSON.stringify(r,null,2);}catch(e){document.getElementById('sync-status').textContent=({fr:'La synchronisation a renvoyé un résultat illisible : ',en:'Sync returned an unreadable result: ',ar:'أعادت المزامنة نتيجة غير مقروءة: '})["{{lang}}"]+e;}};
</script>""")
@app.get("/audit")
def audit_log():
 if super_guard(): return super_guard()
 c=db(); rows=c.execute("SELECT * FROM audit ORDER BY id DESC LIMIT 300").fetchall(); c.close()
 return page("""<section><h3>{{t('audit')}}</h3><div class="wrap"><table><tr><th>{{t('date')}}</th><th>{{t('actor')}}</th><th>{{t('action')}}</th><th>{{t('details')}}</th></tr>{% for r in rows %}<tr><td>{{r.created_at}}</td><td>{{r.actor}}</td><td>{{r.action}}</td><td>{{r.details}}</td></tr>{% else %}<tr><td colspan="4">{{t('none')}}</td></tr>{% endfor %}</table></div></section>""",rows=rows)


def render_sync_devices_page(rows,issued):
 return page("""<section><h3>{{L('Appareils de synchronisation','Sync devices','أجهزة المزامنة')}}</h3><p>{{L('N’enrôlez que des appareils de confiance. Le jeton donne accès aux données des résidents via l’API de synchronisation ; copiez-le une seule fois et conservez-le en lieu sûr. L’accès distant exige HTTPS.','Enroll only trusted devices. The token grants access to resident data over the sync API; copy it once and store it securely. Remote API access requires HTTPS.','لا تسجّل إلا الأجهزة الموثوقة. يمنح الرمز الوصول إلى بيانات المقيمين عبر واجهة المزامنة؛ انسخه مرة واحدة واحفظه بأمان. يتطلب الوصول البعيد HTTPS.')}}</p>{% if issued %}<div style="padding:12px;border:1px solid #16a085;border-radius:12px"><strong>{{L('Copiez ce jeton maintenant : il ne sera plus affiché.','Copy this token now — it will not be shown again:','انسخ هذا الرمز الآن، فلن يظهر مرة أخرى:')}}</strong><p><code>{{issued.token}}</code></p><small>Device ID: {{issued.device_id}}</small></div>{% endif %}<form method="post"><input type="hidden" name="_csrf_token" value="{{csrf_token}}"><input type="hidden" name="action" value="enroll"><label>{{L('Nom de l’appareil','Device label','اسم الجهاز')}}</label><input name="label" maxlength="100" required placeholder="{{L('ex. téléphone résident 01','e.g. Resident phone 01','مثال: هاتف المقيم 01')}}"><label>{{L('Accès','Access','الوصول')}}</label><select name="mode"><option value="read_only">{{L('Lecture seule','Read-only pull','قراءة فقط')}}</option><option value="read_write">{{L('Lire et envoyer les modifications de synchronisation autorisées','Read and submit approved sync changes','قراءة وإرسال تغييرات المزامنة المعتمدة')}}</option></select><button>{{L('Générer un jeton d’enrôlement','Generate enrollment token','إنشاء رمز تسجيل الجهاز')}}</button></form></section><section><h3>{{L('Appareils enregistrés','Enrolled devices','الأجهزة المسجلة')}}</h3><div class="wrap"><table><tr><th>{{L('Nom','Label','الاسم')}}</th><th>{{L('Identifiant de l’appareil','Device ID','معرّف الجهاز')}}</th><th>{{L('Actif','Active','نشط')}}</th><th>{{L('Mode','Mode','الوضع')}}</th><th>{{L('Créé le','Created','تاريخ الإنشاء')}}</th><th>{{L('Dernière activité','Last seen','آخر ظهور')}}</th><th>{{L('Action','Action','الإجراء')}}</th></tr>{% for d in rows %}<tr><td>{{d.label}}</td><td><code>{{d.device_id}}</code></td><td>{{L('Oui','Yes','نعم') if d.active else L('Révoqué','Revoked','ملغى')}}</td><td>{{d.mode}}</td><td>{{d.created_at}}</td><td>{{d.last_seen or L('Jamais','Never','أبدًا')}}</td><td>{% if d.active %}<form method="post"><input type="hidden" name="_csrf_token" value="{{csrf_token}}"><input type="hidden" name="action" value="revoke"><input type="hidden" name="device_id" value="{{d.device_id}}"><button>{{L('Révoquer','Revoke','إلغاء')}}</button></form>{% else %}—{% endif %}</td></tr>{% else %}<tr><td colspan="7">{{L('Aucun appareil enregistré.','No enrolled devices.','لا توجد أجهزة مسجلة.')}}</td></tr>{% endfor %}</table></div></section>""",rows=rows,issued=issued)

@app.route("/sync-devices",methods=["GET","POST"])
def manage_sync_devices():
 if super_guard(): return super_guard()
 c=db(); issued=None
 if request.method=="POST":
  action=request.form.get("action","")
  if action=="enroll":
   label=request.form.get("label","").strip()
   if not label or len(label)>100:
    flash("Device label is required (maximum 100 characters).")
   else:
    device_id=secrets.token_hex(12)
    token=secrets.token_urlsafe(32)
    token_hash=hashlib.sha256(token.encode("utf-8")).hexdigest()
    mode=request.form.get("mode","read_only")
    if mode not in ("read_only","read_write"):
     c.close(); return ("Invalid sync-device access mode",400)
    c.execute("INSERT INTO sync_devices(device_id,label,token_hash,mode) VALUES(?,?,?,?)",(device_id,label,token_hash,mode))
    audit_event(c,"sync_device.enrolled",f"device_id={device_id}; label={label}")
    c.commit()
    issued={"device_id":device_id,"label":label,"token":token}
    flash("Device enrolled. Copy its bearer token now; it is shown only in this response.")
  elif action=="revoke":
   device_id=request.form.get("device_id","").strip()
   row=c.execute("SELECT label,active FROM sync_devices WHERE device_id=?",(device_id,)).fetchone()
   if not row:
    flash("Sync device not found.")
   elif not row["active"]:
    flash("Device already revoked.")
   else:
    c.execute("UPDATE sync_devices SET active=0 WHERE device_id=?",(device_id,))
    audit_event(c,"sync_device.revoked",f"device_id={device_id}; label={row['label']}")
    c.commit(); flash("Device token revoked.")
  else:
   c.close(); return ("Invalid sync-device action",400)
 rows=c.execute("SELECT device_id,label,active,mode,created_at,last_seen FROM sync_devices ORDER BY created_at DESC").fetchall()
 c.close()
 return render_sync_devices_page(rows,issued)

def _authenticate_sync_device():
 authorization=request.headers.get("Authorization","")
 if not authorization.startswith("Bearer "):
  return None,None,(jsonify(error="bearer_token_required"),401)
 token=authorization[7:].strip()
 if not token or len(token)>256:
  return None,None,(jsonify(error="invalid_token"),401)
 # Plain HTTP is allowed only for loopback development. Remote sync must use TLS.
 if not request.is_secure and request.remote_addr not in ("127.0.0.1","::1","localhost"):
  return None,None,(jsonify(error="https_required"),426)
 token_hash=hashlib.sha256(token.encode("utf-8")).hexdigest()
 c=db()
 device=c.execute("SELECT device_id,label,active,mode FROM sync_devices WHERE token_hash=?",(token_hash,)).fetchone()
 if not device or not device["active"]:
  c.close(); return None,None,(jsonify(error="device_revoked_or_unknown"),401)
 c.execute("UPDATE sync_devices SET last_seen=CURRENT_TIMESTAMP WHERE device_id=?",(device["device_id"],))
 c.commit()
 return device,c,None

@app.get("/api/sync/v1/status")
def sync_api_status():
 device,c,error=_authenticate_sync_device()
 if error:return error
 latest=c.execute("SELECT COALESCE(MAX(seq),0) FROM sync_change_log").fetchone()[0]
 c.close()
 return jsonify(api_version=1,device_id=device["device_id"],server_cursor=latest,
                entities=sorted(_SYNC_ENTITY_FIELDS.keys()),push_enabled=(device["mode"]=="read_write"),
                note="Authenticated bootstrap/pull; push is limited to validated entity types.")

@app.get("/api/sync/v1/bootstrap")
def sync_api_bootstrap():
 device,c,error=_authenticate_sync_device()
 if error:return error
 entity=request.args.get("entity","")
 if entity not in _SYNC_ENTITY_FIELDS:
  c.close(); return jsonify(error="unknown_entity",allowed=sorted(_SYNC_ENTITY_FIELDS.keys())),400
 try:
  after_id=int(request.args.get("after_id","0"))
  limit=int(request.args.get("limit","200"))
  if after_id<0 or limit<1 or limit>500: raise ValueError
 except (TypeError,ValueError):
  c.close(); return jsonify(error="invalid_pagination"),400
 fields=_SYNC_ENTITY_FIELDS[entity]
 columns=", ".join('"'+field+'"' for field in fields)
 # Keep the row page and its journal cursor in one SQLite read snapshot.
 c.execute("BEGIN")
 latest=c.execute("SELECT COALESCE(MAX(seq),0) FROM sync_change_log").fetchone()[0]
 rows=c.execute(f"SELECT {columns} FROM {entity} WHERE id>? ORDER BY id LIMIT ?",(after_id,limit+1)).fetchall()
 has_more=len(rows)>limit
 rows=rows[:limit]
 payload=[{field:row[field] for field in fields} for row in rows]
 cursor=rows[-1]["id"] if rows else after_id
 c.commit(); c.close()
 return jsonify(entity=entity,rows=payload,next_id=cursor,has_more=has_more,snapshot_cursor=latest)

@app.get("/api/sync/v1/pull")
def sync_api_pull():
 device,c,error=_authenticate_sync_device()
 if error:return error
 try:
  after_seq=int(request.args.get("after_seq","0"))
  limit=int(request.args.get("limit","200"))
  if after_seq<0 or limit<1 or limit>500: raise ValueError
 except (TypeError,ValueError):
  c.close(); return jsonify(error="invalid_pagination"),400
 rows=c.execute("SELECT seq,entity,entity_id,operation,payload_json,changed_at FROM sync_change_log WHERE seq>? ORDER BY seq LIMIT ?",(after_seq,limit+1)).fetchall()
 has_more=len(rows)>limit
 rows=rows[:limit]
 changes=[]
 for row in rows:
  changes.append({"seq":row["seq"],"entity":row["entity"],"entity_id":row["entity_id"],
                  "operation":row["operation"],"payload":json.loads(row["payload_json"]) if row["payload_json"] else None,
                  "changed_at":row["changed_at"]})
 cursor=rows[-1]["seq"] if rows else after_seq
 c.close()
 return jsonify(device_id=device["device_id"],changes=changes,next_seq=cursor,has_more=has_more)



@app.get("/api/sync/v1/photo/<int:rid>")
def sync_api_photo_download(rid):
 device,c,error=_authenticate_sync_device()
 if error:return error
 row=c.execute("SELECT photo_path FROM residents WHERE id=?",(rid,)).fetchone()
 c.close()
 if not row:return jsonify(error="resident_not_found"),404
 photo_path=(row["photo_path"] or "").strip()
 filename=secure_filename(os.path.basename(photo_path))
 if not photo_path.startswith("/photos/") or not filename.startswith("resident_") or os.path.splitext(filename)[1].lower() not in (".png",".jpg",".jpeg",".webp"):
  return jsonify(error="resident_has_no_syncable_photo"),404
 directory=os.path.join(os.path.dirname(os.path.abspath(DB)),"resident_photos")
 full_path=os.path.join(directory,filename)
 if not os.path.isfile(full_path):return jsonify(error="photo_file_missing"),404
 mimetype={".png":"image/png",".jpg":"image/jpeg",".jpeg":"image/jpeg",".webp":"image/webp"}[os.path.splitext(filename)[1].lower()]
 return send_file(full_path,mimetype=mimetype,as_attachment=False,download_name=filename,conditional=True)

@app.put("/api/sync/v1/photo/<int:rid>")
def sync_api_photo_upload(rid):
 device,c,error=_authenticate_sync_device()
 if error:return error
 if device["mode"]!="read_write":
  c.close(); return jsonify(error="device_read_only"),403
 resident=c.execute("SELECT id FROM residents WHERE id=?",(rid,)).fetchone()
 if not resident:
  c.close(); return jsonify(error="resident_not_found"),404
 try:
  base_seq=int(request.args.get("base_seq","0"))
  if base_seq<0:raise ValueError
 except ValueError:
  c.close(); return jsonify(error="invalid_base_seq"),400
 raw=request.get_data(cache=False)
 if not raw or len(raw)>4*1024*1024:
  c.close(); return jsonify(error="photo_must_be_nonempty_and_at_most_4mb"),413
 try:
  with warnings.catch_warnings():
   warnings.simplefilter("error",Image.DecompressionBombWarning)
   with Image.open(io.BytesIO(raw)) as probe:
    if probe.format not in ("PNG","JPEG","WEBP"):raise ValueError("unsupported image format")
    probe.verify()
   with Image.open(io.BytesIO(raw)) as source:
    image=ImageOps.exif_transpose(source); image.load(); image.thumbnail((1200,1200))
    output=io.BytesIO()
    image.convert("RGBA" if "A" in image.getbands() else "RGB").save(output,format="WEBP",quality=90,method=6)
    normalized=output.getvalue()
 except Exception:
  c.close(); return jsonify(error="invalid_png_jpeg_or_webp_image"),400
 photo_digest=hashlib.sha256(normalized).hexdigest()
 current=c.execute("SELECT photo_path FROM residents WHERE id=?",(rid,)).fetchone()
 prior=c.execute("SELECT photo_path,server_seq FROM sync_photo_uploads WHERE device_id=? AND resident_id=? AND sha256=?",(device["device_id"],rid,photo_digest)).fetchone()
 latest=c.execute("SELECT COALESCE(MAX(seq),0) FROM sync_change_log WHERE entity='residents' AND entity_id=?",(str(rid),)).fetchone()[0]
 if prior and current and current["photo_path"]:
  current_file=os.path.join(os.path.dirname(os.path.abspath(DB)),"resident_photos",secure_filename(os.path.basename(current["photo_path"])))
  if os.path.isfile(current_file):
   with open(current_file,"rb") as handle: current_digest=hashlib.sha256(handle.read(4*1024*1024+1)).hexdigest()
   # Treat a byte-identical retry as idempotent only while the resident row
   # has not advanced since that upload. A later edit must still conflict.
   if current_digest==photo_digest and latest<=prior["server_seq"]:
    c.close(); return jsonify(ok=True,resident_id=rid,photo_path=current["photo_path"],server_seq=latest,replayed=True)
 if latest>base_seq:
  current_change=c.execute("SELECT payload_json FROM sync_change_log WHERE entity='residents' AND entity_id=? ORDER BY seq DESC LIMIT 1",(str(rid),)).fetchone()
  c.close(); return jsonify(error="conflict",server_seq=latest,server_payload=json.loads(current_change["payload_json"]) if current_change and current_change["payload_json"] else None),409
 directory=os.path.join(os.path.dirname(os.path.abspath(DB)),"resident_photos")
 os.makedirs(directory,exist_ok=True)
 filename="resident_"+str(rid)+"_sync_"+secrets.token_hex(8)+".webp"
 target=os.path.join(directory,filename); temporary=target+".tmp"
 try:
  with open(temporary,"wb") as handle:handle.write(normalized)
  os.replace(temporary,target)
 finally:
  if os.path.exists(temporary):os.remove(temporary)
 c.execute("UPDATE residents SET photo_path=? WHERE id=?",("/photos/"+filename,rid))
 audit_event(c,"sync.photo_uploaded",f"device_id={device['device_id']}; resident_id={rid}; photo={filename}")
 c.commit()
 server_seq=c.execute("SELECT COALESCE(MAX(seq),0) FROM sync_change_log WHERE entity='residents' AND entity_id=?",(str(rid),)).fetchone()[0]
 c.execute("INSERT OR REPLACE INTO sync_photo_uploads(device_id,resident_id,sha256,photo_path,server_seq) VALUES(?,?,?,?,?)",(device["device_id"],rid,photo_digest,"/photos/"+filename,server_seq))
 c.commit(); c.close()
 return jsonify(ok=True,resident_id=rid,photo_path="/photos/"+filename,server_seq=server_seq)

_SYNC_PUSH_ENTITIES={"residents","terrains","periods","capacities","assignments","history","exam_records","objectives","leave_years","leave_entries","external_rotations"}
_SYNC_PUSH_DELETE_ENTITIES={"objectives","leave_entries","history"}

def _validate_sync_push_row(c,entity,payload):
 fields=_SYNC_ENTITY_FIELDS[entity]
 if not isinstance(payload,dict) or set(payload.keys())!=set(fields):
  raise ValueError("payload must contain the complete whitelisted row shape")
 row=dict(payload)
 try: row["id"]=int(row["id"])
 except (TypeError,ValueError): raise ValueError("row id must be an integer")
 if entity=="residents":
  if not str(row.get("surname") or "").strip() or not str(row.get("firstname") or "").strip(): raise ValueError("resident surname and first name are required")
  if int(row["year"]) not in (1,2,3,4): raise ValueError("resident study year must be 1-4")
  if row["status"] not in ("active","graduated","archived"): raise ValueError("invalid resident status")
 elif entity=="terrains":
  if not str(row.get("name") or "").strip() or row["kind"] not in ("mandatory","basic","optional"): raise ValueError("invalid terrain name/type")
  conflict=c.execute("SELECT id FROM terrains WHERE name=? AND id<>?",(row["name"],row["id"])).fetchone()
  if conflict: raise ValueError("terrain name already exists under another ID")
 elif entity=="periods":
  label=str(row.get("year_label") or "")
  if not re.fullmatch(r"\d{4}-\d{4}",label) or int(label.split("-")[1])!=int(label.split("-")[0])+1: raise ValueError("invalid academic year label")
  if not str(row.get("name") or "").strip(): raise ValueError("period name is required")
  conflict=c.execute("SELECT id FROM periods WHERE year_label=? AND name=? AND id<>?",(label,row["name"],row["id"])).fetchone()
  if conflict: raise ValueError("period already exists under another ID")
 elif entity=="capacities":
  if int(row["places"])<0: raise ValueError("capacity cannot be negative")
  if not c.execute("SELECT 1 FROM terrains WHERE id=?",(row["terrain_id"],)).fetchone() or not c.execute("SELECT 1 FROM periods WHERE id=?",(row["period_id"],)).fetchone(): raise ValueError("capacity references an unknown terrain/period")
  conflict=c.execute("SELECT id FROM capacities WHERE year_label=? AND terrain_id=? AND period_id=? AND id<>?",(row["year_label"],row["terrain_id"],row["period_id"],row["id"])).fetchone()
  if conflict: raise ValueError("capacity already exists for this terrain/period")
 elif entity=="assignments":
  if not c.execute("SELECT 1 FROM residents WHERE id=?",(row["resident_id"],)).fetchone() or not c.execute("SELECT 1 FROM terrains WHERE id=?",(row["terrain_id"],)).fetchone() or not c.execute("SELECT 1 FROM periods WHERE id=?",(row["period_id"],)).fetchone(): raise ValueError("assignment references an unknown resident/terrain/period")
  if row["status"] not in ("pending","confirmed","rejected","cancelled"): raise ValueError("invalid assignment status")
  conflict=c.execute("SELECT id FROM assignments WHERE resident_id=? AND period_id=? AND id<>? AND status IN ('pending','confirmed')",(row["resident_id"],row["period_id"],row["id"])).fetchone()
  if conflict: raise ValueError("resident already has an active assignment for this period")
  # Central approval is required before a remote assignment becomes confirmed.
  row["status"]="pending"
 elif entity=="history":
  if not c.execute("SELECT 1 FROM residents WHERE id=?",(row["resident_id"],)).fetchone() or not c.execute("SELECT 1 FROM terrains WHERE id=?",(row["terrain_id"],)).fetchone(): raise ValueError("history references an unknown resident/terrain")
  if not str(row.get("year_label") or "").strip() or not str(row.get("period") or "").strip(): raise ValueError("history year and period are required")
  if row.get("date_done"):
   date.fromisoformat(row["date_done"])
 elif entity=="exam_records":
  if not c.execute("SELECT 1 FROM residents WHERE id=?",(row["resident_id"],)).fetchone(): raise ValueError("exam references an unknown resident")
  if int(row["study_year"]) not in (1,2,3,4) or int(row["attempt"])<1: raise ValueError("invalid exam year/attempt")
  if row["exam_type"] not in ("virtual","actual","remedial","dems_virtual","dems_actual","dems_remedial","dems_insertion"): raise ValueError("invalid exam session")
  for field in ("theory","practical"):
   if row[field] is not None and not 0<=float(row[field])<=20: raise ValueError("exam scores must be between 0 and 20")
  if str(row["exam_type"]).startswith("dems_") and int(row["study_year"])!=4: raise ValueError("DEMS sessions belong to year 4")
 elif entity=="objectives":
  if not c.execute("SELECT 1 FROM terrains WHERE id=?",(row["terrain_id"],)).fetchone(): raise ValueError("objective references an unknown terrain")
  if int(row["study_year"]) not in (1,2,3,4) or int(row["completed"]) not in (0,1) or not str(row["objective"] or "").strip(): raise ValueError("invalid objective")
 elif entity=="leave_years":
  if not c.execute("SELECT 1 FROM residents WHERE id=?",(row["resident_id"],)).fetchone() or int(row["year"])<2000 or not 0<=int(row["earned_days"])<=30: raise ValueError("invalid annual leave entitlement")
  conflict=c.execute("SELECT id FROM leave_years WHERE resident_id=? AND year=? AND id<>?",(row["resident_id"],row["year"],row["id"])).fetchone()
  if conflict: raise ValueError("leave entitlement already exists for this resident/year")
 elif entity=="leave_entries":
  if not c.execute("SELECT 1 FROM residents WHERE id=?",(row["resident_id"],)).fetchone(): raise ValueError("leave references an unknown resident")
  start=date.fromisoformat(row["start_date"]); end=date.fromisoformat(row["end_date"])
  if end<start or int(row["days"])!=(end-start).days+1: raise ValueError("leave dates/days are inconsistent")
  overlap=c.execute("SELECT id FROM leave_entries WHERE resident_id=? AND id<>? AND start_date<=? AND end_date>=?",(row["resident_id"],row["id"],end.isoformat(),start.isoformat())).fetchone()
  if overlap: raise ValueError("leave overlaps an existing entry")
  # Match the main leave form: unused entitlement carries forward, and a
  # multi-year entry is counted against the cumulative entitlement by its end year.
  earned=c.execute("SELECT COALESCE(SUM(earned_days),0) FROM leave_years WHERE resident_id=? AND year<=?",(row["resident_id"],end.year)).fetchone()[0]
  used=c.execute("SELECT COALESCE(SUM(days),0) FROM leave_entries WHERE resident_id=? AND end_date<=? AND id<>?",(row["resident_id"],date(end.year,12,31).isoformat(),row["id"])).fetchone()[0]
  if used+int(row["days"])>earned: raise ValueError("leave exceeds the recorded cumulative entitlement")
 elif entity=="external_rotations":
  if not c.execute("SELECT 1 FROM residents WHERE id=?",(row["resident_id"],)).fetchone(): raise ValueError("external stage references an unknown resident")
  start=date.fromisoformat(row["start_date"]); end=date.fromisoformat(row["end_date"])
  if end<start: raise ValueError("external stage end date precedes start date")
  reject_external_overlap_with_local_assignment(c,row["resident_id"],start,end)
  if row["request_status"]!="pending" or row["status"]!="requested": raise ValueError("remote devices may submit requests only; approval must happen on the central administrator")
  overlap=c.execute("SELECT id FROM external_rotations WHERE resident_id=? AND id<>? AND request_status NOT IN ('rejected','cancelled') AND start_date<=? AND end_date>=?",(row["resident_id"],row["id"],end.isoformat(),start.isoformat())).fetchone()
  if overlap: raise ValueError("external request overlaps another non-rejected request/stage")
 return row

@app.post("/api/sync/v1/push")
def sync_api_push():
 device,c,error=_authenticate_sync_device()
 if error:return error
 if device["mode"]!="read_write":
  c.close(); return jsonify(error="device_read_only"),403
 raw=request.get_data()
 if len(raw)>1_000_000:
  c.close(); return jsonify(error="batch_too_large"),413
 body=request.get_json(silent=True)
 operations=body.get("operations") if isinstance(body,dict) else None
 if not isinstance(operations,list) or not 1<=len(operations)<=100:
  c.close(); return jsonify(error="operations_must_be_a_list_of_1_to_100"),400
 for op in operations:
  if not isinstance(op,dict) or not all(k in op for k in ("op_id","entity","entity_id","operation","base_seq")):
   c.close(); return jsonify(error="invalid_operation_shape"),400
  if not isinstance(op["op_id"],str) or not 8<=len(op["op_id"])<=128 or op["entity"] not in _SYNC_PUSH_ENTITIES or op["operation"] not in ("upsert","delete"):
   c.close(); return jsonify(error="invalid_operation"),400
  if op["operation"]=="upsert" and not isinstance(op.get("payload"),dict):
   c.close(); return jsonify(error="upsert_requires_payload"),400
 try:
  c.execute("BEGIN IMMEDIATE")
  results=[]
  for op in operations:
   op_id=op["op_id"]; entity=op["entity"]; entity_id=str(op["entity_id"]); operation=op["operation"]
   try: numeric_id=int(entity_id); base_seq=int(op["base_seq"])
   except (TypeError,ValueError): raise ValueError("entity_id/base_seq must be integers")
   if base_seq<0: raise ValueError("base_seq cannot be negative")
   prior=c.execute("SELECT device_id,result_json FROM sync_applied_ops WHERE op_id=?",(op_id,)).fetchone()
   if prior:
    if prior["device_id"]!=device["device_id"]:
     results.append({"op_id":op_id,"status":"rejected","reason":"op_id belongs to another device"})
    else:
     results.append({**json.loads(prior["result_json"]),"replayed":True})
    continue
   latest=c.execute("SELECT COALESCE(MAX(seq),0) FROM sync_change_log WHERE entity=? AND entity_id=?",(entity,entity_id)).fetchone()[0]
   if latest>base_seq:
    current=c.execute("SELECT payload_json FROM sync_change_log WHERE entity=? AND entity_id=? ORDER BY seq DESC LIMIT 1",(entity,entity_id)).fetchone()
    results.append({"op_id":op_id,"status":"conflict","server_seq":latest,"server_payload":json.loads(current["payload_json"]) if current and current["payload_json"] else None})
    continue
   try:
    if operation=="upsert":
     row=_validate_sync_push_row(c,entity,op["payload"])
     if str(row["id"])!=entity_id: raise ValueError("entity_id must match payload.id")
     fields=_SYNC_ENTITY_FIELDS[entity]
     cols=", ".join('"'+field+'"' for field in fields)
     marks=", ".join("?" for _ in fields)
     values=tuple(row[field] for field in fields)
     c.execute(f"INSERT OR REPLACE INTO {entity} ({cols}) VALUES ({marks})",values)
    else:
     if entity not in _SYNC_PUSH_DELETE_ENTITIES: raise ValueError("delete is not allowed for this entity")
     record=c.execute(f"SELECT * FROM {entity} WHERE id=?",(numeric_id,)).fetchone()
     if not record: raise ValueError("record not found")
     label=str(entity)+" #"+str(numeric_id)
     reason="Deleted through authenticated device synchronization; retained for six months"
     c.execute("INSERT INTO trash_items(entity_type,entity_id,label,reason,snapshot_json,deleted_at,deleted_by) VALUES(?,?,?,?,?,?,?)",
       (entity,str(numeric_id),label,reason,json.dumps({"record":dict(record)},ensure_ascii=False),datetime.now().replace(microsecond=0).isoformat(),device["device_id"]))
     c.execute(f"DELETE FROM {entity} WHERE id=?",(numeric_id,))
     audit_event(c,"sync.record_trashed",f"device_id={device['device_id']}; entity={entity}; entity_id={numeric_id}; retained_for_six_months=true")
    # Return the row's own version, not the later audit-log sequence, so
    # unrelated changes cannot mask a concurrent edit to this entity.
    server_seq=c.execute("SELECT COALESCE(MAX(seq),0) FROM sync_change_log WHERE entity=? AND entity_id=?",(entity,entity_id)).fetchone()[0]
    audit_event(c,"sync.push_applied",f"device_id={device['device_id']}; entity={entity}; entity_id={entity_id}; op_id={op_id}")
    result={"op_id":op_id,"status":"applied","server_seq":server_seq}
   except (ValueError,TypeError,sqlite3.IntegrityError) as exc:
    result={"op_id":op_id,"status":"rejected","reason":str(exc)[:240]}
   c.execute("INSERT INTO sync_applied_ops(op_id,device_id,entity,entity_id,result_json) VALUES(?,?,?,?,?)",(op_id,device["device_id"],entity,entity_id,json.dumps(result,ensure_ascii=False)))
   results.append(result)
  c.commit()
 except Exception as exc:
  c.rollback(); c.close(); return jsonify(error="push_batch_failed",details=str(exc)[:200]),400
 c.close()
 return jsonify(results=results,applied_count=sum(r["status"]=="applied" for r in results),
                conflict_count=sum(r["status"]=="conflict" for r in results),
                rejected_count=sum(r["status"]=="rejected" for r in results))

def normalize_header(value):
 value=str(value or "").strip().lower()
 # Accept both machine-readable CSV/XLSX headers and the human-readable labels
 # printed in the Word/PDF templates. Remove parenthetical instructions such as
 # "(AAAA-MM-JJ)" before matching, while preserving the original cell values.
 value=re.sub(r"\([^)]*\)"," ",value)
 value=re.sub(r"[☐☑✓✔]"," ",value)
 value=re.sub(r"[,:;|/]+"," ",value)
 value=re.sub(r"[\s_\-]+"," ",value).strip()
 aliases={
  "surname":"surname","nom":"surname","family name":"surname","last name":"surname","اللقب":"surname",
  "firstname":"firstname","first name":"firstname","prénom":"firstname","prenom":"firstname","given name":"firstname","الاسم":"firstname",
  "year":"year","study year":"year","study_year":"year","année":"year","annee":"year","année d’étude":"year","annee d’étude":"year","niveau":"year","السنة":"year",
  "phone":"phone","phone 1":"phone","telephone":"phone","téléphone":"phone","mobile":"phone","الهاتف":"phone","phone 2":"phone2","telephone 2":"phone2","téléphone 2":"phone2",
  "birth place":"birth_place","lieu de naissance":"birth_place","مكان الميلاد":"birth_place",
  "birth date":"birth_date","date de naissance":"birth_date","تاريخ الميلاد":"birth_date","postal address":"address","adresse postale":"address","email address":"email","adresse e mail":"email","email":"email","resident card number":"resident_card_number","numéro de carte de résident":"resident_card_number","numéro de la carte de résident":"resident_card_number","nin":"national_id_number","national id number":"national_id_number","medical order registered":"medical_order_registered","medical order registered yes":"medical_order_registered_yes","registered at medical order":"medical_order_registered_yes","inscrit ordre des médecins":"medical_order_registered_yes","inscrit à l ordre des médecins":"medical_order_registered_yes","medical order registered no":"medical_order_registered_no","not registered at medical order":"medical_order_registered_no","non inscrit ordre des médecins":"medical_order_registered_no","non inscrit à l ordre des médecins":"medical_order_registered_no","medical order number":"medical_order_number","social security number":"social_security_number","registration date":"registration_date","date de première inscription":"registration_date","resident type":"resident_type","catégorie":"resident_type","origin hospital":"origin_hospital","external start date":"external_start_date","external end date":"external_end_date","additional info":"additional_info",
  "resident":"resident","resident name":"resident","resident full name":"resident","nom du résident":"resident","nom et prénom du résident":"resident","nom resident":"resident","المقيم":"resident",
  "terrain":"terrain","service":"terrain","stage":"terrain","rotation":"terrain","service de stage":"terrain","المصلحة":"terrain","التربص":"terrain",
  "academic year":"year_label","year label":"year_label","academic year label":"year_label","année de post graduation":"year_label","annee de post graduation":"year_label","année universitaire":"year_label","annee universitaire":"year_label","سنة ما بعد التدرج":"year_label",
  "period":"period","rotation period":"period","période":"period","periode":"period","الفترة":"period",
  "date done":"date_done","completion date":"date_done","date completed":"date_done","date de réalisation":"date_done","date d’achèvement":"date_done","date de fin":"date_done","تاريخ الإنجاز":"date_done","notes référence du document justificatif":"notes"
 }
 if "ordre des médecins" in value or "ordre des medecins" in value:
  if "non inscrit" in value: return "medical_order_registered_no"
  if "inscrit" in value and "numéro" not in value: return "medical_order_registered_yes"
  if "numéro" in value: return "medical_order_number"
 if "identification nationale" in value or value=="nin": return "national_id_number"
 if value.startswith("numéro de sécurité sociale"): return "social_security_number"
 if value.startswith("établissement d’origine") or value.startswith("établissement d'origine"): return "origin_hospital"
 if value.startswith("début du stage au chu"): return "external_start_date"
 if value.startswith("fin du stage au chu"): return "external_end_date"
 if value.startswith("catégorie"): return "resident_type"
 if value.startswith("téléphone 1"): return "phone"
 if value.startswith("téléphone 2"): return "phone2"
 if value.startswith("autres informations utiles"): return "additional_info"
 return aliases.get(value,value.replace(" ","_"))
def parse_import_file(path, extension, import_type='residents'):
 rows=[]
 if extension==".csv":
  with open(path,"r",encoding="utf-8-sig",newline="") as handle:
   sample=handle.read(4096); handle.seek(0)
   try: dialect=csv.Sniffer().sniff(sample,delimiters=",;\t")
   except csv.Error: dialect=csv.excel
   rows=list(csv.reader(handle,dialect))
 elif extension in (".xlsx",".xls"):
  if extension==".xlsx":
   from openpyxl import load_workbook
   book=load_workbook(path,read_only=True,data_only=True)
   rows=[list(row) for row in book.active.iter_rows(values_only=True)]
   book.close()
  else:
   import xlrd
   book=xlrd.open_workbook(path)
   sheet=book.sheet_by_index(0)
   rows=[[sheet.cell_value(r,c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
 elif extension==".docx":
  from docx import Document
  doc=Document(path)
  if doc.tables:
   table=doc.tables[0]
   rows=[[cell.text.strip() for cell in row.cells] for row in table.rows]
  else:
   rows=[re.split(r"[,;\t]+",p.text.strip()) for p in doc.paragraphs if p.text.strip()]
 elif extension==".pdf":
  from pypdf import PdfReader
  text_content="\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
  if not text_content.strip(): raise ValueError("The PDF has no extractable text (it may be scanned). Please provide a text-based PDF or spreadsheet.")
  lines=[line.strip() for line in text_content.splitlines() if line.strip()]
  delimiter=";" if sum(x.count(";") for x in lines)>sum(x.count(",") for x in lines) else ("\t" if any("\t" in x for x in lines) else ",")
  rows=[next(csv.reader([line],delimiter=delimiter)) for line in lines]
 elif extension==".doc":
  try:
   result=subprocess.run(["antiword",path],capture_output=True,text=True,check=True,timeout=20)
   document_text=result.stdout
  except FileNotFoundError:
   try:
    with tempfile.TemporaryDirectory(prefix="rotation_doc_") as converted_dir:
     subprocess.run(["libreoffice","--headless","--convert-to","txt:Text","--outdir",converted_dir,path],capture_output=True,text=True,check=True,timeout=30)
     converted_path=os.path.join(converted_dir,os.path.splitext(os.path.basename(path))[0]+".txt")
     with open(converted_path,"r",encoding="utf-8",errors="replace") as handle: document_text=handle.read()
   except (FileNotFoundError,subprocess.CalledProcessError,subprocess.TimeoutExpired) as exc:
    raise ValueError("Legacy .doc requires antiword or LibreOffice; convert it to .docx if neither is installed.") from exc
  lines=[line.strip() for line in document_text.splitlines() if line.strip()]
  rows=[re.split(r"\s*[,;\t]\s*",line) for line in lines]
 if not rows: raise ValueError("No table rows could be read from this file.")
 header=[normalize_header(x) for x in rows[0]]
 required={"resident","terrain","year_label"} if import_type=="history" else {"surname","firstname","year"}
 if not required.issubset(set(header)):
  if import_type=="history":
   raise ValueError("History import requires resident name, terrain/service, and academic year (YYYY-YYYY) columns; period and completion date are optional.")
  raise ValueError("Missing required columns. Include surname/nom/اللقب, firstname/prénom/الاسم, and year/année/السنة in the first row.")
 parsed=[]
 for line_number,values in enumerate(rows[1:],start=2):
  if not any(str(x or "").strip() for x in values): continue
  item={}
  for i,key in enumerate(header):
   if key and i<len(values): item[key]=str(values[i] if values[i] is not None else "").strip()
  registration_columns_present=any(key in header for key in ("medical_order_registered","medical_order_registered_yes","medical_order_registered_no"))
  yes_value=str(item.pop("medical_order_registered_yes","") or "").strip().lower()
  no_value=str(item.pop("medical_order_registered_no","") or "").strip().lower()
  checked_yes=yes_value in ("1","x","yes","oui","true","✓","☑")
  checked_no=no_value in ("1","x","yes","oui","true","✓","☑")
  if checked_yes and checked_no:
   raise ValueError("Choose only one medical-order registration checkbox (registered or not registered).")
  if checked_yes: item["medical_order_registered"]="yes"
  elif checked_no: item["medical_order_registered"]="no"
  item["_medical_order_choice_present"]=registration_columns_present
  item["row_number"]=line_number
  parsed.append(item)
 if not parsed: raise ValueError("The file contains headers but no resident rows.")
 return parsed
@app.route("/templates",methods=["GET","POST"])
def blank_templates():
 if guard(): return redirect("/login")
 if request.method=="POST":
  form_type=request.form.get("form_type","residents_full")
  extension=request.form.get("extension",".xlsx").lower()
  year=request.form.get("year","")
  if form_type not in ("residents_full","residents_by_year","rotation_history","rotation_plan"):
   flash(L("Type de formulaire invalide.","Invalid template type.","نوع الاستمارة غير صالح.")); return redirect("/templates")
  if extension not in (".csv",".xlsx",".docx",".pdf"):
   flash(L("Choisissez CSV, Excel (.xlsx), Word (.docx) ou PDF.","Choose CSV, Excel (.xlsx), Word (.docx), or PDF.","اختر CSV أو Excel (.xlsx) أو Word (.docx) أو PDF.")); return redirect("/templates")
  if form_type=="residents_full" or form_type=="residents_by_year":
   headers=["surname","firstname","birth_date","birth_place","address","email","phone","phone2","resident_card_number","national_id_number","medical_order_registered_yes","medical_order_registered_no","medical_order_number","social_security_number","registration_date","year","resident_type","origin_hospital","external_start_date","external_end_date","additional_info"]
   labels=["Nom","Prénom","Date de naissance (AAAA-MM-JJ)","Lieu de naissance","Adresse postale","Adresse e-mail","Téléphone 1","Téléphone 2","Numéro de carte de résident","NIN / numéro d’identification nationale","☐ Inscrit à l’Ordre des médecins (cocher 1 ou X)","☐ Non inscrit à l’Ordre des médecins (cocher 1 ou X)","Numéro de l’Ordre des médecins (si inscrit)","Numéro de sécurité sociale","Date de première inscription","Année d’étude (1-4)","Catégorie (local/external)","Établissement d’origine (EHS Cardio DBK, CAC DBK ou autre)","Début du stage au CHU","Fin du stage au CHU","Autres informations utiles"]
   if form_type=="residents_by_year":
    headers=["surname","firstname","year"]+[h for h in headers if h not in ("surname","firstname","year")]
    labels=["Nom","Prénom","Année d’étude (1-4)"]+[h for h in labels if h not in ("Nom","Prénom","Année d’étude (1-4)")]
    if year not in ("1","2","3","4"): year="1"
  elif form_type=="rotation_history":
   headers=["resident","terrain","year_label","period","date_done","start_date","end_date","resident_type","origin_hospital","notes"]
   labels=["Nom et prénom du résident (Nom Prénom)","Service de stage","Année de post-graduation (AAAA-AAAA)","Période (nom exact ou vide)","Date d’achèvement","Date de début","Date de fin","Catégorie (local/external)","Établissement d’origine","Notes / référence du document justificatif"]
  else:
   headers=["postgrad_year","period","period_start_date","period_end_date","terrain_name","capacity","resident_surname","resident_firstname","resident_type","origin_hospital","assignment_status","decision_reason"]
   labels=["Année de post-graduation (AAAA-AAAA)","Période","Début de période","Fin de période","Service","Capacité / nombre de postes","Nom du résident","Prénom du résident","Catégorie (local/external)","Établissement d’origine","État (confirmé/en attente)","Motif / décision"]
  filename={"residents_full":"residents_modele_complet","residents_by_year":"residents_modele_annee_"+year,"rotation_history":"modele_historique_stages","rotation_plan":"modele_planification_affectations"}[form_type]+extension
  output=io.BytesIO()
  try:
   if extension==".csv":
    text_output=io.StringIO(newline="")
    writer=csv.writer(text_output); writer.writerow(headers); blank_row=[""]*len(headers)
    if form_type=="residents_by_year": blank_row[headers.index("year")]=year
    writer.writerow(blank_row)
    output.write(("\\ufeff"+text_output.getvalue()).encode("utf-8")); output.seek(0)
    mimetype="text/csv"
   elif extension==".xlsx":
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb=Workbook(); ws=wb.active; ws.title="Formulaire"
    ws.append(headers); blank_row=[""]*len(headers)
    if form_type=="residents_by_year": blank_row[headers.index("year")]=year
    ws.append(blank_row)
    for cell in ws[1]:
     cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="174B68"); cell.alignment=Alignment(wrap_text=True)
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    # Add a human-readable guide without changing the importable first-row keys.
    guide=wb.create_sheet("Guide")
    guide.append(["Colonne à remplir","Libellé / consigne"])
    for key,label in zip(headers,labels):
     guide.append([key,label])
    guide.append(["medical_order_registered","Inscrit : saisir oui/non. Si oui, renseigner medical_order_number."])
    guide.append(["resident_type","Saisir local ou external. Pour un résident externe, remplir origin_hospital, external_start_date et external_end_date."])
    guide.append(["Dates","Utiliser le format AAAA-MM-JJ. Ne pas inventer d’historique de stage."])
    for cell in guide[1]:
     cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="174B68")
    guide.freeze_panes="A2"; guide.column_dimensions["A"].width=34; guide.column_dimensions["B"].width=76
    from openpyxl.worksheet.datavalidation import DataValidation
    if "year" in headers:
     year_col=headers.index("year")+1
     dv=DataValidation(type="list",formula1='"1,2,3,4"',allow_blank=True)
     ws.add_data_validation(dv); dv.add(f"{ws.cell(2,year_col).coordinate}:{ws.cell(501,year_col).coordinate}")
    for registration_column in ("medical_order_registered_yes","medical_order_registered_no"):
     if registration_column in headers:
      reg_col=headers.index(registration_column)+1
      dv=DataValidation(type="list",formula1='"1,X,x,✓,☑"',allow_blank=True)
      ws.add_data_validation(dv); dv.add(f"{ws.cell(2,reg_col).coordinate}:{ws.cell(501,reg_col).coordinate}")
    if "resident_type" in headers:
     type_col=headers.index("resident_type")+1
     dv=DataValidation(type="list",formula1='"local,external"',allow_blank=True)
     ws.add_data_validation(dv); dv.add(f"{ws.cell(2,type_col).coordinate}:{ws.cell(501,type_col).coordinate}")
    for col in ws.columns:
     letter=col[0].column_letter; ws.column_dimensions[letter].width=min(32,max(16,max(len(str(x.value or "")) for x in col)+2))
    wb.save(output); output.seek(0); mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
   elif extension==".docx":
    from docx import Document
    from docx.shared import Inches, Pt
    doc=Document(); doc.add_heading("ResiManage — "+({"residents_full":"Fiche complète des résidents","residents_by_year":"Liste des résidents par année","rotation_history":"Historique des stages documentés","rotation_plan":"Planification et répartition des stages"}[form_type]),0)
    doc.add_paragraph("Formulaire vierge — remplir les champs puis importer le fichier après saisie structurée.")
    table=doc.add_table(rows=1,cols=len(labels)); table.style="Table Grid"
    for i,label in enumerate(labels): table.rows[0].cells[i].text=label
    for _ in range(8):
     cells=table.add_row().cells
     for i,cell in enumerate(cells): cell.text=(year if form_type=="residents_by_year" and headers[i]=="year" else " ")
    doc.add_paragraph("Ordre des médecins : ☐ Inscrit ☐ Non inscrit. À l’importation, inscrire « oui » ou « non » dans la colonne correspondante ; si inscrit, renseigner le numéro. Ne pas inventer les stages antérieurs : utiliser uniquement des données documentées.")
    doc.save(output); output.seek(0); mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
   else:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from xml.sax.saxutils import escape
    doc=SimpleDocTemplate(output,pagesize=landscape(A4),rightMargin=10*mm,leftMargin=10*mm,topMargin=10*mm,bottomMargin=10*mm)
    styles=getSampleStyleSheet(); title=Paragraph("ResiManage — Formulaire vierge",styles["Title"])
    data=[[Paragraph(escape(label),styles["BodyText"]) for label in labels]]
    for _ in range(8):
     data.append([escape(year) if form_type=="residents_by_year" and headers[i]=="year" else "" for i in range(len(headers))])
    table=Table(data,repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#174B68")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),0.4,colors.grey),("VALIGN",(0,0),(-1,-1),"TOP"),("ROWHEIGHT",(0,1),(-1,-1),22)]))
    doc.build([title,Spacer(1,5*mm),table]); output.seek(0); mimetype="application/pdf"
  except ImportError as exc:
   app.logger.warning("Template format unavailable: %s",exc)
   flash(L("Le module nécessaire à ce format n’est pas installé. Essayez CSV ou installez les dépendances de l’application.","The module required for this format is not installed. Try CSV or install the app dependencies.","الوحدة المطلوبة لهذه الصيغة غير مثبتة. جرّب CSV أو ثبّت تبعيات التطبيق.")); return redirect("/templates")
  return send_file(output,mimetype=mimetype,as_attachment=True,download_name=filename)
 return page("""<section><h3>{{L('Formulaires vierges téléchargeables','Downloadable blank forms','استمارات فارغة قابلة للتنزيل')}}</h3><p>{{L('Téléchargez un modèle structuré, remplissez-le puis importez-le dans la rubrique correspondante. Les colonnes sont normalisées pour réduire les erreurs.','Download a structured template, fill it in, then import it in the matching section. Standardized columns help prevent errors.','نزّل نموذجًا منظمًا واملأه ثم استورده في القسم المناسب. الأعمدة الموحدة تقلل الأخطاء.')}}</p><form method="post"><label>{{L('Formulaire','Template','الاستمارة')}}</label><select name="form_type"><option value="residents_full">{{L('Liste complète des résidents — toutes les informations','Complete resident list — all information','قائمة المقيمين الكاملة — جميع المعلومات')}}</option><option value="residents_by_year">{{L('Liste des résidents par année de post-graduation','Residents by postgraduate year','قائمة المقيمين حسب سنة ما بعد التدرج')}}</option><option value="rotation_history">{{L('Historique documenté des stages antérieurs','Documented past rotation history','السجل الموثق للتربصات السابقة')}}</option><option value="rotation_plan">{{L('Planification des capacités et répartition','Capacity and rotation planning','تخطيط المقاعد والتوزيع')}}</option></select><label>{{L('Année de post-graduation (pour la liste par année)','Postgraduate year (for year-specific list)','سنة ما بعد التدرج (للقائمة حسب السنة)')}}</label><select name="year"><option value="1">1</option><option value="2">2</option><option value="3">3</option><option value="4">4</option></select><label>{{L('Format de téléchargement','Download format','صيغة التنزيل')}}</label><select name="extension"><option value=".xlsx">Excel (.xlsx)</option><option value=".docx">Word (.docx)</option><option value=".csv">CSV</option><option value=".pdf">PDF</option></select><button>{{L('Télécharger le formulaire vierge','Download blank template','تنزيل الاستمارة الفارغة')}}</button></form><p><strong>{{L('Conseil :','Tip:','نصيحة:')}}</strong> {{L('Utilisez CSV ou Excel pour l’importation des données. Le modèle Word/PDF convient surtout à la collecte et à la saisie manuelle.','Use CSV or Excel to import data. Word/PDF templates are best for collecting information and manual entry.','استخدم CSV أو Excel لاستيراد البيانات. نماذج Word/PDF مناسبة أكثر لجمع المعلومات والتعبئة اليدوية.')}}</p></section>""")

@app.route("/import",methods=["GET","POST"])
def import_residents():
 if super_guard(): return super_guard()
 c=db()
 c.execute("DELETE FROM import_previews WHERE created_at < datetime('now','-1 day')"); c.commit()
 if request.method=="POST":
  action=request.form.get("action","upload")
  if action in ("confirm","cancel"):
   preview_id=request.form.get("preview_id",type=int)
   preview=c.execute("SELECT * FROM import_previews WHERE id=? AND owner=?",(preview_id,session["user"])).fetchone()
   if not preview:
    c.close(); flash("Import preview expired or not found."); return redirect("/import")
   if action=="cancel":
    c.execute("DELETE FROM import_previews WHERE id=?",(preview_id,)); c.commit(); c.close(); flash("Import cancelled."); return redirect("/import")
   rows=json.loads(preview["rows_json"]); errors=json.loads(preview["errors_json"])
   if errors:
    c.close(); flash("Fix the invalid rows in the source file before importing."); return redirect("/import")
   # Re-check against live data at confirmation: another administrator may have
   # created one of these residents after the preview was generated.
   seen=set(); confirmation_errors=[]
   for row in rows:
    surname=str(row.get("surname","")).strip()
    firstname=str(row.get("firstname","")).strip()
    try:
     year=int(row.get("year"))
     if year not in (1,2,3,4): raise ValueError()
    except (ValueError,TypeError):
     confirmation_errors.append({"row":row.get("row_number"),"message":"Year must be 1, 2, 3 or 4."})
     continue
    row_type_for_key=str(row.get("resident_type","local") or "local").strip().lower()
    origin_for_key=str(row.get("origin_hospital","") or "").strip().casefold() if row_type_for_key=="external" else ""
    key=(surname.casefold(),firstname.casefold(),year,row_type_for_key,origin_for_key)
    if not surname or not firstname:
     confirmation_errors.append({"row":row.get("row_number"),"message":"Surname and first name are required."})
    elif key in seen:
     confirmation_errors.append({"row":row.get("row_number"),"message":"Duplicate resident row in this file."})
    elif c.execute("SELECT 1 FROM residents WHERE lower(surname)=lower(?) AND lower(firstname)=lower(?) AND year=? AND COALESCE(resident_type,'local')=? AND lower(COALESCE(origin_hospital,''))=?",(surname,firstname,year,row_type_for_key,origin_for_key)).fetchone():
     confirmation_errors.append({"row":row.get("row_number"),"message":"Resident with same name and year already exists; upload a fresh preview."})
    seen.add(key)
   if confirmation_errors:
    c.execute("UPDATE import_previews SET errors_json=? WHERE id=?",(json.dumps(confirmation_errors,ensure_ascii=False),preview_id))
    c.commit(); c.close()
    flash(L("La liste des résidents a changé après l’aperçu. Rien n’a été importé ; renvoyez le fichier pour vérifier les données actuelles.","The resident list changed after preview. Nothing was imported; upload the file again to review the current data.","تغيرت قائمة المقيمين بعد المعاينة. لم يتم استيراد أي شيء؛ ارفع الملف مجددًا لمراجعة البيانات الحالية."))
    return redirect("/import")
   inserted=0
   try:
    for row in rows:
     resident_type=str(row.get("resident_type","local") or "local").strip().lower()
     if resident_type not in ("local","external"): resident_type="local"
     c.execute("INSERT INTO residents(surname,firstname,year,phone,birth_place,birth_date,status,phone2,address,email,resident_card_number,national_id_number,medical_order_registered,medical_order_number,social_security_number,registration_date,resident_type,origin_hospital,external_start_date,external_end_date,additional_info) VALUES(?,?,?,?,?,?, 'active',?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
       (row["surname"],row["firstname"],int(row["year"]),row.get("phone",""),row.get("birth_place",""),row.get("birth_date",""),row.get("phone2",""),row.get("address",""),row.get("email",""),row.get("resident_card_number",""),row.get("national_id_number",""),row.get("medical_order_registered","no"),row.get("medical_order_number",""),row.get("social_security_number",""),row.get("registration_date",""),resident_type,row.get("origin_hospital",""),row.get("external_start_date",""),row.get("external_end_date",""),row.get("additional_info","")))
     inserted+=1
   except sqlite3.IntegrityError:
    c.rollback(); c.close()
    flash(L("Importation annulée car les données des résidents ont changé pendant la confirmation. Rien n’a été importé ; créez un nouvel aperçu.","Import cancelled because the resident data changed during confirmation. Nothing was imported; upload a fresh preview.","تم إلغاء الاستيراد لأن بيانات المقيمين تغيرت أثناء التأكيد. لم يتم استيراد أي شيء؛ أنشئ معاينة جديدة."))
    return redirect("/import")
   audit_event(c,"residents.imported",f"filename={preview['filename']}; imported={inserted}; history_not_imported=true")
   c.execute("DELETE FROM import_previews WHERE id=?",(preview_id,)); c.commit(); c.close()
   flash(L(f"{inserted} résidents ont été importés. Aucun stage antérieur n’a été inventé ni importé.",f"Imported {inserted} residents. No past rotations were invented or imported.",f"تم استيراد {inserted} مقيمين. لم يتم اختلاق أو استيراد أي تربص سابق."))
   return redirect("/residents")
  upload=request.files.get("file")
  if not upload or not upload.filename:
   c.close(); flash(L("Choisissez d’abord un fichier.","Choose a file first.","اختر ملفًا أولًا.")); return redirect("/import")
  filename=os.path.basename(upload.filename); extension=os.path.splitext(filename)[1].lower()
  allowed={".csv",".xlsx",".xls",".docx",".doc",".pdf"}
  if extension not in allowed:
   c.close(); flash(L("Formats acceptés : CSV, XLS/XLSX, DOC/DOCX et PDF contenant du texte.","Supported formats: CSV, XLS/XLSX, DOC/DOCX and text-based PDF.","الصيغ المدعومة: CSV وXLS/XLSX وDOC/DOCX وPDF النصي.")); return redirect("/import")
  staged=None
  try:
   with tempfile.NamedTemporaryFile(prefix="rotation_import_",suffix=extension,dir=os.path.dirname(os.path.abspath(DB)),delete=False) as temp:
    staged=temp.name; upload.save(staged)
   rows=parse_import_file(staged,extension)
   errors=[]; seen=set()
   for row in rows:
    row_errors=[]
    surname=row.get("surname","").strip(); firstname=row.get("firstname","").strip()
    if not surname or not firstname: row_errors.append(L("Le nom et le prénom sont obligatoires.","Surname and first name are required.","اللقب والاسم إجباريان."))
    try:
     numeric_year=float(row.get("year",""))
     if not numeric_year.is_integer(): raise ValueError()
     year=int(numeric_year)
     if year not in (1,2,3,4): raise ValueError()
     row["year"]=year
    except (ValueError,TypeError): row_errors.append(L("L’année doit être 1, 2, 3 ou 4.","Year must be 1, 2, 3 or 4.","يجب أن تكون السنة 1 أو 2 أو 3 أو 4."))
    row["surname"]=surname; row["firstname"]=firstname
    row_type_for_key=str(row.get("resident_type","local") or "local").strip().lower()
    origin_for_key=str(row.get("origin_hospital","") or "").strip().casefold() if row_type_for_key=="external" else ""
    key=(surname.casefold(),firstname.casefold(),row.get("year"),row_type_for_key,origin_for_key)
    if key in seen: row_errors.append(L("Cette ligne de résident est dupliquée dans le fichier.","Duplicate resident row in this file.","سطر المقيم هذا مكرر في الملف."))
    seen.add(key)
    row_type=str(row.get("resident_type","local") or "local").strip().lower()
    if row_type not in ("local","external"): row_errors.append(L("La catégorie doit être local ou external.","Category must be local or external.","يجب أن تكون الفئة local أو external."))
    registered=str(row.get("medical_order_registered","") or "").strip().lower()
    if registered in ("1","x","✓","☑","true","inscrit"): registered="yes"
    if registered in ("0","false","non inscrit","not registered"): registered="no"
    if row.get("_medical_order_choice_present") and registered not in ("yes","no","oui","non"): row_errors.append(L("Cochez exactement une option pour l’Ordre des médecins : inscrit ou non inscrit.","Select exactly one medical-order option: registered or not registered.","اختر خيارًا واحدًا فقط بخصوص التسجيل في هيئة الأطباء: مسجل أو غير مسجل."))
    elif registered in ("yes","oui") and not str(row.get("medical_order_number","") or "").strip(): row_errors.append(L("Le numéro de l’Ordre des médecins est obligatoire si le résident est inscrit.","The medical-order number is required when the resident is registered.","رقم هيئة الأطباء إجباري إذا كان المقيم مسجلًا."))
    row["medical_order_registered"]="yes" if registered in ("yes","oui") else "no" if registered in ("no","non") else "no"
    for date_field,date_label in (("birth_date",L("date de naissance","birth date","تاريخ الميلاد")),("registration_date",L("date de première inscription","first registration date","تاريخ أول تسجيل"))):
     raw_date=str(row.get(date_field,"") or "").strip()
     if raw_date:
      try:
       parsed_date=date.fromisoformat(raw_date)
       if parsed_date>date.today(): raise ValueError()
      except ValueError:
       row_errors.append(L(f"La {date_label} doit être une date valide au format AAAA-MM-JJ et ne peut pas être dans le futur.",f"The {date_label} must be a valid YYYY-MM-DD date and cannot be in the future.",f"يجب أن يكون {date_label} تاريخًا صحيحًا بالتنسيق YYYY-MM-DD وألا يكون في المستقبل."))
    if row_type=="external" and (not row.get("origin_hospital","").strip() or not row.get("external_start_date","").strip() or not row.get("external_end_date","").strip()):
     row_errors.append(L("Un résident externe doit avoir un établissement d’origine et des dates de stage.","An external resident must have a home institution and rotation dates.","يجب أن يكون للمقيم الخارجي مؤسسة أصلية وتواريخ تربص."))
    elif row_type=="external":
     try:
      external_start=date.fromisoformat(row.get("external_start_date","").strip())
      external_end=date.fromisoformat(row.get("external_end_date","").strip())
      if external_end<external_start: raise ValueError()
     except ValueError:
      row_errors.append(L("Les dates du stage externe doivent être valides au format AAAA-MM-JJ et la fin ne peut précéder le début.","External rotation dates must be valid YYYY-MM-DD dates and the end cannot precede the start.","يجب أن تكون تواريخ التربص الخارجي صحيحة بالتنسيق YYYY-MM-DD وألا تسبق النهاية البداية."))
    exists=c.execute("SELECT 1 FROM residents WHERE lower(surname)=lower(?) AND lower(firstname)=lower(?) AND year=? AND COALESCE(resident_type,'local')=? AND COALESCE(origin_hospital,'')=?",(surname,firstname,row.get("year",0),row_type,row.get("origin_hospital",""))).fetchone()
    if exists: row_errors.append(L("Un résident portant le même nom, la même année et le même établissement existe déjà.","Resident with same name, year and institution already exists.","يوجد مقيم بالاسم والسنة والمؤسسة نفسها."))
    if row_errors: errors.append({"row":row.get("row_number"),"message":" ".join(row_errors)})
   rows.sort(key=lambda row:(row.get("year",99),row.get("surname","").casefold(),row.get("firstname","").casefold()))
   c.execute("DELETE FROM import_previews WHERE created_at < datetime('now','-1 day')")
   c.execute("INSERT INTO import_previews(owner,rows_json,errors_json,filename) VALUES(?,?,?,?)",(session["user"],json.dumps(rows,ensure_ascii=False),json.dumps(errors,ensure_ascii=False),filename))
   preview_id=c.execute("SELECT last_insert_rowid()").fetchone()[0]; c.commit()
   preview=c.execute("SELECT * FROM import_previews WHERE id=?",(preview_id,)).fetchone(); c.close()
   return page("""<section><h3>{{t('import')}} — {{t('preview')}}</h3><p>{{L('Fichier :','File:','الملف:')}} {{filename}} — {{rows|length}} {{L('lignes. Aucun stage passé ne sera inventé.','rows. Past stages/history will not be fabricated.','صفوف. لن يتم اختلاق أي تربص سابق.')}}</p>{% if errors %}<p class="bad">{{L('Erreurs de validation :','Validation errors:','أخطاء التحقق:')}} {{errors|length}}. {{L('Corrigez le fichier source puis importez-le à nouveau ; aucune ligne ne sera importée tant que des erreurs subsistent.','Correct the source file and upload again; no rows can be imported while errors remain.','صحح الملف المصدر وأعد رفعه؛ لا يمكن استيراد أي صف ما دامت الأخطاء موجودة.')}}</p>{% endif %}<div class="wrap"><table><tr><th>{{L('Ligne','Row','السطر')}}</th><th>{{L('Nom','Surname','اللقب')}}</th><th>{{L('Prénom','First name','الاسم')}}</th><th>{{L('Année','Year','السنة')}}</th><th>{{L('Téléphone','Phone','الهاتف')}}</th><th>{{L('Validation','Validation','التحقق')}}</th></tr>{% for row in rows %}<tr><td>{{row.row_number}}</td><td>{{row.surname}}</td><td>{{row.firstname}}</td><td>{{row.year}}</td><td>{{row.phone or ''}}</td><td>{% for e in errors if e.row==row.row_number %}<span class="bad">{{e.message}}</span>{% else %}<span class="good">{{L('Valide','Valid','صالح')}}</span>{% endfor %}</td></tr>{% endfor %}</table></div><form method="post"><input type="hidden" name="preview_id" value="{{preview_id}}"><button name="action" value="confirm" {% if errors %}disabled{% endif %}>{{t('confirm_import')}}</button> <button name="action" value="cancel">{{t('cancel')}}</button></form></section>""",filename=filename,rows=rows,errors=errors,preview_id=preview_id)
  except Exception as exc:
   c.close(); app.logger.warning("Resident import validation failed: %s",exc); flash(L("Impossible de lire ou valider le fichier. Vérifiez le format et les colonnes.","Could not read or validate the file. Check the format and columns.","تعذرت قراءة الملف أو التحقق منه. تحقق من الصيغة والأعمدة.")); return redirect("/import")
  finally:
   if staged and os.path.exists(staged): os.remove(staged)
 c.close()
 return page("""<section><h3>{{t('import')}}</h3><p>{{t('import_help')}} {{L('Formats acceptés : CSV, XLS/XLSX, DOC/DOCX et PDF contenant du texte.','Supported formats: CSV, XLS/XLSX, DOC/DOCX and text-based PDF.','الصيغ المدعومة: CSV وXLS/XLSX وDOC/DOCX وPDF النصي.')}}</p><form method="post" enctype="multipart/form-data"><label>{{t('file')}}</label><input type="file" name="file" accept=".csv,.xls,.xlsx,.doc,.docx,.pdf" required><button>{{t('upload_preview')}}</button></form><p><a href="/import-history">{{L('Importer l’historique documenté des stages','Import documented stage history','استيراد سجل التربصات السابقة الموثق')}}</a></p></section>""")

@app.route("/import-history",methods=["GET","POST"])
def import_history():
 if super_guard(): return super_guard()
 c=db()
 c.execute("DELETE FROM import_previews WHERE created_at < datetime('now','-1 day')"); c.commit()
 if request.method=="POST":
  action=request.form.get("action","upload")
  if action in ("confirm","cancel"):
   preview_id=request.form.get("preview_id",type=int)
   preview=c.execute("SELECT * FROM import_previews WHERE id=? AND owner=? AND import_type='history'",(preview_id,session["user"])).fetchone()
   if not preview:
    c.close(); flash("History import preview expired or not found."); return redirect("/import-history")
   if action=="cancel":
    c.execute("DELETE FROM import_previews WHERE id=?",(preview_id,)); c.commit(); c.close(); flash("History import cancelled."); return redirect("/import-history")
   rows=json.loads(preview["rows_json"]); errors=json.loads(preview["errors_json"])
   if errors:
    c.close(); flash("Fix all history validation errors before importing."); return redirect("/import-history")
   seen=set(); validation=[]
   for row in rows:
    resident_name=str(row.get("resident","")).strip()
    terrain_name=str(row.get("terrain","")).strip()
    year_label=str(row.get("year_label","")).strip()
    matches=c.execute("SELECT id,surname,firstname FROM residents WHERE lower(trim(surname||' '||firstname))=lower(?) OR lower(trim(firstname||' '||surname))=lower(?) OR lower(trim(surname||', '||firstname))=lower(?)",(resident_name,resident_name,resident_name)).fetchall()
    terrains=c.execute("SELECT id,name FROM terrains WHERE lower(trim(name))=lower(?)",(terrain_name,)).fetchall()
    if len(matches)!=1: validation.append({"row":row.get("row_number"),"message":"Resident no longer matches exactly one existing resident."}); continue
    if len(terrains)!=1: validation.append({"row":row.get("row_number"),"message":"Terrain no longer matches exactly one existing terrain."}); continue
    try:
     if not re.fullmatch(r"\d{4}-\d{4}",year_label) or int(year_label.split("-")[1])!=int(year_label.split("-")[0])+1: raise ValueError(L("L’année de post-graduation doit suivre le format AAAA-AAAA avec deux années consécutives.","Postgraduate year must use YYYY-YYYY with consecutive years.","يجب أن تكون سنة ما بعد التدرج بالتنسيق YYYY-YYYY مع سنتين متتاليتين."))
     period=str(row.get("period","")).strip()
     valid_periods={"01 Fév - 30 Avr","01 Mai - 30 Juillet","01 Août - 30 Oct","01 Nov - 30 Jan"}
     if period and period not in valid_periods: raise ValueError(L("Choisissez l’une des quatre périodes ou laissez ce champ vide.","Choose one of the four official periods or leave period blank.","اختر إحدى الفترات الأربع أو اترك الحقل فارغًا."))
     date_done=str(row.get("date_done","")).strip()
     if date_done:
      parsed=date.fromisoformat(date_done)
      if parsed>date.today(): raise ValueError(L("Une date de stage historique ne peut pas être dans le futur.","A historical rotation date cannot be in the future.","لا يمكن أن يكون تاريخ التربص السابق في المستقبل."))
    except ValueError as exc:
     validation.append({"row":row.get("row_number"),"message":str(exc)}); continue
    rid=matches[0]["id"]; tid=terrains[0]["id"]
    key=(rid,tid,year_label,period,date_done)
    if key in seen: validation.append({"row":row.get("row_number"),"message":"Duplicate history row in this file."})
    elif c.execute("SELECT 1 FROM history WHERE resident_id=? AND terrain_id=? AND year_label=? AND COALESCE(period,'')=? AND COALESCE(date_done,'')=?",(rid,tid,year_label,period,date_done)).fetchone():
     validation.append({"row":row.get("row_number"),"message":"This exact history record already exists; upload a fresh preview."})
    seen.add(key)
   if validation:
    c.execute("UPDATE import_previews SET errors_json=? WHERE id=?",(json.dumps(validation,ensure_ascii=False),preview_id)); c.commit(); c.close()
    flash(L("L’historique a changé après l’aperçu. Rien n’a été importé ; veuillez renvoyer le fichier.","History changed after preview. Nothing was imported; upload the file again.","تغير السجل بعد المعاينة. لم يتم استيراد أي شيء؛ يرجى رفع الملف مجددًا."))
    return redirect("/import-history")
   inserted=0
   try:
    for row in rows:
     resident_name=str(row["resident"]).strip(); terrain_name=str(row["terrain"]).strip()
     match=c.execute("SELECT id FROM residents WHERE lower(trim(surname||' '||firstname))=lower(?) OR lower(trim(firstname||' '||surname))=lower(?) OR lower(trim(surname||', '||firstname))=lower(?)",(resident_name,resident_name,resident_name)).fetchone()
     terrain=c.execute("SELECT id FROM terrains WHERE lower(trim(name))=lower(?)",(terrain_name,)).fetchone()
     c.execute("INSERT INTO history(resident_id,terrain_id,year_label,period,date_done) VALUES(?,?,?,?,?)",(match["id"],terrain["id"],row["year_label"],row.get("period",""),row.get("date_done","")))
     inserted+=1
   except sqlite3.IntegrityError:
    c.rollback(); c.close(); flash(L("Importation annulée car l’historique a changé pendant la confirmation. Rien n’a été importé ; créez un nouvel aperçu.","Import cancelled because history changed during confirmation. Nothing was imported; upload a fresh preview.","تم إلغاء الاستيراد لأن السجل تغير أثناء التأكيد. لم يتم استيراد أي شيء؛ أنشئ معاينة جديدة.")); return redirect("/import-history")
   audit_event(c,"history.imported",f"filename={preview['filename']}; imported={inserted}; documented_rows_only=true")
   c.execute("DELETE FROM import_previews WHERE id=?",(preview_id,)); c.commit(); c.close()
   flash(L(f"{inserted} stages antérieurs documentés ont été importés. Aucun historique n’a été inventé.",f"Imported {inserted} documented past-stage records. No history was invented.",f"تم استيراد {inserted} من سجلات التربصات السابقة الموثقة. لم يتم اختلاق أي سجل."))
   return redirect("/history")
  upload=request.files.get("file")
  if not upload or not upload.filename:
   c.close(); flash(L("Choisissez d’abord un fichier.","Choose a file first.","اختر ملفًا أولًا.")); return redirect("/import-history")
  filename=os.path.basename(upload.filename); extension=os.path.splitext(filename)[1].lower()
  if extension not in {".csv",".xlsx",".xls",".docx",".doc",".pdf"}:
   c.close(); flash(L("Formats acceptés : CSV, XLS/XLSX, DOC/DOCX et PDF contenant du texte.","Supported formats: CSV, XLS/XLSX, DOC/DOCX and text-based PDF.","الصيغ المدعومة: CSV وXLS/XLSX وDOC/DOCX وPDF النصي.")); return redirect("/import-history")
  staged=None
  try:
   with tempfile.NamedTemporaryFile(prefix="rotation_history_import_",suffix=extension,dir=os.path.dirname(os.path.abspath(DB)),delete=False) as temp:
    staged=temp.name; upload.save(staged)
   rows=parse_import_file(staged,extension,"history")
   errors=[]; seen=set(); valid_periods={"01 Fév - 30 Avr","01 Mai - 30 Juillet","01 Août - 30 Oct","01 Nov - 30 Jan"}
   residents_pool=c.execute("SELECT id,surname,firstname FROM residents").fetchall()
   terrain_pool=c.execute("SELECT id,name FROM terrains").fetchall()
   for row in rows:
    row_errors=[]
    resident_name=str(row.get("resident","")).strip(); terrain_name=str(row.get("terrain","")).strip(); year_label=str(row.get("year_label","")).strip()
    row["resident"]=resident_name; row["terrain"]=terrain_name; row["year_label"]=year_label
    if not resident_name: row_errors.append(L("Le nom du résident est obligatoire.","Resident name is required.","اسم المقيم إجباري."))
    if not terrain_name: row_errors.append(L("Le nom du service de stage est obligatoire.","Rotation department name is required.","اسم مصلحة التربص إجباري."))
    if not re.fullmatch(r"\d{4}-\d{4}",year_label) or int(year_label.split("-")[1])!=int(year_label.split("-")[0])+1:
     row_errors.append(L("L’année de post-graduation doit suivre le format AAAA-AAAA avec deux années consécutives.","Postgraduate year must use YYYY-YYYY with consecutive years.","يجب أن تكون سنة ما بعد التدرج بالتنسيق YYYY-YYYY مع سنتين متتاليتين."))
    resident_matches=[r for r in residents_pool if resident_name.casefold() in {(r["surname"]+" "+r["firstname"]).casefold().strip(),(r["firstname"]+" "+r["surname"]).casefold().strip(),(r["surname"]+", "+r["firstname"]).casefold().strip()}]
    terrain_matches=[t for t in terrain_pool if t["name"].casefold().strip()==terrain_name.casefold().strip()]
    if len(resident_matches)!=1: row_errors.append(L("Le résident doit correspondre exactement à un dossier existant ; utilisez « Nom Prénom ».","Resident must match exactly one existing resident; use 'Surname Firstname'.","يجب أن يطابق المقيم ملفًا موجودًا واحدًا بالضبط؛ استعمل «اللقب الاسم»."))
    else: row["resident_id"]=resident_matches[0]["id"]
    if len(terrain_matches)!=1: row_errors.append(L("Le service de stage doit correspondre exactement à un service existant.","Rotation department must match exactly one existing department.","يجب أن تطابق مصلحة التربص مصلحة موجودة واحدة بالضبط."))
    else: row["terrain_id"]=terrain_matches[0]["id"]
    period=str(row.get("period","")).strip(); date_done=str(row.get("date_done","")).strip()
    if period and period not in valid_periods: row_errors.append(L("Choisissez l’une des quatre périodes ou laissez ce champ vide.","Choose one of the four official periods or leave period blank.","اختر إحدى الفترات الأربع أو اترك الحقل فارغًا."))
    if date_done:
     try:
      parsed=date.fromisoformat(date_done)
      if parsed>date.today(): row_errors.append(L("Une date de stage historique ne peut pas être dans le futur.","A historical rotation date cannot be in the future.","لا يمكن أن يكون تاريخ التربص السابق في المستقبل."))
     except ValueError: row_errors.append(L("La date doit être une date réelle au format AAAA-MM-JJ.","Date must be a real calendar date in YYYY-MM-DD format.","يجب أن يكون التاريخ صحيحًا بالتنسيق YYYY-MM-DD."))
    row["period"]=period; row["date_done"]=date_done
    key=(row.get("resident_id"),row.get("terrain_id"),year_label,period,date_done)
    if key in seen: row_errors.append(L("Cette ligne d’historique est présente plusieurs fois dans le fichier.","Duplicate history row in this file.","هذا السطر من السجل مكرر في الملف."))
    seen.add(key)
    if row.get("resident_id") and row.get("terrain_id") and c.execute("SELECT 1 FROM history WHERE resident_id=? AND terrain_id=? AND year_label=? AND COALESCE(period,'')=? AND COALESCE(date_done,'')=?",(row["resident_id"],row["terrain_id"],year_label,period,date_done)).fetchone():
     row_errors.append(L("Cet enregistrement historique existe déjà.","This exact history record already exists.","سجل التاريخ هذا موجود مسبقًا."))
    if row_errors: errors.append({"row":row.get("row_number"),"message":" ".join(row_errors)})
   rows.sort(key=lambda row:(row.get("year_label",""),row.get("resident","").casefold(),row.get("terrain","").casefold()))
   c.execute("INSERT INTO import_previews(owner,rows_json,errors_json,filename,import_type) VALUES(?,?,?,?,?)",(session["user"],json.dumps(rows,ensure_ascii=False),json.dumps(errors,ensure_ascii=False),filename,"history"))
   preview_id=c.execute("SELECT last_insert_rowid()").fetchone()[0]; c.commit(); c.close()
   return page("""<section><h3>{{L('Importer l’historique des stages documentés','Import documented stage history','استيراد سجل التربصات السابقة الموثق')}}</h3><p>{{L('Fichier :','File:','الملف:')}} {{filename}} — {{rows|length}} {{L('lignes. Seuls les stages explicitement documentés sont importés ; aucun stage n’est déduit.','rows. Only explicitly documented stages are imported; no stages are inferred.','صفوف. يتم استيراد التربصات الموثقة صراحة فقط؛ ولا يتم استنتاج أي تربص.')}}</p>{% if errors %}<p class="bad">{{L('Erreurs de validation :','Validation errors:','أخطاء التحقق:')}} {{errors|length}}. {{L('Aucune donnée ne peut être importée avant de corriger le fichier source.','Nothing can be imported until the source is corrected.','لا يمكن استيراد أي بيانات قبل تصحيح الملف المصدر.')}}</p>{% endif %}<div class="wrap"><table><tr><th>{{L('Ligne','Row','السطر')}}</th><th>{{L('Résident','Resident','المقيم')}}</th><th>{{L('Service de stage','Terrain','مصلحة التربص')}}</th><th>{{L('Année de post-graduation','Postgraduate year','سنة ما بعد التدرج')}}</th><th>{{L('Période','Period','الفترة')}}</th><th>{{L('Date de fin','Completion date','تاريخ الإتمام')}}</th><th>{{L('Validation','Validation','التحقق')}}</th></tr>{% for row in rows %}<tr><td>{{row.row_number}}</td><td>{{row.resident}}</td><td>{{row.terrain}}</td><td>{{row.year_label}}</td><td>{{row.period}}</td><td>{{row.date_done}}</td><td>{% for e in errors if e.row==row.row_number %}<span class="bad">{{e.message}}</span>{% else %}<span class="good">{{L('Valide','Valid','صالح')}}</span>{% endfor %}</td></tr>{% endfor %}</table></div><form method="post"><input type="hidden" name="preview_id" value="{{preview_id}}"><button name="action" value="confirm" {% if errors %}disabled{% endif %}>{{L('Confirmer l’importation','Confirm import','تأكيد الاستيراد')}}</button> <button name="action" value="cancel">{{L('Annuler','Cancel','إلغاء')}}</button></form></section>""",filename=filename,rows=rows,errors=errors,preview_id=preview_id)
  except Exception as exc:
   c.close(); app.logger.warning("History import validation failed: %s",exc); flash(L("Impossible de lire ou valider le fichier d’historique. Vérifiez le format et les colonnes.","Could not read or validate the history file. Check the format and columns.","تعذرت قراءة ملف السجل أو التحقق منه. تحقق من الصيغة والأعمدة.")); return redirect("/import-history")
  finally:
   if staged and os.path.exists(staged): os.remove(staged)
 c.close()
 return page("""<section><h3>{{L('Importer l’historique des stages documentés','Import documented stage history','استيراد سجل التربصات السابقة الموثق')}}</h3><p>{{L('Importez un fichier CSV, XLS/XLSX, DOC/DOCX ou PDF texte avec les colonnes résident, service de stage et année de post-graduation ; la période et la date de fin sont facultatives. Le résident et le service doivent correspondre exactement aux dossiers existants.','Upload CSV, XLS/XLSX, DOC/DOCX or text-based PDF with resident, rotation department and postgraduate year columns; period and completion date are optional. Resident and department must exactly match existing records.','ارفع ملف CSV أو XLS/XLSX أو DOC/DOCX أو PDF نصي يحتوي أعمدة المقيم ومصلحة التربص وسنة ما بعد التدرج؛ الفترة وتاريخ الإتمام اختياريان. يجب أن يطابق المقيم والمصلحة السجلات الموجودة تمامًا.')}}</p><form method="post" enctype="multipart/form-data"><label>{{L('Fichier d’historique','History file','ملف السجل')}}</label><input type="file" name="file" accept=".csv,.xls,.xlsx,.doc,.docx,.pdf" required><button>{{L('Importer et prévisualiser','Upload and preview','رفع ومعاينة')}}</button></form><p><a href="/import">{{L('Retour à l’importation des résidents','Back to resident import','العودة إلى استيراد المقيمين')}}</a></p></section>""")

@app.get("/photos/<filename>")
def resident_photo(filename):
 if guard(): return redirect("/login")
 safe=secure_filename(filename)
 if safe!=filename or not safe.startswith("resident_") or os.path.splitext(safe)[1].lower() not in (".png",".jpg",".jpeg",".webp"): return ("Not found",404)
 directory=os.path.join(os.path.dirname(os.path.abspath(DB)),"resident_photos")
 return send_from_directory(directory,safe)
@app.route("/profile/<int:rid>",methods=["GET","POST"])
def resident_profile(rid):
 if guard(): return redirect("/login")
 c=db(); r=c.execute("SELECT * FROM residents WHERE id=?",(rid,)).fetchone()
 if not r: c.close(); return ("Resident not found",404)
 if request.method=="POST":
  if session.get("role")!="super_admin": c.close(); return ("Forbidden — Super Admin only",403)
  surname=request.form.get("surname",r["surname"]).strip()
  firstname=request.form.get("firstname",r["firstname"]).strip()
  birth_place=request.form.get("birth_place",r["birth_place"] or "").strip()
  birth_date=request.form.get("birth_date",r["birth_date"] or "").strip()
  registration_date=request.form.get("registration_date",r["registration_date"] or "").strip()
  phone=request.form.get("phone",r["phone"] or "").strip()
  phone2=request.form.get("phone2",r["phone2"] or "").strip()
  address=request.form.get("address",r["address"] or "").strip()
  email=request.form.get("email",r["email"] or "").strip()
  resident_card_number=request.form.get("resident_card_number",r["resident_card_number"] or "").strip()
  national_id_number=request.form.get("national_id_number",r["national_id_number"] or "").strip()
  medical_order_registered=request.form.get("medical_order_registered",r["medical_order_registered"] or "no").strip()
  medical_order_number=request.form.get("medical_order_number",r["medical_order_number"] or "").strip()
  social_security_number=request.form.get("social_security_number",r["social_security_number"] or "").strip()
  additional_info=request.form.get("additional_info",r["additional_info"] or "").strip()
  origin_hospital=request.form.get("origin_hospital",r["origin_hospital"] or "").strip()
  external_start_date=request.form.get("external_start_date",r["external_start_date"] or "").strip()
  external_end_date=request.form.get("external_end_date",r["external_end_date"] or "").strip()
  if medical_order_registered not in ("yes","no"):
   c.close(); flash(L("Indiquez si le résident est inscrit ou non à l’Ordre des médecins.","Select whether the resident is registered with the Medical Order.","حدد ما إذا كان المقيم مسجلًا في هيئة الأطباء أم لا.")); return redirect("/profile/"+str(rid))
  if r["resident_type"]=="external" and (not origin_hospital or not external_start_date or not external_end_date):
   c.close(); flash(L("L’établissement d’origine et les dates du stage externe sont obligatoires.","Home institution and external rotation dates are required.","المؤسسة الأصلية وتواريخ التربص الخارجي إجبارية.")); return redirect("/profile/"+str(rid))
  if r["resident_type"]=="external":
   try:
    if date.fromisoformat(external_end_date)<date.fromisoformat(external_start_date): raise ValueError()
   except ValueError:
    c.close(); flash(L("Les dates du stage externe sont invalides.","External rotation dates are invalid.","تواريخ التربص الخارجي غير صالحة.")); return redirect("/profile/"+str(rid))
  if len(phone2)>40 or len(email)>254 or len(address)>500 or len(additional_info)>2000:
   c.close(); flash(L("Vérifiez les longueurs des coordonnées et des informations complémentaires.","Check the length of contact and additional information fields.","تحقق من أطوال حقول الاتصال والمعلومات الإضافية.")); return redirect("/profile/"+str(rid))
  if not surname or not firstname or len(surname)>100 or len(firstname)>100:
   c.close(); flash(L("Le nom et le prénom sont obligatoires et ne doivent pas dépasser 100 caractères.","Surname and first name are required and must be at most 100 characters.","اللقب والاسم إجباريان ويجب ألا يتجاوز كل منهما 100 حرف.")); return redirect("/profile/"+str(rid))
  if len(birth_place)>120 or len(phone)>40:
   c.close(); flash(L("Le lieu de naissance ne doit pas dépasser 120 caractères et le téléphone 40 caractères.","Birth place must be at most 120 characters and phone at most 40 characters.","يجب ألا يتجاوز مكان الميلاد 120 حرفًا ورقم الهاتف 40 حرفًا.")); return redirect("/profile/"+str(rid))
  if registration_date:
   try: parsed_registration_date=date.fromisoformat(registration_date)
   except ValueError:
    c.close(); flash(L("Saisissez une date d’inscription en première année valide.","Enter a valid first-year registration date.","أدخل تاريخ تسجيل صحيحًا في السنة الأولى.")); return redirect("/profile/"+str(rid))
   if parsed_registration_date>date.today():
    c.close(); flash(L("La date d’inscription ne peut pas être dans le futur.","Registration date cannot be in the future.","لا يمكن أن يكون تاريخ التسجيل في المستقبل.")); return redirect("/profile/"+str(rid))
  if birth_date:
   try: parsed_birth_date=date.fromisoformat(birth_date)
   except ValueError:
    c.close(); flash(L("Saisissez une date de naissance valide.","Enter a valid birth date.","أدخل تاريخ ميلاد صحيحًا.")); return redirect("/profile/"+str(rid))
   if parsed_birth_date>date.today():
    c.close(); flash(L("La date de naissance ne peut pas être dans le futur.","Birth date cannot be in the future.","لا يمكن أن يكون تاريخ الميلاد في المستقبل.")); return redirect("/profile/"+str(rid))
  # Keep the stored image unless a new valid upload is explicitly selected.
  photo_path=(r["photo_path"] or "").strip()
  photo_file=request.files.get("photo")
  if photo_file and photo_file.filename:
   ext=os.path.splitext(secure_filename(photo_file.filename))[1].lower()
   expected_formats={".png":"PNG",".jpg":"JPEG",".jpeg":"JPEG",".webp":"WEBP"}
   if ext not in expected_formats:
    c.close(); flash(L("La photo doit être au format PNG, JPG ou WEBP.","Photo must be PNG, JPG or WEBP.","يجب أن تكون الصورة بصيغة PNG أو JPG أو WEBP.")); return redirect("/profile/"+str(rid))
   photo_bytes=photo_file.stream.read(4*1024*1024+1); photo_file.stream.seek(0)
   if not photo_bytes or len(photo_bytes)>4*1024*1024:
    c.close(); flash(L("La photo doit être valide et ne pas dépasser 4 Mo.","Photo must be non-empty and 4 MB or smaller.","يجب أن تكون الصورة غير فارغة وألا يتجاوز حجمها 4 ميغابايت.")); return redirect("/profile/"+str(rid))
   try:
    with warnings.catch_warnings():
     warnings.simplefilter("error", Image.DecompressionBombWarning)
     with Image.open(io.BytesIO(photo_bytes)) as probe:
      if probe.format != expected_formats[ext]:
       raise ValueError("Image format does not match its filename.")
      probe.verify()
     with Image.open(io.BytesIO(photo_bytes)) as source:
      image=ImageOps.exif_transpose(source)
      image.load()
      image.thumbnail((1200,1200))
      output=io.BytesIO()
      has_alpha="A" in image.getbands()
      if expected_formats[ext]=="JPEG":
       image.convert("RGB").save(output,format="JPEG",quality=90,optimize=True)
      elif expected_formats[ext]=="PNG":
       image.convert("RGBA" if has_alpha else "RGB").save(output,format="PNG",optimize=True)
      else:
       image.convert("RGBA" if has_alpha else "RGB").save(output,format="WEBP",quality=90,method=6)
      normalized_photo=output.getvalue()
   except Exception:
    c.close(); flash(L("Le fichier choisi n’est pas une image PNG/JPG/WEBP lisible. Choisissez une vraie photo.","The selected file is not a readable PNG/JPG/WEBP image. Please choose a real photo.","الملف المختار ليس صورة PNG/JPG/WEBP قابلة للقراءة. اختر صورة حقيقية.")); return redirect("/profile/"+str(rid))
   photo_dir=os.path.join(os.path.dirname(os.path.abspath(DB)),"resident_photos"); os.makedirs(photo_dir,exist_ok=True)
   photo_name="resident_"+str(rid)+"_"+str(int(time.time()))+"_"+secrets.token_hex(5)+ext
   target=os.path.join(photo_dir,photo_name); temporary=target+".tmp"
   try:
    with open(temporary,"wb") as handle: handle.write(normalized_photo)
    os.replace(temporary,target)
   finally:
    if os.path.exists(temporary): os.remove(temporary)
   photo_path="/photos/"+photo_name
  c.execute("UPDATE residents SET surname=?,firstname=?,birth_place=?,birth_date=?,registration_date=?,photo_path=?,phone=?,phone2=?,address=?,email=?,resident_card_number=?,national_id_number=?,medical_order_registered=?,medical_order_number=?,social_security_number=?,additional_info=?,origin_hospital=?,external_start_date=?,external_end_date=? WHERE id=?",
   (surname,firstname,birth_place,birth_date,registration_date,photo_path,phone,phone2,address,email,resident_card_number,national_id_number,medical_order_registered,medical_order_number,social_security_number,additional_info,origin_hospital,external_start_date,external_end_date,rid))
  audit_event(c,"resident.profile_updated",f"resident_id={rid}; name={surname} {firstname}; photo_uploaded={bool(photo_file and photo_file.filename)}")
  c.commit(); r=c.execute("SELECT * FROM residents WHERE id=?",(rid,)).fetchone(); flash(tr("done"))
 history_rows=c.execute("SELECT h.*,t.name terrain,p.start_date AS period_start_date,p.end_date AS period_end_date FROM history h JOIN terrains t ON t.id=h.terrain_id LEFT JOIN periods p ON p.year_label=h.year_label AND p.name=h.period WHERE h.resident_id=? ORDER BY h.id DESC",(rid,)).fetchall()
 assignment_rows=c.execute("SELECT a.status,p.name period,p.year_label,t.name terrain FROM assignments a JOIN periods p ON p.id=a.period_id JOIN terrains t ON t.id=a.terrain_id WHERE a.resident_id=? ORDER BY p.year_label DESC,p.id",(rid,)).fetchall()
 objectives=c.execute("SELECT o.*,t.name terrain FROM objectives o JOIN terrains t ON t.id=o.terrain_id ORDER BY o.study_year,t.name").fetchall()
 exam_rows=c.execute("SELECT * FROM exam_records WHERE resident_id=? ORDER BY study_year,module",(rid,)).fetchall()
 done=set()
 today=date.today()
 profile_history=c.execute("SELECT terrain_id,year_label,period,date_done FROM history WHERE resident_id=?",(rid,)).fetchall()
 for item in profile_history:
  done_date=str(item["date_done"] or "").strip()
  if done_date:
   try:
    if date.fromisoformat(done_date)<=today: done.add(item["terrain_id"])
   except ValueError:
    continue
  else:
   try:
    _,period_end=rotation_period_bounds(item["year_label"],item["period"],item["period_start_date"],item["period_end_date"])
    if period_end<today: done.add(item["terrain_id"])
   except (ValueError,TypeError):
    continue
 profile_assignments=c.execute("SELECT a.terrain_id,p.year_label,p.name,p.start_date,p.end_date FROM assignments a JOIN periods p ON p.id=a.period_id WHERE a.resident_id=? AND a.status='confirmed'",(rid,)).fetchall()
 for item in profile_assignments:
  try:
   _,period_end=rotation_period_bounds(item["year_label"],item["name"],item["start_date"],item["end_date"])
   if period_end<today: done.add(item["terrain_id"])
  except (ValueError,TypeError):
   continue
 leave_rows=c.execute("SELECT start_date,end_date,days,notes FROM leave_entries WHERE resident_id=? ORDER BY start_date DESC",(rid,)).fetchall()
 external_rows=c.execute("SELECT supervisor,destination,hospital,country,rotation_type,start_date,end_date,notes,status,request_status FROM external_rotations WHERE resident_id=? ORDER BY start_date DESC",(rid,)).fetchall()
 leave_earned=max(accrued_leave_days(r["registration_date"],today),c.execute("SELECT COALESCE(SUM(earned_days),0) FROM leave_years WHERE resident_id=? AND year<=?",(rid,today.year)).fetchone()[0])
 leave_used=c.execute("SELECT COALESCE(SUM(days),0) FROM leave_entries WHERE resident_id=? AND end_date<=?",(rid,today.isoformat())).fetchone()[0]
 leave_remaining=leave_earned-leave_used
 terrains=c.execute("SELECT * FROM terrains ORDER BY CASE WHEN kind='optional' THEN 1 ELSE 0 END,name").fetchall(); c.close()
 return page("""<section class="rt-profile-panel"><div class="rt-profile-head"><div class="rt-avatar-wrap">{% if r.photo_path %}<img id="resident-photo-preview" class="rt-profile-avatar" src="{{r.photo_path}}" alt="{{L('Photo du résident','Resident photo','صورة المقيم')}}">{% else %}<div id="resident-photo-placeholder" class="rt-profile-avatar rt-avatar-placeholder" aria-label="{{L('Aucune photo du résident','No resident photo','لا توجد صورة للمقيم')}}">👤</div><img id="resident-photo-preview" class="rt-profile-avatar" alt="{{L('Aperçu de la photo du résident','Selected resident photo preview','معاينة صورة المقيم')}}" hidden>{% endif %}</div><div class="rt-profile-heading"><span class="rt-eyebrow">{{L('RESIMANAGE · DOSSIER DU RÉSIDENT','RESIMANAGE · RESIDENT FILE','RESIMANAGE · ملف المقيم')}}</span><h2>{{r.surname}} {{r.firstname}}</h2><p><span class="rt-chip">{{L('Année','Year','السنة')}} {{r.year}}</span>{% if r.resident_type=='external' %}<span class="rt-chip">{{L('Résident externe','External resident','مقيم خارجي')}}{% if r.origin_hospital %} — {{r.origin_hospital}}{% endif %}</span>{% endif %}<span class="rt-chip {{'rt-chip-muted' if r.status!='active' else 'rt-chip-active'}}">{{L('Actif' if r.status=='active' else 'Inactif' if r.status=='inactive' else 'Diplômé' if r.status=='graduated' else r.status,'Active' if r.status=='active' else 'Inactive' if r.status=='inactive' else 'Graduated' if r.status=='graduated' else r.status,'نشط' if r.status=='active' else 'غير نشط' if r.status=='inactive' else 'متخرج' if r.status=='graduated' else r.status)}}</span></p><p class="muted">{{L('Profil, avancement des stages, résultats d’examen et objectifs réunis au même endroit.','Profile, rotation progress, exam results and objectives in one place.','الملف وتقدم التربصات ونتائج الامتحانات والأهداف في مكان واحد.')}}</p></div></div><form class="rt-profile-form" method="post" enctype="multipart/form-data" accept-charset="UTF-8"><div class="rt-field-grid"><label>{{t('surname')}}<input name="surname" value="{{r.surname}}" autocomplete="family-name" maxlength="100" required></label><label>{{t('firstname')}}<input name="firstname" value="{{r.firstname}}" autocomplete="given-name" maxlength="100" required></label><label>{{L('Lieu de naissance','Birth place','مكان الميلاد')}}<input name="birth_place" value="{{r.birth_place or ''}}" maxlength="120"></label><label>{{L('Date de naissance','Birth date','تاريخ الميلاد')}}<input type="date" name="birth_date" value="{{r.birth_date or ''}}"></label><label>{{L('Date d’inscription en 1re année','First-year registration date','تاريخ التسجيل في السنة الأولى')}}<input type="date" name="registration_date" value="{{r.registration_date or ''}}" max="{{today_iso}}"></label><label>{{L('Téléphone 1','Phone 1','الهاتف 1')}}<input name="phone" value="{{r.phone or ''}}" type="tel" autocomplete="tel" maxlength="40"></label><label>{{L('Téléphone 2','Phone 2','الهاتف 2')}}<input name="phone2" value="{{r.phone2 or ''}}" type="tel" maxlength="40"></label><label>{{L('Adresse postale','Postal address','العنوان البريدي')}}<input name="address" value="{{r.address or ''}}" maxlength="500"></label><label>{{L('Adresse e-mail','Email address','البريد الإلكتروني')}}<input name="email" value="{{r.email or ''}}" type="email" maxlength="254"></label><label>{{L('Carte de résident — numéro','Resident card number','رقم بطاقة المقيم')}}<input name="resident_card_number" value="{{r.resident_card_number or ''}}" maxlength="100"></label><label>{{L('NIN / identification nationale','National ID / NIN','رقم التعريف الوطني')}}<input name="national_id_number" value="{{r.national_id_number or ''}}" maxlength="100"></label><label>{{L('Inscription à l’Ordre des médecins','Medical Order registration','التسجيل في هيئة الأطباء')}}<select name="medical_order_registered"><option value="yes" {% if r.medical_order_registered=='yes' %}selected{% endif %}>{{L('Inscrit','Registered','مسجل')}}</option><option value="no" {% if r.medical_order_registered!='yes' %}selected{% endif %}>{{L('Non inscrit','Not registered','غير مسجل')}}</option></select></label><label>{{L('Numéro de l’Ordre des médecins','Medical Order number','رقم هيئة الأطباء')}}<input name="medical_order_number" value="{{r.medical_order_number or ''}}" maxlength="100"></label><label>{{L('Numéro de sécurité sociale','Social security number','رقم الضمان الاجتماعي')}}<input name="social_security_number" value="{{r.social_security_number or ''}}" maxlength="100"></label>{% if r.resident_type=='external' %}<label>{{L('Établissement d’origine','Home institution','المؤسسة الأصلية')}}<input name="origin_hospital" value="{{r.origin_hospital or ''}}" required maxlength="200"></label><label>{{L('Début du stage au CHU','Start date at our hospital','بداية التربص بمستشفانا')}}<input type="date" name="external_start_date" value="{{r.external_start_date or ''}}" required></label><label>{{L('Fin du stage au CHU','End date at our hospital','نهاية التربص بمستشفانا')}}<input type="date" name="external_end_date" value="{{r.external_end_date or ''}}" required></label>{% endif %}<label>{{L('Autres informations utiles','Additional useful information','معلومات مفيدة أخرى')}}<textarea name="additional_info" maxlength="2000">{{r.additional_info or ''}}</textarea></label></div><div class="rt-photo-picker"><label for="resident-photo-input">{{L('Photo du résident','Resident photo','صورة المقيم')}}</label><p class="muted">{{L('Choisissez une véritable image PNG, JPG ou WEBP (4 Mo maximum). Un aperçu apparaît avant l’enregistrement.','Choose a real PNG, JPG or WEBP image (maximum 4 MB). A preview appears before saving.','اختر صورة حقيقية بصيغة PNG أو JPG أو WEBP (بحد أقصى 4 ميغابايت). ستظهر المعاينة قبل الحفظ.')}}</p><input id="resident-photo-input" type="file" name="photo" accept=".png,.jpg,.jpeg,.webp,image/png,image/jpeg,image/webp" aria-describedby="resident-photo-help"><small id="resident-photo-help">{{L('La photo existante est conservée si vous n’en choisissez pas une autre.','The existing photo is kept if you do not select a replacement.','سيتم الاحتفاظ بالصورة الحالية إذا لم تختر صورة بديلة.')}}</small><small id="resident-photo-status" aria-live="polite"></small></div><div class="rt-savebar"><button type="submit">{{L('Enregistrer le profil','Save profile','حفظ الملف')}}</button><a class="rt-secondary-action" href="/residents">{{L('Retour aux résidents','Back to residents','العودة إلى المقيمين')}}</a></div></form><p class="rt-profile-links"><a href="/exams?resident={{r.id}}">{{L('Examens','Exams','الامتحانات')}}</a><a href="/dems">{{L('Archive DEMS','DEMS archive','أرشيف DEMS')}}</a></p></section><section><h3>{{L('Avancement des stages','Rotation progress','تقدم التربصات')}}</h3><div class="wrap"><table><tr><th>{{L('Service de stage','Rotation department','مصلحة التربص')}}</th><th>{{L('État','Status','الحالة')}}</th></tr>{% for t in terrains %}<tr><td>{{t.name}}</td><td class="{{'bad' if t.id in done else 'good'}}">{{L('Effectué','Completed','مكتمل') if t.id in done else L('À faire','Remaining','متبقٍ')}}</td></tr>{% endfor %}</table></div></section><section><h3>{{L('Affectations et historique','Assignments and history','التعيينات والسجل')}}</h3>{% for a in assignment_rows %}<p>{{a.year_label}} — {{period_label(a.period)}} — {{a.terrain}} ({{L('Confirmée' if a.status=='confirmed' else 'En attente' if a.status=='pending' else 'Refusée' if a.status=='rejected' else 'Annulée' if a.status=='cancelled' else a.status,'Confirmed' if a.status=='confirmed' else 'Pending' if a.status=='pending' else 'Rejected' if a.status=='rejected' else 'Cancelled' if a.status=='cancelled' else a.status,'مؤكد' if a.status=='confirmed' else 'معلق' if a.status=='pending' else 'مرفوض' if a.status=='rejected' else 'ملغى' if a.status=='cancelled' else a.status)}})</p>{% endfor %}{% for h in history_rows %}<p>{{h.year_label}} — {{period_label(h.period)}} — {{h.terrain}} — {{h.date_done}}</p>{% endfor %}</section><section><h3>{{L('Congés annuels','Annual leave','الإجازات السنوية')}}</h3><p>{{L('Acquis :','Accrued:','المكتسب:')}} {{leave_earned}} {{L('jours · Utilisés :','days · Used:','أيام · المستعمل:')}} {{leave_used}} {{L('jours · Restants :','days · Remaining:','أيام · المتبقي:')}} {{leave_remaining}} {{L('jours','days','أيام')}}</p><div class="wrap"><table><tr><th>{{L('Début','Start','البداية')}}</th><th>{{L('Fin','End','النهاية')}}</th><th>{{L('Jours','Days','الأيام')}}</th><th>{{L('Notes','Notes','ملاحظات')}}</th></tr>{% for l in leave_rows %}<tr><td>{{l.start_date}}</td><td>{{l.end_date}}</td><td>{{l.days}}</td><td>{{l.notes}}</td></tr>{% else %}<tr><td colspan="4">{{t('none')}}</td></tr>{% endfor %}</table></div></section><section><h3>{{L('Stages à l’étranger','External internships','التربصات بالخارج')}}</h3>{% for x in external_rows %}<div style="padding:12px;margin:8px 0;border:1px solid #d6e4ee;border-radius:12px"><strong>{{x.destination}} — {{x.hospital}}, {{x.country}}</strong><p>{{x.rotation_type}} · {{x.start_date}} → {{x.end_date}} · {{L('En attente' if x.request_status=='pending' else 'Approuvé' if x.request_status=='approved' else 'Refusé' if x.request_status=='rejected' else 'Annulé' if x.request_status=='cancelled' else 'Demandé','Pending' if x.request_status=='pending' else 'Approved' if x.request_status=='approved' else 'Rejected' if x.request_status=='rejected' else 'Cancelled' if x.request_status=='cancelled' else 'Requested','معلق' if x.request_status=='pending' else 'موافق عليه' if x.request_status=='approved' else 'مرفوض' if x.request_status=='rejected' else 'ملغى' if x.request_status=='cancelled' else 'مطلوب')}} / {{L('Actif' if x.status=='active' else 'Terminé' if x.status=='returned' else 'Annulé' if x.status=='cancelled' else 'Demandé' if x.status=='requested' else 'Refusé' if x.status=='rejected' else x.status,'Active' if x.status=='active' else 'Returned' if x.status=='returned' else 'Cancelled' if x.status=='cancelled' else 'Requested' if x.status=='requested' else 'Rejected' if x.status=='rejected' else x.status,'نشط' if x.status=='active' else 'عاد' if x.status=='returned' else 'ملغى' if x.status=='cancelled' else 'مطلوب' if x.status=='requested' else 'مرفوض' if x.status=='rejected' else x.status)}}</p><p>{{L('Encadrant :','Supervisor:','المشرف:')}} {{x.supervisor}}</p><p>{{x.notes}}</p></div>{% else %}<p>{{t('none')}}</p>{% endfor %}</section><section><h3>{{L('Notes d’examen','Exam scores','علامات الامتحان')}}</h3><div class="wrap"><table><tr><th>{{L('Année d’étude','Study year','السنة الدراسية')}}</th><th>{{L('Module / insertion DEMS','Module / DEMS insertion','المادة / إدراج DEMS')}}</th><th>{{L('Session','Session','الدورة')}}</th><th>{{L('Tentative','Attempt','المحاولة')}}</th><th>{{L('Partie 1 / Théorie','Part 1 / Theory','الجزء الأول / النظري')}}</th><th>{{L('Partie 2 / Pratique','Part 2 / Practical','الجزء الثاني / التطبيقي')}}</th></tr>{% for e in exam_rows %}<tr><td>{{e.study_year}}</td><td>{{e.module}}</td><td>{{L('Virtuelle' if e.exam_type=='virtual' else 'Réelle' if e.exam_type=='actual' else 'Rattrapage' if e.exam_type=='remedial' else 'DEMS virtuel' if e.exam_type=='dems_virtual' else 'DEMS réel' if e.exam_type=='dems_actual' else 'DEMS rattrapage' if e.exam_type=='dems_remedial' else 'Insertion DEMS (deux parties)' if e.exam_type=='dems_insertion' else e.exam_type,'Virtual' if e.exam_type=='virtual' else 'Actual' if e.exam_type=='actual' else 'Remedial' if e.exam_type=='remedial' else 'DEMS virtual' if e.exam_type=='dems_virtual' else 'DEMS actual' if e.exam_type=='dems_actual' else 'DEMS remedial' if e.exam_type=='dems_remedial' else 'DEMS insertion (two parts)' if e.exam_type=='dems_insertion' else e.exam_type,'افتراضي' if e.exam_type=='virtual' else 'فعلي' if e.exam_type=='actual' else 'استدراكي' if e.exam_type=='remedial' else 'DEMS افتراضي' if e.exam_type=='dems_virtual' else 'DEMS فعلي' if e.exam_type=='dems_actual' else 'DEMS استدراكي' if e.exam_type=='dems_remedial' else 'إدراج DEMS (جزآن)' if e.exam_type=='dems_insertion' else e.exam_type)}}</td><td>{{e.attempt}}</td><td>{{e.theory}}/20</td><td>{{e.practical}}/20</td></tr>{% else %}<tr><td colspan="6">{{t('none')}}</td></tr>{% endfor %}</table></div></section><section><h3>{{L('Objectifs','Objectives','الأهداف')}}</h3>{% for o in objectives %}<div style="padding:10px;margin:8px 0;border-radius:10px;background:{{'#fee2e2' if o.completed else '#dcfce7'}};color:{{'#991b1b' if o.completed else '#166534'}}"><strong>{{L('Année','Year','السنة')}} {{o.study_year}} — {{o.terrain}}</strong><p>{{o.objective}}</p><small>{{L('Effectué','Completed','مكتمل') if o.completed else L('Non effectué','Not done','غير مكتمل')}}</small></div>{% else %}<p>{{t('none')}}</p>{% endfor %}</section><script>(function(){const lang="{{lang}}";const msg=(fr,en,ar)=>lang==="fr"?fr:lang==="ar"?ar:en;const input=document.getElementById('resident-photo-input');const preview=document.getElementById('resident-photo-preview');const placeholder=document.getElementById('resident-photo-placeholder');if(!input||!preview)return;const status=document.getElementById('resident-photo-status');input.addEventListener('change',function(){const file=input.files&&input.files[0];if(!file)return;const ext=(file.name.split('.').pop()||'').toLowerCase();if(!['png','jpg','jpeg','webp'].includes(ext)){if(status)status.textContent=msg('Choisissez une photo PNG, JPG ou WEBP.','Please choose a PNG, JPG or WEBP photo.','اختر صورة بصيغة PNG أو JPG أو WEBP.');input.value='';return;}if(file.size>4*1024*1024){if(status)status.textContent=msg('La photo ne doit pas dépasser 4 Mo.','The photo must be 4 MB or smaller.','يجب ألا يتجاوز حجم الصورة 4 ميغابايت.');input.value='';return;}const reader=new FileReader();reader.onload=function(){if(typeof reader.result!=='string'||!reader.result.startsWith('data:image/')){if(status)status.textContent=msg('Impossible de prévisualiser ce fichier. Essayez une autre photo.','The selected file could not be previewed. Try another photo.','تعذرت معاينة الملف المختار. جرّب صورة أخرى.');return;}preview.src=reader.result;preview.hidden=false;if(placeholder)placeholder.hidden=true;if(status)status.textContent=msg('Aperçu prêt. Enregistrez le profil pour importer la photo.','Preview ready. Save the profile to upload the photo.','المعاينة جاهزة. احفظ الملف لتحميل الصورة.');};reader.onerror=function(){if(status)status.textContent=msg('Le téléphone ne peut pas lire cette photo. Essayez une autre image ou copiez-la dans Téléchargements.','The phone could not read this photo. Try another image or copy it to Downloads.','تعذر على الهاتف قراءة الصورة. جرّب صورة أخرى أو انسخها إلى مجلد التنزيلات.');};reader.readAsDataURL(file);});})();</script>""",r=r,terrains=terrains,done=done,assignment_rows=assignment_rows,history_rows=history_rows,objectives=objectives,exam_rows=exam_rows,leave_rows=leave_rows,leave_earned=leave_earned,leave_used=leave_used,leave_remaining=leave_remaining,external_rows=external_rows,today_iso=date.today().isoformat())
@app.route("/objectives",methods=["GET","POST"])
def manage_objectives():
 if super_guard(): return super_guard()
 c=db()
 if request.method=="POST":
  action=request.form.get("action","create")
  if action=="toggle":
   try:
    oid=int(request.form["objective_id"]); completed=int(request.form["completed"])
    if completed not in (0,1): raise ValueError("Invalid completion status.")
    objective=c.execute("SELECT * FROM objectives WHERE id=?",(oid,)).fetchone()
    if not objective: raise ValueError("Objective not found.")
    c.execute("UPDATE objectives SET completed=? WHERE id=?",(completed,oid))
    audit_event(c,"objective.completed" if completed else "objective.reopened",f"objective_id={oid}; terrain_id={objective['terrain_id']}; study_year={objective['study_year']}")
    c.commit(); flash(tr("done"))
   except (ValueError,TypeError,KeyError,sqlite3.IntegrityError) as exc:
    c.rollback(); flash(str(exc) or "Invalid objective status.")
  else:
   try:
    tid=int(request.form["terrain"]); sy=int(request.form["study_year"]); obj=request.form.get("objective","").strip()
    terrain=c.execute("SELECT id FROM terrains WHERE id=? AND active=1",(tid,)).fetchone()
    if not terrain: raise ValueError(L("Choisissez un service actif.","Choose an active department.","اختر مصلحة نشطة."))
    if sy not in (1,2,3,4): raise ValueError("Objective study year must be 1, 2, 3 or 4.")
    if not obj or len(obj)>1000: raise ValueError("Objective text is required and must be at most 1000 characters.")
    duplicate=c.execute("SELECT 1 FROM objectives WHERE terrain_id=? AND study_year=? AND lower(objective)=lower(?)",(tid,sy,obj)).fetchone()
    if duplicate: raise ValueError("This objective already exists for the selected terrain and year.")
    c.execute("INSERT INTO objectives(terrain_id,study_year,objective) VALUES(?,?,?)",(tid,sy,obj))
    audit_event(c,"objective.created",f"terrain_id={tid}; study_year={sy}")
    c.commit(); flash(tr("done"))
   except (ValueError,TypeError,KeyError,sqlite3.IntegrityError) as exc:
    c.rollback(); flash(str(exc) or "Invalid objective input.")
 ts=c.execute("SELECT * FROM terrains WHERE active=1 ORDER BY name").fetchall()
 rows=c.execute("SELECT o.*,t.name terrain FROM objectives o JOIN terrains t ON t.id=o.terrain_id ORDER BY study_year,t.name,o.id").fetchall(); c.close()
 return page("""<section><h3>{{t('objectives')}}</h3><form method="post"><input type="hidden" name="action" value="create"><label>{{t('terrain')}}</label><select name="terrain">{% for t in ts %}<option value="{{t.id}}">{{t.name}}</option>{% endfor %}</select><label>{{t('study_year')}}</label><select name="study_year">{% for y in range(1,5) %}<option>{{y}}</option>{% endfor %}</select><label>{{t('objective')}}</label><input name="objective" maxlength="1000" required><button>{{t('save')}}</button></form></section><section><h3>{{L('Avancement par service et année d’étude','Progress by department and study year','التقدم حسب المصلحة والسنة الدراسية')}}</h3>{% for o in rows %}<div style="display:flex;flex-wrap:wrap;align-items:center;gap:10px;padding:12px;margin:8px 0;border-radius:12px;background:{{'#fee2e2' if o.completed else '#dcfce7'}};color:{{'#991b1b' if o.completed else '#166534'}};border:1px solid {{'#fca5a5' if o.completed else '#86efac'}}"><div style="flex:1;min-width:180px"><strong>{{L('Année','Year','السنة')}} {{o.study_year}} — {{o.terrain}}</strong><p style="margin:6px 0 0">{{o.objective}}</p><small>{{L('Effectué','Completed','مكتمل') if o.completed else L('Non effectué','Not done','غير مكتمل')}}</small></div><form method="post"><input type="hidden" name="action" value="toggle"><input type="hidden" name="objective_id" value="{{o.id}}"><input type="hidden" name="completed" value="{{0 if o.completed else 1}}"><button type="submit">{{'Mark not done' if o.completed else 'Mark completed'}}</button></form></div>{% else %}<p>{{t('none')}}</p>{% endfor %}</section>""",ts=ts,rows=rows)
@app.route("/exams",methods=["GET","POST"])
def exams():
 if guard(): return redirect("/login")
 c=db()
 if request.method=="POST":
  action=request.form.get("action","record")
  if action=="delete_record":
   if session.get("role")!="super_admin":
    c.close(); return ("Forbidden — Super Admin only",403)
   try:
    exam_id=int(request.form.get("exam_id","0"))
    old_record=c.execute("SELECT * FROM exam_records WHERE id=?",(exam_id,)).fetchone()
    if not old_record: raise ValueError("Exam record not found.")
    reason=L("Suppression demandée par le Super Admin","Deletion requested by Super Admin","طلب المسؤول العام الحذف")
    c.execute("INSERT INTO trash_items(entity_type,entity_id,label,reason,snapshot_json,deleted_at,deleted_by) VALUES(?,?,?,?,?,?,?)",
      ("exam_record",str(exam_id),str(old_record["module"])+" — "+str(old_record["theory"])+"/"+str(old_record["practical"]),reason,json.dumps({"record":dict(old_record)},ensure_ascii=False),datetime.now().replace(microsecond=0).isoformat(),session.get("username","")))
    audit_event(c,"exam.deleted",f"exam_id={exam_id}; resident_id={old_record['resident_id']}; postgraduate_year={old_record['postgrad_year']}; study_year={old_record['study_year']}; module={old_record['module']}; session={old_record['exam_type']}; attempt={old_record['attempt']}; theory={old_record['theory']}; practical={old_record['practical']}; moved_to_trash_for_six_months=true")
    c.execute("DELETE FROM exam_records WHERE id=?",(exam_id,))
    c.commit(); flash(L("Résultat déplacé dans la corbeille pour six mois.","Result moved to the trash for six months.","تم نقل النتيجة إلى سلة المهملات لمدة ستة أشهر."))
   except Exception as exc:
    c.rollback(); flash("Suppression impossible : "+str(exc)[:200])
   c.close(); return redirect("/exams")
  if action=="assign_postgrad_year":
   if session.get("role")!="super_admin":
    c.close(); return ("Forbidden — Super Admin only",403)
   try:
    exam_id=int(request.form.get("exam_id","0"))
    postgrad_year=request.form.get("postgrad_year","").strip()
    if not re.fullmatch(r"\d{4}-\d{4}",postgrad_year) or int(postgrad_year.split("-")[1])!=int(postgrad_year.split("-")[0])+1:
     raise ValueError("Choose a valid postgraduate year.")
    old_record=c.execute("SELECT id,resident_id,study_year,module,postgrad_year FROM exam_records WHERE id=?",(exam_id,)).fetchone()
    if not old_record: raise ValueError("Exam record not found.")
    c.execute("UPDATE exam_records SET postgrad_year=? WHERE id=?",(postgrad_year,exam_id))
    audit_event(c,"exam.postgrad_year_assigned",f"exam_id={exam_id}; resident_id={old_record['resident_id']}; module={old_record['module']}; study_year={old_record['study_year']}; before={old_record['postgrad_year']}; after={postgrad_year}")
    c.commit(); flash("Année de post-graduation associée au résultat.")
   except Exception as exc:
    c.rollback(); flash("Association impossible : "+str(exc)[:200])
   c.close(); return redirect("/exams")
  try:
   rid=int(request.form.get("resident",""))
   sy=int(request.form.get("study_year",""))
  except (TypeError,ValueError):
   c.close(); flash(L("Choisissez un résident et une année d’étude valides.","Choose a valid resident and study year.","اختر مقيمًا وسنة دراسية صالحين.")); return redirect("/exams")
  resident=c.execute("SELECT * FROM residents WHERE id=?",(rid,)).fetchone()
  if not resident:
   flash(L("Résident introuvable.","Resident not found.","لم يتم العثور على المقيم."))
  elif action=="record":
   if resident["resident_type"]=="external":
    flash(L("Les résidents externes ne relèvent pas du processus interne d’examens et de promotion du CHU.","External residents are not part of the hospital's internal exam and progression process.","المقيمون الخارجيون لا يخضعون لمسار الامتحانات والترقية الداخلي للمستشفى."))
    c.close(); return redirect("/exams")
   if session.get("role")!="super_admin":
    c.close(); return ("Forbidden — Super Admin only",403)
   if resident["status"]=="graduated" or sy>int(resident["year"]):
    flash(L("Impossible d’enregistrer des notes pour une année future ou pour un diplômé archivé du DEMS.","Cannot record marks for a future study year or for an archived DEMS graduate.","لا يمكن تسجيل علامات لسنة دراسية مستقبلية أو لخريج DEMS مؤرشف."))
    c.close(); return redirect("/exams")
   try:
    module=request.form.get("module","").strip()
    exam_type=request.form.get("exam_type","actual").strip()
    attempt=int(request.form.get("attempt","1"))
    theory=float(request.form.get("theory",""))
    practical=float(request.form.get("practical",""))
   except (TypeError,ValueError):
    c.close(); flash(L("Le module et les notes numériques de 0 à 20 sont obligatoires ; choisissez une session et une tentative valides.","Module and numeric marks from 0 to 20 are required; choose a valid exam session and attempt.","المادة والعلامات الرقمية من 0 إلى 20 إجبارية؛ اختر دورة ومحاولة صحيحتين.")); return redirect("/exams")
   allowed_exam_types=("virtual","actual","remedial","dems_virtual","dems_actual","dems_remedial","dems_insertion")
   dems_types=("dems_virtual","dems_actual","dems_remedial","dems_insertion")
   if sy not in (1,2,3,4) or not module or exam_type not in allowed_exam_types or (exam_type in dems_types and sy!=4) or attempt not in (1,2,3) or not (0<=theory<=20 and 0<=practical<=20):
    flash(L("Choisissez une session et une année valides, une tentative de 1 à 3 et des notes de 0 à 20. Les sessions DEMS concernent la 4e année.","Choose a valid session/year, attempt 1–3 and marks from 0 to 20. DEMS sessions are for year 4.","اختر دورة وسنة صحيحتين ومحاولة من 1 إلى 3 وعلامات من 0 إلى 20. دورات DEMS تخص السنة الرابعة."))
   else:
    postgrad_year=active_postgrad_year(c)
    existing=c.execute("SELECT id FROM exam_records WHERE resident_id=? AND study_year=? AND lower(module)=lower(?) AND exam_type=? AND attempt=? AND postgrad_year=?",(rid,sy,module,exam_type,attempt,postgrad_year)).fetchone()
    if existing:
     c.execute("UPDATE exam_records SET theory=?,practical=?,created_at=CURRENT_TIMESTAMP WHERE id=?",(theory,practical,existing["id"]))
     audit_event(c,"exam.updated",f"resident_id={rid}; postgraduate_year={postgrad_year}; study_year={sy}; module={module}; session={exam_type}; attempt={attempt}")
    else:
     c.execute("INSERT INTO exam_records(resident_id,study_year,module,theory,practical,exam_type,attempt,postgrad_year) VALUES(?,?,?,?,?,?,?,?)",(rid,sy,module,theory,practical,exam_type,attempt,postgrad_year))
     audit_event(c,"exam.recorded",f"resident_id={rid}; postgraduate_year={postgrad_year}; study_year={sy}; module={module}; session={exam_type}; attempt={attempt}")
    c.commit(); flash(tr("done"))
  elif action=="evaluate_first_year" and request.form.get("all_modules_confirmed")!="yes":
   flash(L("Confirmez que les résultats de tous les modules de première année ont été saisis avant d’évaluer la promotion.","Confirm that all first-year module results have been entered before evaluating progression.","أكد إدخال نتائج جميع مواد السنة الأولى قبل تقييم الترقية."))
  elif action=="evaluate_first_year":
   if session.get("role")!="super_admin": c.close(); return ("Forbidden — Super Admin only",403)
   if resident["year"]!=1:
    flash(L("Seuls les résidents actuellement en première année peuvent être évalués pour le passage en deuxième année.","Only current first-year residents can be evaluated for first-year progression.","يمكن تقييم المقيمين في السنة الأولى حاليًا فقط للانتقال إلى السنة الثانية."))
   elif resident["status"]!="active":
    flash(L("Le résident est inactif ; la promotion n’est pas disponible.","The resident is inactive; progression is not available.","المقيم غير نشط؛ الترقية غير متاحة."))
   else:
    postgrad_year=active_postgrad_year(c)
    raw_records=c.execute("SELECT id,module,theory,practical,exam_type,attempt FROM exam_records WHERE resident_id=? AND study_year=1 AND postgrad_year=? AND exam_type IN ('actual','remedial') ORDER BY id",(rid,postgrad_year)).fetchall()
    chosen={}
    for record in raw_records:
     key=record["module"].strip().casefold()
     previous=chosen.get(key)
     # A remedial result replaces the ordinary actual result for that module;
     # otherwise the highest recorded attempt in the same session is used.
     priority=(1 if record["exam_type"]=="remedial" else 0,record["attempt"],record["id"])
     if previous is None or priority>(1 if previous["exam_type"]=="remedial" else 0,previous["attempt"],previous["id"]):
      chosen[key]=record
    records=list(chosen.values())
    all_modules={row["module"].strip().casefold() for row in c.execute("SELECT DISTINCT module FROM exam_records WHERE resident_id=? AND study_year=1 AND postgrad_year=?",(rid,postgrad_year)).fetchall()}
    missing_actual_or_remedial=all_modules-set(chosen.keys())
    if not records or missing_actual_or_remedial:
     flash(L("Saisissez les notes réelles ou de rattrapage de chaque module de première année avant l’évaluation ; les notes virtuelles seules ne comptent pas.","Enter actual or remedial marks for every first-year module before evaluation; virtual marks alone do not count.","أدخل العلامات الفعلية أو الاستدراكية لكل مادة في السنة الأولى قبل التقييم؛ العلامات الافتراضية وحدها لا تُحتسب."))
    else:
     average=sum((x["theory"]+x["practical"])/2 for x in records)/len(records)
     passed=average>=10 and all(x["theory"]>=10 and x["practical"]>=10 for x in records)
     c.execute("INSERT INTO decisions(details,decision,reason,actor,postgrad_year) VALUES(?,?,?,?,?)",(f"{resident['surname']} {resident['firstname']} — first-year exams — postgraduate year {postgrad_year}","Passed" if passed else "Failed",f"Average={average:.2f}/20; all theory/practical marks >=10: {all(x['theory']>=10 and x['practical']>=10 for x in records)}; postgraduate_year={postgrad_year}",session["user"],postgrad_year))
     audit_event(c,"exam.first_year_evaluated",f"resident_id={rid}; average={average:.2f}; passed={passed}")
     advanced=False
     if passed and resident["year"]==1 and resident["status"]=="active":
      c.execute("UPDATE residents SET year=2 WHERE id=? AND year=1 AND status='active'",(rid,))
      advanced=c.execute("SELECT changes()").fetchone()[0]==1
      if advanced:
       audit_event(c,"resident.year_advanced",f"resident_id={rid}; from=1; to=2")
     c.commit()
     if not passed:
      flash(L("Échec : la moyenne doit être d’au moins 10/20 et chaque note théorique/pratique doit être d’au moins 10/20.","Failed: average must be at least 10/20 and every theory/practical mark must be at least 10/20.","رسوب: يجب أن يكون المعدل 10/20 على الأقل وكل علامة نظرية/تطبيقية 10/20 على الأقل."))
     elif resident["year"]!=1:
      flash(L("Critères d’examen validés ; aucune année modifiée car le résident n’est plus en première année.","Exam criteria passed; no year change because the resident is not currently in year 1.","تم استيفاء شروط الامتحان؛ لم تتغير السنة لأن المقيم ليس في السنة الأولى حاليًا."))
     elif resident["status"]!="active":
      flash(L("Critères validés, mais le résident est inactif ; aucune promotion n’a été appliquée.","Exam criteria passed, but the resident is inactive; no progression was applied.","تم استيفاء الشروط، لكن المقيم غير نشط؛ لم تُطبق الترقية."))
     elif advanced:
      flash(L("Réussite ; le résident est passé en deuxième année.","Passed; resident advanced to year 2.","نجاح؛ انتقل المقيم إلى السنة الثانية."))
     else:
      flash(L("Critères validés, mais la promotion n’a pas été appliquée ; vérifiez le dossier du résident.","Exam criteria passed, but progression was not applied; review the resident record.","تم استيفاء الشروط لكن لم تُطبق الترقية؛ راجع ملف المقيم."))
  elif action=="approve_advance":
   actor=session.get("user","").lower()
   outcome=request.form.get("outcome","approve")
   reason=request.form.get("reason","").strip()
   if session.get("role")!="super_admin" or actor not in ("benhocine","drdris"):
    c.close(); return ("Only the official Super Admins Dr Dris and Pr Benhocine may vote on progression.",403)
   if outcome not in ("approve","reject") or not reason:
    flash(L("Choisissez approuver ou refuser et indiquez le motif de la décision.","Choose approve/reject and enter the decision reason.","اختر الموافقة أو الرفض وأدخل سبب القرار."))
   elif resident["status"]!="active":
    flash(L("Un résident inactif ne peut pas être promu ; aucun vote de progression n’a été enregistré.","Inactive residents cannot be advanced; no progression vote was recorded.","لا يمكن ترقية مقيم غير نشط؛ لم يتم تسجيل تصويت الترقية."))
   elif sy not in (2,3) or resident["year"]!=sy:
    flash(L("Seuls les résidents actifs de deuxième ou troisième année peuvent faire l’objet du vote conjoint des deux responsables.","Only active year 2 or 3 residents can use the dual-official progression vote.","يمكن فقط للمقيمين النشطين في السنة الثانية أو الثالثة الخضوع لتصويت المسؤولين الاثنين."))
   else:
    c.execute("INSERT INTO promotion_votes(resident_id,from_year,decision,reason,actor) VALUES(?,?,?,?,?) ON CONFLICT(resident_id,from_year,actor) DO UPDATE SET decision=excluded.decision,reason=excluded.reason,created_at=CURRENT_TIMESTAMP",(rid,sy,outcome,reason,actor))
    audit_event(c,"promotion.vote_recorded",f"resident_id={rid}; from_year={sy}; decision={outcome}; reason={reason}")
    votes=c.execute("SELECT actor,decision,reason FROM promotion_votes WHERE resident_id=? AND from_year=? AND actor IN ('benhocine','drdris')",(rid,sy)).fetchall()
    if len(votes)<2:
     c.commit(); flash(L("Vote enregistré. En attente du vote de l’autre responsable officiel.","Vote recorded. Waiting for the other official administrator.","تم تسجيل التصويت. في انتظار تصويت المسؤول الرسمي الآخر."))
    else:
     combined_reason=" | ".join(v["actor"]+": "+v["decision"]+" — "+v["reason"] for v in votes)
     actor_names="+".join(sorted(v["actor"] for v in votes))
     approved=all(v["decision"]=="approve" for v in votes)
     if approved:
      c.execute("UPDATE residents SET year=? WHERE id=? AND year=? AND status='active'",(sy+1,rid,sy))
      if c.execute("SELECT changes()").fetchone()[0]:
       c.execute("INSERT INTO decisions(details,decision,reason,actor,postgrad_year) VALUES(?,?,?,?,?)",(f"{resident['surname']} {resident['firstname']} — year {sy}","Approved by both official administrators",combined_reason,actor_names,active_postgrad_year(c)))
       audit_event(c,"resident.year_advanced",f"resident_id={rid}; from={sy}; to={sy+1}; officials={actor_names}")
       flash(L("Approuvé par les deux responsables officiels ; le résident a été promu.","Approved by both official administrators; resident advanced.","وافق المسؤولان الرسميان؛ تمت ترقية المقيم."))
      else:
       flash(L("Le statut du résident a changé avant la finalisation du vote ; aucune promotion n’a été appliquée.","Resident status changed before the vote was finalized; no progression was applied.","تغيرت حالة المقيم قبل إتمام التصويت؛ لم تُطبق الترقية."))
     else:
      c.execute("INSERT INTO decisions(details,decision,reason,actor,postgrad_year) VALUES(?,?,?,?,?)",(f"{resident['surname']} {resident['firstname']} — year {sy}","Rejected by official vote",combined_reason,actor_names,active_postgrad_year(c)))
      audit_event(c,"promotion.vote_rejected",f"resident_id={rid}; from_year={sy}; officials={actor_names}")
      flash(L("Promotion refusée ; les deux votes officiels ont été enregistrés.","Progression not approved; both official votes are recorded.","لم تتم الموافقة على الترقية؛ تم تسجيل التصويتين الرسميين."))
     c.execute("DELETE FROM promotion_votes WHERE resident_id=? AND from_year=?",(rid,sy))
     c.commit()
 selected_year=active_postgrad_year(c)
 rows=c.execute("SELECT e.*,r.surname,r.firstname FROM exam_records e JOIN residents r ON r.id=e.resident_id WHERE e.postgrad_year=? ORDER BY e.created_at DESC,e.id DESC",(selected_year,)).fetchall()
 unassigned_rows=c.execute("SELECT e.*,r.surname,r.firstname FROM exam_records e JOIN residents r ON r.id=e.resident_id WHERE COALESCE(e.postgrad_year,'')='' ORDER BY e.created_at DESC,e.id DESC").fetchall()
 rs=c.execute("SELECT * FROM residents WHERE status='active' AND COALESCE(resident_type,'local')<>'external' ORDER BY year DESC,surname").fetchall()
 grade_table_rows=[]
 for resident_row in rs:
  for study_year in range(1,5):
   grade_records=c.execute("SELECT module,theory,practical,exam_type,attempt,id FROM exam_records WHERE resident_id=? AND study_year=? AND postgrad_year=? AND exam_type IN ('actual','remedial') ORDER BY id",(resident_row["id"],study_year,selected_year)).fetchall()
   chosen_marks={}
   for grade_record in grade_records:
    key=grade_record["module"].strip().casefold()
    old=chosen_marks.get(key)
    rank=(1 if grade_record["exam_type"]=="remedial" else 0,grade_record["attempt"],grade_record["id"])
    old_rank=(1 if old["exam_type"]=="remedial" else 0,old["attempt"],old["id"]) if old else None
    if old is None or rank>old_rank: chosen_marks[key]=grade_record
   marks=list(chosen_marks.values())
   average=sum((m["theory"]+m["practical"])/2 for m in marks)/len(marks) if marks else None
   passed_first=bool(marks) and average>=10 and all(m["theory"]>=10 and m["practical"]>=10 for m in marks)
   grade_table_rows.append({"resident_id":resident_row["id"],"surname":resident_row["surname"],"firstname":resident_row["firstname"],"current_year":resident_row["year"],"study_year":study_year,"modules":len(marks),"average":average,"passed_first":passed_first,"is_current":resident_row["year"]==study_year,"status":resident_row["status"]})
 selected=request.args.get("resident",type=int)
 pending_votes=c.execute("SELECT v.*,r.surname,r.firstname FROM promotion_votes v JOIN residents r ON r.id=v.resident_id ORDER BY v.created_at DESC").fetchall()
 year_options=available_postgrad_years(c)
 c.close()
 return page("""<section><h3>{{t('exams')}} — {{year}}</h3><p>{{L('Les résultats enregistrés ci-dessous sont rattachés à l’année de post-graduation sélectionnée. Les anciens résultats sans année connue restent à associer manuellement.','The results below are linked to the selected postgraduate year. Older results without a known year must be assigned manually.','ترتبط النتائج المسجلة أدناه بسنة ما بعد التدرج المحددة. يجب ربط النتائج القديمة التي لا تعرف سنتها يدويًا.')}}</p>{% if role=='super_admin' and unassigned_rows %}<details><summary>{{L('Résultats historiques à rattacher à une année','Historical results to assign to a year','نتائج تاريخية لربطها بسنة')}} ({{unassigned_rows|length}})</summary><div class="wrap"><table><tr><th>{{L('Résident','Resident','المقيم')}}</th><th>{{L('Année d’études','Study year','السنة الدراسية')}}</th><th>{{L('Module','Module','المادة')}}</th><th>{{L('Notes','Marks','العلامات')}}</th><th>{{L('Année de post-graduation à confirmer','Postgraduate year to confirm','سنة ما بعد التدرج للتأكيد')}}</th></tr>{% for e in unassigned_rows %}<tr><td>{{e.surname}} {{e.firstname}}</td><td>{{e.study_year}}</td><td>{{e.module}}</td><td>{{e.theory}} / {{e.practical}}</td><td><form method="post"><input type="hidden" name="action" value="assign_postgrad_year"><input type="hidden" name="exam_id" value="{{e.id}}"><select name="postgrad_year" required>{% for yy in year_options %}<option value="{{yy}}" {% if yy==year %}selected{% endif %}>{{yy}}</option>{% endfor %}</select><button type="submit">{{L('Associer','Assign year','ربط السنة')}}</button></form></td></tr>{% endfor %}</table></div></details>{% endif %}<form method="post"><input type="hidden" name="action" value="record"><label>{{L('Résident','Resident','المقيم')}}</label><select name="resident">{% for r in rs %}<option value="{{r.id}}" {% if selected==r.id %}selected{% endif %}>{{r.year}} — {{r.surname}} {{r.firstname}}</option>{% endfor %}</select><label>{{L('Année d’études','Study year','السنة الدراسية')}}</label><select id="exam-study-year" name="study_year">{% for y in range(1,5) %}<option>{{y}}</option>{% endfor %}</select><label>{{L('Session d’examen','Exam session','دورة الامتحان')}}</label><select id="exam-session-type" name="exam_type"><option value="virtual">{{L('Virtuelle','Virtual','افتراضي')}}</option><option value="actual" selected>{{L('Réelle','Actual','فعلي')}}</option><option value="remedial">{{L('Rattrapage','Remedial','استدراكي')}}</option><option value="dems_virtual">{{L('DEMS virtuel (4e année)','DEMS virtual (year 4)','DEMS افتراضي (السنة الرابعة)')}}</option><option value="dems_actual">{{L('DEMS réel (4e année)','DEMS actual (year 4)','DEMS فعلي (السنة الرابعة)')}}</option><option value="dems_remedial">{{L('DEMS rattrapage (4e année)','DEMS remedial (year 4)','DEMS استدراكي (السنة الرابعة)')}}</option><option value="dems_insertion">{{L('Insertion DEMS — deux parties (4e année)','DEMS insertion — two parts (year 4)','إدراج DEMS — جزآن (السنة الرابعة)')}}</option></select><label>{{L('Tentative','Attempt','المحاولة')}}</label><select name="attempt"><option>1</option><option>2</option><option>3</option></select><label>{{t('module')}}</label><input name="module" required><div class="grid"><div><label>{{t('theory')}}</label><input type="number" name="theory" min="0" max="20" step="0.01" required></div><div><label>{{t('practical')}}</label><input type="number" name="practical" min="0" max="20" step="0.01" required></div></div><button>{{L('Enregistrer la note','Record exam','تسجيل العلامة')}}</button></form><script>(function(){const y=document.getElementById("exam-study-year"),t=document.getElementById("exam-session-type");if(!y||!t)return;function syncExamTypes(){const isYearFour=y.value==="4";Array.from(t.options).forEach(function(o){if(o.value.indexOf("dems_")===0)o.disabled=!isYearFour;});if(!isYearFour&&t.value.indexOf("dems_")===0)t.value="actual";}y.addEventListener("change",syncExamTypes);syncExamTypes();})();</script></section><section><h3>{{L('Tableau des notes et validation du passage','Marks and promotion validation table','جدول العلامات والتحقق من الترقية')}} — {{year}}</h3><p>{{L('Les notes et les décisions sont regroupées par année d’études et année de post-graduation. Les décisions de passage se prennent dans ce tableau. Pour la première année, seules les notes réelles ou de rattrapage comptent et tous les modules doivent être renseignés.','Marks and decisions are grouped by study year and postgraduate year. Promotion decisions are made in this table. For first year, only actual or remedial marks count and every module must be entered.','تُنظم العلامات والقرارات حسب السنة الدراسية وسنة ما بعد التدرج. تُتخذ قرارات الترقية في هذا الجدول. في السنة الأولى تُحتسب العلامات الفعلية أو الاستدراكية فقط ويجب إدخال جميع المواد.')}}</p><div class="wrap"><table><tr><th>{{L('Résident','Resident','المقيم')}}</th><th>{{L('Année évaluée','Year evaluated','السنة المُقيّمة')}}</th><th>{{L('Modules notés','Modules graded','المواد المسجلة')}}</th><th>{{L('Moyenne /20','Average /20','المعدل /20')}}</th><th>{{L('Critères','Criteria','المعايير')}}</th><th>{{L('Validation / décision','Validation / decision','التحقق / القرار')}}</th></tr>{% for g in grade_table_rows if g.is_current %}<tr><td>{{g.surname}} {{g.firstname}}</td><td>{{g.study_year}}e année</td><td>{{g.modules}}</td><td>{% if g.average is not none %}{{'%.2f'|format(g.average)}}{% else %}—{% endif %}</td><td>{% if g.study_year==1 %}{% if g.passed_first %}<span style="color:#15803d;font-weight:bold">{{L('Critères de notes atteints','Mark criteria met','تم استيفاء معايير العلامات')}}</span>{% elif g.modules %}<span style="color:#b91c1c">{{L('Critères non atteints / vérifier les notes manquantes','Criteria not met / check missing marks','المعايير غير مستوفاة / تحقق من العلامات الناقصة')}}</span>{% else %}{{L('Notes non saisies','No marks entered','لم تُدخل العلامات')}}{% endif %}{% elif g.study_year in [2,3] %}{{L('Décision conjointe Dr Dris + Pr Benhocine','Joint decision by Dr Dris + Pr Benhocine','قرار مشترك للدكتور دريس والأستاذ بن حسين')}}{% else %}{{L('Validation DEMS dans la rubrique dédiée','DEMS validation in its dedicated section','التحقق من DEMS في القسم المخصص')}}{% endif %}</td><td>{% if role=='super_admin' and g.current_year==1 and g.study_year==1 and g.status=='active' %}<form method="post"><input type="hidden" name="action" value="evaluate_first_year"><input type="hidden" name="resident" value="{{g.resident_id}}"><input type="hidden" name="study_year" value="1"><label><input type="checkbox" name="all_modules_confirmed" value="yes" required> {{L('Toutes les notes sont saisies','All marks entered','تم إدخال جميع العلامات')}}</label><button type="submit">{{L('Valider le passage','Validate promotion','تأكيد الترقية')}}</button></form>{% elif role=='super_admin' and g.study_year in [2,3] and g.current_year==g.study_year and g.status=='active' %}<form method="post"><input type="hidden" name="action" value="approve_advance"><input type="hidden" name="resident" value="{{g.resident_id}}"><input type="hidden" name="study_year" value="{{g.study_year}}"><select name="outcome"><option value="approve">Approuver</option><option value="reject">Refuser</option></select><input name="reason" maxlength="500" placeholder="{{L('Motif obligatoire','Required reason','السبب إلزامي')}}" required><button type="submit">{{L('Enregistrer mon vote','Record my vote','تسجيل تصويتي')}}</button></form>{% else %}—{% endif %}</td></tr>{% else %}<tr><td colspan="6">{{L('Aucun résident actif à évaluer pour cette année.','No active resident to evaluate for this year.','لا يوجد مقيم نشط لتقييمه في هذه السنة.')}}</td></tr>{% endfor %}</table></div></section><section><h4>{{L('Évaluer la première année','Evaluate first year','تقييم السنة الأولى')}}</h4><form method="post"><input type="hidden" name="action" value="evaluate_first_year"><label>{{L('Résident','Resident','المقيم')}}</label><select name="resident">{% for r in rs if r.year==1 %}<option value="{{r.id}}">{{r.surname}} {{r.firstname}}</option>{% endfor %}</select><input type="hidden" name="study_year" value="1"><label><input type="checkbox" name="all_modules_confirmed" value="yes" required> {{L('Confirmer la saisie de tous les résultats des modules','Confirm all module results are entered','تأكيد إدخال نتائج جميع المواد')}}</label><button>{{L('Évaluer et promouvoir si les critères sont remplis','Evaluate and advance if eligible','تقييم الترقية إذا استوفيت الشروط')}}</button></form><h4>{{L('Vote de promotion — années 2 et 3 (Dr Dris + Pr Benhocine)','Progression vote — years 2–3 (Dr Dris + Pr Benhocine)','تصويت الترقية — السنتان الثانية والثالثة (الدكتور دريس والأستاذ بن حسين)')}}</h4><form method="post"><input type="hidden" name="action" value="approve_advance"><label>{{L('Résident','Resident','المقيم')}}</label><select name="resident" required>{% for r in rs if r.year in [2,3] %}<option value="{{r.id}}">{{r.year}} — {{r.surname}} {{r.firstname}}</option>{% endfor %}</select><label>{{L('Année d’études','Study year','السنة الدراسية')}}</label><select name="study_year"><option>2</option><option>3</option></select><label>{{L('Vote','Vote','التصويت')}}</label><select name="outcome"><option value="approve">{{L('Approuver','Approve','موافقة')}}</option><option value="reject">{{L('Refuser','Reject','رفض')}}</option></select><label>{{L('Motif','Reason','السبب')}}</label><input name="reason" required maxlength="500"><button>{{L('Enregistrer le vote officiel','Record official vote','تسجيل التصويت الرسمي')}}</button></form><h4>{{L('Votes officiels en attente','Pending official votes','التصويتات الرسمية المعلقة')}}</h4><div class="wrap"><table><tr><th>{{L('Résident','Resident','المقيم')}}</th><th>{{L('Année de départ','From year','سنة الانطلاق')}}</th><th>{{L('Responsable','Official','المسؤول')}}</th><th>{{L('Vote','Vote','التصويت')}}</th><th>{{L('Motif','Reason','السبب')}}</th></tr>{% for v in pending_votes %}<tr><td>{{v.surname}} {{v.firstname}}</td><td>{{v.from_year}}</td><td>{{v.actor}}</td><td>{{v.decision}}</td><td>{{v.reason}}</td></tr>{% else %}<tr><td colspan="5">{{L('Aucun vote officiel en attente','No pending official votes','لا توجد تصويتات رسمية معلقة')}}</td></tr>{% endfor %}</table></div></section><section><div class="wrap"><table><tr><th>{{L('Année de post-graduation','Postgraduate year','سنة ما بعد التدرج')}}</th><th>{{L('Résident','Resident','المقيم')}}</th><th>{{L('Année d’études','Study year','السنة الدراسية')}}</th><th>{{L('Module / insertion DEMS','Module / DEMS insertion','المادة / إدراج DEMS')}}</th><th>{{L('Session','Session','الدورة')}}</th><th>{{L('Tentative','Attempt','المحاولة')}}</th><th>{{L('Partie 1 / théorie','Part 1 / theory','الجزء الأول / النظري')}}</th><th>{{L('Partie 2 / pratique','Part 2 / practical','الجزء الثاني / التطبيقي')}}</th><th>{{L('Actions','Actions','الإجراءات')}}</th></tr>{% for e in rows %}<tr><td>{{e.postgrad_year}}</td><td>{{e.surname}} {{e.firstname}}</td><td>{{e.study_year}}</td><td>{{e.module}}</td><td>{{e.exam_type}}</td><td>{{e.attempt}}</td><td>{{e.theory}}</td><td>{{e.practical}}</td><td>{% if role=='super_admin' %}<form method="post" onsubmit="return confirm('{{L('Déplacer ce résultat dans la corbeille pendant six mois ?','Move this result to the trash for six months?','نقل هذه النتيجة إلى سلة المهملات لمدة ستة أشهر؟')}}');"><input type="hidden" name="action" value="delete_record"><input type="hidden" name="exam_id" value="{{e.id}}"><button type="submit">{{L('Mettre à la corbeille','Move to trash','نقل إلى سلة المهملات')}}</button></form>{% endif %}</td></tr>{% else %}<tr><td colspan="9">{{t('none')}}</td></tr>{% endfor %}</table></div></section>""",rows=rows,unassigned_rows=unassigned_rows,year_options=year_options,rs=rs,selected=selected,pending_votes=pending_votes)
@app.route("/dems",methods=["GET","POST"])
def dems():
 if super_guard(): return super_guard()
 c=db()
 if request.method=="POST":
  actor=session.get("user","").lower()
  if actor not in ("benhocine","drdris"):
   c.close(); return ("Only Dr Dris and Pr Benhocine may vote on DEMS success.",403)
  rid=request.form.get("resident",type=int)
  outcome=request.form.get("outcome","approve")
  reason=request.form.get("reason","").strip()
  try: success_date=date.fromisoformat(request.form.get("success_date","").strip()).isoformat()
  except ValueError:
   c.close(); flash(L("Saisissez une date d’examen DEMS valide.","Enter a valid DEMS exam date.","أدخل تاريخ امتحان DEMS صالحًا.")); return redirect("/dems")
  if date.fromisoformat(success_date)>date.today():
   c.close(); flash(L("La date de l’examen DEMS ne peut pas être dans le futur.","The DEMS exam date cannot be in the future.","لا يمكن أن يكون تاريخ امتحان DEMS في المستقبل.")); return redirect("/dems")
  r=c.execute("SELECT * FROM residents WHERE id=? AND year=4 AND status='active' AND COALESCE(resident_type,'local')<>'external'",(rid,)).fetchone()
  if not r or outcome not in ("approve","reject") or not reason:
   flash(L("Choisissez un résident actif de quatrième année, un vote valide, une date et un motif.","Choose an active year-4 resident, a valid vote, a date and a reason.","اختر مقيمًا نشطًا في السنة الرابعة وتصويتًا صالحًا وتاريخًا وسببًا."))
  else:
   c.execute("INSERT INTO dems_votes(resident_id,success_date,decision,reason,actor) VALUES(?,?,?,?,?) ON CONFLICT(resident_id,success_date,actor) DO UPDATE SET decision=excluded.decision,reason=excluded.reason,created_at=CURRENT_TIMESTAMP",(rid,success_date,outcome,reason,actor))
   audit_event(c,"dems.vote_recorded",f"resident_id={rid}; success_date={success_date}; decision={outcome}; reason={reason}")
   votes=c.execute("SELECT actor,decision,reason FROM dems_votes WHERE resident_id=? AND success_date=? AND actor IN ('benhocine','drdris')",(rid,success_date)).fetchall()
   if len(votes)<2:
    c.commit(); flash(L("Vote DEMS enregistré. En attente du vote de l’autre responsable officiel.","DEMS vote recorded. Waiting for the other official administrator.","تم تسجيل تصويت DEMS. في انتظار تصويت المسؤول الرسمي الآخر."))
   else:
    combined_reason=" | ".join(v["actor"]+": "+v["decision"]+" — "+v["reason"] for v in votes)
    actor_names="+".join(sorted(v["actor"] for v in votes))
    approved=all(v["decision"]=="approve" for v in votes)
    if approved:
     hist=c.execute("SELECT h.*,t.name terrain FROM history h JOIN terrains t ON t.id=h.terrain_id WHERE h.resident_id=?",(rid,)).fetchall()
     asgs=c.execute("SELECT a.*,p.name period,p.year_label,t.name terrain FROM assignments a JOIN periods p ON p.id=a.period_id JOIN terrains t ON t.id=a.terrain_id WHERE a.resident_id=?",(rid,)).fetchall()
     exams_snapshot=c.execute("SELECT * FROM exam_records WHERE resident_id=? ORDER BY study_year,module",(rid,)).fetchall()
     profile={k:r[k] for k in r.keys()}
     c.execute("INSERT INTO dems_archive(resident_id,success_date,profile_json,history_json) VALUES(?,?,?,?)",(rid,success_date,json.dumps(profile,ensure_ascii=False),json.dumps({"history":[dict(x) for x in hist],"assignments":[dict(x) for x in asgs],"exams":[dict(x) for x in exams_snapshot]},ensure_ascii=False)))
     c.execute("UPDATE residents SET status='graduated',dems_status='success' WHERE id=?",(rid,))
     c.execute("INSERT INTO decisions(details,decision,reason,actor,postgrad_year) VALUES(?,?,?,?,?)",(f"{r['surname']} {r['firstname']} — DEMS","Success approved by both official administrators",combined_reason,actor_names,active_postgrad_year(c)))
     audit_event(c,"dems.archived",f"resident_id={rid}; success_date={success_date}; officials={actor_names}")
     flash(L("Réussite au DEMS approuvée par les deux responsables ; le dossier et l’historique sont archivés.","DEMS success approved by both officials; profile and history archived.","وافق المسؤولان على نجاح DEMS؛ تمت أرشفة الملف والسجل."))
    else:
     c.execute("INSERT INTO decisions(details,decision,reason,actor,postgrad_year) VALUES(?,?,?,?,?)",(f"{r['surname']} {r['firstname']} — DEMS","Not approved by official vote",combined_reason,actor_names,active_postgrad_year(c)))
     audit_event(c,"dems.vote_rejected",f"resident_id={rid}; success_date={success_date}; officials={actor_names}")
     flash(L("Réussite au DEMS refusée ; les deux votes officiels sont enregistrés.","DEMS success was not approved; both official votes are recorded.","لم تتم الموافقة على نجاح DEMS؛ تم تسجيل التصويتين الرسميين."))
    c.execute("DELETE FROM dems_votes WHERE resident_id=? AND success_date=?",(rid,success_date))
    c.commit()
 residentsrow_numbers=c.execute("SELECT * FROM residents WHERE year=4 AND status='active' AND COALESCE(resident_type,'local')<>'external' ORDER BY surname").fetchall()
 archiverow_numbers=c.execute("SELECT d.*,r.surname,r.firstname FROM dems_archive d JOIN residents r ON r.id=d.resident_id ORDER BY d.id DESC").fetchall()
 pending_votes=c.execute("SELECT v.*,r.surname,r.firstname FROM dems_votes v JOIN residents r ON r.id=v.resident_id ORDER BY v.created_at DESC").fetchall()
 c.close()
 return page("""<section><h3>{{t('dems')}}</h3><p>{{L('La réussite au DEMS est définitive uniquement après l’approbation du Dr Dris et du Pr Benhocine.','DEMS success is finalized only after both Dr Dris and Pr Benhocine approve.','لا تُعتمد نتيجة النجاح في DEMS نهائيًا إلا بعد موافقة الدكتور دريس والأستاذ بن حسين.')}}</p><form method="post"><label>{{L('Résident de 4e année','Year-4 resident','مقيم السنة الرابعة')}}</label><select name="resident" required>{% for r in residentsrow_numbers %}<option value="{{r.id}}">{{r.surname}} {{r.firstname}}</option>{% endfor %}</select><label>{{t('success_date')}}</label><input type="date" name="success_date" required><label>{{L('Vote','Vote','التصويت')}}</label><select name="outcome"><option value="approve">{{L('Approuver la réussite','Approve success','الموافقة على النجاح')}}</option><option value="reject">{{L('Refuser la réussite','Reject success','رفض النجاح')}}</option></select><label>{{L('Motif','Reason','السبب')}}</label><input name="reason" required maxlength="500"><button>{{L('Enregistrer le vote officiel','Record official vote','تسجيل التصويت الرسمي')}}</button></form></section><section><h4>{{L('Votes officiels en attente','Pending official votes','التصويتات الرسمية المعلقة')}}</h4><div class="wrap"><table><tr><th>{{L('Résident','Resident','المقيم')}}</th><th>{{L('Date de l’examen','Exam date','تاريخ الامتحان')}}</th><th>{{L('Responsable','Official','المسؤول')}}</th><th>{{L('Vote','Vote','التصويت')}}</th><th>{{L('Motif','Reason','السبب')}}</th></tr>{% for v in pending_votes %}<tr><td>{{v.surname}} {{v.firstname}}</td><td>{{v.success_date}}</td><td>{{v.actor}}</td><td>{{L('Approuvé' if v.decision=='approve' else 'Refusé','Approved' if v.decision=='approve' else 'Rejected','موافق عليه' if v.decision=='approve' else 'مرفوض')}}</td><td>{{v.reason}}</td></tr>{% else %}<tr><td colspan="5">{{L('Aucun vote officiel en attente','No pending official votes','لا توجد تصويتات رسمية معلقة')}}</td></tr>{% endfor %}</table></div></section><section><h4>{{L('Diplômés archivés','Archived graduates','المتخرجون المؤرشفون')}}</h4>{% for a in archiverow_numbers %}<p>{{a.surname}} {{a.firstname}} — {{L('Réussite au DEMS :','DEMS success:','النجاح في DEMS:')}} {{a.success_date}} — {{L('archivé le','archived','تمت الأرشفة في')}} {{a.archived_at}}</p>{% else %}<p>{{t('none')}}</p>{% endfor %}</section>""",residentsrow_numbers=residentsrow_numbers,archiverow_numbers=archiverow_numbers,pending_votes=pending_votes)
@app.get("/export")
def export():
 if super_guard():return super_guard()
 c=db(); out=io.StringIO(); w=csv.writer(out)
 for table in ["residents","terrains","periods","capacities","assignments","history","decisions","exam_records","objectives","dems_archive","leave_years","leave_entries","external_rotations","promotion_votes","dems_votes","audit"]:
  w.writerow([table])
  w.writerow([x[1] for x in c.execute('PRAGMA table_info("'+table+'")')])
  for row in c.execute('SELECT * FROM "'+table+'"'):w.writerow(list(row))
 c.close(); f=io.BytesIO(("\ufeff"+out.getvalue()).encode("utf-8")); return send_file(f,mimetype="text/csv",as_attachment=True,download_name="rotation_export.csv")

def _sqlite_backup_file(source_path, destination_path):
 """Copy a SQLite snapshot through the backup API (safe for Windows file locks)."""
 source=sqlite3.connect(source_path,timeout=30)
 destination=sqlite3.connect(destination_path,timeout=30)
 try:
  source.backup(destination)
  destination.commit()
 finally:
  destination.close(); source.close()

def create_full_backup_bytes(directory):
 with tempfile.TemporaryDirectory(prefix="resimanage_backup_",dir=directory) as workdir:
  snapshot=os.path.join(workdir,"rotation.db")
  source=sqlite3.connect(DB,timeout=20)
  destination=sqlite3.connect(snapshot)
  try:
   source.backup(destination); destination.commit()
  finally:
   destination.close(); source.close()
  output=io.BytesIO()
  with zipfile.ZipFile(output,"w",compression=zipfile.ZIP_DEFLATED) as bundle:
   bundle.write(snapshot,"rotation.db")
   bundle.writestr("resident_photos/","")
   photo_dir=os.path.join(directory,"resident_photos")
   if os.path.isdir(photo_dir):
    for name in sorted(os.listdir(photo_dir)):
     path=os.path.join(photo_dir,name)
     if os.path.isfile(path) and not os.path.islink(path) and os.path.basename(name)==name:
      bundle.write(path,"resident_photos/"+name)
  return output.getvalue()

@app.get("/backup")
def backup():
 if super_guard(): return super_guard()
 directory=os.path.dirname(os.path.abspath(DB))
 data=create_full_backup_bytes(directory)
 return send_file(io.BytesIO(data),mimetype="application/zip",as_attachment=True,download_name="resimanage_full_backup.zip")

@app.route("/trash",methods=["GET","POST"])
def trash():
 if guard(): return redirect("/login")
 if session.get("role")!="super_admin": return ("Forbidden — Super Admin only",403)
 c=db(); now=datetime.now()
 for item in c.execute("SELECT id,deleted_at FROM trash_items").fetchall():
  try:
   deleted=datetime.fromisoformat(item["deleted_at"])
   month=deleted.month+6; year=deleted.year+(month-1)//12; month=(month-1)%12+1
   day=min(deleted.day,__import__("calendar").monthrange(year,month)[1])
   if now>=deleted.replace(year=year,month=month,day=day):
    c.execute("DELETE FROM trash_items WHERE id=?",(item["id"],))
    audit_event(c,"trash.expired",f"trash_id={item['id']}; purged_after_six_calendar_months")
  except (ValueError,TypeError): continue
 c.commit()
 if request.method=="POST":
  action=request.form.get("action",""); tid=int(request.form.get("trash_id","0"))
  item=c.execute("SELECT * FROM trash_items WHERE id=?",(tid,)).fetchone()
  if not item: flash(L("Élément introuvable ou délai de conservation écoulé.","Item not found or retention period expired.","العنصر غير موجود أو انتهت مدة الاحتفاظ."))
  elif action=="restore":
   try:
    snap=json.loads(item["snapshot_json"])
    if item["entity_type"]=="resident":
     resident=snap["resident"]; cols=[x[1] for x in c.execute("PRAGMA table_info(residents)") if x[1] in resident]
     c.execute("INSERT INTO residents("+",".join(cols)+") VALUES("+",".join("?" for _ in cols)+")",[resident[x] for x in cols])
     for table,records in snap.get("related",{}).items():
      allowed={x[1] for x in c.execute("PRAGMA table_info("+table+")")}
      for record in records:
       cols=[k for k in record if k in allowed]
       c.execute("INSERT INTO "+table+"("+",".join(cols)+") VALUES("+",".join("?" for _ in cols)+")",[record[k] for k in cols])
    elif item["entity_type"]=="exam_record":
     record=snap["record"]; cols=[x[1] for x in c.execute("PRAGMA table_info(exam_records)") if x[1] in record]
     c.execute("INSERT INTO exam_records("+",".join(cols)+") VALUES("+",".join("?" for _ in cols)+")",[record[x] for x in cols])
    elif item["entity_type"]=="leave_entry":
     record=snap["record"]; cols=[x[1] for x in c.execute("PRAGMA table_info(leave_entries)") if x[1] in record]
     c.execute("INSERT INTO leave_entries("+",".join(cols)+") VALUES("+",".join("?" for _ in cols)+")",[record[x] for x in cols])
    elif item["entity_type"]=="assignment":
     record=snap["record"]
     current=c.execute("SELECT id,status,terrain_id,period_id FROM assignments WHERE id=?",(record["id"],)).fetchone()
     if not current or current["status"]!="cancelled" or current["terrain_id"]!=record["terrain_id"] or current["period_id"]!=record["period_id"]:
      raise ValueError("Assignment changed after it was removed; cannot safely restore the old version.")
     conflict=c.execute("SELECT id FROM assignments WHERE resident_id=? AND period_id=? AND id<>? AND status IN ('confirmed','pending')",(record["resident_id"],record["period_id"],record["id"])).fetchone()
     if conflict: raise ValueError("Resident already has another active assignment in this period.")
     c.execute("UPDATE assignments SET status=? WHERE id=?",(record["status"],record["id"]))
    elif item["entity_type"] in ("history","objectives","leave_entries"):
     table={"history":"history","objectives":"objectives","leave_entries":"leave_entries"}[item["entity_type"]]
     record=snap["record"]
     columns={x[1] for x in c.execute("PRAGMA table_info("+table+")")}
     cols=[key for key in record if key in columns]
     c.execute("INSERT INTO "+table+"("+",".join(cols)+") VALUES("+",".join("?" for _ in cols)+")",[record[key] for key in cols])
    else: raise ValueError("Unsupported trash type")
    c.execute("DELETE FROM trash_items WHERE id=?",(tid,))
    audit_event(c,"trash.restored",f"trash_id={tid}; resident_id={item['entity_id']}")
    c.commit(); flash(L("Dossier et données associées restaurés.","Resident record and related data restored.","تمت استعادة ملف المقيم والبيانات المرتبطة."))
   except (ValueError,TypeError,KeyError,sqlite3.Error) as exc:
    c.rollback(); app.logger.warning("Trash restore failed: %s",exc); flash(L("Restauration impossible. Vérifiez que ce dossier n’existe pas déjà.","Restore failed. Check whether the record already exists.","تعذرت الاستعادة. تحقق من أن الملف غير موجود مسبقًا."))
  elif action=="purge":
   flash(L("La suppression définitive est automatique après six mois ; l’élément ne peut pas être purgé avant l’échéance.","Permanent deletion happens automatically after six months; items cannot be purged early.","يتم الحذف النهائي تلقائيًا بعد ستة أشهر؛ لا يمكن حذف العنصر نهائيًا قبل انتهاء المدة."))
  return redirect("/trash")
 rows=c.execute("SELECT * FROM trash_items ORDER BY deleted_at DESC").fetchall(); c.close()
 return page("""<section><h3>{{L('Corbeille — conservation de six mois','Trash — six-month retention','سلة المهملات — الاحتفاظ ستة أشهر')}}</h3><p>{{L('Les dossiers restent restaurables pendant six mois calendaires. Après échéance, leur contenu est supprimé définitivement automatiquement lors de la prochaine ouverture de cette rubrique.','Records can be restored for six calendar months. After expiry, their contents are permanently purged automatically when this section is next opened.','يمكن استعادة الملفات خلال ستة أشهر تقويمية. بعد انتهاء المدة تُحذف محتوياتها نهائيًا تلقائيًا عند فتح هذا القسم مرة أخرى.')}}</p><div class="wrap"><table><tr><th>{{L('Élément','Item','العنصر')}}</th><th>{{L('Motif','Reason','السبب')}}</th><th>{{L('Supprimé le','Deleted at','تاريخ الحذف')}}</th><th>{{L('Supprimé par','Deleted by','حذف بواسطة')}}</th><th>{{L('Actions','Actions','الإجراءات')}}</th></tr>{% for x in rows %}<tr><td>{{x.label}} ({{x.entity_type}})</td><td>{{x.reason}}</td><td>{{x.deleted_at}}</td><td>{{x.deleted_by}}</td><td><form method="post"><input type="hidden" name="trash_id" value="{{x.id}}"><button name="action" value="restore">{{L('Restaurer','Restore','استعادة')}}</button></form><small>{{L('Suppression définitive automatique à l’échéance des six mois','Automatically purged when six months expire','الحذف النهائي تلقائيًا عند انتهاء الأشهر الستة')}}</small></td></tr>{% else %}<tr><td colspan="5">{{L('La corbeille est vide.','Trash is empty.','سلة المهملات فارغة.')}}</td></tr>{% endfor %}</table></div></section>""",rows=rows)

@app.route("/restore",methods=["GET","POST"])
def restore():
 if super_guard(): return super_guard()
 if request.method=="POST":
  uploaded=request.files.get("backup_file")
  if not uploaded or not uploaded.filename:
   flash("Choose a .zip full backup or a legacy .db/.sqlite backup.")
   return redirect("/restore")
  extension=os.path.splitext(uploaded.filename.lower())[1]
  if extension not in (".zip",".db",".sqlite",".sqlite3"):
   flash("Unsupported backup type. Choose .zip, .db or .sqlite.")
   return redirect("/restore")
  directory=os.path.dirname(os.path.abspath(DB))
  workdir=None; old_db_path=None; old_photo_path=None; staged_db=None; staged_photos=None
  db_replaced=False; photos_replaced=False
  try:
   workdir=tempfile.mkdtemp(prefix="resimanage_restore_",dir=directory)
   uploaded_path=os.path.join(workdir,"uploaded"+extension)
   uploaded.save(uploaded_path)
   staged_db=os.path.join(workdir,"rotation.db")
   restore_photos=False
   if extension==".zip":
    with zipfile.ZipFile(uploaded_path,"r") as bundle:
     infos=bundle.infolist()
     names=[info.filename for info in infos]
     if names.count("rotation.db")!=1:
      raise ValueError("Full backup must contain exactly one rotation.db.")
     if sum(info.file_size for info in infos)>100*1024*1024:
      raise ValueError("Backup expands beyond the 100 MB safety limit.")
     for info in infos:
      name=info.filename
      if name=="rotation.db" or name=="resident_photos/" or (name.startswith("resident_photos/") and not name.endswith("/")):
       if ".." in name.split("/") or name.startswith("/"):
        raise ValueError("Backup contains an unsafe path.")
      else:
       raise ValueError("Backup contains an unsupported file path.")
     with bundle.open("rotation.db") as source, open(staged_db,"wb") as destination:
      shutil.copyfileobj(source,destination)
     restore_photos="resident_photos/" in names or any(name.startswith("resident_photos/") for name in names)
     if restore_photos:
      staged_photos=os.path.join(workdir,"resident_photos")
      os.makedirs(staged_photos,exist_ok=True)
      for info in infos:
       name=info.filename
       if name.startswith("resident_photos/") and not name.endswith("/"):
        photo_name=name[len("resident_photos/"):]
        if not photo_name or "/" in photo_name or photo_name in (".","..") or info.file_size>8*1024*1024:
         raise ValueError("Backup contains an invalid or oversized resident photo.")
        with bundle.open(info) as source, open(os.path.join(staged_photos,photo_name),"wb") as destination:
         shutil.copyfileobj(source,destination)
   else:
    shutil.copyfile(uploaded_path,staged_db)
   check=sqlite3.connect("file:"+staged_db+"?mode=ro",uri=True,timeout=10)
   integrity=check.execute("PRAGMA integrity_check").fetchone()[0]
   tables={row[0] for row in check.execute("SELECT name FROM sqlite_master WHERE type='table'")}
   check.close()
   required={"residents","terrains","periods","assignments","history","settings"}
   if integrity!="ok" or not required.issubset(tables):
    raise ValueError("The uploaded file is not a valid ResiManage database backup.")
   timestamp=datetime.now().strftime("%Y%m%d_%H%M%S")
   full_before=os.path.join(directory,"rotation_backup_before_restore_"+timestamp+".zip")
   with open(full_before,"wb") as handle:
    handle.write(create_full_backup_bytes(directory))
   old_db_path=os.path.join(directory,"rotation_backup_before_restore_"+timestamp+".db")
   _sqlite_backup_file(DB,old_db_path)
   photo_dir=os.path.join(directory,"resident_photos")
   if restore_photos and os.path.exists(photo_dir):
    old_photo_path=os.path.join(directory,"resident_photos_before_restore_"+timestamp)
    os.replace(photo_dir,old_photo_path)
   # Restore pages into the existing DB through SQLite's online backup API.
   # Windows may deny os.replace() when any process still has the database open.
   _sqlite_backup_file(staged_db,DB); db_replaced=True
   if restore_photos:
    os.replace(staged_photos,photo_dir); staged_photos=None; photos_replaced=True
   app._ready=False
   init()
   c=db(); audit_event(c,"database.restored",f"Restored from {extension} backup; pre-restore full backup={os.path.basename(full_before)}"); c.commit(); c.close()
   flash("Restore completed. Pre-restore full backup saved as "+os.path.basename(full_before))
  except Exception as exc:
   if db_replaced and old_db_path and os.path.exists(old_db_path):
    try: _sqlite_backup_file(old_db_path,DB); app._ready=False
    except Exception: pass
   if photos_replaced:
    photo_dir=os.path.join(directory,"resident_photos")
    try:
     if os.path.isdir(photo_dir): shutil.rmtree(photo_dir)
    except Exception: pass
   if old_photo_path and os.path.exists(old_photo_path):
    try: os.replace(old_photo_path,os.path.join(directory,"resident_photos"))
    except Exception: pass
   flash("Restore failed; a pre-restore full backup was saved if validation had completed. Details: "+str(exc)[:220])
  finally:
   if workdir and os.path.exists(workdir): shutil.rmtree(workdir,ignore_errors=True)
  return redirect("/restore")
 return page("""<section><h3>{{t('restore')}}</h3><p>{{t('restore_help')}} Upload a full .zip backup (database and resident photos) or a legacy .db/.sqlite database backup.</p><form method="post" enctype="multipart/form-data"><label>{{t('file')}} (.zip/.db/.sqlite)</label><input type="file" name="backup_file" accept=".zip,.db,.sqlite,.sqlite3" required><button>{{t('restore')}}</button></form></section>""")
def rotation_period_bounds(year_label, period_name, start_date="", end_date=""):
    try:
        y = int(str(year_label).split("-")[0])
    except (ValueError, IndexError):
        y = date.today().year
    n = str(period_name).lower()
    defaults = {
        "période 1": (date(y,2,1), date(y,4,30)),
        "periode 1": (date(y,2,1), date(y,4,30)),
        "période 2": (date(y,5,1), date(y,7,30)),
        "periode 2": (date(y,5,1), date(y,7,30)),
        "période 3": (date(y,8,1), date(y,10,30)),
        "periode 3": (date(y,8,1), date(y,10,30)),
        "période 4": (date(y,11,1), date(y+1,1,30)),
        "periode 4": (date(y,11,1), date(y+1,1,30)),
    }
    if any(x in n for x in ("fév", "fev", "feb")): fallback=(date(y,2,1), date(y,4,30))
    elif any(x in n for x in ("mai", "may")): fallback=(date(y,5,1), date(y,7,30))
    elif any(x in n for x in ("août", "aout", "aug")): fallback=(date(y,8,1), date(y,10,30))
    elif "nov" in n: fallback=(date(y,11,1), date(y+1,1,30))
    else: fallback=defaults.get(n,(date(y,2,1), date(y+1,1,30)))
    try: start=date.fromisoformat(start_date) if start_date else fallback[0]
    except (ValueError,TypeError): start=fallback[0]
    try: end=date.fromisoformat(end_date) if end_date else fallback[1]
    except (ValueError,TypeError): end=fallback[1]
    return start,end

def reject_external_overlap_with_local_assignment(c, resident_id, start, end):
    local_rows = c.execute(
        """SELECT a.id,t.name terrain,p.year_label,p.name period,p.start_date,p.end_date
           FROM assignments a
           JOIN terrains t ON t.id=a.terrain_id
           JOIN periods p ON p.id=a.period_id
           WHERE a.resident_id=? AND a.status IN ('confirmed','pending')""",
        (resident_id,),
    ).fetchall()
    for assignment in local_rows:
        local_start, local_end = rotation_period_bounds(
            assignment["year_label"], assignment["period"], assignment["start_date"], assignment["end_date"]
        )
        if start <= local_end and end >= local_start:
            raise ValueError(
                L("Le stage à l’étranger chevauche une affectation locale existante : ","External rotation overlaps an existing local assignment: ","يتداخل التربص بالخارج مع تعيين محلي موجود: ") + assignment["terrain"] + " / " + assignment["period"] + "."
            )


def accrued_leave_days(registration_date, as_of):
    """30 days accrue for each completed anniversary of first-year registration."""
    try:
        start=date.fromisoformat(str(registration_date or "").strip())
    except (TypeError,ValueError):
        return 0
    if start>as_of: return 0
    years=as_of.year-start.year
    try: anniversary=date(as_of.year,start.month,start.day)
    except ValueError: anniversary=date(as_of.year,2,28)
    if anniversary>as_of: years-=1
    return max(0,years)*30

@app.route("/leave", methods=["GET","POST"])
def leave_management():
    if guard(): return redirect("/login")
    c=db()
    if request.method=="POST":
        action=request.form.get("action","create_entry")
        if super_guard(): c.close(); return ("Forbidden — Super Admin only",403)
        if action in ("delete_entry","update_entry"):
            try:
                entry_id=int(request.form.get("entry_id","0"))
                existing=c.execute("SELECT * FROM leave_entries WHERE id=?",(entry_id,)).fetchone()
                if not existing: raise ValueError("Leave entry does not exist.")
                if action=="delete_entry":
                    reason=L("Suppression demandée par le Super Admin","Deletion requested by Super Admin","طلب المسؤول العام الحذف")
                    c.execute("INSERT INTO trash_items(entity_type,entity_id,label,reason,snapshot_json,deleted_at,deleted_by) VALUES(?,?,?,?,?,?,?)",
                      ("leave_entry",str(entry_id),L("Congé du résident","Resident leave entry","إجازة المقيم")+" #"+str(entry_id),reason,json.dumps({"record":dict(existing)},ensure_ascii=False),datetime.now().replace(microsecond=0).isoformat(),session.get("username","")))
                    audit_event(c,"leave.entry_deleted",f"entry_id={entry_id}; resident_id={existing['resident_id']}; start={existing['start_date']}; end={existing['end_date']}; days={existing['days']}; notes={existing['notes']}; moved_to_trash_for_six_months=true")
                    c.execute("DELETE FROM leave_entries WHERE id=?",(entry_id,))
                    c.commit(); flash(L("Congé déplacé dans la corbeille pour six mois.","Leave entry moved to the trash for six months.","تم نقل الإجازة إلى سلة المهملات لمدة ستة أشهر."))
                else:
                    start=date.fromisoformat(request.form["start_date"]); end=date.fromisoformat(request.form["end_date"])
                    if end<start: raise ValueError(L("La date de fin ne peut pas précéder la date de début.","End date must not precede start date.","لا يمكن أن يسبق تاريخ النهاية تاريخ البداية."))
                    overlap=c.execute("SELECT id FROM leave_entries WHERE resident_id=? AND id<>? AND start_date<=? AND end_date>=?",(existing["resident_id"],entry_id,end.isoformat(),start.isoformat())).fetchone()
                    if overlap: raise ValueError(L("Les dates du congé chevauchent une autre entrée.","Leave dates overlap another entry.","تتداخل تواريخ الإجازة مع سجل آخر."))
                    days=(end-start).days+1
                    notes=request.form.get("notes","").strip()
                    c.execute("UPDATE leave_entries SET start_date=?,end_date=?,days=?,notes=? WHERE id=?",(start.isoformat(),end.isoformat(),days,notes,entry_id))
                    audit_event(c,"leave.entry_updated",f"entry_id={entry_id}; resident_id={existing['resident_id']}; before={existing['start_date']}..{existing['end_date']} ({existing['days']} days); after={start}..{end} ({days} days); notes={notes}")
                    c.commit(); flash(L("Congé corrigé par le Super Admin ; les anciennes et nouvelles valeurs sont consignées dans le journal.","Leave corrected by Super Admin; old and new values recorded in the audit log.","صحح المسؤول العام الإجازة؛ تم تسجيل القيم القديمة والجديدة في سجل التدقيق."))
            except Exception as exc:
                c.rollback(); app.logger.warning("Leave update failed: %s",exc); flash(L("La correction du congé a échoué : ","Leave update failed: ","فشل تصحيح الإجازة: ")+str(exc)[:220])
            c.close(); return redirect("/leave")
        try:
            rid=int(request.form["resident_id"])
            if not c.execute("SELECT 1 FROM residents WHERE id=?",(rid,)).fetchone():
                raise ValueError(L("Ce résident n’existe pas.","Resident does not exist.","هذا المقيم غير موجود."))
            start=date.fromisoformat(request.form["start_date"]); end=date.fromisoformat(request.form["end_date"])
            if end<start: raise ValueError(L("La date de fin ne peut pas précéder la date de début.","End date must not precede start date.","لا يمكن أن يسبق تاريخ النهاية تاريخ البداية."))
            days=(end-start).days+1
            overlap=c.execute("SELECT id FROM leave_entries WHERE resident_id=? AND start_date<=? AND end_date>=?",(rid,end.isoformat(),start.isoformat())).fetchone()
            if overlap: raise ValueError(L("Les dates du congé chevauchent une entrée existante.","Leave dates overlap an existing entry","تتداخل تواريخ الإجازة مع سجل موجود."))
            historical=request.form.get("historical_record")=="1"
            resident=c.execute("SELECT registration_date FROM residents WHERE id=?",(rid,)).fetchone()
            if historical and end>date.today():
                raise ValueError(L("Un congé historique doit se terminer aujourd’hui ou avant.","Historical leave entries must end today or earlier.","يجب أن تنتهي الإجازة السابقة اليوم أو قبله."))
            if not historical:
                # Maintain explicit annual quotas. If a legacy resident has no registration
                # date yet, record the 30-day entitlement per calendar year being entered;
                # admins can later correct the first-year registration date in the profile.
                if resident["registration_date"]:
                    first_year=date.fromisoformat(resident["registration_date"]).year
                    last_year=end.year
                    earned=accrued_leave_days(resident["registration_date"],end)
                    eligible_years=[]
                    for entitlement_year in range(first_year,last_year+1):
                        anniversary_year=date(entitlement_year,1,1)
                        if entitlement_year<first_year: continue
                        c.execute("INSERT OR IGNORE INTO leave_years(resident_id,year,earned_days,created_by) VALUES(?,?,30,?)",(rid,entitlement_year,session.get("user","")))
                    earned_rows=c.execute("SELECT COALESCE(SUM(earned_days),0) FROM leave_years WHERE resident_id=? AND year<=?",(rid,last_year)).fetchone()
                    earned=max(earned,int(earned_rows[0] or 0))
                else:
                    first_year=start.year
                    for entitlement_year in range(first_year,end.year+1):
                        c.execute("INSERT OR IGNORE INTO leave_years(resident_id,year,earned_days,created_by) VALUES(?,?,30,?)",(rid,entitlement_year,session.get("user","")))
                    earned=c.execute("SELECT COALESCE(SUM(earned_days),0) FROM leave_years WHERE resident_id=? AND year<=?",(rid,end.year)).fetchone()[0]
                used=c.execute("SELECT COALESCE(SUM(days),0) FROM leave_entries WHERE resident_id=? AND end_date<=?",(rid,end.isoformat())).fetchone()[0]
                if used+days>earned:
                    raise ValueError(L(f"Solde de congés acquis insuffisant : {max(0,earned-used)} jours disponibles au {end.isoformat()}.",f"Insufficient accrued leave balance: {max(0,earned-used)} days available on {end.isoformat()}.",f"رصيد الإجازات المكتسب غير كافٍ: المتاح {max(0,earned-used)} يومًا بتاريخ {end.isoformat()}."))
            c.execute("INSERT INTO leave_entries(resident_id,start_date,end_date,days,notes,created_by) VALUES(?,?,?,?,?,?)",(rid,start.isoformat(),end.isoformat(),days,request.form.get("notes","").strip(),session.get("user","")))
            audit_event(c,"leave.entry_created",f"resident_id={rid}; start={start}; end={end}; days={days}")
            c.commit(); flash(L(f"Congé enregistré : {days} jours calendaires. Les jours non utilisés sont reportés.",f"Leave recorded: {days} calendar days. Unused entitlement carries forward.",f"تم تسجيل الإجازة: {days} أيام تقويمية. تُرحّل الأيام غير المستعملة."))
        except Exception as exc: c.rollback(); app.logger.warning("Leave save failed: %s",exc); flash(L("Le congé n’a pas été enregistré : ","Leave was not saved: ","لم يتم حفظ الإجازة: ")+str(exc)[:220])
        c.close(); return redirect("/leave")
    residents=c.execute("SELECT id,surname,firstname,registration_date FROM residents ORDER BY year DESC,surname").fetchall()
    selected_year=active_postgrad_year(c)
    try:
        year_start=date(int(selected_year.split("-")[0]),2,1)
        year_end=date(int(selected_year.split("-")[1]),1,31)
    except (TypeError,ValueError,IndexError):
        year_start=date.today().replace(month=1,day=1); year_end=date.today()
    as_of=min(year_end,date.today())
    balances=[]
    for r in residents:
        earned=accrued_leave_days(r["registration_date"],as_of)
        used=c.execute("SELECT COALESCE(SUM(days),0) FROM leave_entries WHERE resident_id=? AND end_date<=?",(r["id"],as_of.isoformat())).fetchone()[0]
        balances.append({"name":r["surname"]+" "+r["firstname"],"earned":earned,"used":used,"remaining":earned-used,"registration_date":r["registration_date"] or ""})
    entries=c.execute("SELECT l.*,r.surname,r.firstname FROM leave_entries l JOIN residents r ON r.id=l.resident_id WHERE l.start_date<=? AND l.end_date>=? ORDER BY l.start_date DESC",(as_of.isoformat(),year_start.isoformat())).fetchall()
    c.commit(); c.close()
    return page("""<section><h3>{{L('Congés annuels','Annual leave','الإجازات السنوية')}} — {{year}}</h3><p>{{L('Solde calculé jusqu’au','Balance calculated through','الرصيد محسوب حتى')}} {{as_of_label}} ({{L('année de post-graduation sélectionnée','selected postgraduate year','سنة ما بعد التدرج المحددة')}}). {{L('30 jours par année complète à partir de la date d’inscription en première année. Les jours non utilisés sont reportés.','30 days per completed year from first-year registration. Unused days carry forward.','30 يومًا عن كل سنة كاملة ابتداءً من تاريخ التسجيل في السنة الأولى. تُرحّل الأيام غير المستعملة.')}}</p><form method="post"><label>{{L('Résident','Resident','المقيم')}} <select name="resident_id" required>{% for r in residents %}<option value="{{r.id}}">{{r.surname}} {{r.firstname}}</option>{% endfor %}</select></label><div class="grid"><label>{{L('Début','Start','البداية')}} <input type="date" name="start_date" required></label><label>{{L('Fin','End','النهاية')}} <input type="date" name="end_date" required></label></div><label>{{L('Notes','Notes','ملاحظات')}} <input name="notes" maxlength="500"></label><label><input type="checkbox" name="historical_record" value="1"> {{L('Congé historique déjà pris (à partir de 2022 ; sans validation du quota)','Historical leave already taken (2022 onward; bypass quota validation)','إجازة سابقة مأخوذة (ابتداءً من 2022؛ دون التحقق من الحصة)')}}</label><button>{{L('Enregistrer le congé','Record leave','تسجيل الإجازة')}}</button></form></section><section><h3>{{L('Soldes des congés','Leave balances','أرصدة الإجازات')}}</h3><div class="wrap"><table><tr><th>{{L('Résident','Resident','المقيم')}}</th><th>{{L('Acquis','Accrued','المكتسب')}}</th><th>{{L('Utilisés','Used','المستعمل')}}</th><th>{{L('Restants','Remaining','المتبقي')}}</th><th>{{L('Date d’inscription','Registration date','تاريخ التسجيل')}}</th></tr>{% for b in balances %}<tr><td>{{b.name}}</td><td>{{b.earned}}</td><td>{{b.used}}</td><td>{{b.remaining}}</td><td>{{b.registration_date or L('Non renseignée — modifier le profil du résident','Not set — edit resident profile','غير محدد — عدّل ملف المقيم')}}</td></tr>{% endfor %}</table></div></section><section><h3>{{L('Historique des congés','Leave history','سجل الإجازات')}}</h3><div class="wrap"><table><tr><th>{{L('Résident','Resident','المقيم')}}</th><th>{{L('Début','Start','البداية')}}</th><th>{{L('Fin','End','النهاية')}}</th><th>{{L('Jours','Days','الأيام')}}</th><th>{{L('Notes','Notes','ملاحظات')}}</th><th>{{L('Actions Super Admin','Super Admin actions','إجراءات المسؤول العام')}}</th></tr>{% for e in entries %}<tr><td>{{e.surname}} {{e.firstname}}</td><td colspan="4"><form method="post" style="display:grid;grid-template-columns:repeat(4,minmax(90px,1fr));gap:6px"><input type="hidden" name="action" value="update_entry"><input type="hidden" name="entry_id" value="{{e.id}}"><label>{{L('Début','Start','البداية')}}<input type="date" name="start_date" value="{{e.start_date}}" required></label><label>{{L('Fin','End','النهاية')}}<input type="date" name="end_date" value="{{e.end_date}}" required></label><label>{{L('Jours calculés','Calculated days','الأيام المحسوبة')}}<input value="{{e.days}}" readonly></label><label>{{L('Notes','Notes','ملاحظات')}}<input name="notes" value="{{e.notes}}" maxlength="500"></label><button type="submit">{{L('Enregistrer la correction','Save correction','حفظ التصحيح')}}</button></form></td><td><form method="post" onsubmit="return confirm('{{L('Supprimer ce congé ? Cette action sera journalisée.','Delete this leave record? This action will be logged.','حذف هذه الإجازة؟ سيتم تسجيل العملية.')}}');"><input type="hidden" name="action" value="delete_entry"><input type="hidden" name="entry_id" value="{{e.id}}"><button type="submit">{{L('Supprimer','Delete','حذف')}}</button></form></td></tr>{% endfor %}</table></div></section>""",residents=residents,balances=balances,entries=entries,as_of_label=as_of.isoformat())

@app.route("/external-rotations", methods=["GET","POST"])
def external_rotations():
    if guard(): return redirect("/login")
    c=db()
    if request.method=="POST":
        is_admin=session.get("role")=="super_admin"
        try:
            rid=int(request.form["resident_id"])
            resident=c.execute("SELECT id FROM residents WHERE id=?",(rid,)).fetchone()
            if not resident: raise ValueError(L("Ce résident n’existe pas.","Resident does not exist.","هذا المقيم غير موجود."))
            supervisor=request.form["supervisor"].strip(); destination=request.form["destination"].strip()
            hospital=request.form["hospital"].strip(); country=request.form["country"].strip(); kind=request.form["rotation_type"].strip()
            start=date.fromisoformat(request.form["start_date"]); end=date.fromisoformat(request.form["end_date"])
            if not all([supervisor,destination,hospital,country,kind]): raise ValueError("All external rotation fields are required.")
            if end<start: raise ValueError("Return date must not precede start date.")
            overlap=c.execute("SELECT id FROM external_rotations WHERE resident_id=? AND request_status NOT IN ('rejected','cancelled') AND start_date<=? AND end_date>=?",(rid,end.isoformat(),start.isoformat())).fetchone()
            if overlap: raise ValueError("External rotation requests for this resident cannot overlap.")
            if is_admin: reject_external_overlap_with_local_assignment(c,rid,start,end)
            request_status="approved" if is_admin else "pending"
            status=("returned" if end<date.today() else "active") if is_admin else "requested"
            c.execute("INSERT INTO external_rotations(resident_id,supervisor,destination,hospital,country,rotation_type,start_date,end_date,notes,status,request_status,created_by) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",(rid,supervisor,destination,hospital,country,kind,start.isoformat(),end.isoformat(),request.form.get("notes","").strip(),status,request_status,session.get("user","")))
            audit_event(c,"external_rotation.created" if is_admin else "external_rotation.requested",f"resident_id={rid}; hospital={hospital}; country={country}; start={start}; end={end}; request_status={request_status}")
            c.commit(); flash(L("Stage externe enregistré et actif." if is_admin else "Demande de stage à l’étranger envoyée ; elle ne bloque pas les affectations avant son approbation.","External rotation saved and active." if is_admin else "External internship request submitted; it will not block assignments until approved.","تم تسجيل التربص الخارجي وتفعيله." if is_admin else "تم إرسال طلب التربص بالخارج؛ لن يمنع التعيينات قبل الموافقة عليه."))
        except Exception as exc: c.rollback(); app.logger.warning("External rotation save failed: %s",exc); flash(L("Le stage externe n’a pas été enregistré : ","External rotation was not saved: ","لم يتم حفظ التربص الخارجي: ")+str(exc)[:220])
        c.close(); return redirect("/external-rotations")
    residents=c.execute("SELECT id,surname,firstname,year FROM residents ORDER BY year DESC,surname").fetchall()
    selected_year=active_postgrad_year(c)
    year_start=date(int(selected_year.split("-")[0]),2,1)
    year_end=date(int(selected_year.split("-")[1]),1,31)
    today=date.today().isoformat()
    c.execute("UPDATE external_rotations SET status='returned',updated_at=CURRENT_TIMESTAMP WHERE request_status='approved' AND status='active' AND end_date<?",(today,))
    rows=c.execute("SELECT x.*,r.surname,r.firstname FROM external_rotations x JOIN residents r ON r.id=x.resident_id WHERE x.start_date<=? AND x.end_date>=? ORDER BY x.start_date DESC",(year_end.isoformat(),year_start.isoformat())).fetchall()
    c.commit(); c.close()
    return page("""<section><h3>{{L('Stages à l’étranger','External rotations','التربصات بالخارج')}} — {{year}}</h3><p>{{L('Les utilisateurs peuvent soumettre une demande. Seuls les stages approuvés bloquent les affectations locales qui se chevauchent ; les Super Admins peuvent enregistrer directement un stage approuvé.','Standard users may submit an internship request. Only approved internships block overlapping local assignments; Super Admins may register approved stages directly.','يمكن للمستخدمين إرسال طلب تربص. التربصات المعتمدة فقط تمنع التعيينات المحلية المتداخلة؛ ويمكن للمسؤولين العامين تسجيل التربص المعتمد مباشرة.')}}</p><form method="post"><label>{{L('Résident','Resident','المقيم')}} <select name="resident_id" required>{% for r in residents %}<option value="{{r.id}}">{{r.surname}} {{r.firstname}} ({{r.year}})</option>{% endfor %}</select></label><div class="grid"><label>{{L('Responsable / Encadrant','Responsible employee / Supervisor','المسؤول / المشرف')}} <input name="supervisor" required maxlength="160"></label><label>{{L('الوجهة' if lang=='ar' else 'Destination','Destination','الوجهة')}} <input name="destination" required maxlength="160"></label><label>{{L('Hôpital','Hospital','المستشفى')}} <input name="hospital" required maxlength="160"></label><label>{{L('Pays','Country','البلد')}} <input name="country" required maxlength="100"></label><label>{{L('Type de stage','Rotation type','نوع التربص')}} <input name="rotation_type" required maxlength="100"></label><label>{{L('Date de début','Start date','تاريخ البداية')}} <input type="date" name="start_date" required></label><label>{{L('Date de retour','Return date','تاريخ العودة')}} <input type="date" name="end_date" required></label></div><label>{{L('Notes','Notes','ملاحظات')}} <input name="notes" maxlength="1000"></label><button>{% if role=='super_admin' %}{{L('Enregistrer le stage approuvé','Save approved rotation','تسجيل التربص المعتمد')}}{% else %}{{L('Soumettre la demande de stage','Submit internship request','إرسال طلب التربص')}}{% endif %}</button></form></section><section><h3>{{L('Historique des stages à l’étranger','External rotation history','سجل التربصات بالخارج')}}</h3><div class="wrap"><table><tr><th>{{L('Résident','Resident','المقيم')}}</th><th>{{L('Encadrant','Supervisor','المشرف')}}</th><th>{{L('Destination','Destination','الوجهة')}}</th><th>{{L('Hôpital','Hospital','المستشفى')}}</th><th>{{L('Pays','Country','البلد')}}</th><th>{{L('Type','Type','النوع')}}</th><th>{{L('Dates','Dates','التواريخ')}}</th><th>{{L('Approbation / état','Approval / status','الموافقة / الحالة')}}</th><th>{{L('Actions','Actions','الإجراءات')}}</th></tr>{% for x in rows %}<tr><td>{{x.surname}} {{x.firstname}}</td><td>{{x.supervisor}}</td><td>{{x.destination}}</td><td>{{x.hospital}}</td><td>{{x.country}}</td><td>{{x.rotation_type}}</td><td>{{x.start_date}} – {{x.end_date}}</td><td>{{L('En attente' if x.request_status=='pending' else 'Approuvé' if x.request_status=='approved' else 'Refusé' if x.request_status=='rejected' else 'Annulé' if x.request_status=='cancelled' else 'Demandé','Pending' if x.request_status=='pending' else 'Approved' if x.request_status=='approved' else 'Rejected' if x.request_status=='rejected' else 'Cancelled' if x.request_status=='cancelled' else 'Requested','معلق' if x.request_status=='pending' else 'موافق عليه' if x.request_status=='approved' else 'مرفوض' if x.request_status=='rejected' else 'ملغى' if x.request_status=='cancelled' else 'مطلوب')}} / {{L('Actif' if x.status=='active' else 'Terminé' if x.status=='returned' else 'Annulé' if x.status=='cancelled' else 'Demandé' if x.status=='requested' else 'Refusé' if x.status=='rejected' else x.status,'Active' if x.status=='active' else 'Returned' if x.status=='returned' else 'Cancelled' if x.status=='cancelled' else 'Requested' if x.status=='requested' else 'Rejected' if x.status=='rejected' else x.status,'نشط' if x.status=='active' else 'عاد' if x.status=='returned' else 'ملغى' if x.status=='cancelled' else 'مطلوب' if x.status=='requested' else 'مرفوض' if x.status=='rejected' else x.status)}}</td><td>{% if role=='super_admin' and x.request_status=='pending' %}<form method="post" action="/external-rotations/{{x.id}}/decision" style="display:inline"><input type="hidden" name="decision" value="approve"><button>{{L('Approuver','Approve','موافقة')}}</button></form><form method="post" action="/external-rotations/{{x.id}}/decision" style="display:inline"><input type="hidden" name="decision" value="reject"><button>{{L('Refuser','Reject','رفض')}}</button></form>{% endif %}{% if role=='super_admin' and x.request_status!='rejected' and x.status!='cancelled' %}<a href="/external-rotations/{{x.id}}/edit">{{L('Modifier','Edit','تعديل')}}</a><form method="post" action="/external-rotations/{{x.id}}/cancel" style="display:inline"><button>{{L('Annuler','Cancel','إلغاء')}}</button></form>{% endif %}</td></tr>{% endfor %}</table></div></section>""",residents=residents,rows=rows,year=selected_year)

@app.post("/external-rotations/<int:xid>/decision")
def decide_external_rotation(xid):
    if guard(): return redirect("/login")
    denied=super_guard()
    if denied: return denied
    c=db()
    row=c.execute("SELECT * FROM external_rotations WHERE id=?",(xid,)).fetchone()
    if not row:
        c.close(); return ("External rotation request not found",404)
    if row["request_status"]!="pending":
        c.close(); flash(L("Seules les demandes de stage en attente peuvent faire l’objet d’une décision.","Only pending external internship requests can be decided.","يمكن اتخاذ قرار بشأن طلبات التربص المعلقة فقط.")); return redirect("/external-rotations")
    decision=request.form.get("decision","")
    if decision not in ("approve","reject"):
        c.close(); return ("Invalid request decision",400)
    try:
        if decision=="approve":
            start=date.fromisoformat(row["start_date"]); end=date.fromisoformat(row["end_date"])
            overlap=c.execute("SELECT id FROM external_rotations WHERE resident_id=? AND id<>? AND request_status='approved' AND status<>'cancelled' AND start_date<=? AND end_date>=?",(row["resident_id"],xid,end.isoformat(),start.isoformat())).fetchone()
            if overlap: raise ValueError("This date range overlaps another approved external internship.")
            reject_external_overlap_with_local_assignment(c,row["resident_id"],start,end)
            status="returned" if end<date.today() else "active"
            c.execute("UPDATE external_rotations SET request_status='approved',status=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(status,xid))
        else:
            c.execute("UPDATE external_rotations SET request_status='rejected',status='rejected',updated_at=CURRENT_TIMESTAMP WHERE id=?",(xid,))
        audit_event(c,"external_rotation.request_approved" if decision=="approve" else "external_rotation.request_rejected",f"id={xid}; resident_id={row['resident_id']}; decision={decision}; reviewer={session.get('user','')}")
        c.commit(); flash(L("Demande de stage à l’étranger "+("approuvée." if decision=="approve" else "refusée."),"External internship request "+("approved." if decision=="approve" else "rejected."),"طلب التربص بالخارج "+("تمت الموافقة عليه." if decision=="approve" else "تم رفضه.")))
    except Exception as exc:
        c.rollback(); app.logger.warning("External internship decision failed: %s",exc); flash(L("La décision n’a pas été enregistrée : ","Request decision was not saved: ","لم يتم حفظ القرار: ")+str(exc)[:220])
    c.close(); return redirect("/external-rotations")

@app.route("/external-rotations/<int:xid>/edit", methods=["GET","POST"])
def edit_external_rotation(xid):
    if guard(): return redirect("/login")
    denied=super_guard()
    if denied: return denied
    c=db()
    row=c.execute("SELECT * FROM external_rotations WHERE id=?",(xid,)).fetchone()
    if not row:
        c.close(); return ("External rotation not found",404)
    if row["status"]=="cancelled":
        c.close(); return ("Cancelled external rotations are immutable",409)
    if request.method=="POST":
        try:
            rid=int(request.form["resident_id"])
            if not c.execute("SELECT 1 FROM residents WHERE id=?",(rid,)).fetchone(): raise ValueError(L("Ce résident n’existe pas.","Resident does not exist.","هذا المقيم غير موجود."))
            supervisor=request.form["supervisor"].strip(); destination=request.form["destination"].strip()
            hospital=request.form["hospital"].strip(); country=request.form["country"].strip(); kind=request.form["rotation_type"].strip()
            start=date.fromisoformat(request.form["start_date"]); end=date.fromisoformat(request.form["end_date"])
            if not all([supervisor,destination,hospital,country,kind]): raise ValueError("All external rotation fields are required.")
            if end<start: raise ValueError("Return date must not precede start date.")
            overlap=c.execute("SELECT id FROM external_rotations WHERE resident_id=? AND id<>? AND request_status NOT IN ('rejected','cancelled') AND start_date<=? AND end_date>=?",(rid,xid,end.isoformat(),start.isoformat())).fetchone()
            if overlap: raise ValueError("External rotations for this resident cannot overlap.")
            if row["request_status"]=="approved": reject_external_overlap_with_local_assignment(c,rid,start,end)
            old=f"resident_id={row['resident_id']}; hospital={row['hospital']}; country={row['country']}; start={row['start_date']}; end={row['end_date']}"
            status="returned" if end<date.today() else "active"
            c.execute("UPDATE external_rotations SET resident_id=?,supervisor=?,destination=?,hospital=?,country=?,rotation_type=?,start_date=?,end_date=?,notes=?,status=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(rid,supervisor,destination,hospital,country,kind,start.isoformat(),end.isoformat(),request.form.get("notes","").strip(),status,xid))
            audit_event(c,"external_rotation.updated",f"id={xid}; old=({old}); new_resident_id={rid}; hospital={hospital}; country={country}; start={start}; end={end}")
            c.commit(); c.close(); flash("External rotation updated and audited."); return redirect("/external-rotations")
        except Exception as exc:
            c.rollback(); flash(L("Le stage à l’étranger n’a pas été modifié : ","External rotation not updated: ","لم يتم تعديل التربص بالخارج: ")+str(exc)[:220])
    residents=c.execute("SELECT id,surname,firstname,year FROM residents ORDER BY year DESC,surname").fetchall()
    c.close()
    return page("""<section><h3>{{L('Modifier le stage à l’étranger','Edit external rotation','تعديل التربص بالخارج')}}</h3><form method="post"><label>{{L('Résident','Resident','المقيم')}} <select name="resident_id" required>{% for r in residents %}<option value="{{r.id}}" {% if r.id==row.resident_id %}selected{% endif %}>{{r.surname}} {{r.firstname}} ({{r.year}})</option>{% endfor %}</select></label><div class="grid"><label>{{L('Responsable / Encadrant','Responsible employee / Supervisor','المسؤول / المشرف')}} <input name="supervisor" required maxlength="160" value="{{row.supervisor}}"></label><label>{{L('الوجهة' if lang=='ar' else 'Destination','Destination','الوجهة')}} <input name="destination" required maxlength="160" value="{{row.destination}}"></label><label>{{L('Hôpital','Hospital','المستشفى')}} <input name="hospital" required maxlength="160" value="{{row.hospital}}"></label><label>{{L('Pays','Country','البلد')}} <input name="country" required maxlength="100" value="{{row.country}}"></label><label>{{L('Type de stage','Rotation type','نوع التربص')}} <input name="rotation_type" required maxlength="100" value="{{row.rotation_type}}"></label><label>{{L('Date de début','Start date','تاريخ البداية')}} <input type="date" name="start_date" required value="{{row.start_date}}"></label><label>{{L('Date de retour','Return date','تاريخ العودة')}} <input type="date" name="end_date" required value="{{row.end_date}}"></label></div><label>{{L('Notes','Notes','ملاحظات')}} <input name="notes" maxlength="1000" value="{{row.notes}}"></label><button>{{L('Enregistrer les modifications','Save changes','حفظ التعديلات')}}</button></form><p><a href="/external-rotations">{{L('Retour aux stages à l’étranger','Back to external rotations','العودة إلى التربصات بالخارج')}}</a></p></section>""",residents=residents,row=row)

@app.post("/external-rotations/<int:xid>/cancel")
def cancel_external_rotation(xid):
    if guard(): return redirect("/login")
    denied=super_guard()
    if denied: return denied
    c=db()
    row=c.execute("SELECT status FROM external_rotations WHERE id=?",(xid,)).fetchone()
    if not row:
        c.close(); return ("External rotation not found",404)
    if row["status"]=="cancelled":
        c.close(); flash(L("Ce stage externe est déjà annulé.","External rotation is already cancelled.","هذا التربص الخارجي ملغى بالفعل.")); return redirect("/external-rotations")
    c.execute("UPDATE external_rotations SET status='cancelled',request_status='cancelled',updated_at=CURRENT_TIMESTAMP WHERE id=?",(xid,))
    audit_event(c,"external_rotation.cancelled",f"id={xid}; previous_status={row['status']}")
    c.commit(); c.close(); flash(L("Stage externe annulé ; son historique reste conservé dans le journal d’audit.","External rotation cancelled; its audit history is retained.","تم إلغاء التربص الخارجي؛ يبقى سجله محفوظًا في سجل التدقيق."))
    return redirect("/external-rotations")

if __name__=="__main__":
 init(); app.run(host="127.0.0.1",port=int(os.environ.get("PORT","5000")),debug=False)
